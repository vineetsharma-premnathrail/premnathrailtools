"""Technical Offer by UPLOAD (R&D rework 2026-10-30): R&D uploads a Word / PDF offer on an accepted
requirement, the R&D head releases or returns it, BD can upload its own offer (quotable at once),
and quotations pick released offers. Word creation / "Edit in Word" no longer exist."""
import io
from datetime import date, datetime, timezone

import pytest
from docx import Document

from app.core.config import settings
from app.modules.crm.models.inquiry import Quotation
from app.modules.rnd.models.technical_offer import RndTechnicalOffer
from app.modules.rnd.services import technical_offers as to
from app.tests.quote_lines import quote_product  # noqa: F401 — autouse fixture
from app.tests.test_cov_crm_quotations_rnd import ITEM, auth, make_inquiry, make_user, send_tor, tor_env  # noqa: F401 — fixture
from app.tests.test_crm_quotation_documents import sp  # noqa: F401 — fixture
from app.tests.test_rnd_intake import _accept, _setup

RND = "/api/v1/rnd/requirement-work"
CRM = "/api/v1/crm"
TODAY = date.today().isoformat()


def join_offer_team(db, user, create=True):
    """R&D › Requirements › Technical Offer team, given through a role in Role & Permissions."""
    from app.modules.main.models.role import Role
    grants = ["rnd:requirements.offer_team:view"] + (["rnd:requirements.offer_team:create"] if create else [])
    role = Role(name=f"Offer team {user.id}", permissions=grants, is_active=True)
    db.add(role)
    user.roles.append(role)
    db.commit()


def _docx(*lines: str) -> bytes:
    d = Document()
    for line in lines or ("Technical offer",):
        d.add_paragraph(line)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


@pytest.fixture
def store(monkeypatch):
    """A fake file store behind the offer upload / download."""
    files: dict[str, bytes] = {}

    async def up(site, folder, filename, content_type, data):
        path = f"{folder}/{filename}"
        files[path] = data
        return {"name": filename, "path": path, "webUrl": "https://sp/x", "size": len(data)}

    async def down(site, path):
        return files[path], "application/octet-stream"

    monkeypatch.setattr(settings, "SHAREPOINT_SITE_ID", "site-1")
    monkeypatch.setattr(to, "upload_bytes_to_sharepoint", up)
    monkeypatch.setattr(to, "download_file_content", down)
    return files


def _upload(client, inq, who, data=None, name="offer.docx", **form):
    data = _docx() if data is None else data
    fields = {"title": "Rail Grinding Machine", "offer_date": TODAY, **form}
    fields = {k: v for k, v in fields.items() if v is not None}
    return client.post(f"{RND}/inquiry/{inq.id}/technical-offers", data=fields, files={"file": (name, data, "application/octet-stream")}, headers=auth(who))


def _env(client, db):
    """Accepted requirement + an offer-team uploader."""
    bd, head, member, inq = _setup(client, db)
    assert _accept(client, inq, head).status_code == 200
    uploader = make_user(db, "to.up@premnathrail.com", apps=("rnd",))
    join_offer_team(db, uploader)
    return bd, head, member, uploader, inq


# ── R&D upload ────────────────────────────────────────────────────────────

def test_upload_waits_for_head_release_and_bd_sees_it_only_after(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    resp = _upload(client, inq, uploader, notes="For the RDSO trial")
    assert resp.status_code == 201, resp.text
    offer = resp.json()
    assert offer["offer_number"] == f"TO-{inq.universal_id}" and offer["status"] == "submitted" and offer["status_label"] == "Pending release"
    assert offer["origin"] == "rnd" and offer["requirement_revision_no"] == 0 and offer["notes"] == "For the RDSO trial"
    assert offer["file_name"].endswith(".docx") and offer["can_delete"] is True and offer["based_on_earlier_revision"] is False
    assert list(store.values())[0].startswith(b"PK")
    # Pending / returned offers are not in BD's picker and can't be quoted (AC-30).
    assert client.get(f"{CRM}/inquiries/{inq.id}/technical-offers", headers=auth(bd)).json() == []
    bad = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": offer["id"]}, headers=auth(bd))
    assert bad.status_code == 422 and "waiting for the R&D head's release" in bad.json()["detail"]

    released = client.post(f"{RND}/technical-offers/{offer['id']}/release", headers=auth(head))
    assert released.status_code == 200 and released.json()["status"] == "released"
    listed = client.get(f"{CRM}/inquiries/{inq.id}/technical-offers", headers=auth(bd)).json()
    assert [(o["offer_number"], o["origin"], o["stacked_in_quotation"]) for o in listed] == [(f"TO-{inq.universal_id}", "rnd", True)]
    quot = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": offer["id"]}, headers=auth(bd))
    assert quot.status_code == 201, quot.text
    assert quot.json()["technical_offer_number"] == f"TO-{inq.universal_id}"


def test_upload_rejects_bad_files_with_exact_reasons(client, db, tor_env, store, monkeypatch):
    bd, head, member, uploader, inq = _env(client, db)
    only = "Only Word (.docx) or PDF files can be uploaded as a Technical Offer."
    resp = _upload(client, inq, uploader, data=b"hello", name="offer.txt")
    assert resp.status_code == 422 and resp.json()["detail"] == only
    resp = _upload(client, inq, uploader, data=b"", name="offer.pdf")
    assert resp.status_code == 422 and resp.json()["detail"] == "The file is empty."
    resp = _upload(client, inq, uploader, data=_docx(), name="offer.pdf")        # a Word file named .pdf
    assert resp.status_code == 422 and "isn't a real PDF" in resp.json()["detail"]
    resp = _upload(client, inq, uploader, data=PDF, name="offer.docx")
    assert resp.status_code == 422 and "isn't a real DOCX" in resp.json()["detail"]
    monkeypatch.setattr(to, "MAX_BYTES", 100)
    resp = _upload(client, inq, uploader, data=_docx("x" * 500))
    assert resp.status_code == 413 and "the limit is" in resp.json()["detail"] and "MB" in resp.json()["detail"]
    monkeypatch.undo()
    assert _upload(client, inq, uploader, title="").status_code == 422
    assert _upload(client, inq, uploader, offer_date=None).status_code == 422
    resp = client.post(f"{RND}/inquiry/{inq.id}/technical-offers", data={"title": "T", "offer_date": TODAY}, headers=auth(uploader))
    assert resp.status_code == 422 and only in resp.json()["detail"]
    assert db.query(RndTechnicalOffer).count() == 0


def test_permissions_upload_release_and_bd_refused(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    resp = _upload(client, inq, member)                      # R&D member, not on the offer team
    assert resp.status_code == 403 and "Technical Offer team" in resp.json()["detail"]
    assert _upload(client, inq, bd).status_code == 403       # BD (AC-35)
    offer = _upload(client, inq, uploader).json()
    url = f"{RND}/technical-offers/{offer['id']}"
    for who in (bd, member, uploader):
        resp = client.post(f"{url}/release", headers=auth(who))
        assert resp.status_code == 403, who.email
    assert client.get(f"{RND}/technical-offers/{offer['id']}/file", headers=auth(bd)).status_code == 403
    assert db.get(RndTechnicalOffer, offer["id"]).status == "submitted"


def test_return_needs_reason_then_corrected_upload_is_next_revision(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    first = _upload(client, inq, uploader).json()
    url = f"{RND}/technical-offers/{first['id']}"
    assert client.post(f"{url}/return", json={"reason": "  "}, headers=auth(head)).status_code == 422
    assert client.post(f"{url}/return", json={}, headers=auth(head)).status_code == 422
    resp = client.post(f"{url}/return", json={"reason": "Add the brake calculation"}, headers=auth(head))
    assert resp.status_code == 200 and resp.json()["status_label"] == "Returned" and resp.json()["returned_reason"] == "Add the brake calculation"
    assert client.post(f"{url}/release", headers=auth(head)).status_code == 409
    second = _upload(client, inq, uploader, data=_docx("Corrected"), revise_of_id=str(first["id"]))
    assert second.status_code == 201, second.text
    assert second.json()["offer_number"] == f"TO-{inq.universal_id}-R1" and second.json()["revision"] == 1
    # While one is pending, another upload is refused.
    resp = _upload(client, inq, uploader)
    assert resp.status_code == 409 and "waiting for the R&D head's release" in resp.json()["detail"]
    assert client.post(f"{RND}/technical-offers/{second.json()['id']}/release", headers=auth(head)).status_code == 200
    third = _upload(client, inq, uploader).json()
    assert third["offer_number"].endswith("-R2")
    client.post(f"{RND}/technical-offers/{third['id']}/release", headers=auth(head))
    status = {o.offer_number: o.status for o in db.query(RndTechnicalOffer).all()}
    assert status[f"TO-{inq.universal_id}-R1"] == "superseded" and status[f"TO-{inq.universal_id}-R2"] == "released"
    # A stale revise_of_id is refused with what to do.
    resp = _upload(client, inq, uploader, revise_of_id=str(first["id"]))
    assert resp.status_code == 409 and "latest" in resp.json()["detail"]


def test_upload_completes_a_migrated_draft(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    draft = RndTechnicalOffer(source="inquiry", record_id=inq.id, origin="rnd", offer_number_base=f"TO-{inq.universal_id}", revision=0,
                              offer_number=f"TO-{inq.universal_id}", offer_date=date.today(), title="Old draft", status="draft",
                              returned_reason="Offer creation was replaced by upload — upload the final file", created_by_id=uploader.id)
    db.add(draft)
    db.commit()
    resp = _upload(client, inq, uploader)
    assert resp.status_code == 201 and resp.json()["id"] == draft.id and resp.json()["status"] == "submitted"
    assert db.query(RndTechnicalOffer).count() == 1


def test_patch_and_delete_only_before_release(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    offer = _upload(client, inq, uploader).json()
    url = f"{RND}/technical-offers/{offer['id']}"
    assert client.patch(url, json={"title": ""}, headers=auth(uploader)).status_code == 422
    assert client.patch(url, json={"title": "New title", "notes": "n"}, headers=auth(uploader)).json()["title"] == "New title"
    assert client.patch(url, json={"title": "x"}, headers=auth(member)).status_code == 403
    client.post(f"{url}/release", headers=auth(head))
    assert client.patch(url, json={"title": "Late"}, headers=auth(uploader)).status_code == 409
    resp = client.delete(url, headers=auth(uploader))
    assert resp.status_code == 409 and "Upload a new revision" in resp.json()["detail"]
    pending = _upload(client, inq, uploader).json()
    assert client.delete(f"{RND}/technical-offers/{pending['id']}", headers=auth(member)).status_code == 403
    assert client.delete(f"{RND}/technical-offers/{pending['id']}", headers=auth(uploader)).json() == {"ok": True}


def test_offer_file_follows_requirement_access(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    offer = _upload(client, inq, uploader, data=_docx("Secret spec")).json()
    url = f"{RND}/technical-offers/{offer['id']}/file"
    ok = client.get(url, headers=auth(member))
    assert ok.status_code == 200 and ok.content.startswith(b"PK") and ".docx" in ok.headers["content-disposition"]
    designer = make_user(db, "to.design@premnathrail.com", apps=("design",))
    assert client.get(url, headers=auth(designer)).status_code == 403          # no access, no file (IDOR)
    due = date.today().isoformat()
    client.post(f"{RND}/inquiry/{inq.id}/tasks", data={"title": "Check", "assigned_to_id": str(designer.id), "due_date": due}, headers=auth(head))
    assert client.get(url, headers=auth(designer)).status_code == 200            # a task gives View + Download
    assert client.get(f"{RND}/technical-offers/99999/file", headers=auth(head)).status_code == 404


def test_offer_marks_earlier_revision_after_a_resend(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    offer = _upload(client, inq, uploader).json()
    client.post(f"{RND}/technical-offers/{offer['id']}/release", headers=auth(head))
    inq.product = "Rail cutter Mk2"
    inq.product_spec = "BG 1676, revised"  # a re-send needs a changed Customer Requirement
    db.commit()
    send_tor(client, inq.id, bd)
    rows = client.get(f"{RND}/inquiry/{inq.id}/technical-offers", headers=auth(head)).json()
    assert rows[0]["based_on_earlier_revision"] is True and rows[0]["requirement_revision_no"] == 0
    picker = client.get(f"{CRM}/inquiries/{inq.id}/technical-offers", headers=auth(bd)).json()
    assert picker[0]["based_on_earlier_revision"] is True


# ── BD's own offer ────────────────────────────────────────────────────────

def test_bd_offer_without_requirement_is_quotable_at_once(client, db, tor_env, store):
    bd = make_user(db, "bto.bd@premnathrail.com")
    inq = make_inquiry(db, bd)              # never sent to R&D
    resp = client.post(f"{CRM}/inquiries/{inq.id}/technical-offers", data={"title": "Our offer", "offer_date": TODAY},
                       files={"file": ("bd.pdf", PDF, "application/pdf")}, headers=auth(bd))
    assert resp.status_code == 201, resp.text
    offer = resp.json()
    assert offer["offer_number"] == f"BTO-{inq.universal_id}" and offer["status"] == "released" and offer["origin"] == "bd"
    assert offer["label"] == "Uploaded by BD, not reviewed by R&D" and offer["file_type"] == "pdf" and offer["stacked_in_quotation"] is False
    quot = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": offer["id"]}, headers=auth(bd))
    assert quot.status_code == 201, quot.text
    # A newer BD upload supersedes the first but it stays quotable for the existing quotation.
    second = client.post(f"{CRM}/inquiries/{inq.id}/technical-offers", data={"title": "Our offer v2", "offer_date": TODAY},
                         files={"file": ("bd2.docx", _docx(), "application/octet-stream")}, headers=auth(bd)).json()
    assert second["offer_number"].endswith("-R1")
    listed = client.get(f"{CRM}/inquiries/{inq.id}/technical-offers", headers=auth(bd)).json()
    assert [(o["offer_number"], o["status"]) for o in listed] == [(second["offer_number"], "released"), (offer["offer_number"], "superseded")]
    assert listed[0]["can_delete"] is True and listed[1]["can_delete"] is False and listed[1]["quoted_in"]


def test_bd_offer_delete_blocked_while_a_quotation_uses_it(client, db, tor_env, store):
    bd = make_user(db, "bto2.bd@premnathrail.com")
    inq = make_inquiry(db, bd)
    offer = client.post(f"{CRM}/inquiries/{inq.id}/technical-offers", data={"title": "Our offer", "offer_date": TODAY},
                        files={"file": ("bd.pdf", PDF, "application/pdf")}, headers=auth(bd)).json()
    quot = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": offer["id"]}, headers=auth(bd)).json()
    resp = client.delete(f"{CRM}/technical-offers/{offer['id']}", headers=auth(bd))
    assert resp.status_code == 409 and resp.json()["detail"] == f"Used by quotation {quot['quot_number']} — remove it from the quotation first."
    client.patch(f"{CRM}/inquiries/{inq.id}/quotations/{quot['id']}", json={"technical_offer_id": None}, headers=auth(bd))
    assert client.delete(f"{CRM}/technical-offers/{offer['id']}", headers=auth(bd)).status_code == 200
    assert db.query(RndTechnicalOffer).count() == 0


def test_bd_cannot_delete_rnd_offer_or_upload_outside_scope(client, db, tor_env, store):
    bd, head, member, uploader, inq = _env(client, db)
    offer = _upload(client, inq, uploader).json()
    resp = client.delete(f"{CRM}/technical-offers/{offer['id']}", headers=auth(bd))
    assert resp.status_code == 403 and "R&D's offer" in resp.json()["detail"]
    # An R&D-only user has no CRM access: the BD upload endpoints refuse.
    resp = client.post(f"{CRM}/inquiries/{inq.id}/technical-offers", data={"title": "x", "offer_date": TODAY},
                       files={"file": ("a.pdf", PDF, "application/pdf")}, headers=auth(member))
    assert resp.status_code == 403
    assert client.delete(f"{CRM}/technical-offers/{offer['id']}", headers=auth(member)).status_code == 403


def test_bd_offer_on_a_tender_and_scope_check(client, db, tor_env, store):
    from app.modules.crm.models.tender import Tender
    bd = make_user(db, "bto3.bd@premnathrail.com")
    from app.modules.crm.models.organization import Organization
    org = Organization(name="Org for tender", org_type="Customer")
    db.add(org)
    db.flush()
    tender = Tender(universal_id="TND-2026-0001", org_id=org.id, created_by_id=bd.id, tender_name="Tender X")
    db.add(tender)
    db.commit()
    resp = client.post(f"{CRM}/tenders/{tender.id}/technical-offers", data={"title": "Tender offer", "offer_date": TODAY},
                       files={"file": ("t.pdf", PDF, "application/pdf")}, headers=auth(bd))
    assert resp.status_code == 201 and resp.json()["offer_number"] == "BTO-TND-2026-0001"
    listed = client.get(f"{CRM}/tenders/{tender.id}/technical-offers", headers=auth(bd)).json()
    assert len(listed) == 1
    got = client.get(f"{CRM}/tenders/{tender.id}/technical-offers/{listed[0]['id']}/file", headers=auth(bd))
    assert got.status_code == 200 and got.content.startswith(b"%PDF")
    assert client.get(f"{CRM}/tenders/99999/technical-offers", headers=auth(bd)).status_code == 404


# ── quotation document ────────────────────────────────────────────────────

def _text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_docx_offer_is_stacked_pdf_offer_is_attached_separately(client, db, tor_env, store, sp):
    bd, head, member, uploader, inq = _env(client, db)
    word = _upload(client, inq, uploader, data=_docx("UPLOADED TECHNICAL OFFER", "Gauge 1676 mm")).json()
    client.post(f"{RND}/technical-offers/{word['id']}/release", headers=auth(head))
    quot = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": word["id"]}, headers=auth(bd)).json()
    assert client.post(f"/api/v1/managed-documents/crm_quotation/{quot['id']}", headers=auth(bd)).status_code == 200
    text = _text(sp.files[max(sp.files)][-1])
    assert "UPLOADED TECHNICAL OFFER" in text and text.index("UPLOADED TECHNICAL OFFER") < text.index("Amount in words")

    # BD's PDF offer can't be stacked into Word: the quotation Word is commercial only; the PDF is served separately.
    pdf_offer = client.post(f"{CRM}/inquiries/{inq.id}/technical-offers", data={"title": "BD pdf", "offer_date": TODAY},
                            files={"file": ("bd.pdf", PDF, "application/pdf")}, headers=auth(bd)).json()
    quot2 = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM], "technical_offer_id": pdf_offer["id"]}, headers=auth(bd)).json()
    assert client.post(f"/api/v1/managed-documents/crm_quotation/{quot2['id']}", headers=auth(bd)).status_code == 200
    assert _text(sp.files[max(sp.files)][-1]).startswith("QUOTATION")
    sep = client.get(f"{CRM}/inquiries/{inq.id}/quotations/{quot2['id']}/technical-offer-file", headers=auth(bd))
    assert sep.status_code == 200 and sep.content == PDF


def test_quotation_rejects_an_offer_of_another_inquiry(client, db, tor_env, store):
    bd = make_user(db, "bto4.bd@premnathrail.com")
    a = make_inquiry(db, bd, uid="INQ-20261007-0001")
    b = make_inquiry(db, bd, uid="INQ-20261007-0002")
    offer = client.post(f"{CRM}/inquiries/{a.id}/technical-offers", data={"title": "A", "offer_date": TODAY},
                        files={"file": ("a.pdf", PDF, "application/pdf")}, headers=auth(bd)).json()
    resp = client.post(f"{CRM}/inquiries/{b.id}/quotations", json={"items": [ITEM], "technical_offer_id": offer["id"]}, headers=auth(bd))
    assert resp.status_code == 422 and "doesn't belong to this inquiry" in resp.json()["detail"]
    assert client.get(f"{CRM}/inquiries/{b.id}/technical-offers/{offer['id']}/file", headers=auth(bd)).status_code == 422
