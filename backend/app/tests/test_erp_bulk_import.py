"""Bulk import / export for Machines and Service Requests — duplicates are
skipped and reported, bad rows come back with a reason, good rows are saved."""
import io

from openpyxl import load_workbook

from app.modules.erp.models.service_request import ServiceRequest
from app.tests.test_erp_projects import auth_header, make_user


def _csv(text: str):
    return {"file": ("data.csv", io.BytesIO(text.encode("utf-8")), "text/csv")}


def _admin(db, email="bulkadmin@premnathrail.com"):
    return make_user(db, email, role="admin")


def test_machine_import_skips_and_reports_duplicates(client, db):
    h = auth_header(_admin(db))
    assert client.post("/api/v1/erp/projects", json={"serial_number": "SN-100", "engine_number": "ENG-1"}, headers=h).status_code == 201

    body = (
        "serial_number,model_name,engine_number,chassis_number,status,po_date,is_export\n"
        "SN-200,RRV,ENG-2,CH-2,Active,15-08-2025,No\n"      # new
        "sn-100 ,RRV,,,,,\n"                                # same serial as DB (case/space-insensitive)
        "SN-200,RRV,,,,,\n"                                 # repeats row 2
        "SN-300,RRV,ENG-1,,,,\n"                            # engine belongs to SN-100
        "SN-400,RRV,,CH-2,,,\n"                             # chassis repeats row 2
        "SN-500,RRV,,,Flying,,\n"                           # bad status
        "SN-600,RRV,,,,31-31-2025,\n"                       # bad date
        ",RRV,,,,,\n"                                       # no serial
    )
    res = client.post("/api/v1/erp/projects/import", files=_csv(body), headers=h)
    assert res.status_code == 200, res.text
    data = res.json()
    assert [c["ref"] for c in data["created"]] == ["SN-200"]
    assert {e["row"] for e in data["existing"]} == {3, 4}
    reasons = {e["row"]: e["reason"] for e in data["errors"]}
    assert set(reasons) == {5, 6, 7, 8, 9}
    assert "SN-100" in reasons[5]
    assert "row 2" in reasons[6]
    assert "Status" in reasons[7] and "Active" in reasons[7]
    assert "DD-MM-YYYY" in reasons[8]


def test_machine_in_recycle_bin_counts_as_existing(client, db):
    h = auth_header(_admin(db, "bulkadmin2@premnathrail.com"))
    pid = client.post("/api/v1/erp/projects", json={"serial_number": "SN-BIN"}, headers=h).json()["id"]
    assert client.delete(f"/api/v1/erp/projects/{pid}", headers=h).status_code in (200, 204)
    data = client.post("/api/v1/erp/projects/import", files=_csv("serial_number\nSN-BIN\n"), headers=h).json()
    assert not data["created"]
    assert "Recycle Bin" in data["existing"][0]["reason"]


def test_machine_import_needs_create_permission(client, db):
    user = make_user(db, "bulk_noperm@premnathrail.com", assigned_apps=["erp"])
    res = client.post("/api/v1/erp/projects/import", files=_csv("serial_number\nSN-X\n"), headers=auth_header(user))
    assert res.status_code == 403


def test_machine_export_round_trips_into_import(client, db):
    h = auth_header(_admin(db, "bulkadmin3@premnathrail.com"))
    client.post("/api/v1/erp/projects", json={"serial_number": "SN-EXP", "model_name": "M1", "is_export": True}, headers=h)
    res = client.get("/api/v1/erp/projects/export", headers=h)
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws["A1"].value == "Serial Number"
    data = client.post("/api/v1/erp/projects/import",
                       files={"file": ("machines.xlsx", io.BytesIO(res.content), "application/octet-stream")}, headers=h).json()
    assert not data["created"] and not data["errors"]
    assert any(e["key"] == "SN-EXP" for e in data["existing"])


def test_sr_import_new_and_historical_with_duplicates(client, db):
    h = auth_header(_admin(db, "bulkadmin4@premnathrail.com"))
    client.post("/api/v1/erp/projects", json={"serial_number": "SN-SR"}, headers=h)

    body = (
        "machine_serial_number,old_sr_number,issue_title,service_type,priority,status,reported_date,corrective_action,service_cost,tax_percentage,closed_date\n"
        "SN-SR,,Oil leak,Warranty,Urgent,Open,05-10-2026,,,,\n"                                  # 2 new open
        "SN-SR,OLD-1,Engine not starting,Repair / Breakdown,,Closed,14-03-2025,Battery replaced,1000,18,17-03-2025\n"  # 3 historical closed
        "sn-sr,,oil  leak,AMC,,,05-10-2026,,,,\n"                                                 # 4 same machine+title+date as row 2
        "SN-SR,OLD-1,Something else,Warranty,,,01-01-2025,,,,\n"                                 # 5 old number repeats row 3
        "SN-SR,,Gearbox noise,Warranty,,Closed,01-02-2025,,,,\n"                                 # 6 closed without corrective action
        "SN-NONE,,Brake issue,Warranty,,,,,,,\n"                                                  # 7 unknown machine
        "SN-SR,,Brake issue,Teleport,,,,,,,\n"                                                    # 8 bad service type
        "SN-SR,,Late close,Warranty,,Closed,10-02-2025,Fixed,,,01-02-2025\n"                     # 9 closed before reported
    )
    res = client.post("/api/v1/erp/service-requests/import", files=_csv(body), headers=h)
    assert res.status_code == 200, res.text
    data = res.json()
    assert [c["row"] for c in data["created"]] == [2, 3]
    assert {e["row"] for e in data["existing"]} == {4, 5}
    assert {e["row"] for e in data["errors"]} == {6, 7, 8, 9}

    old = db.query(ServiceRequest).filter(ServiceRequest.legacy_sr_number == "OLD-1").one()
    assert old.status == "closed" and old.closed_at is not None
    assert old.total_bill == 1180
    assert old.created_notification_sent and old.closed_notification_sent

    # Same file again → everything that was saved is now a duplicate.
    again = client.post("/api/v1/erp/service-requests/import", files=_csv(body), headers=h).json()
    assert not again["created"]
    assert {2, 3}.issubset({e["row"] for e in again["existing"]})


def test_sr_export_round_trips_into_import(client, db):
    h = auth_header(_admin(db, "bulkadmin5@premnathrail.com"))
    pid = client.post("/api/v1/erp/projects", json={"serial_number": "SN-SRX"}, headers=h).json()["id"]
    assert client.post("/api/v1/erp/service-requests", json={"project_id": pid, "issue_title": "Hose burst", "service_type": "warranty"},
                       headers=h).status_code == 201
    res = client.get("/api/v1/erp/service-requests/export", headers=h)
    assert res.status_code == 200
    data = client.post("/api/v1/erp/service-requests/import",
                       files={"file": ("srs.xlsx", io.BytesIO(res.content), "application/octet-stream")}, headers=h).json()
    assert not data["created"] and not data["errors"]
    assert data["existing"]


def test_templates_download(client, db):
    h = auth_header(_admin(db, "bulkadmin6@premnathrail.com"))
    for url in ("/api/v1/erp/projects/import/template", "/api/v1/erp/service-requests/import/template"):
        res = client.get(url, headers=h)
        assert res.status_code == 200 and res.text.splitlines()[0]
