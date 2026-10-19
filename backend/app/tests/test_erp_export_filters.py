"""Service exports: selected rows (ids) and From–To date ranges."""
import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.auth.jwt_handler import create_access_token
from app.modules.erp.models.project import Project
from app.modules.erp.models.service_request import ServiceRequest
from app.modules.main.models.user import User

API = "/api/v1/erp"


def _admin(db):
    u = User(email="exp_admin@premnathrail.com", name="exp", role="admin", is_active=True, assigned_apps=[], erp_permissions=[])
    db.add(u)
    db.commit()
    return {"Authorization": f"Bearer {create_access_token({'sub': str(u.id), 'email': u.email, 'role': u.role})}"}


def _first_col(resp) -> list:
    assert resp.status_code == 200, resp.text
    ws = load_workbook(io.BytesIO(resp.content)).active
    return [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def _seed(db):
    a = Project(serial_number="EXP-A", delivery_date=datetime(2026, 1, 10).date())
    b = Project(serial_number="EXP-B", delivery_date=datetime(2026, 3, 5).date())
    c = Project(serial_number="EXP-C")
    db.add_all([a, b, c])
    db.flush()
    db.add_all([
        ServiceRequest(request_number="SR-EXP-1", project_id=a.id, issue_title="x", status="open",
                       opened_at=datetime(2026, 2, 1, 6, 0, tzinfo=timezone.utc)),
        ServiceRequest(request_number="SR-EXP-2", project_id=b.id, issue_title="y", status="closed",
                       opened_at=datetime(2026, 4, 1, 6, 0, tzinfo=timezone.utc),
                       closed_at=datetime(2026, 4, 9, 20, 0, tzinfo=timezone.utc)),  # 10-04 in IST
    ])
    db.commit()
    return a, b, c


def test_machine_export_selected_and_dispatch_range(client, db):
    h = _admin(db)
    a, b, c = _seed(db)
    assert _first_col(client.get(f"{API}/projects/export", params={"ids": f"{a.id},{c.id}"}, headers=h)) == ["EXP-A", "EXP-C"]
    # From–To defaults to the dispatch date
    rows = _first_col(client.get(f"{API}/projects/export", params={"date_from": "2026-01-01", "date_to": "2026-01-31"}, headers=h))
    assert rows == ["EXP-A"]
    # only one end filled
    assert _first_col(client.get(f"{API}/projects/export", params={"date_from": "2026-02-01"}, headers=h)) == ["EXP-B"]
    # both blank = everything
    assert _first_col(client.get(f"{API}/projects/export", headers=h)) == ["EXP-A", "EXP-B", "EXP-C"]
    r = client.get(f"{API}/projects/export", params={"date_from": "2026-05-01", "date_to": "2026-01-01"}, headers=h)
    assert r.status_code == 422 and "after To date" in r.json()["detail"]
    assert client.get(f"{API}/projects/export", params={"ids": "1,abc"}, headers=h).status_code == 422


def test_sr_export_selected_and_ranges(client, db):
    h = _admin(db)
    _seed(db)
    sr1 = db.query(ServiceRequest).filter_by(request_number="SR-EXP-1").one()
    assert _first_col(client.get(f"{API}/service-requests/export", params={"ids": str(sr1.id)}, headers=h)) == ["SR-EXP-1"]
    rows = _first_col(client.get(f"{API}/service-requests/export", params={"date_from": "2026-01-15", "date_to": "2026-02-15"}, headers=h))
    assert rows == ["SR-EXP-1"]
    rows = _first_col(client.get(f"{API}/service-requests/export", params={"date_field": "closed", "date_from": "2026-04-10", "date_to": "2026-04-10"}, headers=h))
    assert rows == ["SR-EXP-2"]
    assert client.get(f"{API}/service-requests/export", params={"date_field": "bogus"}, headers=h).status_code == 422
