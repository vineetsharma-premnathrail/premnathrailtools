"""Opening Stock Import (admin only, stock-in only, all-or-nothing) and the
Stock page Excel export. Free-form Receipt / Issue / Damage rows were
removed with manual stock entry on 2026-10-12 (SAP standard)."""
import io

import pytest
from openpyxl import load_workbook

from app.auth.jwt_handler import create_access_token
from app.modules.main.models.module_setting import ModuleSetting
from app.modules.main.models.user import User
from app.modules.store.models.item import StoreItem
from app.modules.store.models.location import StoreLocation
from app.modules.store.models.stock_transaction import StoreStockTransaction
from app.modules.store.services.stock_ledger import post_stock_transaction

API = "/api/v1/store/stock"


def _headers(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}


@pytest.fixture
def users(db):
    admin = User(email="store.admin@premnathrail.com", name="store.admin", role="admin", is_active=True, assigned_apps=["store"])
    keeper = User(email="store.keeper@premnathrail.com", name="store.keeper", role="user", is_active=True, assigned_apps=["store"])
    manager = User(email="store.manager@premnathrail.com", name="store.manager", role="user", is_active=True,
                   assigned_apps=["store"], is_store_manager=True)
    db.add_all([
        admin, keeper, manager,
        StoreItem(item_code="MT-GEN-0001", item_name="Hex Bolt", item_type="material", category="Hardware", uom="NOS"),
        StoreItem(item_code="MT-GEN-0002", item_name="Gloves", item_type="material", category="PPE", uom="PAIR"),
        StoreLocation(name="Main Store", code="MAIN"),
    ])
    db.commit()
    db.query(StoreLocation).filter(StoreLocation.code == "MAIN").update({"manager_user_id": manager.id})
    db.commit()
    return {"admin": _headers(admin), "keeper": _headers(keeper), "manager": _headers(manager)}


@pytest.fixture
def h(users):
    return users["admin"]


def upload(client, h, text):
    return client.post(f"{API}/import", files={"file": ("stock.csv", text, "text/csv")}, headers=h)


def test_import_posts_opening_stock_rows(client, h, db):
    r = upload(client, h,
               "item_code,warehouse,quantity,transaction_date,reference_number,remarks,entry_type\n"
               "mt-gen-0001,MAIN,100,2026-10-01,OPEN-1,Opening stock,Issue\n"   # entry_type column is ignored
               "MT-GEN-0001,Main Store,30,,OPEN-1,Opening stock,\n"
               "MT-GEN-0002,MAIN,5,,OPEN-1,Opening stock,\n")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["errors"] == [] and len(body["posted"]) == 3
    assert all(p["stock_effect"] == "in" for p in body["posted"])
    bal = {b["item_code"]: b["on_hand_qty"] for b in client.get(f"{API}/balances", headers=h).json()}
    assert bal == {"MT-GEN-0001": 130, "MT-GEN-0002": 5}
    rows = db.query(StoreStockTransaction).all()
    assert {(t.transaction_type, t.entry_type, t.reference_type) for t in rows} == {("manual_in", "opening_stock", "opening_stock")}
    assert {t.reference_number for t in rows} == {"OPEN-1"}


def test_any_bad_row_posts_nothing(client, h, db):
    r = upload(client, h,
               "item_code,warehouse,quantity,reference_number,remarks\n"
               "MT-GEN-0001,MAIN,10,OPEN-1,Opening stock\n"   # row 2 fine
               "NOPE,MAIN,1,OPEN-1,Opening stock\n"           # 3 unknown item
               "MT-GEN-0002,NOWHERE,5,OPEN-1,Opening stock\n"  # 4 unknown store
               "MT-GEN-0002,MAIN,0,OPEN-1,Opening stock\n"     # 5 zero
               "MT-GEN-0002,MAIN,-3,OPEN-1,Opening stock\n"    # 6 negative
               "MT-GEN-0002,MAIN,3,,Opening stock\n"           # 7 no reference
               "MT-GEN-0002,MAIN,3,OPEN-1,\n")                 # 8 no remark
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["posted"] == [] and body["would_post"] == 1
    assert [e["row"] for e in body["errors"]] == [3, 4, 5, 6, 7, 8]
    reasons = [e["reason"] for e in body["errors"]]
    assert "isn't in the Item Master" in reasons[0]
    assert "not found" in reasons[1]
    assert "greater than zero" in reasons[2] and "greater than zero" in reasons[3]
    assert "reference" in reasons[4] and "reference" in reasons[5]
    assert db.query(StoreStockTransaction).count() == 0


def test_rate_becomes_unit_cost_when_valuation_is_on(client, h, db):
    db.add(ModuleSetting(module_key="store", key="valuation_enabled", value=True))
    db.commit()
    r = upload(client, h, "item_code,warehouse,quantity,rate,reference_number,remarks\nMT-GEN-0001,MAIN,10,42.50,OPEN-1,Opening stock\n")
    assert r.status_code == 200 and r.json()["errors"] == [], r.text
    txn = db.query(StoreStockTransaction).one()
    assert txn.unit_cost == pytest.approx(42.5) and txn.cost_source == "opening"
    bad = upload(client, h, "item_code,warehouse,quantity,rate,reference_number,remarks\nMT-GEN-0002,MAIN,1,-1,OPEN-1,Opening stock\n")
    assert bad.json()["posted"] == [] and "negative" in bad.json()["errors"][0]["reason"]


@pytest.mark.parametrize("who", ["keeper", "manager"])
def test_import_and_template_are_admin_only(client, users, db, who):
    r = upload(client, users[who], "item_code,warehouse,quantity,reference_number,remarks\nMT-GEN-0001,MAIN,10,OPEN-1,Opening stock\n")
    assert r.status_code == 403 and "admin only" in r.json()["detail"]
    assert client.get(f"{API}/import/template", headers=users[who]).status_code == 403
    assert db.query(StoreStockTransaction).count() == 0


def test_template_has_opening_stock_columns(client, h):
    r = client.get(f"{API}/import/template", headers=h)
    assert r.status_code == 200
    header = r.text.splitlines()[0].split(",")
    assert header == ["item_code", "warehouse", "quantity", "rate", "batch_number", "reference_number", "transaction_date", "remarks"]


def test_export_has_balances_and_movements_with_filters(client, h):
    upload(client, h, "item_code,warehouse,quantity,reference_number,remarks\nMT-GEN-0001,MAIN,10,OPEN-1,Opening stock\nMT-GEN-0002,MAIN,4,OPEN-1,Opening stock\n")
    r = client.get(f"{API}/export", params={"category": "PPE"}, headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert [row[0] for row in wb["Balances"].iter_rows(min_row=2, values_only=True)] == ["MT-GEN-0002"]
    assert [row[1] for row in wb["Movements"].iter_rows(min_row=2, values_only=True)] == ["MT-GEN-0002"]


def test_vendor_names_on_earlier_entries_are_suggested(client, users, db):
    # Manual entry (where the vendor name used to be typed) is removed; names
    # already on the ledger are still offered as suggestions.
    item = db.query(StoreItem).filter_by(item_code="MT-GEN-0001").one()
    loc = db.query(StoreLocation).filter_by(code="MAIN").one()
    for name in ("Sharma Steels", "Bharat Bearings"):
        post_stock_transaction(db, item_id=item.id, location_id=loc.id, transaction_type="manual_in", quantity=1,
                               reference_type="stock_entry", reference_number="SE-OLD", vendor_name=name)
    db.commit()
    assert client.get(f"{API}/vendor-options", headers=users["keeper"]).json() == ["Bharat Bearings", "Sharma Steels"]
