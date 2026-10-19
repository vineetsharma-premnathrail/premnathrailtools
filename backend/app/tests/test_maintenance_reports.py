"""
Tests for Maintenance Phase 2 reporting + Excel asset import:
MTTR / MTBF / availability / PM compliance maths on fixed fixture data,
downtime clamping to the report window, cost roll-up, the export workbook,
the dashboard additions, and the import dry-run / all-or-nothing rules.
"""
import io
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from app.auth.jwt_handler import create_access_token
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch
from app.modules.store.models.item import StoreItem
from app.modules.store.models.location import StoreLocation
from app.modules.production.models.workstation import ProductionWorkstation
from app.modules.maintenance.models.asset import MaintenanceAsset, MaintenanceAssetSpare
from app.modules.maintenance.models.work_order import MaintenanceWorkOrder, MaintenanceWorkOrderSpare, MaintenanceLabourLog
from app.modules.maintenance.reporting import compute_summary, _parse_hours

BASE = "/api/v1/maintenance"
IST = ZoneInfo("Asia/Kolkata")


def ist(y, m, d, hh=0, mm=0):
    """IST wall-clock → UTC (SQLite drops tzinfo, so store UTC)."""
    return datetime(y, m, d, hh, mm, tzinfo=IST).astimezone(timezone.utc)


def make_user(db, email, apps=("maintenance",)):
    u = User(email=email, name=email.split("@")[0], role="user", is_active=True, assigned_apps=list(apps))
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def hdr(u):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(u.id), 'email': u.email, 'role': u.role})}"}


_seq = iter(range(1, 1000))


def wo(db, asset, **kw):
    n = next(_seq)
    w = MaintenanceWorkOrder(wo_number=f"MWO-2026-{n:04d}", asset_id=asset.id, title=f"Job {n}", **kw)
    db.add(w)
    db.flush()
    return w


def seed(db):
    """August 2026: Aug 1 is a Saturday → 5 Sundays → 26 Mon–Sat working days."""
    branch = Branch(name="Faridabad Plant", code="FBD")  # no working days/hours → Mon–Sat × 8 h
    ws = ProductionWorkstation(code="VMC-01", name="VMC 1", capacity_hours_per_day=16)
    item = StoreItem(item_code="SP-1", item_name="Bearing", uom="NOS", item_type="spare_part", minimum_stock=5)
    loc = StoreLocation(name="Main Store", code="MS")
    db.add_all([branch, ws, item, loc])
    db.flush()
    a = MaintenanceAsset(asset_code="CNC-01", name="CNC Lathe", branch_id=branch.id, category="production_machine", criticality="A")
    b = MaintenanceAsset(asset_code="VMC-01", name="VMC", branch_id=branch.id, category="production_machine", workstation_id=ws.id)
    db.add_all([a, b])
    db.flush()

    # A: 4 h breakdown inside August, with costs.
    w1 = wo(db, a, wo_type="breakdown", status="closed", machine_down=True, failure_category="mechanical",
            downtime_start=ist(2026, 8, 10, 10), downtime_end=ist(2026, 8, 10, 14), downtime_minutes=240,
            actual_end=ist(2026, 8, 10, 14), external_cost=500, created_at=ist(2026, 8, 10, 10))
    db.add(MaintenanceWorkOrderSpare(work_order_id=w1.id, store_item_id=item.id, location_id=loc.id, qty_issued=3, qty_returned=1, unit_cost=100))
    db.add(MaintenanceLabourLog(work_order_id=w1.id, technician_id=1, hours=2, hourly_rate=300))
    # A: started 31 Jul 22:00, ended 1 Aug 02:00 — only 120 of 240 min fall in August.
    wo(db, a, wo_type="breakdown", status="closed", machine_down=True, failure_category="mechanical",
       downtime_start=ist(2026, 7, 31, 22), downtime_end=ist(2026, 8, 1, 2), downtime_minutes=240,
       actual_end=ist(2026, 8, 1, 2), created_at=ist(2026, 7, 31, 22))
    # B: 1 h electrical breakdown.
    wo(db, b, wo_type="breakdown", status="completed", machine_down=True, failure_category="electrical",
       downtime_start=ist(2026, 8, 20, 8), downtime_end=ist(2026, 8, 20, 9), downtime_minutes=60,
       actual_end=ist(2026, 8, 20, 9), created_at=ist(2026, 8, 20, 8))
    # PM: on time, late, missed (still open after due date), cancelled (skipped — counts against compliance).
    wo(db, a, wo_type="preventive", status="closed", pm_due_date=date(2026, 8, 15), actual_end=ist(2026, 8, 14, 16), created_at=ist(2026, 8, 12))
    wo(db, a, wo_type="preventive", status="closed", pm_due_date=date(2026, 8, 25), actual_end=ist(2026, 8, 27, 11), created_at=ist(2026, 8, 22))
    wo(db, b, wo_type="preventive", status="assigned", pm_due_date=date(2026, 8, 28), created_at=ist(2026, 8, 25))
    wo(db, b, wo_type="preventive", status="cancelled", pm_due_date=date(2026, 8, 5), created_at=ist(2026, 8, 1))
    # Critical spare: 2 on hand vs minimum 5.
    db.add(MaintenanceAssetSpare(asset_id=a.id, store_item_id=item.id, is_critical=True))
    from app.modules.store.models.stock_balance import StoreStockBalance
    db.add(StoreStockBalance(item_id=item.id, location_id=loc.id, on_hand_qty=2, reserved_qty=0))
    db.commit()
    return branch, a, b


def test_metrics_maths(db):
    _, a, b = seed(db)
    s = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31))
    k = s["kpis"]
    assert k["downtime_minutes"] == 240 + 120 + 60          # July overlap clamped to 120
    assert k["breakdowns"] == 2                               # the July-started one belongs to July
    assert k["mttr_minutes"] == 150                           # (240 + 60) / 2 full repairs
    # A: 26 days × 8 h (default); B: 26 × 16 h (workstation capacity)
    assert k["available_hours"] == 26 * 8 + 26 * 16 == 624
    assert k["availability_pct"] == round((624 - 7) / 624 * 100, 1)
    assert k["mtbf_hours"] == round((624 - 7) / 2, 1)
    basis = s["availability_basis"]
    assert basis["assets_from_workstation"] == 1 and basis["assets_default_hours"] == 1

    pm = s["pm_compliance"]
    assert (pm["due"], pm["on_time"], pm["late"], pm["missed"], pm["skipped"]) == (4, 1, 1, 1, 1)
    assert k["pm_compliance_pct"] == 25.0                     # 1 on time / (1 + 1 + 1 + 1 skipped)

    # Cost: spares (3-1)×100 + labour 2×300 + external 500 on w1.
    assert s["cost"]["total"]["spares"] == 200 and s["cost"]["total"]["labour"] == 600
    assert s["cost"]["total"]["external"] == 500 and k["maintenance_cost"] == 1300
    assert s["cost"]["by_asset"][0]["asset_code"] == "CNC-01"

    assert s["top_assets"]["by_downtime"][0]["asset_code"] == "CNC-01"
    assert s["top_assets"]["by_downtime"][0]["downtime_minutes"] == 360
    pareto = s["failure_pareto"]
    assert pareto[0]["failure_category"] == "mechanical" and pareto[0]["count"] == 1
    assert pareto[-1]["cumulative_pct"] == 100.0
    assert s["backlog"]["total"] == 1  # the missed PM job is still open


def test_downtime_clamping_single_day(db):
    seed(db)
    s = compute_summary(db, date(2026, 8, 1), date(2026, 8, 1))
    assert s["kpis"]["downtime_minutes"] == 120
    assert s["kpis"]["breakdowns"] == 0 and s["kpis"]["mttr_minutes"] is None
    jul = compute_summary(db, date(2026, 7, 1), date(2026, 7, 31))
    assert jul["kpis"]["downtime_minutes"] == 120 and jul["kpis"]["breakdowns"] == 1
    assert jul["downtime"]["by_month"] == [{"month": "2026-07", "downtime_minutes": 120}]
    two = compute_summary(db, date(2026, 7, 1), date(2026, 8, 31))
    assert [m["downtime_minutes"] for m in two["downtime"]["by_month"]] == [120, 420]


def test_filters_and_plant_hours(db):
    branch, a, b = seed(db)
    branch.working_days = "Mon-Fri"
    branch.working_hours = "09:00–18:00"
    db.commit()
    s = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31), asset_id=a.id)
    # August 2026 has 21 Mon–Fri days; plant hours 9 h/day.
    assert s["kpis"]["available_hours"] == 21 * 9
    assert s["availability_basis"]["assets_from_plant_hours"] == 1
    assert s["kpis"]["downtime_minutes"] == 360
    assert _parse_hours("8") == 8 and _parse_hours("9am - 5:30pm") == 8.5 and _parse_hours("shift") is None


def test_report_endpoints_and_dashboard(client, db):
    seed(db)
    h = hdr(make_user(db, "rep@premnathrail.com"))
    r = client.get(f"{BASE}/reports/summary", params={"from": "2026-08-01", "to": "2026-08-31"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["kpis"]["breakdowns"] == 2
    r = client.get(f"{BASE}/reports/summary", params={"from": "2026-08-31", "to": "2026-08-01"}, headers=h)
    assert r.status_code == 400 and "after" in r.json()["detail"]

    r = client.get(f"{BASE}/reports/export", params={"from": "2026-08-01", "to": "2026-08-31", "format": "xlsx"}, headers=h)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert "Summary" in wb.sheetnames and "Failure Pareto" in wb.sheetnames and "Backlog" in wb.sheetnames

    other = make_user(db, "crm@premnathrail.com", apps=("crm",))
    assert client.get(f"{BASE}/reports/summary", headers=hdr(other)).status_code == 403

    d = client.get(f"{BASE}/dashboard", headers=h).json()
    assert "counts" in d and "machines_down" in d  # Phase 1 keys intact
    assert d["pm_overdue"]["count"] == 1 and d["pm_overdue"]["items"][0]["days_overdue"] > 0
    assert d["critical_spares_below_min"]["count"] == 1
    assert d["critical_spares_below_min"]["items"][0]["shortfall"] == 3
    assert "availability_pct" in d["this_month"] and "pm_due_this_week" in d


def test_corrective_machine_down_counts_as_breakdown(db):
    _, a, _ = seed(db)
    # 90-minute corrective (machine-down) job on A, finished in August.
    wo(db, a, wo_type="corrective", status="closed", machine_down=True, failure_category="hydraulic",
       downtime_start=ist(2026, 8, 12, 10), downtime_end=ist(2026, 8, 12, 11, 30), downtime_minutes=90,
       actual_end=ist(2026, 8, 12, 11, 30), created_at=ist(2026, 8, 12, 10))
    # A PM job never records downtime, so it stays out of the breakdown KPIs.
    wo(db, a, wo_type="preventive", status="closed", created_at=ist(2026, 8, 13), actual_end=ist(2026, 8, 13, 12))
    db.commit()
    k = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31))["kpis"]
    assert k["downtime_minutes"] == 240 + 120 + 60 + 90
    assert k["breakdowns"] == 3
    assert k["mttr_minutes"] == round((240 + 60 + 90) / 3)
    assert k["mtbf_hours"] == round((624 - (420 + 90) / 60) / 3, 1)


def test_corrective_without_machine_down_is_not_a_breakdown(db):
    _, a, _ = seed(db)
    base = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31))
    # Corrective job raised from an inspection finding; machine kept running.
    wo(db, a, wo_type="corrective", status="closed", machine_down=False, failure_category="hydraulic",
       actual_end=ist(2026, 8, 18, 15), created_at=ist(2026, 8, 18, 10))
    db.commit()
    s = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31))
    k, bk = s["kpis"], base["kpis"]
    assert k["breakdowns"] == bk["breakdowns"] == 2
    assert k["mtbf_hours"] == bk["mtbf_hours"] and k["mttr_minutes"] == bk["mttr_minutes"]
    assert k["downtime_minutes"] == bk["downtime_minutes"]
    assert s["failure_pareto"] == base["failure_pareto"]
    assert all(r["failure_category"] != "hydraulic" for r in s["failure_pareto"])
    assert s["top_assets"]["by_breakdowns"] == base["top_assets"]["by_breakdowns"]


def test_cancelled_pm_not_due_or_retired_is_not_skipped(db):
    from datetime import timedelta
    from app.modules.maintenance.reporting import today_ist

    branch, a, b = seed(db)
    t = today_ist()
    # (a) PM due next week, cancelled today → not yet due, left out.
    wo(db, a, wo_type="preventive", status="cancelled", pm_due_date=t + timedelta(days=5), created_at=ist(2026, 8, 1))
    # (b) PM on an asset decommissioned 20 Aug, cancelled by retire_asset_pm, due 26 Aug.
    c = MaintenanceAsset(asset_code="DRL-01", name="Drill", branch_id=branch.id, category="production_machine",
                         decommissioned_on=date(2026, 8, 20))
    db.add(c)
    db.flush()
    wo(db, c, wo_type="preventive", status="cancelled", pm_due_date=date(2026, 8, 26), created_at=ist(2026, 8, 1),
       cancel_reason="Asset decommissioned")
    db.commit()

    pm = compute_summary(db, date(2026, 8, 1), date(2026, 8, 31))["pm_compliance"]
    assert (pm["due"], pm["on_time"], pm["late"], pm["missed"], pm["skipped"]) == (4, 1, 1, 1, 1)
    assert pm["compliance_pct"] == 25.0

    fut = compute_summary(db, t, t + timedelta(days=10))["pm_compliance"]
    assert fut["skipped"] == 0 and fut["due"] == 0 and fut["compliance_pct"] is None


def test_dashboard_last_30_days_matches_compute_summary(client, db):
    from datetime import timedelta
    from app.modules.maintenance.reporting import today_ist

    _, a, b = seed(db)
    now = datetime.now(timezone.utc)
    wo(db, a, wo_type="breakdown", status="closed", machine_down=True,
       downtime_start=now - timedelta(days=3, hours=2), downtime_end=now - timedelta(days=3), downtime_minutes=120,
       actual_end=now - timedelta(days=3), created_at=now - timedelta(days=3, hours=2))
    wo(db, b, wo_type="corrective", status="completed", machine_down=True,
       downtime_start=now - timedelta(days=2, minutes=45), downtime_end=now - timedelta(days=2), downtime_minutes=45,
       actual_end=now - timedelta(days=2), created_at=now - timedelta(days=2, minutes=45))
    db.commit()
    h = hdr(make_user(db, "dash30@premnathrail.com"))
    d = client.get(f"{BASE}/dashboard", headers=h).json()["last_30_days"]
    t = today_ist()
    k = compute_summary(db, t - timedelta(days=29), t)["kpis"]
    assert d["mttr_minutes"] == k["mttr_minutes"] == round((120 + 45) / 2)
    assert d["breakdowns"] == k["breakdowns"] == 2
    assert d["downtime_minutes"] == k["downtime_minutes"]
    assert d["maintenance_cost"] == k["maintenance_cost"]


def test_pm_due_this_week_completed_job_not_listed_as_not_generated(client, db):
    from datetime import timedelta
    from app.modules.maintenance.models.pm_plan import MaintenancePmPlan
    from app.modules.maintenance.reporting import today_ist

    _, a, b = seed(db)
    due = today_ist() + timedelta(days=2)
    done_plan = MaintenancePmPlan(plan_code="PM-0001", asset_id=a.id, title="Monthly lube", next_due_date=due)
    open_plan = MaintenancePmPlan(plan_code="PM-0002", asset_id=b.id, title="Monthly clean", next_due_date=due)
    db.add_all([done_plan, open_plan])
    db.flush()
    # Done early, awaiting close — the plan hasn't rolled forward yet.
    wo(db, a, wo_type="preventive", status="completed", pm_plan_id=done_plan.id, pm_due_date=due,
       actual_end=datetime.now(timezone.utc))
    db.commit()
    h = hdr(make_user(db, "pmweek@premnathrail.com"))
    items = client.get(f"{BASE}/dashboard", headers=h).json()["pm_due_this_week"]["items"]
    not_generated = {i["number"] for i in items if i["kind"] == "plan"}
    assert "PM-0001" not in not_generated
    assert "PM-0002" in not_generated


def test_parse_hours_shorthand_and_overnight():
    assert _parse_hours("9-6") == 9
    assert _parse_hours("09:00-18:00") == 9
    assert _parse_hours("22:00-06:00") == 8
    assert _parse_hours("8:30 to 5") == 8.5
    assert _parse_hours("9am-6pm") == 9


def test_parse_hours_dotted_hhmm_and_mixed_ampm():
    assert _parse_hours("9.30 am to 6.00 pm") == 8.5
    assert _parse_hours("9.00-6.00") == 9
    assert _parse_hours("0900-1800") == 9
    assert _parse_hours("2200-0600") == 8
    assert _parse_hours("9am-6") == 9
    assert _parse_hours("9pm-6") == 9          # bare end after a pm start = overnight
    assert _parse_hours("9-30") is None        # minute-looking end: unreadable, not 21 h
    assert _parse_hours("6am-11pm") == 17      # both ends explicit am/pm: trusted


def test_lookup_assets_ids_always_included(client, db):
    branch = Branch(name="Faridabad Plant", code="FBD")
    db.add(branch)
    db.flush()
    db.add_all([MaintenanceAsset(asset_code=f"A-{i:03d}", name=f"Machine {i}", branch_id=branch.id) for i in range(105)])
    far = MaintenanceAsset(asset_code="Z-LAST", name="Last machine", branch_id=branch.id)
    old = MaintenanceAsset(asset_code="B-OLD", name="Scrapped press", branch_id=branch.id, status="decommissioned")
    gone = MaintenanceAsset(asset_code="C-DEL", name="Deleted", branch_id=branch.id, is_deleted=True)
    db.add_all([far, old, gone])
    db.commit()
    h = hdr(make_user(db, "look@premnathrail.com"))

    rows = client.get(f"{BASE}/lookups/assets", headers=h).json()
    assert len(rows) == 100 and far.id not in {r["id"] for r in rows} and old.id not in {r["id"] for r in rows}

    rows = client.get(f"{BASE}/lookups/assets", params={"ids": f"{far.id}, {old.id},{gone.id}"}, headers=h).json()
    got = {r["id"] for r in rows}
    assert far.id in got and old.id in got and gone.id not in got
    assert [r["asset_code"] for r in rows] == sorted(r["asset_code"] for r in rows)
    assert set(rows[0]) == {"id", "asset_code", "name", "criticality", "status", "branch_id", "location_text"}

    rows = client.get(f"{BASE}/lookups/assets", params={"search": "machine 10"}, headers=h).json()
    assert {r["asset_code"] for r in rows} == {"A-010", "A-100", "A-101", "A-102", "A-103", "A-104"}

    r = client.get(f"{BASE}/lookups/assets", params={"ids": "12,abc"}, headers=h)
    assert r.status_code == 400 and "comma-separated asset numbers" in r.json()["detail"]


# ---------------------------------------------------------------------------
# Excel import
# ---------------------------------------------------------------------------

HEAD = ["asset_code", "name", "category", "plant", "department", "location", "make", "model",
        "serial_number", "year_of_manufacture", "criticality", "commissioned_on", "workstation_code", "remarks"]


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(HEAD)
    for r in rows:
        ws.append([r.get(c) for c in HEAD])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, h, content, dry_run):
    return client.post(
        f"{BASE}/assets/import", params={"dry_run": str(dry_run).lower()}, headers=h,
        files={"file": ("machines.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_import_dry_run_errors(client, db):
    branch = Branch(name="Faridabad Plant", code="FBD")
    db.add(branch)
    db.flush()
    db.add(MaintenanceAsset(asset_code="OLD-1", name="Old", branch_id=branch.id))
    db.commit()
    h = hdr(make_user(db, "imp@premnathrail.com"))
    content = xlsx([
        {"asset_code": "CNC-1", "name": "Lathe", "category": "Production Machine", "plant": "FBD"},     # row 2 ok (label + plant code)
        {"asset_code": "CNC-2", "name": "Mill", "category": "production_machine", "plant": "Plant 3"},  # row 3 bad plant
        {"asset_code": "cnc-1", "name": "Dup", "category": "utility", "plant": "Faridabad Plant"},       # row 4 dup in file
        {"asset_code": "OLD-1", "name": "Dup DB", "category": "utility", "plant": "Faridabad Plant"},    # row 5 dup in DB
        {"name": "Compressor", "category": "air", "plant": "Faridabad Plant", "criticality": "Z",
         "year_of_manufacture": 1800, "commissioned_on": "31-02-2020"},                                   # row 6 several
    ])
    r = upload(client, h, content, True)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dry_run"] and body["total_rows"] == 5 and body["valid_rows"] == 1 and body["error_rows"] == 4
    errs = "\n".join(body["errors"])
    assert "Row 3: plant 'Plant 3' not found — use one of: Faridabad Plant" in errs
    assert "Row 4: asset_code 'CNC-1' repeats row 2" in errs
    assert "Row 5: asset_code 'OLD-1' already exists" in errs
    assert "Row 6: category 'air' isn't valid" in errs and "criticality 'Z'" in errs
    assert "year_of_manufacture '1800'" in errs and "commissioned_on '31-02-2020'" in errs
    assert body["rows"][4]["asset_code"] == "EQ-0001" and body["rows"][4]["auto_code"]
    assert db.query(MaintenanceAsset).count() == 1  # dry run wrote nothing

    # A real import with errors refuses and writes nothing.
    r = upload(client, h, content, False)
    assert r.status_code == 422 and "Nothing was imported" in r.json()["detail"]
    assert db.query(MaintenanceAsset).count() == 1


def test_import_success_and_template(client, db):
    branch = Branch(name="Faridabad Plant", code="FBD")
    ws = ProductionWorkstation(code="CNC-WS", name="CNC station")
    db.add_all([branch, ws])
    db.commit()
    h = hdr(make_user(db, "imp2@premnathrail.com"))

    t = client.get(f"{BASE}/assets/import-template", headers=h)
    assert t.status_code == 200
    twb = load_workbook(io.BytesIO(t.content))
    assert twb.sheetnames == ["Assets", "Instructions"]
    assert [c.value for c in twb["Assets"][1]][:4] == ["asset_code", "name", "category", "plant"]

    content = xlsx([
        {"asset_code": "cnc-10", "name": "Lathe", "category": "production_machine", "plant": "Faridabad Plant",
         "criticality": "a", "year_of_manufacture": 2018, "commissioned_on": datetime(2018, 6, 15), "workstation_code": "cnc-ws"},
        {"name": "DG Set", "category": "Utility", "plant": "fbd", "commissioned_on": "01-04-2020"},
        {"name": "Forklift", "category": "material_handling", "plant": "FBD"},
    ])
    r = upload(client, h, content, False)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 3
    assets = {x.asset_code: x for x in db.query(MaintenanceAsset).all()}
    assert set(assets) == {"CNC-10", "EQ-0001", "EQ-0002"}
    assert assets["CNC-10"].criticality == "A" and assets["CNC-10"].workstation_id == ws.id
    assert assets["CNC-10"].commissioned_on == date(2018, 6, 15)
    assert assets["EQ-0001"].category == "utility" and assets["EQ-0001"].commissioned_on == date(2020, 4, 1)

    # Re-importing the same file now clashes with the register.
    r = upload(client, h, content, True)
    assert "already exists" in "\n".join(r.json()["errors"]) and "already linked" in "\n".join(r.json()["errors"])

    bad = client.post(f"{BASE}/assets/import", headers=h, files={"file": ("m.csv", b"a,b", "text/csv")})
    assert bad.status_code == 422 and ".xlsx" in bad.json()["detail"]
