"""Store item master: auto item codes (<TYPE>-<CATEGORY>-<NNNN>), the fixed
UOM list, and the meta/preview endpoints the Add Item form uses."""
import pytest

from app.auth.jwt_handler import create_access_token
from app.modules.main.models.user import User
from app.modules.store.models.category import StoreItemCategory

API = "/api/v1/store/items"


@pytest.fixture
def h(db):
    user = User(email="store.keeper@premnathrail.com", name="store.keeper", role="user", is_active=True, assigned_apps=["store"])
    db.add(user)
    db.add_all([StoreItemCategory(name="Hydraulics", code="HYD"), StoreItemCategory(name="Bearings", code="BRG")])
    db.commit()
    db.refresh(user)
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}


def create(client, h, **body):
    return client.post(API, json={"item_name": "Item", "uom": "NOS", "item_type": "material", **body}, headers=h)


def test_codes_are_generated_per_type_and_category(client, h):
    first = create(client, h, item_name="Hydraulic cylinder 80 bore", category="Hydraulics")
    assert first.status_code == 200, first.text
    assert first.json()["item_code"] == "MT-HYD-0001"
    assert create(client, h, item_name="Hose 1/2in", category="Hydraulics").json()["item_code"] == "MT-HYD-0002"
    assert create(client, h, item_name="Bearing 6205", category="Bearings", item_type="tool_equipment").json()["item_code"] == "TE-BRG-0001"
    assert create(client, h, item_name="Misc part").json()["item_code"] == "MT-GEN-0001"


def test_preview_matches_the_next_saved_code(client, h):
    create(client, h, category="Hydraulics")
    preview = client.get(f"{API}/next-code", params={"item_type": "material", "category": "Hydraulics"}, headers=h).json()
    assert preview["item_code"] == "MT-HYD-0002"
    assert create(client, h, category="Hydraulics").json()["item_code"] == preview["item_code"]


def test_manual_code_still_accepted_but_must_be_unique(client, h):
    assert create(client, h, item_code="legacy-001").json()["item_code"] == "LEGACY-001"
    dup = create(client, h, item_code="LEGACY-001")
    assert dup.status_code == 409 and "leave the code blank" in dup.json()["detail"]


def test_uom_is_required_and_must_come_from_the_list(client, h):
    missing = client.post(API, json={"item_name": "No unit"}, headers=h)
    assert missing.status_code == 400 and "unit of measure" in missing.json()["detail"]
    bad = create(client, h, uom="Nos.")
    assert bad.status_code == 400 and "UOM dropdown" in bad.json()["detail"]
    ok = create(client, h, uom="kg")
    assert ok.status_code == 200 and ok.json()["uom"] == "KG"

    item_id = ok.json()["id"]
    assert client.patch(f"{API}/{item_id}", json={"uom": "furlong"}, headers=h).status_code == 400
    assert client.patch(f"{API}/{item_id}", json={"uom": ""}, headers=h).status_code == 400
    assert client.patch(f"{API}/{item_id}", json={"uom": "mtr"}, headers=h).json()["uom"] == "MTR"


def test_invalid_item_type_is_refused(client, h):
    r = create(client, h, item_type="gadget")
    assert r.status_code == 400 and "valid item type" in r.json()["detail"]


def test_meta_lists_types_with_prefixes_and_uoms(client, h):
    meta = client.get(f"{API}/meta", headers=h).json()
    assert {"value": "material", "label": "Material", "prefix": "MT"} in meta["item_types"]
    assert any(u["value"] == "NOS" for u in meta["uoms"])


def test_uoms_can_be_added_and_renamed(client, h):
    added = client.post("/api/v1/store/uoms", json={"code": "rft", "label": "Running foot"}, headers=h)
    assert added.status_code == 200 and any(u["value"] == "RFT" for u in added.json())
    assert client.post("/api/v1/store/uoms", json={"code": "RFT"}, headers=h).status_code == 409
    item = create(client, h, uom="rft").json()
    renamed = client.put("/api/v1/store/uoms/RFT", json={"code": "RFOOT", "label": "Running foot"}, headers=h)
    assert renamed.status_code == 200 and any(u["value"] == "RFOOT" for u in renamed.json())
    assert client.get(f"{API}/{item['id']}", headers=h).json()["uom"] == "RFOOT"


def test_categories_filter_by_type_and_subcategory_rename_follows_items(client, h):
    cat = client.post("/api/v1/store/categories", json={"name": "Electrical Parts", "item_type": "material"}, headers=h).json()
    client.post("/api/v1/store/categories", json={"name": "Hand Kit", "item_type": "tool_equipment"}, headers=h)
    names = [c["name"] for c in client.get("/api/v1/store/categories", params={"item_type": "material"}, headers=h).json()]
    assert "Electrical Parts" in names and "Hand Kit" not in names
    sub = client.post("/api/v1/store/categories", json={"name": "Cables", "parent_id": cat["id"]}, headers=h)
    assert sub.status_code == 200 and sub.json()["code"] and sub.json()["item_type"] == "material"
    item = create(client, h, category="Electrical Parts", subcategory="Cables").json()
    client.patch(f"/api/v1/store/categories/{sub.json()['id']}", json={"name": "Wires & Cables"}, headers=h)
    assert client.get(f"{API}/{item['id']}", headers=h).json()["subcategory"] == "Wires & Cables"


def test_bulk_import_skips_existing_and_reports_both(client, h):
    create(client, h, item_name="Hex Bolt M12")
    client.post("/api/v1/store/categories", json={"name": "PPE", "item_type": "consumable"}, headers=h)
    csv_text = (
        "item_name,uom,item_type,category,subcategory,description\n"
        "Hex Bolt M12,NOS,Material,,,\n"
        "Safety Gloves,PAIR,Consumable,PPE,Leather,\n"
        "safety gloves,PAIR,Consumable,,,\n"
        "Bad Unit,FURLONG,Material,,,\n"
        "Hex Bolt M12,NOS,Material,,,Zinc plated\n"
    )
    r = client.post(f"{API}/import", files={"file": ("items.csv", csv_text, "text/csv")}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [c["item_name"] for c in body["created"]] == ["Safety Gloves", "Hex Bolt M12"]
    assert {e["row"] for e in body["existing"]} == {2, 4}
    assert body["errors"][0]["row"] == 5 and "FURLONG" in body["errors"][0]["reason"]
    assert client.get(f"{API}/import/template", headers=h).status_code == 200


def test_export_respects_filters(client, h):
    create(client, h, item_name="Export Me", item_type="asset")
    create(client, h, item_name="Not Me", item_type="material")
    r = client.get(f"{API}/export", params={"item_type": "asset"}, headers=h)
    assert r.status_code == 200
    from openpyxl import load_workbook
    import io
    names = [row[4] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2, values_only=True)]
    assert names == ["Export Me"]


def test_search_matches_words_across_name_and_description(client, h):
    create(client, h, item_name="Bearing 6205-2RS", description="Make: Bosch")
    create(client, h, item_name="Bearing 6205-2RS", description="Make: SKF")
    for q in ["Bearing 6205-2RS make bosch", "  bosch   bearing ", "6205 BOSCH"]:
        names = [i["description"] for i in client.get(API, params={"search": q}, headers=h).json()]
        assert names == ["Make: Bosch"], q
    assert len(client.get(API, params={"search": "Bearing 6205-2RS "}, headers=h).json()) == 2


def test_item_types_can_be_added_and_renamed(client, h):
    added = client.post("/api/v1/store/item-types", json={"label": "Spare Kit", "prefix": "sk"}, headers=h)
    assert added.status_code == 200 and {"value": "spare_kit", "label": "Spare Kit", "prefix": "SK"} in added.json()
    assert client.post("/api/v1/store/item-types", json={"label": "spare kit"}, headers=h).status_code == 409
    assert client.post("/api/v1/store/item-types", json={"label": "Other", "prefix": "MT"}, headers=h).status_code == 409
    item = create(client, h, item_type="spare_kit").json()
    assert item["item_code"].startswith("SK-")
    renamed = client.put("/api/v1/store/item-types/spare_kit", json={"label": "Spare Kits", "prefix": "SPK"}, headers=h)
    assert {"value": "spare_kit", "label": "Spare Kits", "prefix": "SPK"} in renamed.json()
    assert create(client, h, item_type="spare_kit").json()["item_code"].startswith("SPK-")


def test_item_photo_upload_replace_and_remove(client, h, monkeypatch):
    from app.modules.store.routes import item_photo
    store = {}

    async def fake_upload(site, folder, f):
        f.file.seek(0)
        store[f"{folder}/{f.filename}"] = f.file.read()
        return {"name": f.filename, "path": f"{folder}/{f.filename}", "size": 3}

    async def fake_download(site, path):
        return store[path], "image/png"

    async def fake_delete(site, path):
        store.pop(path, None)

    monkeypatch.setattr(item_photo.settings, "SHAREPOINT_SITE_ID", "site")
    monkeypatch.setattr(item_photo, "upload_file_to_sharepoint", fake_upload)
    monkeypatch.setattr(item_photo, "download_file_content", fake_download)
    monkeypatch.setattr(item_photo, "delete_file_from_sharepoint", fake_delete)

    item = create(client, h, item_name="Photo Item").json()
    assert item["has_photo"] is False
    bad = client.post(f"{API}/{item['id']}/photo", files={"file": ("a.pdf", b"%PDF", "application/pdf")}, headers=h)
    assert bad.status_code == 400 and "isn't an image" in bad.json()["detail"]
    assert client.post(f"{API}/{item['id']}/photo", files={"file": ("a.png", b"PNG", "image/png")}, headers=h).status_code == 200
    assert client.get(f"{API}/{item['id']}", headers=h).json()["has_photo"] is True
    assert client.get(f"{API}/{item['id']}/photo", headers=h).content == b"PNG"
    client.post(f"{API}/{item['id']}/photo", files={"file": ("b.jpg", b"JPG", "image/jpeg")}, headers=h)
    assert client.get(f"{API}/{item['id']}/photo", headers=h).content == b"JPG" and len(store) == 1
    client.delete(f"{API}/{item['id']}/photo", headers=h)
    assert client.get(f"{API}/{item['id']}", headers=h).json()["has_photo"] is False and not store


def test_search_ranks_name_matches_first(client, h):
    create(client, h, item_name="A4 Paper", description="Demo office item")
    create(client, h, item_name="Bolt", description="demo bolt")
    create(client, h, item_name="demo-2")
    names = [i["item_name"] for i in client.get(API, params={"search": "demo"}, headers=h).json()]
    assert names[0] == "demo-2" and set(names) == {"demo-2", "A4 Paper", "Bolt"}


def test_part_code_is_saved_editable_and_searchable(client, h):
    item = create(client, h, item_name="Brake Shoe", part_number="DRG-4471").json()
    assert item["part_number"] == "DRG-4471"
    assert [i["item_name"] for i in client.get(API, params={"search": "drg-4471"}, headers=h).json()] == ["Brake Shoe"]
    assert client.patch(f"{API}/{item['id']}", json={"part_number": None}, headers=h).json()["part_number"] is None


def test_part_code_is_optional_in_every_category(client, h):
    client.post("/api/v1/store/categories", json={"name": "Brake Parts", "item_type": "material"}, headers=h)
    item = create(client, h, item_name="Brake Shoe", category="Brake Parts")
    assert item.status_code == 200 and item.json()["part_number"] is None
    r = client.patch(f"{API}/{item.json()['id']}", json={"part_number": "DRG-1"}, headers=h)
    assert r.status_code == 200 and r.json()["part_number"] == "DRG-1"
