"""The quotation sent to the customer = the Technical Offer on top + the
commercial quotation below, one Word file editable in Word; Download PDF
without a Word document joins the two PDFs (owner, 2026-10-27). The stacking of an
uploaded offer is tested in test_rnd_technical_offers.py."""
import io

from docx import Document

from app.tests.quote_lines import quote_product  # noqa: F401 — autouse fixture
from app.tests.test_cov_crm_quotations_rnd import ITEM, auth, make_user
from app.tests.test_crm_quotation_documents import sp  # noqa: F401 — fixture
from app.tests.test_cov_crm_quotations_rnd import make_inquiry

CRM = "/api/v1/crm"


def _sent_inquiry(db, owner):
    return make_inquiry(db, owner, uid="INQ-2026-0042")
from app.utils.document_merge import merge_docx


def _text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def _doc(*lines: str) -> bytes:
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_merge_docx_keeps_order_and_both_parts():
    text = _text(merge_docx(_doc("TECHNICAL PART"), _doc("COMMERCIAL PART")))
    assert text.index("TECHNICAL PART") < text.index("COMMERCIAL PART")


def test_quotation_without_technical_offer_is_commercial_only(client, db, sp):
    bd = make_user(db, "qto2.bd@premnathrail.com")
    inq = _sent_inquiry(db, bd)
    quot = client.post(f"{CRM}/inquiries/{inq.id}/quotations", json={"items": [ITEM]}, headers=auth(bd)).json()
    r = client.post(f"/api/v1/managed-documents/crm_quotation/{quot['id']}", headers=auth(bd))
    assert r.status_code == 200, r.text
    text = _text(sp.files[max(sp.files)][-1])
    assert text.startswith("QUOTATION") and "Amount in words" in text


def test_merge_pdf_puts_technical_offer_pages_first():
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    from app.utils.document_merge import merge_pdf

    def pdf(*pages: str) -> bytes:
        out = io.BytesIO()
        c = canvas.Canvas(out)
        for text in pages:
            c.drawString(72, 720, text)
            c.showPage()
        c.save()
        return out.getvalue()

    merged = PdfReader(io.BytesIO(merge_pdf(pdf("TECH 1", "TECH 2"), pdf("COMMERCIAL"))))
    assert [p.extract_text().strip() for p in merged.pages] == ["TECH 1", "TECH 2", "COMMERCIAL"]
