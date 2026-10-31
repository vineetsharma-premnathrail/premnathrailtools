"""'Add New Project' = machine form: POST /projects/with-machine saves the Project
Management project and its machine (erp_projects) in one transaction."""
from datetime import date, timedelta

import pytest

from app.auth.jwt_handler import create_access_token
from app.modules.erp.models.project import Project as Machine
from app.modules.main.models.user import User
from app.modules.projects.models.project import PmProject

API = "/api/v1"
TODAY = date.today()


def make_user(db, email, apps=("projects",), **flags):
    flags.setdefault("is_active", True)
    user = User(email=email, name=email.split("@")[0], role="user", assigned_apps=list(apps), **flags)
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def auth(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}


def payload(**machine_extra):
    machine = {"client_company": "RVNL", "po_number": "PO-77", "model_name": "TM-1"}
    machine.update(machine_extra)
    return {
        "name": "Tamping machine", "sales_order_ref": "SO-1",
        "start_date": str(TODAY), "target_handover_date": str(TODAY + timedelta(days=60)),
        "machine": machine,
    }


@pytest.fixture
def partner(db):
    return make_user(db, "wm.partner@x.com", is_partner=True)


def test_create_with_only_name(client, db, partner):
    # The form asks only the project name: no SO, dates, number, manager or description.
    body = {k: v for k, v in payload().items() if k not in ("start_date", "target_handover_date", "sales_order_ref")}
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=body)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["start_date"] is None and out["target_handover_date"] is None
    assert out["sales_order_ref"] is None
    assert out["project_code"].startswith("PEW-")
    assert db.get(PmProject, out["id"]).project_manager_id == partner.id


def test_create_registers_project_and_machine(client, db, partner):
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=payload())
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["machine_id"] and body["machine_serial"] == body["project_code"]
    assert body["client_name"] == "RVNL" and body["customer_po_ref"] == "PO-77"
    db.expire_all()
    machine = db.get(Machine, body["machine_id"])
    assert machine.pm_project_id == body["id"]
    assert machine.serial_number == body["project_code"]
    assert machine.created_by_id == partner.id
    # the list and the detail both carry the machine
    rows = client.get(f"{API}/projects", headers=auth(partner)).json()
    assert [x["machine_serial"] for x in rows if x["id"] == body["id"]] == [body["project_code"]]
    assert client.get(f"{API}/projects/{body['id']}", headers=auth(partner)).json()["machine_id"] == body["machine_id"]


def test_given_serial_is_kept(client, db, partner):
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=payload(serial_number="PNR-RR-9"))
    assert r.status_code == 200, r.text
    assert r.json()["machine_serial"] == "PNR-RR-9" and r.json()["project_code"] != "PNR-RR-9"


def test_serial_clash_saves_nothing(client, db, partner):
    db.add(Machine(serial_number="PRT-1", created_by_id=partner.id))
    db.commit()
    before = db.query(PmProject).count()
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=payload(serial_number="prt-1"))
    assert r.status_code == 409 and "PRT-1" in r.json()["detail"]
    db.expire_all()
    assert db.query(PmProject).count() == before
    assert db.query(Machine).count() == 1


def test_missing_client_and_po_have_real_messages(client, db, partner):
    p = payload(client_company=None)
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=p)
    assert r.status_code == 422 and "Client & Site tab" in r.json()["detail"]
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=payload(po_number=None))
    assert r.status_code == 422 and "Timeline & Warranty" in r.json()["detail"]
    assert db.query(PmProject).count() == 0


def test_non_management_is_refused(client, db):
    plain = make_user(db, "wm.plain@x.com")
    r = client.post(f"{API}/projects/with-machine", headers=auth(plain), json=payload())
    assert r.status_code == 403
    assert db.query(PmProject).count() == 0 and db.query(Machine).count() == 0
    noapp = make_user(db, "wm.noapp@x.com", apps=())
    assert client.post(f"{API}/projects/with-machine", headers=auth(noapp), json=payload()).status_code == 403


def test_code_change_moves_machine_serial_and_delete_is_blocked(client, db, partner):
    body = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=payload()).json()
    r = client.patch(f"{API}/projects/{body['id']}", headers=auth(partner), json={"project_code": "PEW-900"})
    assert r.status_code == 200, r.text
    assert r.json()["machine_serial"] == "PEW-900"
    d = client.delete(f"{API}/projects/{body['id']}", headers=auth(partner))
    assert d.status_code == 409 and "PEW-900" in d.json()["detail"]


def test_erp_machine_create_unchanged(client, db):
    admin = User(email="wm.admin@x.com", name="adm", role="admin", is_active=True, assigned_apps=["erp", "projects"])
    db.add(admin)
    db.commit()
    r = client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "ERP-1", "client_company": "X"})
    assert r.status_code == 201, r.text
    assert r.json()["serial_number"] == "ERP-1" and r.json()["pm_project_id"] is None
    r2 = client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "erp-1"})
    assert r2.status_code == 409


# ---- linking existing machines by number -------------------------------------------------

def _admin(db, email="wm.adm2@x.com"):
    admin = User(email=email, name="adm", role="admin", is_active=True, assigned_apps=["erp", "projects"])
    db.add(admin)
    db.commit()
    return admin


def _plain_project(client, partner, code, **extra):
    p = {"project_code": code, "name": "P " + code, "client_name": "C", "customer_po_ref": "PO", "sales_order_ref": "SO",
         "start_date": str(TODAY), "target_handover_date": str(TODAY + timedelta(days=30))}
    p.update(extra)
    r = client.post(f"{API}/projects", headers=auth(partner), json=p)
    assert r.status_code == 200, r.text
    return r.json()


def test_pm_project_for_serial_rule(client, db, partner):
    from app.core.pew_numbers import pm_project_for_serial
    a = _plain_project(client, partner, "PEW-601")
    b = _plain_project(client, partner, "PEW-601-X")
    assert pm_project_for_serial(db, "pew-601").id == a["id"]
    assert pm_project_for_serial(db, "PEW-601-01").id == a["id"]
    assert pm_project_for_serial(db, "PEW-601-X-2").id == b["id"]  # longest code wins
    assert pm_project_for_serial(db, "PEW-601A") is None  # PEW-601A is not PEW-601
    assert pm_project_for_serial(db, "") is None


def test_erp_create_links_by_number_but_not_over_explicit_link(client, db, partner):
    a = _plain_project(client, partner, "PEW-602")
    admin = _admin(db)
    r = client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "PEW-602-A1"})
    assert r.status_code == 201, r.text
    assert r.json()["pm_project_id"] == a["id"]
    r = client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "PEW-602A"})
    assert r.json()["pm_project_id"] is None
    # PATCH serial links an unlinked machine...
    mid = r.json()["id"]
    p = client.patch(f"{API}/erp/projects/{mid}", headers=auth(admin), json={"serial_number": "PEW-602-B"})
    assert p.status_code == 200 and p.json()["pm_project_id"] == a["id"]
    # ...but never moves an explicit link
    b = _plain_project(client, partner, "PEW-603")
    p = client.patch(f"{API}/erp/projects/{mid}", headers=auth(admin), json={"serial_number": "PEW-603-B"})
    assert p.json()["pm_project_id"] == a["id"] != b["id"]


def test_import_links_and_export_has_project_column(client, db, partner):
    import io
    from openpyxl import load_workbook
    a = _plain_project(client, partner, "PEW-604")
    admin = _admin(db)
    csv = b"Serial Number,Model Name\nPEW-604-01,M1\nOTHER-1,M2\n"
    r = client.post(f"{API}/erp/projects/import", headers=auth(admin), files={"file": ("m.csv", csv, "text/csv")})
    assert r.status_code == 200, r.text
    assert len(r.json()["created"]) == 2, r.text
    db.expire_all()
    assert db.query(Machine).filter(Machine.serial_number == "PEW-604-01").one().pm_project_id == a["id"]
    assert db.query(Machine).filter(Machine.serial_number == "OTHER-1").one().pm_project_id is None
    x = client.get(f"{API}/erp/projects/export", headers=auth(admin))
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    header = [c.value for c in ws[1]]
    assert header[-1] == "Project"
    rows = {row[0].value: row[-1].value for row in ws.iter_rows(min_row=2)}
    assert rows["PEW-604-01"].startswith("PEW-604") and rows["OTHER-1"] in (None, "")


def test_new_project_number_held_by_a_machine_is_refused(client, db, partner):
    """A PM project can't take a number a machine already carries (so nothing needs linking at create)."""
    db.add_all([Machine(serial_number="PEW-605-A1", created_by_id=partner.id), Machine(serial_number="PEW-605A", created_by_id=partner.id)])
    db.commit()
    p = {"project_code": "PEW-605", "name": "x", "client_name": "c", "customer_po_ref": "PO", "sales_order_ref": "SO",
         "start_date": str(TODAY), "target_handover_date": str(TODAY + timedelta(days=30))}
    r = client.post(f"{API}/projects", headers=auth(partner), json=p)
    assert r.status_code == 409 and "PEW-605-A1" in r.json()["detail"]
    assert client.post(f"{API}/projects", headers=auth(partner), json={**p, "project_code": "PEW-605A"}).status_code == 409


def test_pew_serials_are_normalised_everywhere(client, db, partner):
    from app.core.pew_numbers import next_pew_number, pew_number_taken
    admin = _admin(db)
    r = client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "  pew-777 "})
    assert r.status_code == 201 and r.json()["serial_number"] == "PEW-777"
    assert client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "PEW-777"}).status_code == 409
    assert pew_number_taken(db, "pew-777") is not None
    assert next_pew_number(db) == "PEW-778"
    p = client.patch(f"{API}/erp/projects/{r.json()['id']}", headers=auth(admin), json={"serial_number": " pew-780"})
    assert p.status_code == 200 and p.json()["serial_number"] == "PEW-780"
    csv = b"Serial Number\n pew-790 \nother-790\n"
    imp = client.post(f"{API}/erp/projects/import", headers=auth(admin), files={"file": ("m.csv", csv, "text/csv")}).json()
    assert sorted(c["ref"] for c in imp["created"]) == ["PEW-790", "other-790"]
    db.expire_all()
    stored = db.query(Machine).filter(Machine.serial_number == "PEW-790").one()
    assert stored is not None


def test_machine_side_link_is_audited_on_the_project(client, db, partner):
    from app.modules.main.models.audit_log import AuditLog
    a = _plain_project(client, partner, "PEW-607")
    admin = _admin(db)
    assert client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "PEW-607-A1"}).status_code == 201
    rows = db.query(AuditLog).filter(AuditLog.entity_type == "pm_project", AuditLog.entity_id == a["id"]).all()
    assert any("PEW-607-A1" in (x.summary or "") and "linked by its number" in (x.summary or "") for x in rows)


def test_linked_machine_visible_to_project_viewers_only_inside_scope(client, db, partner):
    """A machine linked to a PM project is read through the project's access; PM view alone
    (outside the project's data scope) is still refused, and unlinked machines keep today's rules."""
    own = {"projects": "own"}  # test DB default is organization-wide; narrow it so scope matters
    manager = make_user(db, "wm.mgr@x.com", data_access_scopes=own)
    outsider = make_user(db, "wm.outsider@x.com", data_access_scopes=own)
    body = payload()
    body["project_manager_id"] = manager.id
    p = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=body).json()
    mid = p["machine_id"]

    ok = client.get(f"{API}/erp/projects", headers=auth(manager), params={"pm_project_id": p["id"]})
    assert ok.status_code == 200 and [m["id"] for m in ok.json()] == [mid]
    assert client.get(f"{API}/erp/projects/{mid}", headers=auth(manager)).status_code == 200
    assert client.get(f"{API}/erp/projects/{mid}/attachments", headers=auth(manager)).status_code == 200
    # the project manager may edit the machine from the project without the ERP edit grant
    e = client.patch(f"{API}/erp/projects/{mid}", headers=auth(manager), json={"site_name": "New site"})
    assert e.status_code == 200 and e.json()["site_name"] == "New site"

    for call in (
        lambda: client.get(f"{API}/erp/projects", headers=auth(outsider), params={"pm_project_id": p["id"]}),
        lambda: client.get(f"{API}/erp/projects/{mid}", headers=auth(outsider)),
        lambda: client.get(f"{API}/erp/projects/{mid}/attachments", headers=auth(outsider)),
        lambda: client.patch(f"{API}/erp/projects/{mid}", headers=auth(outsider), json={"site_name": "x"}),
    ):
        r = call()
        assert r.status_code in (403, 404), (r.request.method, r.request.url, r.text[:200])
    # an unlinked machine is not opened by project view
    loose = Machine(serial_number="LOOSE-1", created_by_id=partner.id)
    db.add(loose)
    db.commit()
    assert client.get(f"{API}/erp/projects/{loose.id}", headers=auth(manager)).status_code == 403


def test_export_project_column_blank_without_project_access(client, db, partner):
    from app.modules.erp.routes.project_import import _machines_xlsx
    a = _plain_project(client, partner, "PEW-608")
    m = Machine(serial_number="PEW-608-1", created_by_id=partner.id, pm_project_id=a["id"])
    db.add(m)
    db.commit()
    nobody = make_user(db, "wm.nobody@x.com", data_access_scopes={"projects": "own"})  # PM project outside scope
    res = _machines_xlsx(db, [m], nobody)
    assert _project_cell(res) in (None, "")
    res = _machines_xlsx(db, [m], partner)
    assert (_project_cell(res) or "").startswith("PEW-608")


def _body(resp) -> bytes:
    import asyncio

    async def collect():
        return b"".join([c async for c in resp.body_iterator])
    return asyncio.run(collect())


def _project_cell(resp):
    import io
    from openpyxl import load_workbook
    buf = _body(resp)
    ws = load_workbook(io.BytesIO(buf)).active
    return ws.cell(row=2, column=ws.max_column).value


def test_xlsx_text_that_looks_like_a_formula_stays_text():
    import io
    from openpyxl import load_workbook
    from app.modules.erp.importing import xlsx_response
    res = xlsx_response("T", ["a", "b", "c", "d"], [["=1+1", "@SUM(1)", 5, "-x"]], "t.xlsx")
    ws = load_workbook(io.BytesIO(_body(res))).active
    assert [c.data_type for c in ws[2]] == ["s", "s", "n", "s"]
    assert [c.value for c in ws[2]] == ["=1+1", "@SUM(1)", 5, "-x"]


def test_machine_list_filter_by_project(client, db, partner):
    a = _plain_project(client, partner, "PEW-606")
    admin = _admin(db)
    client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "PEW-606-1"})
    client.post(f"{API}/erp/projects", headers=auth(admin), json={"serial_number": "ZZ-1"})
    rows = client.get(f"{API}/erp/projects?pm_project_id={a['id']}", headers=auth(admin)).json()
    assert [r["serial_number"] for r in rows] == ["PEW-606-1"]
    assert len(client.get(f"{API}/erp/projects", headers=auth(admin)).json()) == 2
    # a user without the machines tab is still refused
    plain = make_user(db, "wm.plain2@x.com", apps=())
    assert client.get(f"{API}/erp/projects?pm_project_id={a['id']}", headers=auth(plain)).status_code == 403


def test_demo_project_full_details_then_milestone_and_task(client, db, partner):
    """Owner's demo: add a project with every machine detail filled, then a
    department milestone inside it, then a task inside that milestone."""
    from app.modules.organization.models.department import Department

    head = make_user(db, "wm.head@x.com", apps=("design",), department="Design")
    member = make_user(db, "wm.member@x.com", apps=("design",), department="Design")
    dept = Department(name="Design", code="DES", head_user_id=head.id)
    db.add(dept)
    db.commit()

    body = payload(
        model_name="DEMO RRV", machine_type="Road Rail", application_type="OHE", year_of_manufacture="2026-27",
        client_company="DEMO Client Pvt Ltd", client_name="Demo Contact", client_email="demo@example.com",
        client_address="Demo address", site_name="Demo Site", site_location="Demo City", site_state="Haryana",
        specifications="Demo specs", po_number="PO-DEMO-001", po_date=str(TODAY),
    )
    body.update(name="DEMO - Test Project", sales_order_ref="SO-DEMO-001", description="Demo created by live test")
    r = client.post(f"{API}/projects/with-machine", headers=auth(partner), json=body)
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["client_name"] == "DEMO Client Pvt Ltd" and p["customer_po_ref"] == "PO-DEMO-001"
    assert p["machine_serial"] == p["project_code"]

    # Overview lists the project's machines with every detail that was filled.
    machines = client.get(f"{API}/erp/projects", headers=auth(partner), params={"pm_project_id": p["id"]}).json()
    assert isinstance(machines, list), machines
    assert len(machines) == 1
    m = machines[0]
    assert (m["model_name"], m["site_name"], m["site_state"], m["client_email"], m["specifications"]) == (
        "DEMO RRV", "Demo Site", "Haryana", "demo@example.com", "Demo specs")

    ms = client.post(f"{API}/projects/{p['id']}/milestones", headers=auth(partner), json={
        "title": "Design release", "department_id": dept.id, "lead_user_id": head.id,
        "target_date": str(TODAY + timedelta(days=30)), "deliverable": "GA drawing released",
    })
    assert ms.status_code == 200, ms.text
    task = client.post(f"{API}/pm-work/milestones/{ms.json()['id']}/tasks", headers=auth(head), json={
        "title": "Prepare GA drawing", "assignee_id": member.id, "due_date": str(TODAY + timedelta(days=14)),
    })
    assert task.status_code == 200, task.text

    after = client.get(f"{API}/projects/{p['id']}", headers=auth(partner)).json()
    assert after["status"] == "active" and after["milestones_total"] == 1 and after["machine_id"] == m["id"]
    print(f"\nDEMO {after['project_code']}: machine {after['machine_serial']}, status {after['status']}, "
          f"milestone '{ms.json()['title']}' ({dept.name}), task '{task.json()['title']}' -> {task.json().get('assignee_name')}")
