"""Design BOM "Import from Excel": the file is read into the form (nothing
saved) — Item Master by item_code, new-item rows by name, drawings by number,
per-row errors for bad rows."""
import io

from openpyxl import Workbook, load_workbook

from app.tests.test_cov_design_support import API, auth, create_bom, item, make_team, set_team


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEAD = ["position", "item_code", "new_item_name", "new_item_spec", "quantity", "uom", "drawing_no", "make_buy", "remarks"]


def test_bom_lines_import_reads_rows_and_reports_errors(client, db):
    team = make_team(db)
    set_team(client, team, {"author": "designer", "approver": "head"})
    plate = item(db, "MT-PL-0001", "Base plate", uom="NOS")
    item(db, "MT-OLD-0001", "Old bolt", status="inactive")
    h = auth(team["author"])

    tpl = client.get(f"{API}/boms/import-lines/template", headers=h)
    assert tpl.status_code == 200
    assert [c.value for c in load_workbook(io.BytesIO(tpl.content)).active[1]] == HEAD

    data = _xlsx([
        HEAD,
        [20, "mt-pl-0001 ", "", "", 4, "", "", "", "Weld side"],
        ["", "", "Side frame weldment", "IS 2062", 1, "", "", "make", ""],
        [30, "NOPE-1", "", "", 1, "", "", "", ""],
        [40, "MT-OLD-0001", "", "", 1, "", "", "", ""],
        [50, "MT-PL-0001", "", "", 0, "", "", "", ""],
        [20, "MT-PL-0001", "", "", 1, "", "", "", ""],
        [60, "MT-PL-0001", "", "", 1, "", "DWG-X", "", ""],
        [70, "MT-PL-0001", "", "", 1, "", "", "maybe", ""],
    ])
    r = client.post(f"{API}/boms/import-lines", files={"file": ("bom.xlsx", data)}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    lines = body["lines"]
    assert [(ln["position"], ln["store_item_id"], ln["new_item_name"], ln["quantity"], ln["uom"], ln["make_buy"]) for ln in lines] == [
        (20, plate.id, None, 4, "NOS", "buy"),
        (30, None, "Side frame weldment", 1, "NOS", "make"),
    ]
    reasons = {e["row"]: e["reason"] for e in body["errors"]}
    assert "isn't in the Item Master" in reasons[4]
    assert "inactive" in reasons[5]
    assert "greater than 0" in reasons[6]
    assert "used twice" in reasons[7]
    assert "Drawing 'DWG-X'" in reasons[8]
    assert "make or buy" in reasons[9]

    # Nothing was saved — the lines go into the normal create call.
    product = item(db, "MT-ASSY-0001", "Frame assembly")
    created = create_bom(client, team, product, [{k: ln[k] for k in ("position", "store_item_id", "new_item_name", "new_item_spec", "quantity", "uom", "make_buy")} for ln in lines])
    assert created.status_code == 201, created.text


def test_bom_lines_import_needs_quantity_column(client, db):
    team = make_team(db)
    set_team(client, team, {"author": "designer", "approver": "head"})
    r = client.post(f"{API}/boms/import-lines", files={"file": ("bom.xlsx", _xlsx([["item_code"], ["X"]]))}, headers=auth(team["author"]))
    assert r.status_code == 422 and "quantity" in r.json()["detail"]
