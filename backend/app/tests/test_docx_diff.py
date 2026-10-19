import io

from docx import Document
from docx.shared import Inches
from PIL import Image

from app.utils.docx_diff import compare_docx


def _docx(paragraphs, rows=(), image=False):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if rows:
        table = doc.add_table(rows=0, cols=len(rows[0]))
        for row in rows:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text
    if image:
        buf = io.BytesIO()
        Image.new("RGB", (10, 10), "red").save(buf, format="PNG")
        buf.seek(0)
        doc.add_picture(buf, width=Inches(1))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_identical_documents_have_no_changes():
    d = _docx(["Dear Sir,", "Thanks."])
    result = compare_docx(d, d)
    assert result["changed"] is False and result["lines"] == []


def test_edited_added_removed_lines_images_and_amounts():
    old = _docx(["QUOTATION", "Delivery: 7 weeks", "Old note"], rows=[["1", "RRV", "47,214.16"]])
    new = _docx(["QUOTATION", "Delivery: 6 weeks", "Warranty: 24 months"], rows=[["1", "RRV", "45,000.00"]], image=True)
    result = compare_docx(old, new)
    assert result["changed"] is True
    assert {"type": "changed", "old": "Delivery: 7 weeks", "new": "Delivery: 6 weeks"} in result["lines"]
    assert {"type": "changed", "old": "Old note", "new": "Warranty: 24 months"} in result["lines"]
    assert {"type": "changed", "old": "1 | RRV | 47,214.16", "new": "1 | RRV | 45,000.00"} in result["lines"]
    assert result["images_added"] == 1 and result["images_removed"] == 0
    assert result["amounts_removed"] == ["47,214.16"] and result["amounts_added"] == ["45,000.00"]


def test_pure_insert_is_reported_as_added():
    result = compare_docx(_docx(["A", "C"]), _docx(["A", "B", "C"]))
    assert result["lines"] == [{"type": "added", "text": "B"}]
