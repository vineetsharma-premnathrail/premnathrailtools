"""Procurement redesign (2026-10-07; approvals on the central engine since
2026-10-08, seeded from the old Settings): PR approvers (all
approve), two-level PO approval (Settings PO approvers → any one Partner),
vendor master + catalog import, RFQ/PO vendor link, PR export/import."""
import io
from itertools import count

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import object_session

from app.auth.jwt_handler import create_access_token
from app.modules.accounts.models.vendor import Vendor
from app.modules.main.models.notification import Notification
from app.modules.main.models.user import User
from app.tests.p2p_po_helpers import make_ship_to, post_upload_po, upload_po_for_pr
from app.tests.p2p_workflow_helpers import link_items, master_item

BASE = "/api/v1/p2p/requests"
_seq = count(1)


@pytest.fixture(autouse=True)
def _sqlite_number_series(monkeypatch):
    from app.modules.p2p import service
    monkeypatch.setattr(service, "_lock_number_series", lambda db, prefix: None)


# RBAC (2026-10-08): a Purchase-team user holds the `purchase` desk AND the
# Procurement tabs (p2p) AND Store › GRN & Inspection — the `purchase` app
# alone no longer opens requisitions, RFQs, POs or GRNs.
def _apps(apps):
    apps = list(apps)
    if "purchase" in apps and "p2p" not in apps:
        apps.append("p2p")
    return apps


def _grants(apps):
    if not apps:
        # A named PR / PO approver from another department: their role gives
        # them only 'View' on the two tabs; the routing lets them approve.
        return ["p2p:purchase_requisitions:view", "p2p:purchase_orders:view"]
    return ["store:grn:view", "store:grn:create", "store:grn:approve"] if "purchase" in apps else []


def _user(db, email, apps=("p2p",), role="user", **flags):
    u = User(email=email, name=email.split("@")[0], role=role, is_active=True, assigned_apps=_apps(apps), granular_permissions=_grants(apps), department="R&D")
    for k, v in flags.items():
        setattr(u, k, v)
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _h(u):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(u.id), 'email': u.email, 'role': u.role})}"}


def _setup(client, db):
    """Procurement › Settings as it stood before the approval engine: PR
    approvers and level-1 PO approvers stored as module settings (the keys are
    no longer editable — the engine seeds its workflows from them on first use),
    plus the company terms that still are a setting."""
    from app.modules.main.models.module_setting import ModuleSetting

    n = next(_seq)
    admin = _user(db, f"ps.admin{n}@premnathrail.com", role="admin")
    # Settings approvers from other departments: their role adds 'View' on
    # Requisitions / Purchase Orders (RBAC 2026-10-08) — routing does the rest.
    pr_a = _user(db, f"ps.pra{n}@premnathrail.com", apps=("hr",), granular_permissions=_grants(()))
    pr_b = _user(db, f"ps.prb{n}@premnathrail.com", apps=("store",), granular_permissions=_grants(()))
    po_a = _user(db, f"ps.poa{n}@premnathrail.com", apps=("hr",), granular_permissions=_grants(()))
    partner1 = _user(db, f"ps.partner1{n}@premnathrail.com", apps=(), is_partner=True)
    partner2 = _user(db, f"ps.partner2{n}@premnathrail.com", apps=(), is_partner=True)
    db.add_all([ModuleSetting(module_key="p2p", key="pr_approvers", value=[pr_a.id, pr_b.id]),
                ModuleSetting(module_key="p2p", key="po_approvers", value=[po_a.id])])
    db.commit()
    r = client.put("/api/v1/p2p/settings", json={"values": {"company_terms": "Prices firm till delivery.\nGST extra."}},
                   headers=_h(admin))
    assert r.status_code == 200, r.text
    return {"admin": admin, "pr": [pr_a, pr_b], "po": [po_a], "partners": [partner1, partner2]}


def _create(client, requester):
    body = {"category_code": "RAW", "project_label": "Project Alpha", "project_type": "new", "priority": "standard",
            "approvers": {"design_manager": 999999},  # ignored once Settings has approvers
            "items": [{"item_name": "Hydraulic Cylinder", "quantity": 2, "unit": "pcs", "project_inhouse": "Project"}]}
    r = client.post(BASE, json=link_items(object_session(requester), body), headers=_h(requester))
    assert r.status_code == 200, r.text
    return r.json()


def test_pr_needs_every_settings_approver(client, db):
    from app.modules.organization.services import approval_engine as engine

    s = _setup(client, db)
    req = _user(db, "ps.req1@premnathrail.com")
    pr = _create(client, req)
    # The engine's PR workflow was seeded from the old Settings list.
    steps = engine.latest_request(db, "p2p_request", pr["id"]).steps
    assert sorted(st.approver_id for st in steps) == sorted(u.id for u in s["pr"])
    assert client.post(f"{BASE}/{pr['id']}/approve", json={"comment": "ok"}, headers=_h(s["pr"][0])).json()["status"] == "submitted"
    assert client.post(f"{BASE}/{pr['id']}/approve", json={"comment": "ok"}, headers=_h(s["pr"][1])).json()["status"] == "approved"


def test_po_two_levels_then_any_one_partner(client, db):
    s = _setup(client, db)
    req = _user(db, "ps.req2@premnathrail.com")
    buyer = _user(db, "ps.buyer@premnathrail.com", apps=("purchase", "p2p"))
    pr = _create(client, req)
    for a in s["pr"]:
        client.post(f"{BASE}/{pr['id']}/approve", json={"comment": "ok"}, headers=_h(a))
    po = upload_po_for_pr(client, db, pr["id"], buyer, tally_po_no="PO-T-1")
    assert po["status"] == "issued"
    p1, p2 = s["partners"]
    # Partner can't jump the queue: the PO approvers' stage comes first.
    r = client.post(f"{BASE}/{pr['id']}/approve-po", json={"comment": "go"}, headers=_h(p1))
    assert r.status_code == 403 and "PO approvers" in r.json()["detail"]
    r = client.post(f"{BASE}/{pr['id']}/approve-po", json={"comment": "fine"}, headers=_h(s["po"][0]))
    assert r.status_code == 200 and r.json()["status"] == "po_raised"
    assert db.query(Notification).filter(Notification.user_id == p2.id, Notification.notification_type == "approval").count() >= 1
    r = client.post(f"{BASE}/{pr['id']}/approve-po", json={"comment": "approved"}, headers=_h(p2))
    assert r.status_code == 200 and r.json()["status"] == "po_approved"
    r = client.post(f"{BASE}/{pr['id']}/approve-po", json={"comment": "late"}, headers=_h(p1))
    assert r.status_code == 409


def test_vendor_master_catalog_and_exports(client, db):
    buyer = _user(db, "ps.vbuyer@premnathrail.com", apps=("purchase",))
    r = client.post("/api/v1/p2p/vendors", json={"name": "Shree Hydraulics", "contact_name": "Rakesh", "contact_phone": "9876543210",
                                                 "city": "Pune", "terms_conditions": "30 days credit"}, headers=_h(buyer))
    assert r.status_code == 201, r.text
    v = r.json()
    assert v["code"].startswith("V-") and v["gst_registration_type"] == "unregistered"
    assert db.get(Vendor, v["id"]).name == "Shree Hydraulics"  # same master Finance uses

    csv_body = "Item Name *,Part Code,Unit,Price,Dimensions,Technical Specification,Lead Time (days)\n" \
               "Hose 1/2,HH-12,Mtr,450,12.7 mm,275 bar,7\nFitting,FT-1,Nos,abc,,,\n"
    r = client.post(f"/api/v1/p2p/vendors/{v['id']}/catalog/import", files={"file": ("cat.csv", csv_body, "text/csv")}, headers=_h(buyer))
    assert r.status_code == 200, r.text
    assert r.json()["added"] == 1 and len(r.json()["errors"]) == 1
    # Vendors for an item are found by Item Master id (info record), never by name (2026-10-14).
    from app.modules.accounts.models.vendor_catalog_item import VendorCatalogItem
    from app.tests.p2p_workflow_helpers import master_item
    hose = master_item(db, "Hose 1/2 inch", uom="Mtr")
    line = db.query(VendorCatalogItem).filter_by(vendor_id=v["id"]).one()
    line.store_item_id = hose.id
    db.commit()
    hits = client.get("/api/v1/p2p/vendors/for-item", params={"store_item_id": hose.id}, headers=_h(buyer)).json()
    assert hits[0]["vendor_name"] == "Shree Hydraulics" and hits[0]["unit_price"] == 450

    imp = "Company Name *,City,GSTIN\nShree Hydraulics,Pune,\nNew Steel Co,Raipur,\n"
    r = client.post("/api/v1/p2p/vendors/import", files={"file": ("v.csv", imp, "text/csv")}, headers=_h(buyer))
    assert r.json()["created"] == ["New Steel Co"] and r.json()["existing"] == ["Shree Hydraulics"]
    wb = load_workbook(io.BytesIO(client.get("/api/v1/p2p/vendors/export", headers=_h(buyer)).content))
    assert sorted(row[1] for row in wb.active.iter_rows(min_row=2, values_only=True)) == ["New Steel Co", "Shree Hydraulics"]


def test_pr_import_and_export(client, db):
    _setup(client, db)
    req = _user(db, "ps.req3@premnathrail.com")
    # Lines are identified by their Item Master code; name / unit come from
    # the master (2026-10-11).
    master_item(db, "MS plate", code="RM-PLATE-01", uom="KG")
    master_item(db, "MS angle", code="RM-ANGLE-01", uom="KG")
    csv_body = ("PR Ref *,Category Code *,Project Type *,Project *,Item Code *,Quantity *,Project / Inhouse *\n"
                "A,RAW,existing,RRV-1,RM-PLATE-01,10,Project\nA,RAW,existing,RRV-1,rm-angle-01,5,Project\n"
                "B,XXX,new,Proto,RM-PLATE-01,1,Inhouse\n"
                "C,RAW,new,Proto,NOPE-999,1,Inhouse\n")
    r = client.post(f"{BASE}/import", files={"file": ("pr.csv", csv_body, "text/csv")}, headers=_h(req))
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body["created"]) == 1
    errors = {e["ref"]: e["reason"] for e in body["errors"]}
    assert set(errors) == {"B", "C"}
    assert "NOPE-999" in errors["C"] and "isn't in the Item Master" in errors["C"]
    r = client.get(f"{BASE}/export", headers=_h(req))
    assert r.status_code == 200, r.text
    sheet = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in sheet[1]]
    assert header[11:13] == ["Item Code", "Item"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert sorted((row[11], row[12], row[15]) for row in rows) == [
        ("RM-ANGLE-01", "MS angle", "KG"), ("RM-PLATE-01", "MS plate", "KG")]


def test_upload_po_vendor_choice_and_word_document(client, db):
    from datetime import date
    from app.modules.p2p.models.p2p_request import P2PRequest
    from app.modules.p2p.models.p2p_request_item import P2PRequestItem
    from app.modules.p2p.models.rfq import RFQ
    from app.modules.p2p.models.rfq_attachment import RFQAttachment
    from app.modules.p2p.reports.purchase_order import build_po_docx, po_context
    from app.modules.p2p.models.purchase_order import P2PPurchaseOrder

    buyer = _user(db, "ps.pobuyer@premnathrail.com", apps=("purchase", "p2p"))
    _user(db, "ps.pobuyer.partner@premnathrail.com", apps=("p2p",), is_partner=True)  # the PO goes to a Partner for approval
    v1 = Vendor(code="V-9001", name="Cheap Seals", status="active", gst_registration_type="unregistered")
    v2 = Vendor(code="V-9002", name="Fast Seals", status="active", gst_registration_type="unregistered", terms_conditions="Ex-works Pune")
    db.add_all([v1, v2])
    db.flush()
    pr = P2PRequest(p2p_number="P2P-RAW-2026-9001", category_code="RAW", status="vendor_quotations", project_label="RRV-1",
                    project_type="new", request_date=date.today(), requested_by_id=buyer.id, priority="standard")
    db.add(pr)
    db.flush()
    db.add(P2PRequestItem(p2p_request_id=pr.id, item_name="O-ring", quantity=10, unit="Nos", project_inhouse="Project",
                          fulfillment_status="sent_to_procurement"))
    # Company terms come from Procurement → Settings (default text here), not from the RFQ.
    rfq = RFQ(rfq_number="RFQ-2026-9001", p2p_request_id=pr.id, status="locked",
              late_delivery_clause="0.5% per week of delay, max 10%")
    db.add(rfq)
    db.flush()
    for tier, v in (("L1", v1), ("L2", v2)):
        db.add(RFQAttachment(rfq_id=rfq.id, vendor_tier=tier, vendor_name=v.name, vendor_id=v.id, filename=f"{tier}.pdf",
                             sharepoint_path=f"x/{tier}.pdf", created_by_id=buyer.id))
    db.commit()

    def _details(**over):
        return {"ship_to_location_id": make_ship_to(db).id,
                "lines": [{"pr_item_id": it.id, "unit_price": 50, "tax_rate": 18} for it in pr.items], **over}

    # The PO goes to the quote the buyer picks: L1 by default, any other tier needs a reason.
    r = post_upload_po(client, buyer, rfq.id, _details(vendor_tier="L2"), "TALLY-V-1")
    assert r.status_code == 400 and "instead of L1" in r.json()["detail"]
    r = post_upload_po(client, buyer, rfq.id, _details(vendor_tier="L2", selection_reason="L1 can't deliver in time"), "TALLY-V-1")
    assert r.status_code == 200, r.text
    po = r.json()
    assert po["vendor_name"] == "Fast Seals" and po["vendor_id"] == v2.id
    # The vendor's own terms are for the RFQ email only; the PO carries the company terms.
    assert "Ex-works Pune" not in (po["terms_conditions"] or "") and "quote the Purchase Order" in (po["terms_conditions"] or "")
    assert "Late delivery clause:\n0.5% per week of delay, max 10%" in po["terms_conditions"]
    # PO lines are the PR's own lines, priced from the Tally PO (never free-typed).
    assert [(i["item_name"], i["unit_price"], i["price_source"]) for i in po["items"]] == [("O-ring", 50, "manual")]

    # A line spec (printed under the line on the PO) is kept on the line.
    db.get(P2PPurchaseOrder, po["id"]).items[0].specification = "Size: 20 mm\nNBR 70"
    db.commit()

    row = db.get(P2PPurchaseOrder, po["id"])
    db.refresh(row)
    ctx = po_context(db, row)
    assert ctx["po_number"] == row.po_number and ctx["reference"].startswith("P2P-RAW-2026-9001")
    assert ctx["other_references"] == "FOR RRV-1" and ctx["lines"][0]["sub"] == ["Size: 20 mm", "NBR 70"]
    # 10 x 50 + 18% GST = 590; vendor state unknown -> one plain GST line.
    assert (ctx["total"], ctx["gst"], ctx["tax_split"]) == (590, 90.0, "unknown")
    assert ctx["amount_in_words"] == "INR Five Hundred Ninety Only"
    assert len(build_po_docx(ctx).getvalue()) > 10000


def test_item_rate_lookups_follow_the_vendor_prices_switch(client, db):
    """Security fix: /vendors/for-item and the info-record lookups returned
    vendor rates to a Vendors viewer whose role has Prices & rates off."""
    from app.modules.accounts.models.vendor_catalog_item import VendorCatalogItem
    v = Vendor(code="V-SEC-1", name="Price Secret Co", gst_registration_type="regular", status="active")
    db.add(v)
    db.flush()
    item = master_item(db, "Secret priced hose", uom="Mtr")
    db.add(VendorCatalogItem(vendor_id=v.id, store_item_id=item.id, item_name=item.item_name, unit_price=450.0,
                             currency="INR", is_active=True))
    db.commit()

    def lookups(user):
        h = {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}
        q = {"store_item_id": item.id}
        return (client.get("/api/v1/p2p/vendors/for-item", params=q, headers=h).json()[0]["unit_price"],
                client.get("/api/v1/p2p/info-records/for-item", params=q, headers=h).json()["info_records"][0]["unit_price"],
                client.get("/api/v1/p2p/info-records/default-price", params={**q, "vendor_id": v.id}, headers=h).json()["unit_price"],
                client.get("/api/v1/p2p/info-records/price-history", params=q, headers=h).status_code)

    def user(email, grants):
        u = User(email=email, name=email.split("@")[0], role="user", is_active=True, granular_permissions=grants)
        db.add(u)
        db.commit()
        return u

    # Customised role: Vendors view + item rates, Prices & rates not picked.
    blind = user("vend.blind@premnathrail.com", ["p2p:vendors:view", "p2p:vendors.info_records:view"])
    assert lookups(blind)[:3] == (None, None, None)
    sees = user("vend.sees@premnathrail.com", ["p2p:vendors:view"])  # inherits every Vendors feature
    assert lookups(sees)[:3] == (450, 450, 450)
