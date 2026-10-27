"""Service reports (modules/erp/reporting.py): derived figures, scope, filters,
engineer performance and drill consistency, tested on the data layer the
framework engine calls (loaders, measures, filters, composite)."""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace

from app.core.permission_registry import MODULES
from app.core.reporting import get_dataset, get_report, module_reports, token
from app.modules.erp import reporting as rep
from app.modules.erp.models.project import Project
from app.modules.erp.models.service_request import ServiceRequest
from app.modules.erp.models.service_visit import ServiceVisit
from app.modules.main.models.user import User

UTC = timezone.utc
TODAY = date(2026, 10, 26)
_seq = {"n": 0}


def mk_user(db, email, role="user", scope=None):
    u = User(email=email, name=email.split("@")[0], role=role, is_active=True, assigned_apps=["erp"],
             data_access_scopes=scope or {})
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def mk_project(db, serial, **kw):
    p = Project(serial_number=serial, model_name="RRV-X", **kw)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


def mk_sr(db, project, *, opened, status="open", assignee=None, creator=None, **kw):
    _seq["n"] += 1
    sr = ServiceRequest(request_number=f"SR-T-{_seq['n']:04d}", project_id=project.id, issue_title=kw.pop("title", "Leak"),
                        status=status, opened_at=opened, assigned_service_person_id=assignee.id if assignee else None,
                        created_by_id=creator.id if creator else None, **kw)
    db.add(sr)
    db.commit()
    db.refresh(sr)
    return sr


def mk_visit(db, sr, engineer, day, **kw):
    v = ServiceVisit(service_request_id=sr.id, engineer_id=engineer.id, visit_date=day, work_done=kw.pop("work_done", "Fixed"), **kw)
    db.add(v)
    db.commit()
    return v


def ctx_for(db, user, d_from=date(2026, 10, 1), d_to=date(2026, 10, 31), **kw):
    base = dict(db=db, user=user, today=TODAY, date_field="opened", d_from=d_from, d_to=d_to, filters={}, q=None,
                group_by="engineer", group_by2=None, measures=[], sees_money=True, sees_people=True, self_only=False)
    base.update(kw)
    return SimpleNamespace(**base)


def sr_rows(db, user):
    return rep.load_sr_rows(ctx_for(db, user))


def by_no(rows, sr):
    return next(r for r in rows if r["sr_id"] == sr.id)


# ── Registration ────────────────────────────────────────────────────────────

def test_ten_reports_and_two_datasets_registered_with_valid_defaults():
    keys = [r.key for r in module_reports("erp")]
    assert keys == ["engineer_performance", "pending_ageing", "turnaround", "volume", "customers_machines",
                    "repeat_failures", "billing_cost", "visit_log", "sr_explorer", "visit_explorer"]
    for r in module_reports("erp"):
        ds = get_dataset("erp", r.dataset)
        assert ds is not None
        dims = {d.key for d in ds.dimensions}
        measures = {m.key for m in (r.measures or ds.measures)}
        for g in (r.defaults.get("group_by"), r.defaults.get("group_by2")):
            assert g is None or g in dims or g in dict(r.group_bys) or g == "period_month"
        for m in r.defaults.get("measures", []):
            assert m in measures, f"{r.key}: unknown default measure {m}"
        for f in list(r.defaults.get("filters", {})) + list(r.locked_filters):
            assert f in {x.key for x in ds.filters}, f"{r.key}: unknown filter {f}"
        assert r.defaults["date_field"] in {d.key for d in ds.date_fields}
        for text in (r.label, r.description) + tuple(x for pair in r.definitions for x in pair):
            assert "sla" not in text.lower().split(), f"{r.key}: the word SLA must not appear"


def test_ratio_measures_reference_existing_measures_and_money_is_flagged():
    ds = get_dataset("erp", "sr")
    keys = {m.key for m in ds.measures}
    for m in ds.measures:
        if m.ratio:
            assert m.ratio[0] in keys and m.ratio[1] in keys
    money = {m.key for m in ds.measures if m.money}
    assert {"total_bill", "gross_margin", "actual_cost", "to_bill", "paid"} <= money
    assert ds.money_feature == ("erp", "service_requests", "charges")
    assert get_report("erp", "billing_cost").money_report is True
    assert get_report("erp", "engineer_performance").self_only_without_management is True
    assert any(c.money for c in ds.columns) and any(f.money for f in ds.filters)
    assert next(d for d in ds.dimensions if d.key == "billing_status").money is True


def test_management_feature_is_registered_explicit_only_on_erp_reports():
    feat = MODULES["erp"]["features"]["reports"]["management"]
    assert feat["label"] == "People & performance"
    assert feat["inherit"].get("view") is None


# ── Derived figures ─────────────────────────────────────────────────────────

def test_response_resolution_and_imported_sr_with_closed_at_only(db):
    admin, eng = mk_user(db, "r1admin@premnathrail.com", "admin"), mk_user(db, "r1eng@premnathrail.com")
    p = mk_project(db, "SN-R1")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    a = mk_sr(db, p, opened=t0, status="closed", assignee=eng, first_response_at=t0 + timedelta(hours=3),
              resolution_at=t0 + timedelta(days=2), closed_at=t0 + timedelta(days=3))
    imported = mk_sr(db, p, opened=t0 + timedelta(days=40), status="closed", assignee=eng,
                     closed_at=t0 + timedelta(days=44), service_type="amc")
    rows = sr_rows(db, admin)
    ra, ri = by_no(rows, a), by_no(rows, imported)
    assert ra["response_h"] == 3.0
    assert ra["resolution_days"] == 2.0            # resolution_at wins over closed_at
    assert ri["resolution_days"] == 4.0            # imported: closed_at stands in
    assert ri["response_h"] is None and ri["is_done"] is True


def test_cancelled_sr_is_not_done_and_has_no_timing(db):
    admin = mk_user(db, "r2admin@premnathrail.com", "admin")
    p = mk_project(db, "SN-R2")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    c = mk_sr(db, p, opened=t0, status="cancelled", first_response_at=t0 + timedelta(hours=1),
              closed_at=t0 + timedelta(days=1))
    r = by_no(sr_rows(db, admin), c)
    assert r["is_done"] is False and r["response_h"] is None and r["resolution_days"] is None
    assert r["status_group"] == "cancelled" and r["overdue"] is False and r["age_days"] is None


def test_on_time_uses_expected_date_then_internal_due_time(db):
    admin = mk_user(db, "r3admin@premnathrail.com", "admin")
    p = mk_project(db, "SN-R3")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    done = t0 + timedelta(days=3)  # 8 Oct
    ok = mk_sr(db, p, opened=t0, status="closed", resolution_at=done, expected_completion_date=date(2026, 10, 8))
    late = mk_sr(db, p, opened=t0, status="closed", resolution_at=done, expected_completion_date=date(2026, 10, 7))
    fallback = mk_sr(db, p, opened=t0, status="closed", resolution_at=done, sla_resolution_due_at=t0 + timedelta(days=2))
    none = mk_sr(db, p, opened=t0, status="closed", resolution_at=done)
    rows = sr_rows(db, admin)
    assert by_no(rows, ok)["on_time"] is True
    assert by_no(rows, late)["on_time"] is False
    assert by_no(rows, fallback)["on_time"] is False and by_no(rows, fallback)["target_date"] == date(2026, 10, 7)
    assert by_no(rows, none)["on_time"] is None and by_no(rows, none)["on_time_label"] == "no_target"


def test_ist_midnight_boundary_decides_the_reported_day(db):
    admin = mk_user(db, "r4admin@premnathrail.com", "admin")
    p = mk_project(db, "SN-R4")
    late_oct = mk_sr(db, p, opened=datetime(2026, 10, 31, 18, 0, tzinfo=UTC))   # 23:30 IST on 31 Oct
    early_nov = mk_sr(db, p, opened=datetime(2026, 10, 31, 19, 0, tzinfo=UTC))  # 00:30 IST on 1 Nov
    naive = mk_sr(db, p, opened=datetime(2026, 10, 31, 18, 0))                 # naive = UTC
    rows = sr_rows(db, admin)
    assert by_no(rows, late_oct)["reported_on"] == date(2026, 10, 31)
    assert by_no(rows, early_nov)["reported_on"] == date(2026, 11, 1)
    assert by_no(rows, naive)["reported_on"] == date(2026, 10, 31)
    ds = get_dataset("erp", "sr")
    opened = next(d for d in ds.date_fields if d.key == "opened").value
    assert opened(by_no(rows, late_oct)).month == 10 and opened(by_no(rows, early_nov)).month == 11


def test_overdue_age_and_bucket(db):
    admin = mk_user(db, "r5admin@premnathrail.com", "admin")
    p = mk_project(db, "SN-R5")
    young = mk_sr(db, p, opened=datetime(2026, 10, 25, 5, 0, tzinfo=UTC), expected_completion_date=date(2026, 10, 30))
    old = mk_sr(db, p, opened=datetime(2026, 8, 1, 5, 0, tzinfo=UTC), expected_completion_date=date(2026, 10, 20))
    waiting = mk_sr(db, p, opened=datetime(2026, 10, 1, 5, 0, tzinfo=UTC), status="review",
                    expected_completion_date=date(2026, 10, 2))
    rows = sr_rows(db, admin)
    assert by_no(rows, young)["overdue"] is False and by_no(rows, young)["age_bucket"] == "b1"
    ro = by_no(rows, old)
    assert ro["overdue"] is True and ro["days_overdue"] == 6 and ro["age_bucket"] == "b6"
    rw = by_no(rows, waiting)  # awaiting closure is never "overdue" (work is done) but is aged
    assert rw["overdue"] is False and rw["status_group"] == "awaiting_closure" and rw["age_bucket"] == "b4"
    assert {bucket: rep._bucket_age(d) for bucket, d in {"a": 2, "b": 3, "c": 15, "d": 31, "e": 61}.items()} == \
        {"a": "b1", "b": "b2", "c": "b3", "d": "b5", "e": "b6"}


def test_comeback_within_30_days_ignores_planned_work_and_other_machines(db):
    admin = mk_user(db, "r6admin@premnathrail.com", "admin")
    m1, m2 = mk_project(db, "SN-R6A"), mk_project(db, "SN-R6B")
    t0 = datetime(2026, 9, 1, 5, 0, tzinfo=UTC)
    first = mk_sr(db, m1, opened=t0, status="closed", resolution_at=t0 + timedelta(days=2), service_type="repair_breakdown")
    back = mk_sr(db, m1, opened=t0 + timedelta(days=20), service_type="repair_breakdown")            # 18 days after done
    amc = mk_sr(db, m1, opened=t0 + timedelta(days=21), service_type="amc")                          # planned: ignored
    far = mk_sr(db, m1, opened=t0 + timedelta(days=100), service_type="repair_breakdown")            # too late
    other = mk_sr(db, m2, opened=t0 + timedelta(days=21), service_type="repair_breakdown")           # other machine
    rows = sr_rows(db, admin)
    assert by_no(rows, first)["comeback"] is False
    assert by_no(rows, back)["comeback"] is True
    assert by_no(rows, amc)["comeback"] is False
    assert by_no(rows, far)["comeback"] is False and by_no(rows, other)["comeback"] is False


def test_visit_sums_per_sr_and_per_visit_engineer_and_arrival_hour(db):
    admin = mk_user(db, "r7admin@premnathrail.com", "admin")
    e1, e2 = mk_user(db, "r7e1@premnathrail.com"), mk_user(db, "r7e2@premnathrail.com")
    p = mk_project(db, "SN-R7")
    sr = mk_sr(db, p, opened=datetime(2026, 10, 5, 4, 0, tzinfo=UTC), assignee=e1, material_cost_actual=100.0, total_bill=1000.0,
               tax_amount=100.0, billing_status="invoiced")
    mk_visit(db, sr, e1, date(2026, 10, 6), work_hours=4, travel_hours=1.5, km=40, travel_cost=50, stay_cost=20, misc_cost=5,
             arrive_at=datetime(2026, 10, 6, 4, 0, tzinfo=UTC), depart_at=datetime(2026, 10, 6, 9, 0, tzinfo=UTC),
             customer_sign_name="Mr Rao")
    mk_visit(db, sr, e2, date(2026, 10, 7), work_hours=2, travel_hours=0.5, km=10, travel_cost=10)
    srs, visits = rep._load_all(db, admin, TODAY)
    r = by_no(srs, sr)
    assert (r["visit_count"], r["visit_work_h"], r["visit_travel_h"], r["visit_km"]) == (2, 6.0, 2.0, 50.0)
    assert r["visit_expenses"] == 85.0 and r["actual_cost"] == 185.0
    assert r["gross_margin"] == 1000.0 - 100.0 - 185.0 and r["invoiced"] == 1000.0 and r["to_bill"] == 0.0
    v1 = next(v for v in visits if v["engineer_id"] == e1.id)
    v2 = next(v for v in visits if v["engineer_id"] == e2.id)
    assert v1["arrival_hour"] == 9.5 and v1["on_site_h"] == 5.0 and v1["signed"] is True
    assert v2["arrival_hour"] is None and v2["signed"] is False and v2["engineer"] == "r7e2"
    assert v1["request_number"] == sr.request_number and v1["engineer_day"] != v2["engineer_day"]


# ── Scope ───────────────────────────────────────────────────────────────────

def test_assigned_scope_sees_only_own_and_assigned_srs_in_both_datasets(db):
    admin = mk_user(db, "s1admin@premnathrail.com", "admin")
    eng = mk_user(db, "s1eng@premnathrail.com", scope={"erp": "assigned"})
    other = mk_user(db, "s1other@premnathrail.com")
    p = mk_project(db, "SN-S1")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    mine = mk_sr(db, p, opened=t0, assignee=eng)
    created = mk_sr(db, p, opened=t0, creator=eng)
    theirs = mk_sr(db, p, opened=t0, assignee=other, creator=other)
    mk_visit(db, mine, eng, date(2026, 10, 6))
    mk_visit(db, theirs, other, date(2026, 10, 6))
    ids = {r["sr_id"] for r in sr_rows(db, eng)}
    assert ids == {mine.id, created.id}
    visits = rep.load_visit_rows(ctx_for(db, eng))
    assert [v["sr_id"] for v in visits] == [mine.id]
    assert {r["sr_id"] for r in sr_rows(db, admin)} >= {mine.id, created.id, theirs.id}
    # Deleted SRs never appear.
    mine.is_deleted = True
    db.commit()
    assert mine.id not in {r["sr_id"] for r in sr_rows(db, eng)}


def test_options_without_management_return_only_me(db):
    eng = mk_user(db, "s2eng@premnathrail.com", scope={"erp": "assigned"})
    boss = mk_user(db, "s2boss@premnathrail.com", "admin")
    p = mk_project(db, "SN-S2")
    mk_sr(db, p, opened=datetime(2026, 10, 5, 4, 0, tzinfo=UTC), assignee=eng)
    assert rep._people_options(ctx_for(db, eng, sees_people=False), "") == [{"value": str(eng.id), "label": "Me (s2eng)"}]
    team = rep._people_options(ctx_for(db, boss, sees_people=True), "")
    assert {"value": str(eng.id), "label": "s2eng"} in team
    assert rep._people_options(ctx_for(db, boss), "zzz") == []
    assert rep._raised_by_options(ctx_for(db, eng, sees_people=False), "")[0]["label"].startswith("Me (")


# ── Filters ─────────────────────────────────────────────────────────────────

def test_filters_and_semantics_and_not_assigned_token(db):
    admin = mk_user(db, "f1admin@premnathrail.com", "admin")
    eng = mk_user(db, "f1eng@premnathrail.com")
    p = mk_project(db, "SN-F1", zone="North")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    a = mk_sr(db, p, opened=t0, assignee=eng, priority="urgent", status="in_progress", service_type="repair_breakdown")
    b = mk_sr(db, p, opened=t0, priority="standard", status="closed", resolution_at=t0 + timedelta(days=1))
    rows = sr_rows(db, admin)
    flt = lambda **f: {r["sr_id"] for r in rep.filter_rows(rep.SR_FILTERS, rows, f)}  # noqa: E731
    assert a.id in flt(engineer=[eng.id]) and b.id not in flt(engineer=[eng.id])
    assert flt(engineer=["__none__"]) >= {b.id} and a.id not in flt(engineer=["__none__"])
    assert flt(priority=["urgent"], zone=["North"]) >= {a.id} and b.id not in flt(priority=["urgent"])
    assert b.id not in flt(priority=["urgent"], status=["closed"])           # AND across filters
    assert {a.id, b.id} <= flt(priority=["urgent", "standard"])              # OR inside one filter
    assert b.id in flt(completed=True) and a.id not in flt(completed=True)
    assert a.id in flt(status_group=["open"]) and b.id in flt(status_group=["closed"])
    assert a.id in flt(age_min=3, age_max=30) and flt(age_min=100) == set()
    assert token(None) == "__none__"


# ── Engineer performance (composite) ───────────────────────────────────────

def _perf_data(db):
    admin = mk_user(db, "e1admin@premnathrail.com", "admin")
    ravi, amit = mk_user(db, "ravi@premnathrail.com"), mk_user(db, "amit@premnathrail.com")
    m = mk_project(db, "SN-E1")
    n = mk_project(db, "SN-E2")
    t0 = datetime(2026, 10, 5, 4, 0, tzinfo=UTC)
    a1 = mk_sr(db, m, opened=t0, status="closed", assignee=ravi, first_response_at=t0 + timedelta(hours=2),
               resolution_at=t0 + timedelta(days=2), expected_completion_date=date(2026, 10, 8), total_bill=500.0,
               billing_status="paid", service_type="repair_breakdown")
    a2 = mk_sr(db, n, opened=t0, status="closed", assignee=ravi, first_response_at=t0 + timedelta(hours=4),
               resolution_at=t0 + timedelta(days=4), expected_completion_date=date(2026, 10, 6), service_type="repair_breakdown")
    p1 = mk_sr(db, m, opened=t0 + timedelta(days=1), status="in_progress", assignee=ravi, expected_completion_date=date(2026, 10, 10))
    u1 = mk_sr(db, n, opened=t0, status="open")
    c1 = mk_sr(db, n, opened=t0, status="cancelled", assignee=amit)
    mk_visit(db, a1, ravi, date(2026, 10, 6), work_hours=3, travel_hours=1, km=20, travel_cost=30)
    mk_visit(db, a1, ravi, date(2026, 10, 6), work_hours=2, travel_hours=0, km=0)      # same day: one day on site
    mk_visit(db, p1, amit, date(2026, 10, 9), work_hours=5, travel_hours=2, km=60)
    return admin, ravi, amit, dict(a1=a1, a2=a2, p1=p1, u1=u1, c1=c1)


def test_engineer_performance_figures_not_assigned_row_and_cancelled_excluded(db):
    admin, ravi, amit, s = _perf_data(db)
    res = rep.engineer_performance(ctx_for(db, admin))
    rows = {r["label"]["engineer"]: r["values"] for r in res.rows}
    assert list(rows) == ["amit", "ravi", "Not assigned"] or list(rows) == ["amit", "ravi", rep.NOT_ASSIGNED]
    r = rows["ravi"]
    assert r["assigned"] == 3 and r["completed"] == 2
    assert r["on_time_n"] == 1 and r["with_target_done"] == 2 and r["on_time_pct"] == 0.5
    assert r["pending_now"] == 1 and r["overdue_now"] == 1 and r["awaiting_closure"] == 0
    assert r["response_h.avg"] == 3.0
    assert r["resolution_days.avg"] == 3.0 and r["resolution_days.median"] == 3.0
    assert r["visit_count"] == 2 and r["days_on_site"] == 1 and r["work_h.sum"] == 5.0 and r["km.sum"] == 20.0
    assert r["billed"] == 500.0 and r["expenses.sum"] == 30.0
    # Visit figures follow the visit engineer, request figures the assignee.
    assert rows["amit"]["visit_count"] == 1 and rows["amit"]["assigned"] == 0 and rows["amit"]["work_h.sum"] == 5.0
    na = rows[rep.NOT_ASSIGNED]
    assert na["assigned"] == 1 and na["pending_now"] == 1 and na["visit_count"] == 0
    assert res.totals["completed"] == 2 and res.group_by == "engineer"
    assert [k.key for k in res.kpis][:3] == ["engineers_with_work", "completed", "on_time_pct"]
    assert res.series["workload"]["labels"][0] == "ravi"
    assert next(k for k in res.kpis if k.key == "engineers_with_work").value == 2
    assert any("Cancelled" in n for n in res.notes)


def test_engineer_performance_without_management_shows_only_own_row(db):
    admin, ravi, amit, s = _perf_data(db)
    res = rep.engineer_performance(ctx_for(db, admin, sees_people=False, self_only=True))
    assert res.rows == []                                  # the admin has no work of their own
    res = rep.engineer_performance(ctx_for(db, ravi, sees_people=False, self_only=True))
    assert [r["label"]["engineer"] for r in res.rows] == ["ravi"]
    assert any("People & performance" in n for n in res.notes)
    # Even if the engine did not force the filter, the report never shows another person.
    res2 = rep.engineer_performance(ctx_for(db, ravi, sees_people=False, self_only=False))
    assert [r["label"]["engineer"] for r in res2.rows] == ["ravi"]


def test_engineer_performance_month_children_and_month_first_grouping(db):
    admin, ravi, amit, s = _perf_data(db)
    res = rep.engineer_performance(ctx_for(db, admin, group_by2="period_month", d_from=date(2026, 9, 15), d_to=date(2026, 10, 31)))
    r = next(x for x in res.rows if x["label"]["engineer"] == "ravi")
    assert [c["key"]["period_month"] for c in r["children"]] == ["2026-10"]
    assert r["children"][0]["values"]["completed"] == r["values"]["completed"]
    res = rep.engineer_performance(ctx_for(db, admin, group_by="period_month", d_from=date(2026, 9, 15), d_to=date(2026, 10, 31)))
    assert res.group_by == "period_month" and [x["key"]["period_month"] for x in res.rows] == ["2026-10"]
    assert {c["label"]["engineer"] for c in res.rows[0]["children"]} == {"ravi", "amit", rep.NOT_ASSIGNED}


def _apply_drill(db, user, drill, d_from, d_to):
    """Resolve a drill the way the engine does: dataset, date field window, f_ filters."""
    ds = get_dataset("erp", "visit" if drill["report"] == "visit_explorer" else "sr")
    rows = ds.load(ctx_for(db, user), None)
    field = next(d for d in ds.date_fields if d.key == drill.get("date_field", ds.default_date_field))
    if field.value is not None:
        lo, hi = date.fromisoformat(drill["date_from"]), date.fromisoformat(drill["date_to"])
        rows = [r for r in rows if field.value(r) is not None and lo <= field.value(r) <= hi]
    filters = {}
    for k, v in drill.items():
        if k.startswith("f_"):
            filters[k[2:]] = [x for x in v.split(",")]
    for k in list(filters):
        fd = next(f for f in ds.filters if f.key == k)
        if fd.type == "bool":
            filters[k] = True
        elif fd.type in ("user", "entity"):
            filters[k] = filters[k]
    return rep.filter_rows(ds.filters, rows, filters)


def test_every_composite_cell_drill_matches_its_figure(db):
    admin, ravi, amit, s = _perf_data(db)
    d_from, d_to = date(2026, 10, 1), date(2026, 10, 31)
    res = rep.engineer_performance(ctx_for(db, admin, d_from=d_from, d_to=d_to))
    counts = {"assigned": None, "completed": None, "on_time_n": None, "with_target_done": None, "pending_now": None,
              "overdue_now": None, "awaiting_closure": None, "comebacks": None, "visit_count": None}
    checked = 0
    for row in res.rows:
        for key in counts:
            drill = row["measure_drill"].get(key) or (row["measure_drill"].get("on_time_pct") if key == "with_target_done" else None)
            if drill is None:
                continue
            got = _apply_drill(db, admin, drill, d_from, d_to)
            assert len(got) == row["values"][key], f"{row['label']['engineer']} / {key}: drill has {len(got)} rows, figure says {row['values'][key]}"
            checked += 1
    assert checked >= 9 * 3 - 3
    # KPI drills carry the same dates.
    k = next(k for k in res.kpis if k.key == "completed")
    assert len(_apply_drill(db, admin, k.drill, d_from, d_to)) == res.totals["completed"]
    k = next(k for k in res.kpis if k.key == "overdue_now")
    assert len(_apply_drill(db, admin, k.drill, d_from, d_to)) == res.totals["overdue_now"]


def test_month_helpers():
    assert rep._months_between(date(2026, 11, 20), date(2027, 2, 1)) == ["2026-11", "2026-12", "2027-01", "2027-02"]
    assert rep._month_bounds("2026-12") == (date(2026, 12, 1), date(2026, 12, 31))
    assert rep._month_bounds("2028-02")[1] == date(2028, 2, 29)


def test_dataset_row_ids_and_links_are_declared(db):
    sr, visit = get_dataset("erp", "sr"), get_dataset("erp", "visit")
    assert sr.id_keys == ("sr_id", "project_id") and "visit_id" in visit.id_keys
    link = {c.key: c.link for c in sr.columns}
    assert link["request_number"] == ("sr_id", "/dashboard/erp/service-requests")
    assert link["serial"] == ("project_id", "/dashboard/projects/service-projects")
    cols = {c.key: c for c in sr.columns}
    assert cols["issue"].cut and cols["total_bill"].money and cols["total_bill"].total == "sum"


# ── Through the framework engine (validation, permissions, drill consistency) ──

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

from app.core.reporting import engine  # noqa: E402

RANGE = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def grant_user(db, email, grants):
    u = User(email=email, name=email.split("@")[0], role="user", is_active=True, granular_permissions=grants)
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def run_report(db, user, key, **qp):
    rep_def = get_report("erp", key)
    params = {**RANGE, **qp}
    return engine.run(engine.prepare(db, user, "erp", rep_def, params))


BASE_GRANTS = ["erp:reports:view", "erp:service_requests:view"]


def test_every_report_runs_for_an_admin_without_error(db):
    admin, ravi, amit, s = _perf_data(db)
    for r in module_reports("erp"):
        out = run_report(db, admin, r.key)
        assert out["report"] == r.key
        assert out["level"] in ("summary", "detail")
    out = run_report(db, admin, "turnaround", group_by="engineer", group_by2="service_type", compare="previous_period")
    assert out["rows"] and out["rows"][0]["children"] is not None and "compare" in out["rows"][0]
    for report in ("sr_explorer", "visit_explorer"):
        out = run_report(db, admin, report, level="detail")
        assert out["row_count"] > 0 and {"sr_id", "project_id"} <= set(out["rows"][0])


def test_summary_cell_drills_return_exactly_the_records_behind_each_number(db):
    admin, ravi, amit, s = _perf_data(db)
    out = run_report(db, admin, "turnaround", group_by="priority", date_field="opened",
                     measures="count,completed,on_time_n,open,overdue,comebacks,urgent,free_count")
    checked = 0
    for row in out["rows"]:
        for m, value in row["values"].items():
            if m not in row["measure_drill"] and m != "count":
                continue
            if m in row["measure_drill"] and row["measure_drill"][m] is None:
                assert value == 0, f"{row['label']} / {m}: no records behind the number but it says {value}"
                continue
            drill = row["measure_drill"].get(m) or row["drill"]
            d = run_report(db, admin, "sr_explorer", level="detail", **{k: v for k, v in drill.items() if k != "report"})
            assert d["row_count"] == value, f"{row['label']} / {m}: drill gives {d['row_count']}, figure {value}"
            checked += 1
    assert checked >= 6
    for k in out["kpis"]:
        drill = {k2: v for k2, v in k["drill"].items() if k2 != "report"}
        assert run_report(db, admin, "sr_explorer", level="detail", **drill)["row_count"] == k["value"], k["key"]


def test_time_bucket_grouping_drill_is_the_month(db):
    admin, ravi, amit, s = _perf_data(db)
    out = run_report(db, admin, "volume", group_by="period_month", measures="count,completed", date_from="2026-09-01")
    row = next(r for r in out["rows"] if r["values"]["count"] > 0)
    d = run_report(db, admin, "sr_explorer", level="detail", **{k: v for k, v in row["drill"].items() if k != "report"})
    assert d["row_count"] == row["values"]["count"]


def test_pending_ageing_default_counts_open_and_awaiting_closure_only(db):
    admin, ravi, amit, s = _perf_data(db)
    out = run_report(db, admin, "pending_ageing", **{"f_status_group": "open,awaiting_closure"})
    assert out["period"]["date_field"] == "none"
    assert out["totals"]["values"]["count"] == 2      # p1 and u1; closed and cancelled are not pending
    labels = {r["label"]["age_bucket"] for r in out["rows"]}
    assert labels <= set(rep.AGE_LABELS.values())


def test_without_management_people_figures_are_refused_but_own_figures_work(db):
    admin, ravi, amit, s = _perf_data(db)
    me = grant_user(db, "plain@premnathrail.com", BASE_GRANTS)
    own = grant_user(db, "ravi2@premnathrail.com", BASE_GRANTS)
    mk_sr(db, mk_project(db, "SN-P1"), opened=datetime(2026, 10, 7, 4, 0, tzinfo=UTC), assignee=own, status="closed",
          resolution_at=datetime(2026, 10, 8, 4, 0, tzinfo=UTC))
    with pytest.raises(HTTPException) as e:
        run_report(db, me, "pending_ageing", group_by="engineer")
    assert e.value.status_code == 403 and "People & performance" in e.value.detail
    with pytest.raises(HTTPException) as e:
        run_report(db, me, "turnaround", **{"f_engineer": str(ravi.id)})
    assert e.value.status_code == 403
    assert run_report(db, me, "turnaround", group_by="priority")["rows"] is not None
    out = run_report(db, own, "engineer_performance")
    assert out["self_only"] is True and [r["label"]["engineer"] for r in out["rows"]] == ["ravi2"]
    assert any("People & performance" in n for n in out["notes"])
    log = run_report(db, own, "visit_log")
    assert log["self_only"] is True and log["filters_applied"][0]["value"] == "Me"
    assert run_report(db, own, "visit_explorer", level="detail")["row_count"] == 0   # own visits only
    assert run_report(db, own, "sr_explorer", level="detail")["row_count"] >= 1
    prep = engine.prepare(db, own, "erp", get_report("erp", "turnaround"), RANGE)
    f = next(f for f in get_dataset("erp", "sr").filters if f.key == "engineer")
    assert engine.options(prep, f, "") == [{"value": own.id, "label": "Me (ravi2)"}]


def test_management_right_unlocks_team_figures(db):
    admin, ravi, amit, s = _perf_data(db)
    boss = grant_user(db, "boss@premnathrail.com", BASE_GRANTS + ["erp:reports.management:view"])
    out = run_report(db, boss, "engineer_performance")
    assert {r["label"]["engineer"] for r in out["rows"]} >= {"ravi", "amit", rep.NOT_ASSIGNED}
    assert run_report(db, boss, "pending_ageing", group_by="engineer")["rows"]
    prep = engine.prepare(db, boss, "erp", get_report("erp", "turnaround"), RANGE)
    f = next(f for f in get_dataset("erp", "sr").filters if f.key == "engineer")
    assert {"ravi", "amit"} <= {o["label"] for o in engine.options(prep, f, "")}


def test_billing_report_and_money_are_denied_without_the_charges_switch(db):
    admin, ravi, amit, s = _perf_data(db)
    blind = grant_user(db, "blind@premnathrail.com", ["erp:reports:view", "erp:service_requests:view",
                                                      "erp:service_requests.visits:view"])  # Charges not picked
    sees = grant_user(db, "sees@premnathrail.com", BASE_GRANTS)
    with pytest.raises(HTTPException) as e:
        run_report(db, blind, "billing_cost")
    assert e.value.status_code == 403 and "Charges, costs & billing amounts" in e.value.detail
    with pytest.raises(HTTPException) as e:
        run_report(db, blind, "turnaround", measures="count,total_bill")
    assert e.value.status_code == 403
    with pytest.raises(HTTPException) as e:
        run_report(db, blind, "turnaround", **{"f_billing_status": "paid"})
    assert e.value.status_code == 403
    det = run_report(db, blind, "sr_explorer", level="detail")
    keys = {c["key"] for c in det["columns"]}
    assert "total_bill" not in keys and "gross_margin" not in keys and "request_number" in keys
    assert all("total_bill" not in r for r in det["rows"])
    ep = run_report(db, blind, "engineer_performance")
    with pytest.raises(HTTPException):
        run_report(db, blind, "engineer_performance", measures="assigned,billed")
    assert "billed" not in {m["key"] for m in ep["measures"]} and all("billed" not in r["values"] for r in ep["rows"])
    ok = run_report(db, sees, "billing_cost")
    assert ok["totals"]["values"]["total_bill"] == 500.0 and ok["totals"]["values"]["paid"] == 500.0
    assert "billed" in {m["key"] for m in run_report(db, sees, "engineer_performance", measures="assigned,billed")["measures"]}


def test_scoped_user_sees_only_own_srs_in_summary_and_drill(db):
    admin, ravi, amit, s = _perf_data(db)
    scoped = User(email="scoped@premnathrail.com", name="scoped", role="user", is_active=True,
                  granular_permissions=BASE_GRANTS, data_access_scopes={"erp": "assigned"})
    db.add(scoped)
    db.commit()
    db.refresh(scoped)
    mine = mk_sr(db, mk_project(db, "SN-SC1"), opened=datetime(2026, 10, 7, 4, 0, tzinfo=UTC), assignee=scoped)
    out = run_report(db, scoped, "turnaround", group_by="status", measures="count", date_field="opened")
    assert out["totals"]["values"]["count"] == 1
    assert run_report(db, admin, "turnaround", group_by="status", measures="count", date_field="opened")["totals"]["values"]["count"] >= 6
    det = run_report(db, scoped, "sr_explorer", level="detail", f_status_group="open,awaiting_closure,closed,cancelled")
    assert [r["sr_id"] for r in det["rows"]] == [mine.id]


def test_bad_inputs_get_real_messages(db):
    admin, ravi, amit, s = _perf_data(db)
    with pytest.raises(HTTPException) as e:
        run_report(db, admin, "turnaround", **{"f_priority": "high"})
    assert e.value.status_code == 422 and "Standard, Urgent" in e.value.detail
    with pytest.raises(HTTPException) as e:
        run_report(db, admin, "turnaround", **{"f_customer": "abc"})
    assert e.value.status_code == 422 and "picked from the list" in e.value.detail
    with pytest.raises(HTTPException) as e:
        run_report(db, admin, "turnaround", **{"f_age_min": "30", "f_age_max": "10"})
    assert e.value.status_code == 422
    with pytest.raises(HTTPException) as e:
        run_report(db, admin, "turnaround", group_by="planet")
    assert e.value.status_code == 422


# ── Through the HTTP router (the way the frontend calls it) ─────────────────

from io import BytesIO  # noqa: E402

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.core.reporting.router import router as reports_router  # noqa: E402
from app.db.session import get_db  # noqa: E402
from app.modules.main.models.audit_log import AuditLog  # noqa: E402
from app.modules.main.routes.auth import get_current_user  # noqa: E402


def _api(db, user):
    app = FastAPI()
    app.include_router(reports_router, prefix="/api/v1")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app)


def test_http_catalogue_run_and_export_for_service(db):
    admin, ravi, amit, s = _perf_data(db)
    c = _api(db, admin)
    cat = c.get("/api/v1/reports/erp")
    assert cat.status_code == 200, cat.text
    body = cat.json()
    assert [r["key"] for r in body["reports"]][:2] == ["engineer_performance", "pending_ageing"]
    assert {"sr", "visit"} <= {d["key"] for d in body["datasets"]}
    assert body["sees_money"] and body["sees_people"]
    run = c.get("/api/v1/reports/erp/turnaround", params={**RANGE, "group_by": "priority"})
    assert run.status_code == 200, run.text
    det = c.get("/api/v1/reports/erp/sr_explorer", params={**RANGE, "level": "detail"}).json()
    assert det["row_count"] >= 5 and "sr_id" in det["rows"][0]
    ex = c.get("/api/v1/reports/erp/engineer_performance/export", params=RANGE)
    assert ex.status_code == 200, ex.text
    wb = load_workbook(BytesIO(ex.content))
    assert "Summary" in wb.sheetnames
    assert db.query(AuditLog).filter(AuditLog.entity_type == "report_export").count() >= 1


def test_http_catalogue_hides_money_and_people_items_without_the_switches(db):
    blind = grant_user(db, "httpblind@premnathrail.com", ["erp:reports:view", "erp:service_requests:view",
                                                          "erp:service_requests.visits:view"])
    body = _api(db, blind).get("/api/v1/reports/erp").json()
    assert "billing_cost" not in {r["key"] for r in body["reports"]}
    sr = next(d for d in body["datasets"] if d["key"] == "sr")
    assert "total_bill" not in {m["key"] for m in sr["measures"]} and "gross_margin" not in {c["key"] for c in sr["columns"]}
    assert "engineer" not in {d["key"] for d in sr["dimensions"]}
    assert "billing_status" not in {f["key"] for f in sr["filters"]}
    r = _api(db, blind).get("/api/v1/reports/erp/billing_cost", params=RANGE)
    assert r.status_code == 403 and "Charges, costs & billing amounts" in r.json()["detail"]
    opts = _api(db, blind).get("/api/v1/reports/erp/turnaround/options/engineer")
    assert opts.status_code == 200 and [o["label"] for o in opts.json()] == ["Me (httpblind)"]


# ── Dashboard pages (ready-made KPI cards and charts) ───────────────────────

from app.core.reporting import dates as rdates  # noqa: E402
from app.core.reporting import module_dashboards  # noqa: E402
from app.core.reporting.engine import all_filters  # noqa: E402

DASH = "/api/v1/reports/erp/dashboard"
PAGES = ["overview", "engineers", "pending", "timing", "customers", "billing", "visits"]
OCT_BODY = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


def dash_post(db, user, page, body=None):
    return _api(db, user).post(f"{DASH}/{page}", json={**OCT_BODY, **(body or {})})


def live_today_rows(db, user):
    return rep._load_all(db, user, rdates.today_ist())[0]


def test_dashboard_definition_is_consistent_with_the_datasets():
    dash = module_dashboards("erp")[0]
    assert [p.key for p in dash.pages] == PAGES
    for page in dash.pages:
        pds = get_dataset("erp", page.dataset)
        assert page.hint
        assert set(page.slicers) <= set(all_filters(pds)), page.key
        assert page.date_mode == ("snapshot" if page.key == "pending" else "range")
        assert page.money_page == (page.key == "billing")
        assert 6 <= len(page.kpis) + 0 <= 8 or page.key in ("timing", "customers"), page.key
        for v in (*page.kpis, *page.visuals):
            ds = get_dataset("erp", v.dataset or page.dataset)
            mm, dm = {m.key for m in ds.measures}, {d.key: d for d in ds.dimensions}
            assert v.definition and v.title, (page.key, v.id)
            assert "sla" not in (v.title + " " + v.definition).lower().split()
            for k in ([v.measure] if v.measure else []) + list(v.measures):
                assert k in mm, (page.key, v.id, k)
            for k in v.columns:
                assert k in {c.key for c in ds.columns}, (page.key, v.id, k)
            if v.kind == "kpi":
                assert v.measure and v.id.startswith("k_")
                continue
            assert v.group_by in dm or v.kind == "table", (page.key, v.id)
            drill = v.drill or v.group_by
            if drill and not dm[drill].time_bucket:
                assert drill in all_filters(ds), f"{page.key}/{v.id}: click filter '{drill}' is not a filter"
            if v.group_by2:
                assert v.group_by2 in dm
            if v.kind != "table":
                assert v.measure or v.measures


def test_every_dashboard_page_evaluates_through_the_batch_endpoint(db):
    admin, ravi, amit, s = _perf_data(db)
    for page in PAGES:
        r = dash_post(db, admin, page)
        assert r.status_code == 200, f"{page}: {r.text}"
        body = r.json()
        pg = next(p for p in module_dashboards("erp")[0].pages if p.key == page)
        assert set(body["kpis"]) == {v.id for v in pg.kpis}, page
        assert set(body["results"]) == {v.id for v in pg.visuals}, page
        assert body["hint"]
    cat = _api(db, admin).get("/api/v1/reports/erp").json()
    assert [p["key"] for p in cat["dashboards"][0]["pages"]] == PAGES


def test_overview_kpis_match_the_underlying_records(db):
    admin, ravi, amit, s = _perf_data(db)
    rows = [r for r in live_today_rows(db, admin) if r["reported_on"] and date(2026, 10, 1) <= r["reported_on"] <= date(2026, 10, 31)]
    k = dash_post(db, admin, "overview").json()["kpis"]
    assert k["k_total"]["value"] == len(rows) >= 5
    assert k["k_open"]["value"] == sum(1 for r in rows if r["status_group"] == "open")
    assert k["k_completed"]["value"] == sum(1 for r in rows if r["is_done"]) == 2
    assert k["k_overdue"]["value"] == sum(1 for r in rows if r["overdue"])
    assert k["k_comebacks"]["value"] == sum(1 for r in rows if r["comeback"])
    ot = [r for r in rows if r["on_time"] is not None]
    assert k["k_ontime"]["value"] == sum(1 for r in ot if r["on_time"]) / len(ot) == 0.5
    assert k["k_response"]["value"] == 3.0 and k["k_resolution"]["value"] == 3.0
    assert k["k_total"]["better"] is None and k["k_open"]["better"] == "down" and k["k_ontime"]["better"] == "up"
    assert k["k_total"]["definition"]


def test_a_bar_click_filter_returns_exactly_the_records_behind_it(db):
    admin, ravi, amit, s = _perf_data(db)
    body = dash_post(db, admin, "overview").json()
    checked = 0
    for vid in ("status", "priority", "categories", "service_type", "customers", "cover"):
        res = body["results"][vid]
        assert res["drill"]["dimension"] and len(res["tokens"]) == len(res["labels"]) == len(res["values"])
        for tok, val in zip(res["tokens"], res["values"]):
            d = run_report(db, admin, "sr_explorer", level="detail", date_field="opened", **{f"f_{res['drill']['dimension']}": tok})
            assert d["row_count"] == val, f"{vid} / {tok}: click gives {d['row_count']} records, bar says {val}"
            checked += 1
    assert checked >= 10


def test_cross_filter_changes_other_cards_but_a_card_can_ignore_its_own_selection(db):
    admin, ravi, amit, s = _perf_data(db)
    urgent = len([r for r in live_today_rows(db, admin) if r["priority"] == "urgent" and r["reported_on"]
                  and date(2026, 10, 1) <= r["reported_on"] <= date(2026, 10, 31)])
    r = dash_post(db, admin, "overview", {"filters": {"f_priority": "urgent"},
                                          "visuals": [{"id": "k_total"}, {"id": "priority", "filters": {}}, {"id": "status"}]}).json()
    assert r["kpis"]["k_total"]["value"] == urgent
    assert sum(r["results"]["status"]["values"]) == urgent
    assert sum(r["results"]["priority"]["values"]) >= urgent and len(r["results"]["priority"]["labels"]) >= 1
    assert set(r["results"]) == {"priority", "status"}
    assert {f["key"] for f in r["filters_applied"]} == {"priority"}
    # Two values of one field are OR, two fields are AND.
    both = dash_post(db, admin, "overview", {"filters": {"f_priority": "urgent,standard"}, "visuals": [{"id": "k_total"}]}).json()
    assert both["kpis"]["k_total"]["value"] >= urgent
    none = dash_post(db, admin, "overview", {"filters": {"f_priority": "urgent", "f_status": "closed"}, "visuals": [{"id": "k_total"}]}).json()
    assert none["kpis"]["k_total"]["value"] == 0


def test_pending_page_is_a_snapshot_and_ignores_the_date_range(db):
    admin, ravi, amit, s = _perf_data(db)
    a = dash_post(db, admin, "pending").json()
    b = dash_post(db, admin, "pending", {"date_from": "2019-01-01", "date_to": "2019-01-31"}).json()
    assert a["period"]["mode"] == "snapshot" and a["period"]["from"] is None
    assert a["kpis"] == b["kpis"] and a["results"] == b["results"]
    pend = [r for r in live_today_rows(db, admin) if r["status_group"] in ("open", "awaiting_closure")]
    assert a["kpis"]["k_open"]["value"] == sum(1 for r in pend if r["status_group"] == "open") >= 2
    assert a["kpis"]["k_unassigned"]["value"] == sum(1 for r in pend if r["engineer_id"] is None) >= 1
    assert a["kpis"]["k_oldest"]["value"] == max(r["age_days"] for r in pend)
    assert sum(a["results"]["status"]["values"]) == len(pend)         # closed and cancelled never appear
    bad = dash_post(db, admin, "pending", {"compare": "previous_period"})
    assert bad.status_code == 422 and "as of today" in bad.json()["detail"]
    over = a["results"]["overdue_list"]
    assert over["rows"] and all(r["days_overdue"] is not None for r in over["rows"])
    days = [r["days_overdue"] for r in over["rows"]]
    assert days == sorted(days, reverse=True)


def test_age_buckets_are_in_age_order(db):
    admin, ravi, amit, s = _perf_data(db)
    labels = dash_post(db, admin, "pending").json()["results"]["age"]["labels"]
    order = list(rep.AGE_LABELS.values())
    assert labels == sorted(labels, key=order.index), f"age buckets out of order: {labels}"


def test_people_and_money_pages_and_visuals_follow_the_rights(db):
    admin, ravi, amit, s = _perf_data(db)
    own = grant_user(db, "dash_own@premnathrail.com", BASE_GRANTS)
    blind = grant_user(db, "dash_blind@premnathrail.com", ["erp:reports:view", "erp:service_requests:view",
                                                           "erp:service_requests.visits:view"])
    boss = grant_user(db, "dash_boss@premnathrail.com", BASE_GRANTS + ["erp:reports.management:view"])
    mk_sr(db, mk_project(db, "SN-D1"), opened=datetime(2026, 10, 7, 4, 0, tzinfo=UTC), assignee=own, status="closed",
          resolution_at=datetime(2026, 10, 8, 4, 0, tzinfo=UTC), total_bill=100.0, billing_status="invoiced")
    # catalogue: the Billing page is absent without the charges switch, the pending-by-engineer card without management
    cat = lambda u: _api(db, u).get("/api/v1/reports/erp").json()["dashboards"][0]["pages"]  # noqa: E731
    assert "billing" not in {p["key"] for p in cat(blind)} and "billing" in {p["key"] for p in cat(own)}
    pend = lambda u: next(p for p in cat(u) if p["key"] == "pending")  # noqa: E731
    ids = lambda p: {v["id"] for v in p.get("visuals", [])}  # noqa: E731
    assert "engineer" not in ids(pend(own)) and "engineer" in ids(pend(boss))
    # running a denied page / card names the right to ask for
    r = dash_post(db, blind, "billing")
    assert r.status_code == 403 and "Charges, costs & billing amounts" in r.json()["detail"]
    r = dash_post(db, own, "pending", {"visuals": [{"id": "engineer"}]})
    assert r.status_code == 403 and "People & performance" in r.json()["detail"]
    r = dash_post(db, own, "pending").json()
    assert "engineer" not in r["results"] and "overdue_list" in r["results"]
    # engineers page: own figures only, with the note
    e = dash_post(db, own, "engineers").json()
    assert e["self_only"] is True and any("People & performance" in n for n in e["notes"])
    assert [row["label"] for row in e["results"]["table"]["rows"]] == ["dash_own"]
    assert e["kpis"]["k_engineers"]["value"] == 1
    team = dash_post(db, boss, "engineers").json()
    assert {row["label"] for row in team["results"]["table"]["rows"]} >= {"ravi", "amit"}
    assert team["self_only"] is False
    # visits page: own visits only without management
    v = dash_post(db, own, "visits").json()
    assert v["self_only"] is True and v["kpis"]["k_visits"]["value"] == 0
    assert dash_post(db, boss, "visits").json()["kpis"]["k_visits"]["value"] == 3
    # money cards on the visits page are absent without the charges switch
    assert "k_expenses" not in dash_post(db, blind, "visits").json()["kpis"]
    assert "k_expenses" in dash_post(db, boss, "visits").json()["kpis"]
    # billing page works with the switch and its numbers match
    b = dash_post(db, own, "billing").json()
    assert b["kpis"]["k_billed"]["value"] >= 100.0 + 500.0 - 500.0 and b["kpis"]["k_paid"]["value"] == 500.0


def test_dashboard_data_scope_applies(db):
    admin, ravi, amit, s = _perf_data(db)
    scoped = User(email="dash_scoped@premnathrail.com", name="dash_scoped", role="user", is_active=True,
                  granular_permissions=BASE_GRANTS, data_access_scopes={"erp": "assigned"})
    db.add(scoped)
    db.commit()
    db.refresh(scoped)
    mk_sr(db, mk_project(db, "SN-D2"), opened=datetime(2026, 10, 7, 4, 0, tzinfo=UTC), assignee=scoped)
    assert dash_post(db, scoped, "overview").json()["kpis"]["k_total"]["value"] == 1
    assert dash_post(db, admin, "overview").json()["kpis"]["k_total"]["value"] >= 6


def test_dashboard_bad_requests_get_real_messages_and_export_works(db):
    admin, ravi, amit, s = _perf_data(db)
    r = dash_post(db, admin, "nope")
    assert r.status_code == 404 and "Reload the Reports page" in r.json()["detail"]
    r = dash_post(db, admin, "overview", {"filters": {"f_priority": "high"}})
    assert r.status_code == 422 and "Standard, Urgent" in r.json()["detail"]
    r = dash_post(db, admin, "overview", {"visuals": [{"id": "ghost"}]})
    assert r.status_code == 422
    ex = _api(db, admin).get(f"{DASH}/overview/export", params=OCT_BODY)
    assert ex.status_code == 200, ex.text
    assert "Summary" in load_workbook(BytesIO(ex.content)).sheetnames
