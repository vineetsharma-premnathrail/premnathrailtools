"""CO: KP06 cost-centre planning (versions, grid, Excel, plan vs actual),
KL01/KP26 activity rates used by production costing, KSU5/KSV5 assessment
and distribution cycles posting a reversible allocation JV, and KGI2/KGI4
costing-sheet overhead on work orders (order cost, WIP, BOM estimate)."""
import io
from datetime import date

import pytest

from app.modules.accounts.models.co_activity import CoActivityRate
from app.modules.accounts.models.co_overhead import CoCostingSheetRate
from app.modules.accounts.models.gl_account import GLAccount
from app.modules.accounts.models.journal_entry import JournalEntry, JournalEntryLine
from app.modules.accounts.services import co_activity, co_allocation, co_overhead, co_planning
from app.modules.accounts.services.controlling import cost_center_actuals
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch
from app.modules.organization.models.cost_center import CostCenter
from app.modules.organization.models.department import Department
from app.modules.production import costing
from app.modules.production.models.bom import ProductionBom
from app.modules.production.models.work_order import ProductionOrderCostPosting, ProductionTimeLog, ProductionWorkOrder
from app.modules.production.models.workstation import ProductionWorkstation
from app.modules.production.service import work_order_costing, workstation_rates
from app.modules.store.models.item import StoreItem
from app.tests.test_accounts_finance import books, _jv, _capital, _headers  # noqa: F401  (fixture)

FY = "2026-27"


@pytest.fixture
def co(books):
    db, gl = books["db"], books["gl"]
    plant = Branch(name="Plant 1", code="P1")
    db.add(plant)
    db.flush()
    d_prod = Department(name="Production", code="PROD")
    d_qc = Department(name="Quality", code="QC")
    db.add_all([d_prod, d_qc])
    db.flush()
    admin = CostCenter(code="CC-ADM", name="Admin", status="active")
    prod = CostCenter(code="CC-PRD", name="Production", status="active", department_id=d_prod.id, branch_id=plant.id)
    qc = CostCenter(code="CC-QC", name="Quality", status="active", department_id=d_qc.id)
    db.add_all([admin, prod, qc])
    db.flush()
    rent = gl["5200"]
    assess = GLAccount(code="5990", name="Overhead assessed", account_type="expense", group_id=rent.group_id,
                       is_posting_account=True, status="active")
    db.add(assess)
    db.flush()
    gl["5990"] = assess
    db.commit()
    _capital(books)
    return {**books, "plant": plant, "admin": admin, "prod": prod, "qc": qc, "d_prod": d_prod, "d_qc": d_qc}


def _cost(b, d, gl_code, cc, amount):
    return _jv(b, d, gl_code, "1121", amount, dr={"cost_center_id": cc.id})


# ── KP06 planning ───────────────────────────────────────────────────────

def test_plan_grid_versions_and_plan_vs_actual(co):
    db, gl = co["db"], co["gl"]
    v = co_planning.create_version(db, code="0", name="Original", financial_year=FY)
    assert v.is_primary   # the first version of a year becomes primary
    db.commit()
    g = co_planning.save_grid(db, v.id, [
        {"cost_center_id": co["admin"].id, "gl_account_id": gl["5200"].id, "amounts": {"2026-04": 1000, "2026-05": 1000}},
        {"cost_center_id": co["prod"].id, "gl_account_id": gl["5220"].id, "amounts": {"2026-04": 500}},
    ])
    assert g["periods"][0] == "2026-04" and len(g["periods"]) == 12 and g["grand_total"] == 2500
    db.commit()
    with pytest.raises(ValueError, match="isn't a P&L posting ledger"):
        co_planning.save_grid(db, v.id, [{"cost_center_id": co["admin"].id, "gl_account_id": gl["1121"].id, "amounts": {}}])
    db.rollback()

    _cost(co, date(2026, 4, 10), "5200", co["admin"], 1200)
    _cost(co, date(2026, 6, 10), "5200", co["admin"], 999)   # outside the selected months
    pva = co_planning.plan_vs_actual(db, financial_year=FY, from_period="2026-04", to_period="2026-05")
    row = next(r for r in pva["rows"] if r["cost_center_id"] == co["admin"].id)
    assert (row["plan"], row["actual"], row["variance"], row["variance_pct"]) == (2000, 1200, -800, -40)
    assert pva["total_plan"] == 2500 and pva["total_actual"] == 1200

    # Revised version copied +10%, made primary; the original gets locked.
    r1 = co_planning.create_version(db, code="R1", name="Revised", financial_year=FY, version_type="revised",
                                    copy_from_version_id=v.id, uplift_pct=10, is_primary=True)
    db.refresh(v)
    assert r1.is_primary and not v.is_primary
    assert co_planning.grid(db, r1.id)["grand_total"] == 2750
    co_planning.set_locked(db, v.id, locked=True, user_id=co["maker"].id)
    with pytest.raises(ValueError, match="locked"):
        co_planning.save_grid(db, v.id, [])
    with pytest.raises(ValueError, match="already exists"):
        co_planning.create_version(db, code="0", name="Dup", financial_year=FY)


def test_plan_from_last_years_actuals(co):
    db = co["db"]
    _cost(co, date(2026, 5, 3), "5200", co["admin"], 400)
    nxt = co_planning.create_version(db, code="0", name="Next year", financial_year="2027-28", copy_actuals_from_fy=FY, uplift_pct=5)
    g = co_planning.grid(db, nxt.id)
    assert g["rows"][0]["amounts"]["2027-05"] == 420


def test_plan_excel_round_trip_and_row_errors(co):
    from openpyxl import load_workbook

    db, gl = co["db"], co["gl"]
    v = co_planning.create_version(db, code="0", name="Original", financial_year=FY)
    co_planning.save_grid(db, v.id, [{"cost_center_id": co["admin"].id, "gl_account_id": gl["5200"].id, "amounts": {"2026-04": 100}}])
    db.commit()
    content, filename = co_planning.export_xlsx(db, v.id)
    assert filename == f"CO-plan-{FY}-0.xlsx"
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Plan"]
    assert ws.cell(1, 5).value == "Apr-2026" and ws.cell(2, 5).value == 100
    ws.cell(2, 6).value = 250                      # May
    ws.append(["CC-PRD", "", "5220", "", 10, "1,000"])
    buf = io.BytesIO()
    wb.save(buf)
    g = co_planning.import_xlsx(db, v.id, buf.getvalue())
    assert g["grand_total"] == 100 + 250 + 10 + 1000
    db.commit()

    ws.append(["CC-NOPE", "", "5200", "", "abc"])
    bad = io.BytesIO()
    wb.save(bad)
    with pytest.raises(ValueError, match=r"Row 4: cost centre code 'CC-NOPE'.*Row 4: 'abc' under Apr-2026 isn't a number"):
        co_planning.import_xlsx(db, v.id, bad.getvalue())
    db.rollback()
    assert co_planning.grid(db, v.id)["grand_total"] == 1360   # nothing changed


# ── KL01 / KP26 activity rates ──────────────────────────────────────────

def _workstation(db, cc=None, rate=100.0):
    ws = ProductionWorkstation(code="WS-CNC", name="CNC", hourly_rate=rate, cost_center_id=cc.id if cc else None)
    db.add(ws)
    db.flush()
    return ws


def test_activity_rate_lookup_validity_and_fallback(co):
    db = co["db"]
    ws = _workstation(db, co["prod"])
    on = date(2026, 7, 1)
    assert co_activity.workstation_hour_rate(db, ws, on) == (100.0, "hourly_rate")
    db.add(CoActivityRate(**co_activity.validate_rate(db, {
        "activity_type": "mch", "cost_center_id": co["prod"].id, "valid_from": date(2026, 4, 1), "valid_to": date(2027, 3, 31),
        "fixed_rate": 300, "variable_rate": 200})))
    db.add(CoActivityRate(**co_activity.validate_rate(db, {
        "activity_type": "LAB", "workstation_id": ws.id, "valid_from": date(2026, 4, 1), "fixed_rate": 0, "variable_rate": 250})))
    db.flush()
    assert co_activity.workstation_hour_rate(db, ws, on, "MCH") == (500.0, "co_rate")   # via its cost centre
    assert co_activity.workstation_hour_rate(db, ws, on) == (750.0, "co_rate")          # plain hour = MCH + LAB
    assert co_activity.workstation_hour_rate(db, ws, on, "SET") == (100.0, "hourly_rate")
    assert workstation_rates(db, {ws.id})[ws.id] == 750.0   # production costing uses it
    with pytest.raises(ValueError, match="overlaps"):
        co_activity.validate_rate(db, {"activity_type": "LAB", "workstation_id": ws.id, "valid_from": date(2026, 9, 1), "variable_rate": 1})
    with pytest.raises(ValueError, match="exactly one"):
        co_activity.validate_rate(db, {"activity_type": "LAB", "valid_from": date(2026, 9, 1), "variable_rate": 1})
    with pytest.raises(ValueError, match="one financial year"):
        co_activity.validate_rate(db, {"activity_type": "SET", "workstation_id": ws.id, "valid_from": date(2026, 9, 1),
                                       "valid_to": date(2027, 6, 1), "variable_rate": 1})
    with pytest.raises(ValueError, match="isn't an activity type"):
        co_activity.validate_rate(db, {"activity_type": "XYZ", "workstation_id": ws.id, "valid_from": date(2026, 9, 1), "variable_rate": 1})


# ── KSV5 / KSU5 cycles ──────────────────────────────────────────────────

def _cycle(co, method="distribution", basis="fixed_percent", receivers=None, **kw):
    db = co["db"]
    receivers = receivers or [{"cost_center_id": co["prod"].id, "percent": 60}, {"cost_center_id": co["qc"].id, "percent": 40}]
    return co_allocation.save_cycle(db, {
        "code": kw.get("code", "DIST-ADM"), "name": "Admin to plant", "method": method, "status": "active",
        "assessment_gl_account_id": co["gl"]["5990"].id if method == "assessment" else None,
        "segments": [{"name": "Admin", "sender_cost_center_id": co["admin"].id, "basis": basis, "receivers": receivers}],
    })


def test_distribution_moves_cost_keeps_pl_and_is_reversible(co):
    db = co["db"]
    _cost(co, date(2026, 7, 5), "5200", co["admin"], 1000)
    _cost(co, date(2026, 7, 6), "5230", co["admin"], 333.33)
    cycle = _cycle(co)
    before = cost_center_actuals(db, from_date=date(2026, 7, 1), to_date=date(2026, 7, 31))
    preview = co_allocation.run_cycle(db, cycle.id, "2026-07", test_run=True)
    assert preview["total_amount"] == 1333.33 and db.query(JournalEntry).filter_by(posting_type="co_allocation").count() == 0
    rec = {r["cost_center_id"]: r["amount"] for r in preview["segments"][0]["receivers"]}
    assert rec == {co["prod"].id: 800.0, co["qc"].id: 533.33}   # last receiver takes the paisa

    out = co_allocation.run_cycle(db, cycle.id, "2026-07", test_run=False, user_id=co["maker"].id)
    db.commit()
    after = cost_center_actuals(db, from_date=date(2026, 7, 1), to_date=date(2026, 7, 31))
    assert after["column_totals"].get(str(co["admin"].id), 0) == 0
    assert after["column_totals"][str(co["prod"].id)] == 800
    assert after["grand_total"] == before["grand_total"]       # P&L unchanged
    with pytest.raises(ValueError, match="already posted for 2026-07"):
        co_allocation.run_cycle(db, cycle.id, "2026-07", test_run=False)
    # A re-preview still sees the original sender cost (own runs are excluded).
    assert co_allocation.run_cycle(db, cycle.id, "2026-07")["total_amount"] == 1333.33

    run = co_allocation.reverse_run(db, out["run_id"], reason="wrong shares", user_id=co["checker"].id)
    db.commit()
    assert run.status == "reversed"
    assert cost_center_actuals(db, from_date=date(2026, 7, 1), to_date=date(2026, 7, 31))["column_totals"][str(co["admin"].id)] == 1333.33
    again = co_allocation.run_cycle(db, cycle.id, "2026-07", test_run=False, user_id=co["maker"].id)
    assert again["total_amount"] == 1333.33


def test_assessment_and_headcount_basis(co):
    db, gl = co["db"], co["gl"]
    _cost(co, date(2026, 8, 5), "5200", co["admin"], 900)
    for i, dept in enumerate([co["d_prod"], co["d_prod"], co["d_qc"]]):
        db.add(User(email=f"u{i}@premnathrail.com", name=f"U{i}", role="user", is_active=True, department_id=dept.id))
    db.commit()
    cycle = _cycle(co, method="assessment", basis="headcount", code="ASSESS",
                   receivers=[{"cost_center_id": co["prod"].id}, {"cost_center_id": co["qc"].id}])
    out = co_allocation.run_cycle(db, cycle.id, "2026-08", test_run=False, user_id=co["maker"].id)
    lines = db.query(JournalEntryLine).filter_by(journal_entry_id=out["journal_entry_id"]).all()
    assert {l.gl_account_id for l in lines} == {gl["5990"].id}
    by_cc = {l.cost_center_id: (l.debit_amount, l.credit_amount) for l in lines}
    assert by_cc == {co["admin"].id: (0, 900), co["prod"].id: (600, 0), co["qc"].id: (300, 0)}


def test_cycle_definition_errors(co):
    with pytest.raises(ValueError, match="can't exceed 100%"):
        _cycle(co, receivers=[{"cost_center_id": co["prod"].id, "percent": 70}, {"cost_center_id": co["qc"].id, "percent": 40}])
    with pytest.raises(ValueError, match="is the sender"):
        _cycle(co, receivers=[{"cost_center_id": co["admin"].id, "percent": 10}])
    stores = CostCenter(code="CC-STR", name="Stores", status="active")
    co["db"].add(stores)
    co["db"].flush()
    with pytest.raises(ValueError, match="has no department"):
        _cycle(co, basis="headcount", receivers=[{"cost_center_id": stores.id}, {"cost_center_id": co["prod"].id}], code="X")
    with pytest.raises(ValueError, match="assessment ledger"):
        co_allocation.save_cycle(co["db"], {"code": "A", "name": "A", "method": "assessment", "segments": [
            {"sender_cost_center_id": co["admin"].id, "basis": "fixed_percent", "receivers": [{"cost_center_id": co["prod"].id, "percent": 100}]}]})


def test_activity_hours_basis(co):
    db = co["db"]
    ws_p = _workstation(db, co["prod"])
    ws_q = ProductionWorkstation(code="WS-QC", name="Test bench", hourly_rate=50, cost_center_id=co["qc"].id)
    db.add(ws_q)
    wo = _work_order(db)
    db.add_all([ProductionTimeLog(work_order_id=wo.id, workstation_id=ws_p.id, log_date=date(2026, 9, 3), hours=3),
                ProductionTimeLog(work_order_id=wo.id, workstation_id=ws_q.id, log_date=date(2026, 9, 4), hours=1)])
    _cost(co, date(2026, 9, 5), "5200", co["admin"], 400)
    cycle = _cycle(co, basis="activity_hours", code="HRS",
                   receivers=[{"cost_center_id": co["prod"].id}, {"cost_center_id": co["qc"].id}])
    out = co_allocation.run_cycle(db, cycle.id, "2026-09")
    assert {r["cost_center_id"]: r["amount"] for r in out["segments"][0]["receivers"]} == {co["prod"].id: 300, co["qc"].id: 100}


# ── KGI2 / KGI4 overhead ────────────────────────────────────────────────

def _work_order(db, branch=None):
    fg = StoreItem(item_code="FG-RRV-9", item_name="RRV", uom="NOS", item_type="finished_good")
    db.add(fg)
    db.flush()
    bom = ProductionBom(bom_number="BOM-2026-0099", product_item_id=fg.id, version=1, status="active")
    db.add(bom)
    db.flush()
    wo = ProductionWorkOrder(wo_number="WO-2026-0099", product_item_id=fg.id, bom_id=bom.id, quantity_planned=1,
                             status="in_progress", branch_id=branch.id if branch else None)
    db.add(wo)
    db.flush()
    return wo


def test_overhead_run_on_work_orders_flows_into_order_cost_and_wip(co):
    db = co["db"]
    wo = _work_order(db, co["plant"])
    ws = _workstation(db, co["prod"], rate=200)
    db.add(ProductionOrderCostPosting(work_order_id=wo.id, posting_type="material_issue", posting_date=date(2026, 9, 10),
                                      quantity=1, unit_cost=10000, amount=10000))
    db.add(ProductionTimeLog(work_order_id=wo.id, workstation_id=ws.id, log_date=date(2026, 9, 11), hours=5, activity_rate=200))
    db.add(CoCostingSheetRate(**co_overhead.validate_sheet_rate(db, {"branch_id": None, "valid_from": date(2026, 4, 1),
                                                                    "material_overhead_pct": 5, "labour_overhead_pct": 50})))
    db.add(CoCostingSheetRate(**co_overhead.validate_sheet_rate(db, {"branch_id": co["plant"].id, "valid_from": date(2026, 4, 1),
                                                                    "material_overhead_pct": 10, "labour_overhead_pct": 100})))
    db.commit()
    with pytest.raises(ValueError, match="overlaps"):
        co_overhead.validate_sheet_rate(db, {"branch_id": co["plant"].id, "valid_from": date(2026, 6, 1)})

    preview = co_overhead.run_overhead(db, "2026-09")
    assert preview["rows"][0]["amount"] == 1000 + 1000   # plant row: 10% of 10,000 + 100% of 1,000
    assert costing.order_cost(db, wo)["overhead"] == 0     # a test run writes nothing
    out = co_overhead.run_overhead(db, "2026-09", test_run=False, user_id=co["maker"].id)
    db.commit()
    oc = costing.order_cost(db, wo)
    assert (oc["overhead"], oc["open"]) == (2000, 13000)
    assert costing.wip_snapshot(db, date(2026, 9, 30))[0]["wip_value"] == 13000
    assert costing.wip_snapshot(db, date(2026, 8, 31)) == []
    wc = work_order_costing(db, wo)
    assert wc["actual_overhead_cost"] == 2000
    assert wc["actual_total_cost"] == wc["actual_material_cost"] + wc["actual_labour_cost"] + 2000
    with pytest.raises(ValueError, match="already applied for 2026-09"):
        co_overhead.run_overhead(db, "2026-09", test_run=False)

    co_overhead.reverse_overhead_run(db, out["run_id"], reason="sheet was wrong", user_id=co["checker"].id)
    db.commit()
    assert costing.order_cost(db, wo)["open"] == 11000
    assert co_overhead.overhead_for(db, None, date(2026, 9, 1), 1000, 100)["overhead"] == 100   # company-wide row


# ── API + RBAC ──────────────────────────────────────────────────────────

def test_api_routes_and_deny_by_default(client, co):
    db, gl = co["db"], co["gl"]
    h = _headers(co["maker"])
    r = client.post("/api/v1/accounts/co/plan-versions", headers=h, json={"code": "0", "name": "Original", "financial_year": FY})
    assert r.status_code == 201, r.text
    vid = r.json()["id"]
    r = client.put(f"/api/v1/accounts/co/plan-versions/{vid}/grid", headers=h, json={"rows": [
        {"cost_center_id": co["admin"].id, "gl_account_id": gl["5200"].id, "amounts": {"2026-04": 10}}]})
    assert r.status_code == 200 and r.json()["grand_total"] == 10
    x = client.get(f"/api/v1/accounts/co/plan-versions/{vid}/export", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    up = client.post(f"/api/v1/accounts/co/plan-versions/{vid}/import", headers=h,
                     files={"file": ("plan.xlsx", x.content, "application/octet-stream")})
    assert up.status_code == 200 and up.json()["grand_total"] == 10
    bad = client.post(f"/api/v1/accounts/co/plan-versions/{vid}/grid", headers=h, json={})
    assert bad.status_code == 405
    assert client.post(f"/api/v1/accounts/co/plan-versions/{vid}/lock", headers=h).json()["status"] == "locked"
    r = client.put(f"/api/v1/accounts/co/plan-versions/{vid}/grid", headers=h, json={"rows": []})
    assert r.status_code == 400 and "locked" in r.json()["detail"]
    assert client.get("/api/v1/accounts/co/plan-vs-actual", headers=h, params={"financial_year": FY}).status_code == 200
    assert client.get("/api/v1/accounts/co/activity-types", headers=h).json()[0]["value"] == "MCH"
    assert client.post("/api/v1/accounts/co/overhead-runs", headers=h, json={"period": "2026-09"}).json()["rows"] == []
    post = client.post("/api/v1/accounts/co/overhead-runs", headers=h, json={"period": "2026-09", "test_run": False})
    assert post.status_code == 400 and "No overhead to apply" in post.json()["detail"]

    outsider = User(email="nobody@premnathrail.com", name="Nobody", role="user", is_active=True)
    db.add(outsider)
    db.commit()
    r = client.get("/api/v1/accounts/co/plan-versions", headers=_headers(outsider))
    assert r.status_code == 403 and "Cost Centre Planning" in r.json()["detail"]
