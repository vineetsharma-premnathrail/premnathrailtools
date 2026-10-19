"""
Tests for Maintenance Phase 3 — AMC & warranty tracking: the expiry bucket
maths, the /maintenance/contracts list + summary + filters + export, and the
daily expiry reminder (trigger days, specific message, once per user/asset/day,
decommissioned assets skipped).
"""
from datetime import date, timedelta

from app.auth.jwt_handler import create_access_token
from app.modules.accounts.models.vendor import Vendor
from app.modules.main.models.notification import Notification
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch
from app.modules.maintenance.contracts import contract_bucket, reminder_due
from app.modules.maintenance.models.asset import MaintenanceAsset
from app.modules.maintenance.service import today_ist
from app.tasks.maintenance_amc_reminders import _send_maintenance_amc_reminders, AMC_EXPIRING

BASE = "/api/v1/maintenance"


def make_user(db, email, apps=("maintenance",), role="user"):
    user = User(email=email, name=email.split("@")[0], role=role, is_active=True, assigned_apps=list(apps))
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def hdr(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}


def make_asset(db, branch, code, name="Machine", **extra):
    a = MaintenanceAsset(asset_code=code, name=name, branch_id=branch.id, criticality="A", **extra)
    db.add(a)
    db.commit()
    db.refresh(a)
    return a


def setup(db):
    branch = Branch(name="Faridabad Plant", code="FBD")
    vendor = Vendor(code="V-ABC", name="ABC Services")
    supplier = Vendor(code="V-OEM", name="Jyoti CNC")
    db.add_all([branch, vendor, supplier])
    db.commit()
    for o in (branch, vendor, supplier):
        db.refresh(o)
    return branch, vendor, supplier


def test_contract_bucket():
    t = date(2026, 10, 6)
    assert contract_bucket(None, t) == ("none", None)
    assert contract_bucket(date(2026, 10, 5), t) == ("expired", -1)
    assert contract_bucket(t, t) == ("expiring", 0)
    assert contract_bucket(t + timedelta(days=60), t) == ("expiring", 60)
    assert contract_bucket(t + timedelta(days=61), t) == ("active", 61)
    assert contract_bucket(t + timedelta(days=61), t, window_days=90) == ("expiring", 61)
    assert [d for d in range(-3, 40) if reminder_due(t + timedelta(days=d), t)] == [-1, 1, 7, 15, 30]


def test_contracts_list_summary_filters_and_export(client, db):
    branch, vendor, supplier = setup(db)
    today = today_ist()
    h = hdr(make_user(db, "m1@premnathrail.com"))
    make_asset(db, branch, "CNC-04", "Lathe", amc_vendor_id=vendor.id, amc_expiry=today + timedelta(days=7),
               supplier_vendor_id=supplier.id, warranty_expiry=today - timedelta(days=10))
    make_asset(db, branch, "CMP-01", "Compressor", amc_expiry=today + timedelta(days=200))
    make_asset(db, branch, "OLD-01", "Scrapped press", amc_expiry=today + timedelta(days=5), status="decommissioned")

    r = client.get(f"{BASE}/contracts", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    # 2 live assets × 2 kinds; decommissioned skipped.
    assert body["summary"] == {"expired": 1, "expiring": 1, "active": 1, "none": 1, "total": 4}
    first = body["items"][0]
    assert (first["asset_code"], first["kind"], first["status"], first["days_left"], first["vendor_name"]) == \
        ("CNC-04", "warranty", "expired", -10, "Jyoti CNC")
    amc = next(i for i in body["items"] if i["asset_code"] == "CNC-04" and i["kind"] == "amc")
    assert amc["vendor_name"] == "ABC Services" and amc["status"] == "expiring" and amc["branch_name"] == "Faridabad Plant"

    r = client.get(f"{BASE}/contracts", params={"kind": "amc", "status": "active"}, headers=h).json()
    assert [i["asset_code"] for i in r["items"]] == ["CMP-01"]
    assert r["summary"]["total"] == 2  # summary ignores the status filter
    r = client.get(f"{BASE}/contracts", params={"kind": "amc", "window_days": 365}, headers=h).json()
    assert r["summary"]["expiring"] == 2
    r = client.get(f"{BASE}/contracts", params={"search": "compress"}, headers=h).json()
    assert {i["asset_code"] for i in r["items"]} == {"CMP-01"}

    bad = client.get(f"{BASE}/contracts", params={"kind": "insurance"}, headers=h)
    assert bad.status_code == 400 and "warranty, amc or all" in bad.json()["detail"]
    other = make_user(db, "x1@premnathrail.com", apps=("crm",))
    assert client.get(f"{BASE}/contracts", headers=hdr(other)).status_code == 403

    r = client.get(f"{BASE}/contracts/export", headers=h)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(r.content)).active
    assert ws["A4"].value == "Asset Code" and ws.max_row == 4 + 4


def test_amc_reminders_trigger_days_and_dedupe(db):
    branch, vendor, supplier = setup(db)
    today = today_ist()
    planner = make_user(db, "planner@premnathrail.com")
    make_user(db, "crm@premnathrail.com", apps=("crm",))  # not a maintenance user
    a7 = make_asset(db, branch, "CNC-04", "Lathe", amc_vendor_id=vendor.id, amc_expiry=today + timedelta(days=7))
    both = make_asset(db, branch, "CMP-01", "Compressor", amc_expiry=today + timedelta(days=30),
                      supplier_vendor_id=supplier.id, warranty_expiry=today - timedelta(days=1))
    make_asset(db, branch, "PRS-02", "Press", amc_expiry=today + timedelta(days=8))  # not a trigger day
    make_asset(db, branch, "OLD-01", "Scrapped", amc_expiry=today + timedelta(days=7), status="decommissioned")

    stats = _send_maintenance_amc_reminders(db, today)
    assert stats == {"assets": 2, "sent": 2, "failed": 0}
    notes = db.query(Notification).filter(Notification.notification_type == AMC_EXPIRING).all()
    assert {n.user_id for n in notes} == {planner.id}
    assert {n.entity_id for n in notes} == {a7.id, both.id} and all(n.entity_type == "maintenance_asset" for n in notes)
    msg7 = next(n.message for n in notes if n.entity_id == a7.id)
    assert msg7 == (f"AMC for CNC-04 Lathe with ABC Services expires on {today + timedelta(days=7):%d %b %Y} (7 days). "
                    "Renew or plan a replacement contract.")
    msg_both = next(n for n in notes if n.entity_id == both.id)
    assert "AMC & Warranty" in msg_both.title and "Warranty for CMP-01 Compressor from Jyoti CNC expired yesterday" in msg_both.message

    # Second run the same day sends nothing new.
    assert _send_maintenance_amc_reminders(db, today)["sent"] == 0
    assert db.query(Notification).filter(Notification.notification_type == AMC_EXPIRING).count() == 2
