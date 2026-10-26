"""Customer Organization scope: Domestic / International first, scope-specific tax
fields, independent Country, scope switching, duplicates, audit, list filter,
export and import."""
import io

from openpyxl import load_workbook

from app.auth.jwt_handler import create_access_token
from app.modules.crm.models.organization import Organization, OrgContact
from app.modules.main.models.audit_log import AuditLog
from app.modules.main.models.user import User

URL = "/api/v1/crm/organizations"
GSTIN = "27ABCDE1234F1Z5"


def make_user(db, email, role="user", apps=("crm",)):
    u = User(email=email, name=email.split("@")[0], role=role, is_active=True, assigned_apps=list(apps))
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def h(u):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(u.id), 'email': u.email, 'role': u.role})}"}


def body(**kw):
    base = {"org_scope": "domestic", "org_type": "Customer", "name": "Scope Org"}
    base.update(kw)
    return base


def detail(r):
    j = r.json()
    errs = " ".join(str(e.get("msg", "")) + " " + ".".join(str(x) for x in e.get("loc", ())) for e in j.get("errors") or [])
    d = j.get("detail")
    return (d if isinstance(d, str) else "") + " " + errs


# ── create ──────────────────────────────────────────────────────────────

def test_domestic_with_gstin_derives_pan_and_state_code(client, db):
    u = make_user(db, "scope_a@premnathrail.com")
    r = client.post(URL, json=body(name="Dom GST", gst_number=GSTIN.lower()), headers=h(u))
    assert r.status_code == 201, r.text
    d = r.json()
    assert (d["org_scope"], d["gst_number"], d["pan"], d["gst_state_code"]) == ("domestic", GSTIN, "ABCDE1234F", "27")
    assert d["tax_vat_id"] is None and d["company_registration_no"] is None


def test_domestic_with_pan_only(client, db):
    u = make_user(db, "scope_b@premnathrail.com")
    r = client.post(URL, json=body(name="Dom PAN", pan="abcde1234f"), headers=h(u))
    assert r.status_code == 201, r.text
    assert r.json()["pan"] == "ABCDE1234F" and r.json()["gst_number"] is None


def test_international_with_tax_vat_and_registration(client, db):
    u = make_user(db, "scope_c@premnathrail.com")
    r = client.post(URL, json=body(org_scope="international", name="Intl", country="Germany",
                                   tax_vat_id="DE 123.456/789", company_registration_no="HRB-12345"), headers=h(u))
    assert r.status_code == 201, r.text
    d = r.json()
    assert (d["org_scope"], d["tax_vat_id"], d["company_registration_no"], d["country"]) == (
        "international", "DE 123.456/789", "HRB-12345", "Germany")
    assert d["gst_number"] is None and d["pan"] is None and d["gst_state_code"] is None


def test_international_organization_in_india_is_allowed(client, db):
    u = make_user(db, "scope_d@premnathrail.com")
    r = client.post(URL, json=body(org_scope="international", name="Intl India", country="India", state="Delhi",
                                   tax_vat_id="AB-1234"), headers=h(u))
    assert r.status_code == 201, r.text
    d = r.json()
    assert d["org_scope"] == "international" and d["country"] == "India"
    assert d["gst_state_code"] is None  # state name must not be turned into a GST state code


def test_domestic_must_be_in_india(client, db):
    u = make_user(db, "scope_e@premnathrail.com")
    msg = "A Domestic organization is in India. Choose International for a customer outside India, or set the country to India."
    r = client.post(URL, json=body(name="Dom Abroad", country="Nepal"), headers=h(u))
    assert r.status_code == 422 and msg in detail(r)
    r = client.post(URL, json=body(name="Dom Blank", country=""), headers=h(u))
    assert r.status_code == 201 and r.json()["country"] == "India"
    r = client.post(URL, json=body(name="Dom India", country="india"), headers=h(u))
    assert r.status_code == 201
    # editing: domestic can't be moved abroad; International -> Domestic needs India
    org = client.post(URL, json=body(org_scope="international", name="Intl Nepal", country="Nepal"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "domestic"}, headers=h(u))
    assert r.status_code == 422 and r.json()["detail"] == msg
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "domestic", "country": "India"}, headers=h(u))
    assert r.status_code == 200 and r.json()["country"] == "India"
    r = client.patch(f"{URL}/{org['id']}", json={"country": "Nepal"}, headers=h(u))
    assert r.status_code == 422 and r.json()["detail"] == msg


def test_gst_state_code_not_required_and_derived(client, db):
    u = make_user(db, "scope_sc@premnathrail.com")
    r = client.post(URL, json=body(name="No SC", gst_number=GSTIN), headers=h(u))
    assert r.status_code == 201 and r.json()["gst_state_code"] == "27"
    org = client.post(URL, json=body(name="No SC 2"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"gst_number": "07AAACN1234Q1Z2"}, headers=h(u))
    assert r.status_code == 200 and r.json()["gst_state_code"] == "07"


def test_scope_is_required_and_must_be_valid(client, db):
    u = make_user(db, "scope_f@premnathrail.com")
    r = client.post(URL, json={"org_type": "Customer", "name": "No Scope"}, headers=h(u))
    assert r.status_code == 422
    assert "Organization Scope is required" in detail(r) and "Domestic or International" in detail(r)
    r = client.post(URL, json=body(org_scope="foreign"), headers=h(u))
    assert r.status_code == 422 and "Choose Domestic or International" in detail(r)


def test_name_and_type_still_required(client, db):
    u = make_user(db, "scope_g@premnathrail.com")
    assert client.post(URL, json={"org_scope": "domestic", "org_type": "Customer"}, headers=h(u)).status_code == 422
    r = client.post(URL, json={"org_scope": "domestic", "name": "X"}, headers=h(u))
    assert r.status_code == 422


def test_wrong_scope_fields_are_rejected_with_the_field_named(client, db):
    u = make_user(db, "scope_h@premnathrail.com")
    r = client.post(URL, json=body(tax_vat_id="GB123456789"), headers=h(u))
    assert r.status_code == 422
    assert "Tax / VAT ID" in detail(r) and "International" in detail(r) and "clear" in detail(r)
    r = client.post(URL, json=body(company_registration_no="12345678"), headers=h(u))
    assert r.status_code == 422 and "Company Registration Number" in detail(r)
    for field, label, value in (("gst_number", "GSTIN", GSTIN), ("pan", "PAN", "ABCDE1234F"), ("gst_state_code", "GST State Code", "27")):
        r = client.post(URL, json=body(org_scope="international", **{field: value}), headers=h(u))
        assert r.status_code == 422, field
        assert label in detail(r) and "Domestic" in detail(r) and "clear" in detail(r)


def test_gstin_pan_state_code_format_and_mismatch(client, db):
    u = make_user(db, "scope_i@premnathrail.com")
    r = client.post(URL, json=body(gst_number="NOT-A-GSTIN"), headers=h(u))
    assert r.status_code == 422 and "15-character GSTIN" in detail(r)
    r = client.post(URL, json=body(pan="12345"), headers=h(u))
    assert r.status_code == 422 and "PAN" in detail(r) and "5 letters, 4 digits, 1 letter" in detail(r)
    r = client.post(URL, json=body(gst_number=GSTIN, pan="ZZZZZ9999Z"), headers=h(u))
    assert r.status_code == 422 and "doesn't match GSTIN" in detail(r)
    r = client.post(URL, json=body(gst_number=GSTIN, gst_state_code="07"), headers=h(u))
    assert r.status_code == 422 and "state code 07 doesn't match GSTIN".lower() in detail(r).lower()
    r = client.post(URL, json=body(gst_state_code="AB"), headers=h(u))
    assert r.status_code == 422 and "2-digit GST state code" in detail(r)


def test_tax_vat_and_registration_format(client, db):
    u = make_user(db, "scope_j@premnathrail.com")
    for bad in ("A", "X" * 51, "AB#123", "AB_123"):
        r = client.post(URL, json=body(org_scope="international", tax_vat_id=bad), headers=h(u))
        assert r.status_code == 422, bad
        assert "Tax / VAT ID must be 2-50 characters" in detail(r)
    r = client.post(URL, json=body(org_scope="international", company_registration_no="no@way"), headers=h(u))
    assert r.status_code == 422 and "Company Registration Number must be 2-50 characters" in detail(r)


def test_blank_tax_fields_are_treated_as_not_sent(client, db):
    u = make_user(db, "scope_k@premnathrail.com")
    r = client.post(URL, json=body(name="Blank Dom", gst_number="", tax_vat_id="", company_registration_no=" "), headers=h(u))
    assert r.status_code == 201, r.text
    r = client.post(URL, json=body(org_scope="international", name="Blank Intl", gst_number="", pan="", gst_state_code=""), headers=h(u))
    assert r.status_code == 201, r.text
    assert db.query(Organization).filter(Organization.gst_number == "").count() == 0


# ── duplicates ──────────────────────────────────────────────────────────

def test_duplicate_gstin_names_existing_organization(client, db):
    u = make_user(db, "scope_l@premnathrail.com")
    assert client.post(URL, json=body(name="First GST", gst_number=GSTIN), headers=h(u)).status_code == 201
    r = client.post(URL, json=body(name="Second GST", gst_number=GSTIN), headers=h(u))
    assert r.status_code == 409 and "First GST" in r.json()["detail"]


def test_duplicate_tax_vat_id_per_country(client, db):
    u = make_user(db, "scope_m@premnathrail.com")
    first = body(org_scope="international", name="Vat One", country="France", tax_vat_id="FR 99 887766")
    assert client.post(URL, json=first, headers=h(u)).status_code == 201
    r = client.post(URL, json={**first, "name": "Vat Two", "tax_vat_id": "fr 99 887766"}, headers=h(u))
    assert r.status_code == 409
    assert "Tax / VAT ID" in r.json()["detail"] and "Vat One" in r.json()["detail"] and "France" in r.json()["detail"]
    # same ID in another country is a different company
    r = client.post(URL, json={**first, "name": "Vat Three", "country": "Belgium"}, headers=h(u))
    assert r.status_code == 201, r.text


# ── edit / switch scope ─────────────────────────────────────────────────

def test_edit_international_fields(client, db):
    u = make_user(db, "scope_n@premnathrail.com")
    org = client.post(URL, json=body(org_scope="international", name="Edit Intl", country="Japan", tax_vat_id="JP-1"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"tax_vat_id": "JP-2", "company_registration_no": "REG-77"}, headers=h(u))
    assert r.status_code == 200, r.text
    assert (r.json()["tax_vat_id"], r.json()["company_registration_no"]) == ("JP-2", "REG-77")
    r = client.patch(f"{URL}/{org['id']}", json={"gst_number": GSTIN}, headers=h(u))
    assert r.status_code == 422 and "GSTIN" in detail(r) and "Domestic" in detail(r)


def test_edit_domestic_rejects_international_fields(client, db):
    u = make_user(db, "scope_o@premnathrail.com")
    org = client.post(URL, json=body(name="Edit Dom"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"tax_vat_id": "XY-12"}, headers=h(u))
    assert r.status_code == 422 and "Tax / VAT ID" in detail(r)


def test_switch_domestic_to_international_clears_gst_fields(client, db):
    u = make_user(db, "scope_p@premnathrail.com")
    org = client.post(URL, json=body(name="Switch A", gst_number=GSTIN, state="Maharashtra"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "international", "tax_vat_id": "SW-100", "country": "India"}, headers=h(u))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["org_scope"] == "international" and d["tax_vat_id"] == "SW-100" and d["country"] == "India"
    assert d["gst_number"] is None and d["pan"] is None and d["gst_state_code"] is None


def test_switch_international_to_domestic_clears_foreign_fields(client, db):
    u = make_user(db, "scope_q@premnathrail.com")
    org = client.post(URL, json=body(org_scope="international", name="Switch B", country="Italy",
                                     tax_vat_id="IT-5", company_registration_no="REA-5"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "domestic", "country": "India", "gst_number": GSTIN}, headers=h(u))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["org_scope"] == "domestic" and d["gst_number"] == GSTIN and d["pan"] == "ABCDE1234F"
    assert d["tax_vat_id"] is None and d["company_registration_no"] is None and d["country"] == "India"


def test_switch_with_values_of_the_new_wrong_scope_is_rejected(client, db):
    u = make_user(db, "scope_r@premnathrail.com")
    org = client.post(URL, json=body(name="Switch C"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "international", "gst_number": GSTIN}, headers=h(u))
    assert r.status_code == 422 and "GSTIN" in detail(r)
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": None}, headers=h(u))
    assert r.status_code == 422 and "Organization Scope can't be empty" in detail(r)


def test_switch_to_international_checks_tax_vat_duplicate(client, db):
    u = make_user(db, "scope_s@premnathrail.com")
    client.post(URL, json=body(org_scope="international", name="Holder", country="Spain", tax_vat_id="ES-1"), headers=h(u))
    org = client.post(URL, json=body(name="Mover"), headers=h(u)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "international", "country": "Spain", "tax_vat_id": "ES-1"}, headers=h(u))
    assert r.status_code == 409 and "Holder" in r.json()["detail"]


# ── permission and data scope ───────────────────────────────────────────

def test_create_and_patch_require_crm_access(client, db):
    outsider = make_user(db, "scope_t@premnathrail.com", apps=())
    assert client.post(URL, json=body(name="Denied"), headers=h(outsider)).status_code == 403
    owner = make_user(db, "scope_u@premnathrail.com")
    org = client.post(URL, json=body(name="Guarded"), headers=h(owner)).json()
    assert client.patch(f"{URL}/{org['id']}", json={"city": "X"}, headers=h(outsider)).status_code == 403


def test_only_creator_or_admin_can_edit_scope(client, db):
    owner = make_user(db, "scope_v@premnathrail.com")
    other = make_user(db, "scope_w@premnathrail.com")
    org = client.post(URL, json=body(name="Owned"), headers=h(owner)).json()
    r = client.patch(f"{URL}/{org['id']}", json={"org_scope": "international"}, headers=h(other))
    assert r.status_code == 403
    assert db.get(Organization, org["id"]).org_scope == "domestic"


# ── contacts and audit ──────────────────────────────────────────────────

def test_contacts_saved_with_the_organization(client, db):
    u = make_user(db, "scope_x@premnathrail.com")
    org = client.post(URL, json=body(org_scope="international", name="With Contacts", country="Chile", tax_vat_id="CL-1"), headers=h(u)).json()
    for n, m in (("Ana", "111"), ("Ben", "222")):
        r = client.post(f"{URL}/{org['id']}/contacts", json={"name": n, "mobile": m}, headers=h(u))
        assert r.status_code == 201, r.text
    d = client.get(f"{URL}/{org['id']}/detail", headers=h(u)).json()
    assert d["org_scope"] == "international" and d["tax_vat_id"] == "CL-1"
    assert sorted(c["name"] for c in d["contacts"]) == ["Ana", "Ben"]
    assert db.query(OrgContact).filter(OrgContact.org_id == org["id"]).count() == 2


def test_audit_rows_contain_scope_and_new_fields(client, db):
    u = make_user(db, "scope_y@premnathrail.com")
    org = client.post(URL, json=body(org_scope="international", name="Audited", country="Peru",
                                     tax_vat_id="PE-1", company_registration_no="RUC-9"), headers=h(u)).json()
    client.post(f"{URL}/{org['id']}/contacts", json={"name": "Zed"}, headers=h(u))
    client.patch(f"{URL}/{org['id']}", json={"tax_vat_id": "PE-2", "company_registration_no": "RUC-10"}, headers=h(u))
    client.patch(f"{URL}/{org['id']}", json={"org_scope": "domestic", "country": "India"}, headers=h(u))
    rows = db.query(AuditLog).filter(AuditLog.entity_type == "organization", AuditLog.entity_id == org["id"]).order_by(AuditLog.id).all()
    by_action = {}
    for a in rows:
        by_action.setdefault(a.action, []).append(a.summary)
    assert "Scope: international" in by_action["created"][0]
    assert "Tax / VAT ID: PE-1" in by_action["created"][0] and "Company Registration Number: RUC-9" in by_action["created"][0]
    assert "Zed" in by_action["contact_added"][0]
    assert "Tax / VAT ID: PE-1 -> PE-2" in by_action["updated"][0]
    assert "Company Registration Number: RUC-9 -> RUC-10" in by_action["updated"][0]
    switch = by_action["updated"][1]
    assert "Scope: international -> domestic" in switch and "Tax / VAT ID: PE-2 -> -" in switch
    api = client.get(f"{URL}/{org['id']}/audit", headers=h(u))
    assert api.status_code == 200 and len(api.json()) == len(rows)


# ── list filter, export, import ─────────────────────────────────────────

def test_list_filter_by_scope(client, db):
    u = make_user(db, "scope_z@premnathrail.com")
    client.post(URL, json=body(name="List Dom"), headers=h(u))
    client.post(URL, json=body(org_scope="international", name="List Intl", country="Oman", tax_vat_id="OM-1"), headers=h(u))
    names = lambda r: sorted(o["name"] for o in r.json())  # noqa: E731
    r = client.get(URL, params={"org_scope": "international"}, headers=h(u))
    assert names(r) == ["List Intl"] and r.json()[0]["tax_vat_id"] == "OM-1"
    assert names(client.get(URL, params={"org_scope": "domestic"}, headers=h(u))) == ["List Dom"]
    assert names(client.get(URL, headers=h(u))) == ["List Dom", "List Intl"]
    r = client.get(URL, params={"org_scope": "mars"}, headers=h(u))
    assert r.status_code == 422 and "domestic" in r.json()["detail"]
    assert names(client.get(URL, params={"search": "OM-1"}, headers=h(u))) == ["List Intl"]


def test_export_has_scope_columns(client, db):
    admin = make_user(db, "scope_exp@premnathrail.com", role="admin")
    client.post(URL, json=body(name="Exp Intl", org_scope="international", country="Fiji", tax_vat_id="FJ-7", company_registration_no="C-7"), headers=h(admin))
    client.post(URL, json=body(name="Exp Dom", gst_number=GSTIN), headers=h(admin))
    r = client.get("/api/v1/crm/exports/organizations", headers=h(admin))
    assert r.status_code == 200
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    head = list(rows[0])
    for col in ("Scope", "Tax / VAT ID", "Company Registration No", "GST No.", "PAN", "GST State Code"):
        assert col in head, col
    data = {row[1]: dict(zip(head, row)) for row in rows[1:]}
    assert data["Exp Intl"]["Scope"] == "International" and data["Exp Intl"]["Tax / VAT ID"] == "FJ-7"
    assert data["Exp Intl"]["Company Registration No"] == "C-7"
    assert data["Exp Dom"]["Scope"] == "Domestic" and data["Exp Dom"]["PAN"] == "ABCDE1234F"


def _import(client, admin, csv_text):
    return client.post("/api/v1/crm/bulk-import", files={"file": ("orgs.csv", csv_text.encode(), "text/csv")}, headers=h(admin))


def test_import_defaults_to_domestic_and_reads_scope_column(client, db):
    admin = make_user(db, "scope_imp@premnathrail.com", role="admin")
    csv_text = (
        "org_name,org_type,country,gst_number,scope,tax_vat_id,company_registration_no\n"
        "Imp Default,Customer,India,,,,\n"
        "Imp Intl,Customer,Brazil,,international,BR-55,CNPJ-1\n"
        "Imp Bad,Customer,Brazil,,international,#bad#,\n"
    )
    r = _import(client, admin, csv_text)
    assert r.status_code == 200, r.text
    assert r.json()["organizations_created"] == 2
    assert db.query(Organization).filter(Organization.name == "Imp Default").one().org_scope == "domestic"
    o = db.query(Organization).filter(Organization.name == "Imp Intl").one()
    assert (o.org_scope, o.tax_vat_id, o.company_registration_no) == ("international", "BR-55", "CNPJ-1")
    assert any("tax_vat_id" in str(e) for e in r.json()["errors"])


def test_import_rejects_wrong_scope_fields(client, db):
    admin = make_user(db, "scope_imp2@premnathrail.com", role="admin")
    csv_text = (
        "org_name,org_type,gst_number,scope,tax_vat_id\n"
        "Imp Wrong A,Customer,27ABCDE1234F1Z5,international,\n"
        "Imp Wrong B,Customer,,domestic,VAT-1\n"
    )
    r = _import(client, admin, csv_text)
    assert r.status_code == 200, r.text
    assert r.json()["organizations_created"] == 0
    assert len(r.json()["errors"]) == 2
