"""R&D calculator verification pack: Braking (tools/braking).

Reference: docs/02-modules/rnd/overview.md §8.1 and §14 rows 2-4.

(a) Golden values are worked by hand from the documented formulas (units in
    the comments). They pass against the current code.
(b) Every documented discrepancy that can be expressed in backend code is
    written as the CORRECT engineering behaviour and marked
    xfail(strict=True, raises=AssertionError): the test fails today, and the
    day someone fixes the maths it will XPASS and turn the suite red, so the
    marker has to be removed consciously after engineering sign-off.

No calculator maths is changed by this file.
"""
import io
import math

import pytest

from app.modules.rnd.tools.braking.constants import BRAKING_DATA, G, MAX_STOPPING_DISTANCES
from app.modules.rnd.tools.braking.core import (
    calculate_deceleration_and_distances,
    calculate_gbr,
    calculate_max_rail_force,
    calculate_rail_scenario_forces,
)
from app.modules.rnd.tools.braking.service import perform_braking_calculation
from app.modules.rnd.tools.braking.units import calculate_angle, get_compliance

FOLLOWED = "✓ Standard Followed"
EXCEEDED = "✗ Standard Exceeded"


def _inputs(**over):
    base = dict(
        mass_kg=10000.0, reaction_time=1.0, num_wheels=4, calc_mode="Rail",
        rail_speed_input="20", rail_gradient_input="", rail_gradient_type="1 in G",
        road_speed_input="", road_gradient_input="", road_gradient_type="Percentage (%)",
        mu=0.7, doc_no="", made_by="", checked_by="", approved_by="", wheel_dia=0,
    )
    base.update(over)
    return base


def _row(rows, scenario, speed):
    return next(r for r in rows if r["mode"] == "Rail" and r["scenario"] == scenario and r["speed"] == speed)


# ---------------------------------------------------------------------------
# (a) Golden values
# ---------------------------------------------------------------------------

def test_max_rail_force_and_gbr_hand_calc():
    # a = v²/(2d) for every reference pair; the largest wins.
    #   8 km/h → v = 2.2222 m/s, a = 4.9383 / (2·3)   = 0.823045 m/s²
    #  16 km/h → v = 4.4444 m/s, a = 19.7531 / (2·12) = 0.823045 m/s²
    #  40 km/h → v = 11.111 m/s, a = 123.457 / (2·75) = 0.823045 m/s²
    #  (exact: 64 / 3.6² / 6 = 0.8230453 m/s²)
    # m = 10 000 kg → F_max = 8 230.45 N
    # GBR = 8 230.45 / (10 000 · 9.81) · 100 = 8.39 %
    f = calculate_max_rail_force(10000.0, 1.0)
    assert f == pytest.approx(8230.45, abs=0.01)
    assert calculate_gbr(f, 10000.0) == pytest.approx(8.39, abs=0.005)
    # Doc engineering note: GBR is ~8.39 % for ANY mass (force scales with m).
    assert calculate_gbr(calculate_max_rail_force(37000.0, 1.0), 37000.0) == pytest.approx(8.39, abs=0.005)
    # sanity on the reference table itself (no pair needs more than 0.8231 m/s²)
    assert max((s / 3.6) ** 2 / (2 * d) for s, d in BRAKING_DATA.items()) == pytest.approx(0.8230453, abs=1e-6)


def test_flat_track_20kmh_stopping_distance_hand_calc():
    # m = 10 000 kg, v = 20 km/h = 5.5556 m/s, t_r = 1 s, flat.
    #   a        = 0.823045 m/s²
    #   braking  = v²/(2a) = 30.864 / 1.646091 = 18.75 m
    #   reaction = 5.5556 · 1                  =  5.56 m
    #   total    = 24.31 m  ≤ 27 m (EN limit @ 20 km/h) → followed
    rows, ctx = perform_braking_calculation(_inputs(rail_speed_input="20"))
    r = _row(rows, "Straight Track", 20.0)
    assert r["decel"] == pytest.approx(0.82, abs=0.005)
    assert r["dist"] == pytest.approx(18.75, abs=0.01)
    assert r["total"] == pytest.approx(24.31, abs=0.01)
    assert r["status"] == FOLLOWED
    assert ctx["gbr"] == pytest.approx(8.39, abs=0.005)


def test_moving_up_1_in_100_at_40kmh_hand_calc():
    # 1 in 100 → θ = atan(0.01) = 0.5729°, sin θ = 0.0099995
    #   W        = 10 000 · 9.81      = 98 100 N
    #   W·sin θ  = 98 100 · 0.0099995 =    980.95 N
    #   F_net    = 8 230.45 + 980.95  =  9 211.40 N
    #   a        = 9 211.40 / 10 000  =    0.92114 m/s²
    #   v        = 40/3.6             =   11.111 m/s
    #   braking  = 123.457 / (2·0.92114) = 67.01 m
    #   total    = 67.01 + 11.11      =   78.12 m  ≤ 90 m → followed
    rows, _ = perform_braking_calculation(_inputs(rail_speed_input="40", rail_gradient_input="100"))
    r = _row(rows, "Moving up", 40.0)
    assert calculate_angle(100, "1 in G") == pytest.approx(0.5729, abs=1e-4)
    assert r["gravitational_force"] == pytest.approx(980.95, abs=0.05)
    assert r["net_force"] == pytest.approx(9211.40, abs=0.05)
    assert r["dist"] == pytest.approx(67.01, abs=0.01)
    assert r["total"] == pytest.approx(78.12, abs=0.01)
    assert r["status"] == FOLLOWED


def test_moving_down_gentle_gradient_hand_calc():
    # 1 in 100 downhill (vehicle CAN stop: W·sin θ < F_max)
    #   F_net = 8 230.45 − 980.95 = 7 249.50 N → a = 0.72495 m/s²
    #   v = 40 km/h: braking = 123.457 / 1.4499 = 85.15 m, total = 96.26 m > 90 m → exceeded
    rows, _ = perform_braking_calculation(_inputs(rail_speed_input="40", rail_gradient_input="100"))
    r = _row(rows, "Moving down", 40.0)
    assert r["net_force"] == pytest.approx(7249.50, abs=0.05)
    assert r["dist"] == pytest.approx(85.15, abs=0.01)
    assert r["total"] == pytest.approx(96.26, abs=0.01)
    assert r["status"] == EXCEEDED


def test_compliance_lookup_uses_lower_tabulated_speed_as_documented():
    # Documented server rule: limit of the highest tabulated speed ≤ input.
    assert get_compliance(7, 1.0) == "Standard Not Found"      # below 8 km/h
    assert get_compliance(45, 89.9) == FOLLOWED                # 40 km/h limit = 90 m
    assert get_compliance(45, 90.1) == EXCEEDED
    assert get_compliance(100, 620) == FOLLOWED
    assert MAX_STOPPING_DISTANCES[40] == 90 and MAX_STOPPING_DISTANCES[50] == 155


def test_gradient_angle_conversions():
    # degree as is; % → atan(p/100); 1 in G → atan(1/G)
    assert calculate_angle(2, "Degree (°)") == 2
    assert calculate_angle(5, "Percentage (%)") == pytest.approx(math.degrees(math.atan(0.05)))
    assert calculate_angle(20, "1 in G") == pytest.approx(2.8624, abs=1e-4)


# ---------------------------------------------------------------------------
# (b) Discrepancies (pinned, expected to fail until engineering confirms a fix)
# ---------------------------------------------------------------------------

@pytest.mark.xfail(
    reason="DISCREPANCY: abs(f_net/mass) turns a runaway on a steep downhill into a finite stopping "
           "distance, needs engineering confirmation — backend/app/modules/rnd/tools/braking/core.py:50",
    strict=True, raises=AssertionError,
)
def test_steep_downhill_vehicle_that_cannot_stop_is_flagged():
    # 1 in 10 downhill: sin θ = 0.1/√1.01 = 0.0995037
    #   W·sin θ = 98 100 · 0.0995037 = 9 761.3 N  >  F_max = 8 230.45 N
    #   F_net = 8 230.45 − 9 761.3 = −1 530.9 N → the vehicle ACCELERATES downhill.
    # Correct behaviour: no finite stopping distance and not a PASS. Today abs()
    # gives a = 0.153 m/s² and a "finite" 30 km/h distance of ≈ 227 m.
    rows, _ = perform_braking_calculation(_inputs(rail_speed_input="30", rail_gradient_input="10"))
    r = _row(rows, "Moving down", 30.0)
    assert r["net_force"] < 0  # holds today; documents the physics
    dist_is_finite = isinstance(r["dist"], (int, float)) and math.isfinite(r["dist"])
    assert not dist_is_finite, f"runaway reported as finite braking distance {r['dist']} m"
    assert "cannot stop" in str(r["status"]).lower() or r["total"] in ("Inf", math.inf)


@pytest.mark.xfail(
    reason="DISCREPANCY: abs() in the shared distance helper also makes a negative net force look like "
           "braking, needs engineering confirmation — backend/app/modules/rnd/tools/braking/core.py:50",
    strict=True, raises=AssertionError,
)
def test_distance_helper_negative_net_force_is_not_a_deceleration():
    f_net, _ = calculate_rail_scenario_forces(8230.45, 98100.0, calculate_angle(10, "1 in G"), "Moving down")
    decel, braking_dist, _, _ = calculate_deceleration_and_distances(f_net, 10000.0, 30 / 3.6, 1.0)
    assert decel <= 0 or math.isinf(braking_dist)


def _screen_limit(speed_kmh):
    """Python port of getStdDistance() in
    frontend/src/app/dashboard/rnd/braking/page.tsx:51 (linear interpolation)."""
    tbl = MAX_STOPPING_DISTANCES
    if speed_kmh in tbl:
        return tbl[speed_kmh]
    speeds = sorted(tbl)
    for lo, hi in zip(speeds, speeds[1:]):
        if lo < speed_kmh <= hi:
            return tbl[lo] + (speed_kmh - lo) / (hi - lo) * (tbl[hi] - tbl[lo])
    return None


def test_screen_interpolation_port_at_45kmh():
    # 45 km/h lies between 40 (90 m) and 50 (155 m):
    #   limit = 90 + (45−40)/(50−40) · (155−90) = 90 + 0.5·65 = 122.5 m
    assert _screen_limit(45) == pytest.approx(122.5)


@pytest.mark.xfail(
    reason="DISCREPANCY: screen interpolates EN limit (122.5 m @45 km/h) but server/PDF uses lower "
           "tabulated speed (90 m), needs engineering confirmation — "
           "backend/app/modules/rnd/tools/braking/units.py:37 vs frontend/src/app/dashboard/rnd/braking/page.tsx:51",
    strict=True, raises=AssertionError,
)
def test_screen_and_server_compliance_agree_at_45kmh():
    # A total of 100 m at 45 km/h: screen says PASS (100 ≤ 122.5), server says
    # "Standard Exceeded" (100 > 90). Whichever rule engineering picks, the two
    # must give the same verdict.
    total = 100.0
    screen_pass = total <= _screen_limit(45)
    server_pass = get_compliance(45, total) == FOLLOWED
    assert screen_pass == server_pass


@pytest.mark.xfail(
    reason="DISCREPANCY: brake force is not an input, so GBR is ~8.39 % for any vehicle, needs "
           "engineering confirmation — backend/app/modules/rnd/tools/braking/core.py:9",
    strict=True, raises=AssertionError,
)
def test_actual_brake_force_is_an_input_so_gbr_can_differ():
    from app.modules.rnd.tools.braking.schemas import BrakingInput
    fields = set(BrakingInput.model_fields)
    assert fields & {"brake_force_n", "brake_force", "applied_brake_force"}, \
        "a real brake can't be verified without its force as an input"


@pytest.mark.xfail(
    reason="DISCREPANCY: Word report receives LaTeX-escaped header text (BRK\\_001), needs engineering "
           "confirmation — backend/app/modules/rnd/tools/braking/service.py:310",
    strict=True, raises=AssertionError,
)
def test_docx_header_is_not_latex_escaped():
    import docx

    from app.modules.rnd.tools.braking.reports.pdf_builder import create_braking_docx_report

    rows, ctx = perform_braking_calculation(_inputs(doc_no="BRK_001"))
    stream = create_braking_docx_report(rows, ctx)
    d = docx.Document(io.BytesIO(stream.getvalue()))
    text = "\n".join(p.text for p in d.paragraphs)
    text += "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "BRK\\_001" not in text


@pytest.mark.xfail(
    reason="DISCREPANCY (new): above 100 km/h the screen shows no EN limit but the server keeps applying "
           "the 100 km/h limit (620 m), needs engineering confirmation — "
           "backend/app/modules/rnd/tools/braking/units.py:37 vs frontend/src/app/dashboard/rnd/braking/page.tsx:51",
    strict=True, raises=AssertionError,
)
def test_speed_above_table_is_not_silently_given_the_100kmh_limit():
    # EN table ends at 100 km/h. Screen: getStdDistance(120) → null (no verdict).
    # Server: get_compliance(120, 600) → "Standard Followed" using the 100 km/h limit.
    assert _screen_limit(120) is None
    assert get_compliance(120, 600.0) == "Standard Not Found"
