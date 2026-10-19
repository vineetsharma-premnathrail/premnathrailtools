"""Store foundation plan Phases A + B (docs/02-modules/store/foundation-plan.md):
the weighted-average valuation engine inside post_stock_transaction (§5.1),
the valuation_enabled switch (§5.2, off by default), ledger_sort_key (§5.1),
the read-only moving_average_cost mirror (§3.4), and the Item Master fields
lead_time_days / legacy_code (§4.4, §5.3, §5.5).

Engine tests run with the switch ON; the invariant (§9) is checked after each:
valuation.qty = Σ (on hand + quality + blocked) over all stores, and
valuation.value = Σ signed external ledger values + Σ events."""
import io

import pytest
from openpyxl import load_workbook

from app.auth.jwt_handler import create_access_token
from app.modules.main.models.audit_log import AuditLog
from app.modules.main.models.module_setting import ModuleSetting
from app.modules.main.models.user import User
from app.modules.store.models.item import StoreItem
from app.modules.store.models.location import StoreLocation
from app.modules.store.models.stock_balance import StoreStockBalance
from app.modules.store.models.stock_transaction import STORE_STOCK_TYPE_EFFECT, StoreStockTransaction
from app.modules.store.models.valuation import StoreItemValuation, StoreValuationEvent
from app.modules.store.services import valuation
from app.modules.store.services.stock_ledger import post_stock_transaction

API = "/api/v1/store/items"


# ---------------------------------------------------------------- fixtures

@pytest.fixture
def stores(db):
    s1, s2 = StoreLocation(name="Main Store", code="MAIN"), StoreLocation(name="Plant 2 Store", code="P2")
    db.add_all([s1, s2])
    db.commit()
    return s1, s2


@pytest.fixture
def bolt(db):
    item = StoreItem(item_code="MT-HWF-0001", item_name="Hex bolt M12", uom="NOS")
    db.add(item)
    db.commit()
    return item


def enable(db, on: bool = True):
    row = db.query(ModuleSetting).filter(ModuleSetting.module_key == "store", ModuleSetting.key == "valuation_enabled").first()
    if row is None:
        db.add(ModuleSetting(module_key="store", key="valuation_enabled", value=on))
    else:
        row.value = on
    db.commit()


@pytest.fixture
def valued(db):
    enable(db)


def post(db, item, store, ttype, qty, **kw):
    txn = post_stock_transaction(db, item_id=item.id, location_id=store.id, transaction_type=ttype, quantity=qty, **kw)
    db.commit()
    return txn


def val(db, item) -> StoreItemValuation:
    db.expire_all()
    return db.get(StoreItemValuation, item.id)


def assert_invariant(db, item):
    db.expire_all()
    v = db.get(StoreItemValuation, item.id)
    bal = sum((b.on_hand_qty or 0) + (b.quality_qty or 0) + (b.quarantine_qty or 0)
              for b in db.query(StoreStockBalance).filter(StoreStockBalance.item_id == item.id))
    assert v.qty == pytest.approx(bal, abs=1e-6)
    signed = 0.0
    for t in db.query(StoreStockTransaction).filter(StoreStockTransaction.item_id == item.id):
        src, dst = STORE_STOCK_TYPE_EFFECT[t.transaction_type]
        if src is None:
            signed += t.value or 0
        elif dst is None:
            signed -= t.value or 0
    events = sum(e.value_delta for e in db.query(StoreValuationEvent).filter(StoreValuationEvent.item_id == item.id))
    assert v.value == pytest.approx(signed + events, abs=0.005)


# ---------------------------------------------------------------- switch

def test_switch_is_off_by_default_and_ledger_is_quantity_only(db, stores, bolt):
    assert valuation.valuation_enabled(db) is False
    txn = post(db, bolt, stores[0], "receipt", 10, unit_cost=100)
    out = post(db, bolt, stores[0], "issue", 4)
    for t in (txn, out):
        assert t.unit_cost is None and t.value is None and t.cost_source is None
        assert t.valuation_qty_after is None and t.valuation_value_after is None
    assert db.get(StoreItemValuation, bolt.id) is None
    db.refresh(bolt)
    assert bolt.moving_average_cost is None
    bal = db.query(StoreStockBalance).filter_by(item_id=bolt.id).one()
    assert bal.on_hand_qty == 6


# ---------------------------------------------------------------- formula

def test_worked_example_grn_at_120_on_stock_at_100(db, valued, stores, bolt):
    s1 = stores[0]
    first = post(db, bolt, s1, "receipt", 10, unit_cost=100)
    assert first.cost_source == "po_rate" and first.value == 1000
    grn = post(db, bolt, s1, "receipt", 5, unit_cost=120, reference_number="GRN-0002")
    v = val(db, bolt)
    assert v.qty == 15 and v.value == 1600
    assert v.avg_rate == pytest.approx(106.6667)
    assert v.last_receipt_rate == 120 and v.last_receipt_ref == "GRN-0002"
    assert grn.valuation_qty_after == 15 and grn.valuation_value_after == 1600

    issue = post(db, bolt, s1, "issue", 4)
    assert issue.unit_cost == pytest.approx(106.6667) and issue.value == pytest.approx(426.67)
    assert issue.cost_source == "average"
    v = val(db, bolt)
    assert v.qty == 11 and v.value == pytest.approx(1173.33)

    last = post(db, bolt, s1, "issue", 11)
    assert last.value == pytest.approx(1173.33)  # the last unit takes the remaining value
    v = val(db, bolt)
    assert v.qty == 0 and v.value == 0
    # Kept as the last rate: the rate the stock went out at (V ÷ Q = 1,173.33 ÷ 11;
    # avg_rate is always derived from the totals, so it carries the paise rounding).
    assert v.avg_rate == pytest.approx(1173.33 / 11, abs=1e-4)
    assert last.unit_cost == pytest.approx(1173.33 / 11, abs=1e-4)
    assert_invariant(db, bolt)


def test_negative_stock_still_blocked_and_valuation_untouched(db, valued, stores, bolt):
    post(db, bolt, stores[0], "receipt", 3, unit_cost=50)
    with pytest.raises(ValueError, match="Not enough stock"):
        post_stock_transaction(db, item_id=bolt.id, location_id=stores[0].id, transaction_type="issue", quantity=5)
    db.rollback()
    v = val(db, bolt)
    assert v.qty == 3 and v.value == 150
    assert_invariant(db, bolt)


# ---------------------------------------------------------------- internal moves

def test_transfer_between_stores_keeps_value_even_through_zero(db, valued, stores, bolt):
    s1, s2 = stores
    post(db, bolt, s1, "receipt", 3, unit_cost=10)
    post(db, bolt, s1, "receipt", 4, unit_cost=11)  # V = 74, Q = 7
    out = post(db, bolt, s1, "transfer_out", 7, reference_type="stock_transfer")  # entire stock: Q passes 0
    assert out.value == 74
    inn = post(db, bolt, s2, "transfer_in", 7, value=out.value, unit_cost=out.unit_cost, reference_type="stock_transfer")
    assert inn.value == 74
    v = val(db, bolt)
    assert v.qty == 7 and v.value == 74
    assert v.avg_rate == pytest.approx(74 / 7, abs=1e-4)
    assert_invariant(db, bolt)


def test_stock_type_moves_are_value_neutral(db, valued, stores, bolt):
    s1 = stores[0]
    post(db, bolt, s1, "quality_in", 10, unit_cost=100)  # valued at receipt (§5.1.3)
    assert val(db, bolt).value == 1000
    rel = post(db, bolt, s1, "quality_release", 6)
    blk = post(db, bolt, s1, "quality_block", 4)
    assert rel.value == 600 and blk.value == 400  # shown only
    post(db, bolt, s1, "free_block", 1)
    post(db, bolt, s1, "blocked_release", 2)
    post(db, bolt, s1, "blocked_to_quality", 1)
    post(db, bolt, s1, "quality_hold", 1)
    v = val(db, bolt)
    assert v.qty == 10 and v.value == 1000 and v.avg_rate == 100
    # Quarantine release pair (out at average, in at the out row's value).
    out = post(db, bolt, s1, "quarantine_out", 2)
    post(db, bolt, s1, "return_in", 2, value=out.value)
    assert val(db, bolt).value == 1000
    assert_invariant(db, bolt)


# ---------------------------------------------------------------- returns / vendor returns

def test_return_at_issue_cost_and_vendor_return_clamp(db, valued, stores, bolt):
    s1 = stores[0]
    post(db, bolt, s1, "receipt", 10, unit_cost=100)
    issue = post(db, bolt, s1, "issue", 2)
    post(db, bolt, s1, "receipt", 10, unit_cost=130)  # average now (800 + 1300) / 18
    ret = post(db, bolt, s1, "return_in", 2, unit_cost=issue.unit_cost, cost_source="issue_rate")
    assert ret.value == 200 and ret.cost_source == "issue_rate"
    assert_invariant(db, bolt)

    # Vendor return at the GRN rate (reverse of the receipt).
    vr = post(db, bolt, s1, "return_out", 5, unit_cost=130, cost_source="grn_rate")
    assert vr.value == 650 and vr.cost_source == "grn_rate"
    assert_invariant(db, bolt)


def test_vendor_return_above_stock_value_is_clamped_with_revaluation_event(db, valued, stores, bolt):
    s1 = stores[0]
    post(db, bolt, s1, "receipt", 1, unit_cost=10)
    post(db, bolt, s1, "receipt", 1, unit_cost=10)
    vr = post(db, bolt, s1, "return_out", 1, unit_cost=50, cost_source="grn_rate")
    assert vr.value == 20  # min(q × r, V)
    ev = db.query(StoreValuationEvent).filter_by(item_id=bolt.id, event_type="revaluation").one()
    assert "above current average" in ev.reason
    v = val(db, bolt)
    assert v.qty == 1 and v.value == 0
    assert_invariant(db, bolt)


# ---------------------------------------------------------------- zero stock / unpriced

def test_receipt_after_zero_takes_new_rate_and_residual_is_written_off(db, valued, stores, bolt):
    s1 = stores[0]
    post(db, bolt, s1, "receipt", 2, unit_cost=100)
    post(db, bolt, s1, "issue", 2)
    post(db, bolt, s1, "receipt", 4, unit_cost=80)
    v = val(db, bolt)
    assert v.avg_rate == 80 and v.value == 320

    # Residual value with zero quantity (should not happen) → rounding_writeoff on the next receipt.
    post(db, bolt, s1, "issue", 4)
    row = val(db, bolt)
    row.value = 0.03
    db.commit()
    post(db, bolt, s1, "receipt", 1, unit_cost=90)
    ev = db.query(StoreValuationEvent).filter_by(item_id=bolt.id, event_type="rounding_writeoff").one()
    assert ev.value_delta == pytest.approx(-0.03)
    assert val(db, bolt).value == 90


def test_unpriced_user_receipt_refused_system_receipt_unvalued(db, valued, stores, bolt):
    s1 = stores[0]
    with pytest.raises(ValueError) as exc:
        post_stock_transaction(db, item_id=bolt.id, location_id=s1.id, transaction_type="manual_in", quantity=5)
    assert "has no average or standard cost yet" in str(exc.value) and "per NOS" in str(exc.value)
    db.rollback()
    assert db.query(StoreStockBalance).filter_by(item_id=bolt.id).first() is None or \
        db.query(StoreStockBalance).filter_by(item_id=bolt.id).one().on_hand_qty == 0

    sysrow = post(db, bolt, s1, "receipt", 5)  # system document with no rate (e.g. legacy unpriced PO)
    assert sysrow.cost_source == "none" and sysrow.value == 0
    v = val(db, bolt)
    assert v.unvalued_qty == 5 and v.avg_rate is None
    assert_invariant(db, bolt)


def test_unpriced_receipt_uses_average_then_standard(db, valued, stores, bolt):
    s1 = stores[0]
    bolt.standard_cost = 40
    db.commit()
    std = post(db, bolt, s1, "manual_in", 2)
    assert std.cost_source == "standard" and std.value == 80
    avg = post(db, bolt, s1, "adjustment_in", 1)
    assert avg.cost_source == "average" and avg.unit_cost == 40
    assert_invariant(db, bolt)


def test_zero_rate_needs_manual_rate_and_remark(db, valued, stores, bolt):
    with pytest.raises(ValueError, match="zero rate"):
        post_stock_transaction(db, item_id=bolt.id, location_id=stores[0].id, transaction_type="receipt", quantity=1, unit_cost=0)
    db.rollback()
    free = post(db, bolt, stores[0], "manual_in", 1, unit_cost=0, cost_source="manual_rate", remarks="Free replacement under warranty")
    assert free.value == 0


def test_back_dated_receipt_does_not_recost_earlier_issue(db, valued, stores, bolt):
    from datetime import date, timedelta
    s1 = stores[0]
    post(db, bolt, s1, "receipt", 10, unit_cost=100)
    issue = post(db, bolt, s1, "issue", 5)
    post(db, bolt, s1, "receipt", 5, unit_cost=200, transaction_date=date.today() - timedelta(days=30))
    db.refresh(issue)
    assert issue.unit_cost == 100 and issue.value == 500
    assert val(db, bolt).avg_rate == 150


def test_switch_on_mid_life_seeds_opening_from_balances(db, stores, bolt):
    s1 = stores[0]
    bolt.standard_cost = 25
    db.commit()
    post(db, bolt, s1, "receipt", 8)  # switch off: quantity only
    enable(db)
    out = post(db, bolt, s1, "issue", 2)
    assert out.unit_cost == 25 and out.value == 50
    ev = db.query(StoreValuationEvent).filter_by(item_id=bolt.id, event_type="cutover_opening").one()
    assert ev.value_delta == 200 and ev.qty_basis == 8
    assert_invariant(db, bolt)


def test_switch_off_again_returns_to_quantity_only(db, valued, stores, bolt):
    post(db, bolt, stores[0], "receipt", 4, unit_cost=10)
    enable(db, False)
    t = post(db, bolt, stores[0], "issue", 1)
    assert t.value is None and t.cost_source is None
    assert val(db, bolt).qty == 4  # valuation frozen while off (rollback switch)


# ---------------------------------------------------------------- mirror / sort key

def test_mirror_follows_average_without_audit_row_or_updated_at_change(db, valued, stores, bolt):
    db.refresh(bolt)
    before = bolt.updated_at
    audits_before = db.query(AuditLog).filter(AuditLog.entity_type == "store_item").count()
    post(db, bolt, stores[0], "receipt", 10, unit_cost=100)
    post(db, bolt, stores[0], "receipt", 5, unit_cost=120)
    db.expire_all()
    item = db.get(StoreItem, bolt.id)
    assert item.moving_average_cost == pytest.approx(106.6667)
    assert item.updated_at == before
    assert db.query(AuditLog).filter(AuditLog.entity_type == "store_item").count() == audits_before


def test_ledger_sort_key_orders_by_item_then_store():
    lines = [(2, 1), (1, 2), (1, 1), (2, 2)]
    assert sorted(lines, key=lambda l: valuation.ledger_sort_key(*l)) == [(1, 1), (1, 2), (2, 1), (2, 2)]


# ---------------------------------------------------------------- item master fields (Phase A)

@pytest.fixture
def h(db):
    user = User(email="store.keeper@premnathrail.com", name="store.keeper", role="user", is_active=True, assigned_apps=["store"])
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id), 'email': user.email, 'role': user.role})}"}


def test_lead_time_and_legacy_code_on_item_api(client, h):
    r = client.post(API, json={"item_name": "Hose", "uom": "NOS", "lead_time_days": 21, "legacy_code": " HYD-HOS-0007 "}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["lead_time_days"] == 21 and body["legacy_code"] == "HYD-HOS-0007"

    bad = client.post(API, json={"item_name": "Pump", "uom": "NOS", "lead_time_days": 731}, headers=h)
    assert bad.status_code == 400 and "between 0 and 730" in bad.json()["detail"]
    neg = client.patch(f"{API}/{body['id']}", json={"lead_time_days": -1}, headers=h)
    assert neg.status_code == 400

    dup = client.post(API, json={"item_name": "Other", "uom": "NOS", "legacy_code": "hyd-hos-0007"}, headers=h)
    assert dup.status_code == 409 and "already used" in dup.json()["detail"]

    # Searchable by legacy code; "No lead time" filter.
    client.post(API, json={"item_name": "Gasket", "uom": "NOS"}, headers=h)
    found = client.get(API, params={"search": "HOS-0007"}, headers=h).json()
    assert [i["item_name"] for i in found] == ["Hose"]
    no_lt = client.get(API, params={"no_lead_time": True}, headers=h).json()
    assert [i["item_name"] for i in no_lt] == ["Gasket"]


def test_moving_average_cost_is_read_only_in_api(client, h, db):
    r = client.post(API, json={"item_name": "Hose", "uom": "NOS", "moving_average_cost": 99}, headers=h)
    item_id = r.json()["id"]
    assert r.json()["moving_average_cost"] is None
    p = client.patch(f"{API}/{item_id}", json={"moving_average_cost": 55, "standard_cost": 40}, headers=h)
    assert p.status_code == 200
    assert p.json()["moving_average_cost"] is None and p.json()["standard_cost"] == 40


def test_import_export_round_trip_lead_time_and_legacy_code(client, h):
    csv_text = ("item_type,category,subcategory,item_name,uom,technical_specification,lead_time_days,legacy_code\n"
                "Material,,,Hex Bolt M12x50,NOS,Zinc plated,21,V3-0012\n"
                "Material,,,Washer M12,NOS,,800,\n"
                "Material,,,Nut M12,NOS,,,V3-0012\n")
    r = client.post(f"{API}/import", files={"file": ("items.csv", csv_text.encode(), "text/csv")}, headers=h)
    assert r.status_code == 200, r.text
    out = r.json()
    assert [c["item_name"] for c in out["created"]] == ["Hex Bolt M12x50"]
    reasons = {e["item_name"]: e["reason"] for e in out["errors"]}
    assert "between 0 and 730" in reasons["Washer M12"]
    assert "V3-0012" in reasons["Nut M12"]

    x = client.get(f"{API}/export", headers=h)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    row = dict(zip(header, rows[1]))
    assert row["Lead Time (days)"] == 21 and row["Legacy Code"] == "V3-0012"

    # The exported titles are accepted back by the import (as a CSV with the same headers).
    csv2 = ",".join(["item_name", "uom", "Lead Time (days)", "Legacy Code"]) + "\nSpring washer,NOS,14,V3-0099\n"
    r2 = client.post(f"{API}/import", files={"file": ("items.csv", csv2.encode(), "text/csv")}, headers=h)
    assert r2.json()["created"][0]["item_name"] == "Spring washer"
    items = client.get(API, params={"search": "V3-0099"}, headers=h).json()
    assert items[0]["lead_time_days"] == 14
