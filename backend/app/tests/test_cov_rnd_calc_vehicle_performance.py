"""R&D calculator verification pack: Vehicle performance (tools/vehicle_performance).

Reference: docs/02-modules/rnd/overview.md §8.6 and §14 row 9.
The live code path is service.VehiclePerformanceCalculator (core.py and
calculate.py are unused). Golden values hand-computed (forces in N, weight in t).
Discrepancies pinned with xfail(strict=True, raises=AssertionError).
"""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.modules.rnd.tools.vehicle_performance.api import router as vp_router
from app.modules.rnd.tools.vehicle_performance.service import (
    VehiclePerformanceCalculator,
    curvature_resistance,
    gradient_resistance,
    interpolate_torque,
    rolling_resistance_loco,
    rolling_resistance_wagon,
    starting_resistance_loco,
    starting_resistance_wagon,
)


def _calc(**over):
    d = dict(
        loco_gvw_kg=20000.0, max_speed_kmh=100.0, num_axles=2, rear_axle_ratio=5.0,
        gear_ratios=[4.0], shunting_load_t=0.0, peak_power_kw=500.0, friction_mu=0.3,
        wheel_dia_m=0.8, min_rpm=1000, max_rpm=2000, torque_curve={1000: 400.0, 2000: 500.0},
        max_curve=0.0, curve_unit="degree", max_slope=0.0, slope_unit="%",
    )
    d.update(over)
    return VehiclePerformanceCalculator(d)


# ---------------------------------------------------------------------------
# (a) Golden values
# ---------------------------------------------------------------------------

def test_resistance_functions_hand_calc():
    # Loco 20 t, 2 axles, 20 km/h:
    #   A = 0.647 + 13.17/10 = 1.964 ; B·v = 0.1866 ; C·v² = (0.057/20)·400 = 1.14
    #   R = 3.2906 · 20 · 9.81 = 645.62 N
    assert rolling_resistance_loco(20, 20, 2) == pytest.approx(645.62, abs=0.01)
    # Wagon 1 t, 10 km/h: (0.6438797 + 0.1047218 + 0.007323)·9.81 = 0.7559245·9.81 = 7.4156 N
    assert rolling_resistance_wagon(10, 1) == pytest.approx(7.4156, abs=1e-4)
    # Gradient 20 t @ 2 %: 20·1000·9.81·0.02 = 3924 N ; curve 20 t @ 3°: 0.4·20·3·9.81 = 235.44 N
    assert gradient_resistance(20, 2) == pytest.approx(3924.0)
    assert curvature_resistance(20, 3) == pytest.approx(235.44)
    # Starting: loco 6·20·9.81 = 1177.2 N ; wagon 4·1·9.81 = 39.24 N
    assert starting_resistance_loco(20) == pytest.approx(1177.2)
    assert starting_resistance_wagon(1) == pytest.approx(39.24)


def test_torque_interpolation_and_clamping():
    curve = {1000: 400.0, 2000: 500.0}
    assert interpolate_torque(1500, curve) == pytest.approx(450.0)
    assert interpolate_torque(500, curve) == pytest.approx(400.0)   # clamped low
    assert interpolate_torque(3000, curve) == pytest.approx(500.0)  # clamped high


def test_traction_snapshot_not_limited_hand_calc():
    # F = 2·T_max·i_gear·i_axle/D = 2·500·4·5/0.8 = 25 000 N
    # F_adh = 20 t·0.3·1000·9.81 = 58 860 N → 25 000 < 58 860 → not limited
    s = _calc().run_tractive_calculation()
    assert s["max_traction_generated_n"] == pytest.approx(25000.0)
    assert s["max_traction_slipping_n"] == pytest.approx(58860.0)
    assert s["result_message"] == "Not limited by slipping"


def test_traction_snapshot_limited_hand_calc():
    # Highest gear ratio 10: F = 2·500·10·5/0.8 = 62 500 N > 58 860 N → limited
    s = _calc(gear_ratios=[10.0, 4.0]).run_tractive_calculation()
    assert s["max_traction_generated_n"] == pytest.approx(62500.0)
    assert s["result_message"] == "Limited by slipping"


def test_slope_in_degrees_and_curve_as_radius_are_converted():
    # 2° → tan(2°)·100 = 3.4921 % ; R = 175 m → D = 1750/175 = 10°
    c = _calc(max_slope=2.0, slope_unit="degree", max_curve=175.0, curve_unit="m")
    assert c.max_slope_percent == pytest.approx(3.4921, abs=1e-4)
    assert c.max_curve_deg == pytest.approx(10.0)


def test_report_builder_physics_matches_service():
    # Physics is duplicated in reports/pdf_builder.py (doc §8.6 item 6): pin the copies together.
    from app.modules.rnd.tools.vehicle_performance.reports import pdf_builder as rb

    for v in (0, 15, 42.5, 80):
        assert rb._rr_loco(v, 20, 2) == pytest.approx(rolling_resistance_loco(v, 20, 2))
        assert rb._rr_wagon(v, 7) == pytest.approx(rolling_resistance_wagon(v, 7))
    assert rb._grad_res(20, 2.5) == pytest.approx(gradient_resistance(20, 2.5))
    assert rb._curve_res(20, 4) == pytest.approx(curvature_resistance(20, 4))
    assert rb._start_loco(20) == pytest.approx(starting_resistance_loco(20))
    assert rb._start_wagon(3) == pytest.approx(starting_resistance_wagon(3))


# ---------------------------------------------------------------------------
# (b) Discrepancies
# ---------------------------------------------------------------------------

# Two-gear vehicle where the low gear has traction but runs out of rpm early,
# and the top gear reaches high speed but lacks traction:
#   flat torque 100 N·m (800–2000 rpm), axle 5, Ø0.8 m, gears 4 and 1, 20 t, 2 axles, flat, no load
#   gear 4: F = 2·100·4·5/0.8 = 5000 N, but v_max @2000 rpm = 2000·π·0.8/(4·5·60)·3.6 = 15.08 km/h
#   gear 1: F = 2·100·1·5/0.8 = 1250 N, v_max @2000 rpm = 60.32 km/h
#     balance 1250 = (1.964 + 0.00933v + 0.00285v²)·196.2 → 0.00285v² + 0.00933v − 4.407 = 0
#     → v = 37.7 km/h  ← the true maximum on 0 % slope
#   The table scans 6.03…60.32 km/h but always uses gear 4 (5000 N, torque clamped past
#   2000 rpm) → reports 60.32 km/h.
TWO_GEAR = dict(gear_ratios=[4.0, 1.0], min_rpm=800, max_rpm=2000,
                torque_curve={800: 100.0, 2000: 100.0}, peak_power_kw=500.0)


def test_two_gear_case_documented_current_output():
    table = _calc(**TWO_GEAR).calculate_speed_for_shunting_load()
    assert table[0]["slope"] == 0
    assert table[0]["max_speed_kmh"] == pytest.approx(60.32, abs=0.01)


@pytest.mark.xfail(
    reason="DISCREPANCY: speed-vs-slope evaluates every speed in the lowest gear past that gear's rpm limit "
           "(torque clamped), needs engineering confirmation — "
           "backend/app/modules/rnd/tools/vehicle_performance/service.py:280 (range at :264)",
    strict=True, raises=AssertionError,
)
def test_speed_vs_slope_respects_gear_rpm_limit():
    table = _calc(**TWO_GEAR).calculate_speed_for_shunting_load()
    assert table[0]["max_speed_kmh"] <= 38.0  # true max ≈ 37.7 km/h (gear 1)


@pytest.mark.xfail(
    reason="DISCREPANCY: max_speed input is stored but never caps any result, needs engineering "
           "confirmation — backend/app/modules/rnd/tools/vehicle_performance/service.py:121",
    strict=True, raises=AssertionError,
)
def test_speed_table_is_capped_at_vehicle_max_speed():
    table = _calc(max_speed_kmh=30.0, **TWO_GEAR).calculate_speed_for_shunting_load()
    assert all(row["max_speed_kmh"] <= 30.0 for row in table)


@pytest.mark.xfail(
    reason="DISCREPANCY: document metadata (doc no., made/checked/approved by) is collected but not printed "
           "in the DOCX, needs engineering confirmation — "
           "backend/app/modules/rnd/tools/vehicle_performance/reports/pdf_builder.py",
    strict=True, raises=AssertionError,
)
def test_docx_report_prints_document_metadata():
    import docx

    app = FastAPI()
    app.include_router(vp_router, prefix="/vp")
    client = TestClient(app, raise_server_exceptions=False)
    payload = dict(
        max_curve=0, max_slope=1, loco_gvw=20000, max_speed=60, num_axles=2, rear_axle_ratio=5,
        gear_ratios=[4, 1], shunting_load=50, peak_power=100, friction_mu=0.3, wheel_dia=0.8,
        min_rpm=800, max_rpm=2000, torque_curve={"800": 300, "1400": 400, "2000": 350},
        doc_no="VP-DOC-77", made_by="Tester A", checked_by="Checker B", approved_by="Approver C",
    )
    res = client.post("/vp/download-report", json=payload)
    assert res.status_code == 200, res.text
    d = docx.Document(io.BytesIO(res.content))
    text = "\n".join(p.text for p in d.paragraphs)
    text += "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "VP-DOC-77" in text
