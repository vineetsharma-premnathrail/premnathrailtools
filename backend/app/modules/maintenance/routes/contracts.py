"""AMC & Warranty tracker (Phase 3) — one row per (asset, warranty|amc)
with an expiry bucket. Read-only: the dates themselves are edited on the
asset. Bucket rules live in modules/maintenance/contracts.py."""
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.core.permissions import require_app_access, require_tab_access, require_tab_action
from app.db.session import get_db
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch
from app.modules.maintenance.contracts import (
    CONTRACT_KINDS, CONTRACT_STATUSES, DEFAULT_WINDOW_DAYS, KIND_LABELS, STATUS_LABELS,
    build_contract_rows, summarize,
)
from app.modules.maintenance.service import today_ist

router = APIRouter(
    prefix="/maintenance", tags=["Maintenance"],
    dependencies=[Depends(require_app_access("maintenance"))],
)


def _validate(kind: str, status: str | None, window_days: int) -> None:
    if kind not in ("all", *CONTRACT_KINDS):
        raise HTTPException(status_code=400, detail=f"Unknown cover type '{kind}'. Use warranty, amc or all.")
    if status and status not in CONTRACT_STATUSES:
        raise HTTPException(status_code=400, detail=f"Unknown status '{status}'. Use expired, expiring, active or none.")
    if not (1 <= window_days <= 365):
        raise HTTPException(status_code=400, detail=f"'Expiring within' must be between 1 and 365 days (got {window_days}).")


def _query(db: Session, branch_id, kind, window_days, status, search):
    _validate(kind, status, window_days)
    today = today_ist()
    rows = build_contract_rows(db, today, branch_id=branch_id, kind=kind, window_days=window_days, search=search)
    summary = summarize(rows)
    if status:
        rows = [r for r in rows if r["status"] == status]
    return today, rows, summary


@router.get("/contracts")
async def list_contracts(
    branch_id: int | None = None,
    kind: str = "all",
    window_days: int = DEFAULT_WINDOW_DAYS,
    status: str | None = None,
    search: str | None = None,
    db: Session = Depends(get_db),
    _tab: User = Depends(require_tab_access("maintenance", "assets")),
):
    """Summary counts ignore the status filter (so the tiles always show the
    full picture for the chosen plant / cover type / search)."""
    today, rows, summary = _query(db, branch_id, kind, window_days, status, search)
    return {"as_of": today, "window_days": window_days, "summary": summary, "items": rows}


@router.get("/contracts/export")
async def export_contracts(
    branch_id: int | None = None,
    kind: str = "all",
    window_days: int = DEFAULT_WINDOW_DAYS,
    status: str | None = None,
    search: str | None = None,
    db: Session = Depends(get_db),
    _tab: User = Depends(require_tab_access("maintenance", "assets")),
    _exp: User = Depends(require_tab_action("maintenance", "assets", "export")),
):
    today, rows, summary = _query(db, branch_id, kind, window_days, status, search)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    scope = []
    if branch_id:
        scope.append(f"Plant: {db.query(Branch.name).filter(Branch.id == branch_id).scalar() or f'#{branch_id}'}")
    scope.append(f"Cover: {'Warranty + AMC' if kind == 'all' else KIND_LABELS[kind]}")
    if status:
        scope.append(f"Status: {STATUS_LABELS[status]}")
    if search:
        scope.append(f"Search: {search}")
    subtitle = (f"As of {today:%d-%m-%Y} · Expiring = within {window_days} days · " + " · ".join(scope)
                + f" · Expired {summary['expired']}, Expiring {summary['expiring']}, Active {summary['active']}, Not covered {summary['none']}")

    wb = Workbook()
    ws = wb.active
    ws.title = "AMC & Warranty"
    ws.append(["AMC & Warranty Tracker"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    headers = ["Asset Code", "Asset Name", "Plant", "Criticality", "Cover", "Vendor", "Expiry Date", "Days Left", "Status"]
    widths = [16, 34, 20, 11, 11, 30, 14, 11, 14]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FDF1E6")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([
            r["asset_code"], r["asset_name"], r["branch_name"], r["criticality"], KIND_LABELS[r["kind"]],
            r["vendor_name"], r["expiry_date"], r["days_left"], STATUS_LABELS[r["status"]],
        ])
        if r["expiry_date"]:
            ws.cell(row=ws.max_row, column=7).number_format = "DD-MM-YYYY"
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    ws.freeze_panes = "A5"

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="maintenance_amc_warranty_{today:%Y%m%d}.xlsx"'},
    )
