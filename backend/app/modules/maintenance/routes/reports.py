"""Maintenance reports (plan §6): GET /maintenance/reports/summary and the
Excel export of the same numbers. The maths lives in ../reporting.py so the
dashboard reuses it."""
import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.permissions import require_app_access, require_tab_access, require_tab_action
from app.db.session import get_db
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch
from app.modules.maintenance.models.asset import MaintenanceAsset, MAINTENANCE_ASSET_CATEGORIES
from app.modules.maintenance.reporting import compute_summary, today_ist

router = APIRouter(prefix="/maintenance", tags=["Maintenance"])

MAX_RANGE_DAYS = 3 * 366


def _resolve(db: Session, date_from: date | None, date_to: date | None, branch_id: int | None, asset_id: int | None, category: str | None) -> tuple[date, date]:
    today = today_ist()
    d_from = date_from or today.replace(day=1)
    d_to = date_to or today
    if d_from > d_to:
        raise HTTPException(status_code=400, detail=f"'From' date {d_from:%d-%m-%Y} is after 'To' date {d_to:%d-%m-%Y} — swap them.")
    if (d_to - d_from).days > MAX_RANGE_DAYS:
        raise HTTPException(status_code=400, detail="The period is longer than 3 years — pick a shorter range (e.g. one financial year).")
    if category and category not in MAINTENANCE_ASSET_CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Unknown asset category '{category}'. Use one of: {', '.join(MAINTENANCE_ASSET_CATEGORIES)}.")
    if branch_id and not db.query(Branch.id).filter(Branch.id == branch_id).first():
        raise HTTPException(status_code=404, detail=f"Plant #{branch_id} not found — pick the plant again.")
    if asset_id and not db.query(MaintenanceAsset.id).filter(MaintenanceAsset.id == asset_id, MaintenanceAsset.is_deleted == False).first():  # noqa: E712
        raise HTTPException(status_code=404, detail=f"Asset #{asset_id} not found (it may have been deleted) — pick the asset again.")
    return d_from, d_to


@router.get("/reports/summary")
async def report_summary(
    date_from: date | None = Query(None, alias="from"),
    date_to: date | None = Query(None, alias="to"),
    branch_id: int | None = None,
    asset_id: int | None = None,
    category: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("maintenance")),
    _tab: User = Depends(require_tab_access("maintenance", "reports")),
):
    d_from, d_to = _resolve(db, date_from, date_to, branch_id, asset_id, category)
    return compute_summary(db, d_from, d_to, branch_id=branch_id, asset_id=asset_id, category=category, user=user)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def _hours(mins) -> float | None:
    return None if mins is None else round(mins / 60, 2)


def _sheet(wb, title: str, subtitle: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    from openpyxl.styles import Font, PatternFill, Alignment

    ws = wb.create_sheet(title=title[:31])
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FDF1E6")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [18] * len(headers)):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    ws.freeze_panes = "A5"
    return ws


@router.get("/reports/export")
async def report_export(
    date_from: date | None = Query(None, alias="from"),
    date_to: date | None = Query(None, alias="to"),
    branch_id: int | None = None,
    asset_id: int | None = None,
    category: str | None = None,
    export_format: str = Query("xlsx", alias="format"),
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("maintenance")),
    _tab: User = Depends(require_tab_access("maintenance", "reports")),
    _exp: User = Depends(require_tab_action("maintenance", "reports", "export")),
):
    if export_format != "xlsx":
        raise HTTPException(status_code=400, detail=f"Export format '{export_format}' isn't supported — only xlsx (Excel).")
    d_from, d_to = _resolve(db, date_from, date_to, branch_id, asset_id, category)
    s = compute_summary(db, d_from, d_to, branch_id=branch_id, asset_id=asset_id, category=category, user=user)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    scope = []
    if branch_id:
        scope.append(f"Plant: {db.query(Branch.name).filter(Branch.id == branch_id).scalar()}")
    if asset_id:
        scope.append(f"Asset: {db.query(MaintenanceAsset.asset_code).filter(MaintenanceAsset.id == asset_id).scalar()}")
    if category:
        scope.append(f"Category: {category.replace('_', ' ').title()}")
    subtitle = f"Period {d_from:%d-%m-%Y} to {d_to:%d-%m-%Y}" + (f" · {' · '.join(scope)}" if scope else " · All plants")

    k = s["kpis"]
    d = s["definitions"]
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Summary", subtitle, ["Metric", "Value", "Definition"], [
        ["Breakdown downtime (hours)", _hours(k["downtime_minutes"]), d["downtime"]],
        ["Breakdowns (incl. corrective machine-down jobs)", k["breakdowns"],
         "Breakdown + corrective (machine-down) work orders whose downtime started (or, with no downtime recorded, which were raised) in the period."],
        ["MTTR (hours)", _hours(k["mttr_minutes"]), d["mttr"]],
        ["MTBF (hours)", k["mtbf_hours"], d["mtbf"]],
        ["Available hours", k["available_hours"], s["availability_basis"]["description"]],
        ["Availability %", k["availability_pct"], d["availability"]],
        ["PM compliance %", k["pm_compliance_pct"], d["pm_compliance"]],
        ["Maintenance cost (₹)", k["maintenance_cost"], d["cost"]],
        ["Open work orders (now)", k["open_work_orders"], "Jobs not yet completed, as of today."],
    ], [46, 16, 110])
    for row in ws.iter_rows(min_row=5, max_col=3):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)

    _sheet(wb, "Downtime by Month", subtitle, ["Month", "Downtime (hours)"],
           [[m["month"], _hours(m["downtime_minutes"])] for m in s["downtime"]["by_month"]], [14, 18])
    _sheet(wb, "Downtime by Asset", subtitle,
           ["Asset Code", "Asset Name", "Breakdowns", "Downtime (hours)", "Available hours", "Availability %"],
           [[r["asset_code"], r["asset_name"], r["breakdowns"], _hours(r["downtime_minutes"]), r["available_hours"], r["availability_pct"]]
            for r in s["downtime"]["by_asset"]], [16, 34, 12, 16, 16, 14])
    pm = s["pm_compliance"]
    _sheet(wb, "PM Compliance", subtitle, ["Measure", "Count"], [
        ["PM jobs due in period", pm["due"]], ["Done on/before due date", pm["on_time"]], ["Done late", pm["late"]],
        ["Overdue, not done", pm["missed"]], ["Skipped (cancelled)", pm["skipped"]], ["Not yet due", pm["not_yet_due"]], ["Compliance %", pm["compliance_pct"]],
    ], [30, 12])
    mix = s["work_mix"]
    _sheet(wb, "Breakdown vs Planned", subtitle, ["Work order type", "Jobs raised in period"],
           [[t.replace("_", " ").title(), n] for t, n in sorted(mix["by_type"].items())]
           + [["Breakdown + corrective", mix["breakdown"]], ["Preventive + calibration + inspection", mix["planned"]], ["Planned share %", mix["planned_pct"]]],
           [38, 22])
    _sheet(wb, "Cost by Month", subtitle, ["Month", "Spares (₹)", "Labour (₹)", "External (₹)", "Total (₹)"],
           [[m["month"], m["spares"], m["labour"], m["external"], m["total"]] for m in s["cost"]["by_month"]], [14, 14, 14, 14, 14])
    _sheet(wb, "Cost by Asset", subtitle, ["Asset Code", "Asset Name", "Jobs", "Spares (₹)", "Labour (₹)", "External (₹)", "Total (₹)"],
           [[r["asset_code"], r["asset_name"], r["jobs"], r["spares"], r["labour"], r["external"], r["total"]] for r in s["cost"]["by_asset"]],
           [16, 34, 8, 14, 14, 14, 14])
    _sheet(wb, "Top 10 by Breakdowns", subtitle, ["Asset Code", "Asset Name", "Breakdowns", "Downtime (hours)"],
           [[r["asset_code"], r["asset_name"], r["breakdowns"], _hours(r["downtime_minutes"])] for r in s["top_assets"]["by_breakdowns"]], [16, 34, 12, 16])
    _sheet(wb, "Top 10 by Downtime", subtitle, ["Asset Code", "Asset Name", "Downtime (hours)", "Breakdowns"],
           [[r["asset_code"], r["asset_name"], _hours(r["downtime_minutes"]), r["breakdowns"]] for r in s["top_assets"]["by_downtime"]], [16, 34, 16, 12])
    b = s["backlog"]
    _sheet(wb, "Backlog", subtitle.split(" · ", 1)[-1] + " · open jobs as of today",
           ["Age (days)"] + [p.title() for p in b["priorities"]] + ["Total"],
           [[r["age_bucket"]] + [r[p] for p in b["priorities"]] + [r["total"]] for r in b["matrix"]], [12, 10, 10, 10, 10, 10])
    _sheet(wb, "Failure Pareto", subtitle, ["Failure category", "Jobs", "Downtime (hours)", "% of jobs", "Cumulative %"],
           [[r["failure_category"].replace("_", " ").title(), r["count"], _hours(r["downtime_minutes"]), r["pct"], r["cumulative_pct"]] for r in s["failure_pareto"]],
           [26, 8, 16, 12, 14])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"maintenance_report_{d_from:%Y%m%d}_{d_to:%Y%m%d}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})
