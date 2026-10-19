"""Excel import of the plant's existing machine list into the asset register
(plan §11 Phase 2). Used by POST /maintenance/assets/import and
GET /maintenance/assets/import-template in routes/assets.py.

Rules:
- dry_run validates every row and writes nothing.
- A real import is all-or-nothing: if any row has an error, nothing is saved.
- asset_code is the plant's machine number; a duplicate (in the file or in
  the register) is an error. Blank → EQ-NNNN, numbered like the create form.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.organization.models.branch import Branch
from app.modules.organization.models.department import Department
from app.modules.maintenance.models.asset import (
    MaintenanceAsset, MAINTENANCE_ASSET_CATEGORIES, MAINTENANCE_ASSET_CRITICALITIES,
)
from app.modules.maintenance.service import suggest_asset_code

MAX_ROWS = 2000

# (column, required, help text, example)
COLUMNS = [
    ("asset_code", False, "Plant's machine number, e.g. CNC-07. Leave blank to auto-number EQ-NNNN. Must be unique.", "CNC-07"),
    ("name", True, "Machine / equipment name.", "CNC Lathe — LMW LL20"),
    ("category", True, "One of the categories listed below (code or label).", "production_machine"),
    ("plant", True, "Plant name or plant code, exactly as in Organization › Plants.", ""),
    ("department", False, "Department name under that plant.", ""),
    ("location", False, "Where it stands, e.g. Bay 3, Line 2.", "Bay 3, Line 2"),
    ("make", False, "Manufacturer.", "LMW"),
    ("model", False, "Model.", "LL20T L5"),
    ("serial_number", False, "Manufacturer serial number.", "LL20-88123"),
    ("year_of_manufacture", False, "4-digit year.", 2018),
    ("criticality", False, "A = line stops if it fails, B = important, C = low impact. Blank = B.", "A"),
    ("commissioned_on", False, "Date the machine went into service (DD-MM-YYYY or an Excel date).", "15-06-2018"),
    ("workstation_code", False, "Production workstation code to link (blocks the workstation while the machine is down).", ""),
    ("remarks", False, "Anything else.", ""),
]
CATEGORY_LABELS = {
    "production_machine": "Production Machine", "material_handling": "Material Handling", "utility": "Utility",
    "tooling_fixture": "Tooling / Fixture", "instrument": "Instrument / Gauge", "facility": "Facility",
    "vehicle": "Vehicle", "other": "Other",
}
_HEADER_ALIASES = {
    "code": "asset_code", "asset_no": "asset_code", "machine_no": "asset_code", "machine_number": "asset_code",
    "asset_name": "name", "machine_name": "name", "branch": "plant", "plant_name": "plant", "plant_code": "plant",
    "location_text": "location", "serial_no": "serial_number", "year": "year_of_manufacture",
    "workstation": "workstation_code", "commissioned": "commissioned_on", "commissioning_date": "commissioned_on",
}


def _norm_header(h) -> str:
    s = str(h or "").strip().lower()
    s = re.sub(r"\(.*?\)|\*", "", s).strip()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    return _HEADER_ALIASES.get(s, s)


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(s)


def read_rows(filename: str, raw: bytes) -> list[tuple[int, dict]]:
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=422, detail=f"'{filename}' isn't an Excel .xlsx file — download the template, fill it in and upload it as .xlsx.")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=422, detail="Couldn't open the Excel file — it may be corrupt or password-protected. Save it again as .xlsx and retry.")
    ws = wb["Assets"] if "Assets" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm_header(h) for h in rows[0]]
    missing = [c for c, req, *_ in COLUMNS if req and c not in headers]
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"The first row must be the header row; column(s) {', '.join(missing)} are missing. Download the template and keep its header row.",
        )
    out = []
    for idx, r in enumerate(rows[1:], start=2):
        rec = {headers[i]: v for i, v in enumerate(r) if i < len(headers) and headers[i]}
        if all(_cell_str(v) == "" for v in rec.values()):
            continue
        out.append((idx, rec))
    return out


def _category(raw: str) -> str | None:
    key = raw.strip().lower()
    for code, label in CATEGORY_LABELS.items():
        if key in (code, label.lower(), code.replace("_", " "), label.lower().replace(" / ", "/")):
            return code
    return None


def validate_rows(db: Session, rows: list[tuple[int, dict]]) -> list[dict]:
    """One result per row: the cleaned values ready for MaintenanceAsset and
    a list of real-reason errors."""
    branches = db.query(Branch).all()
    plant_names = ", ".join(sorted(b.name for b in branches)) or "(no plants set up — add one in Organization › Plants)"
    by_plant = {}
    for b in branches:
        by_plant[b.name.strip().lower()] = b
        by_plant[b.code.strip().lower()] = b
    departments = db.query(Department).all()
    existing_codes = {c.upper() for (c,) in db.query(MaintenanceAsset.asset_code).all()}

    ws_by_code, linked_ws = {}, {}
    try:
        from app.modules.production.models.workstation import ProductionWorkstation
        ws_by_code = {w.code.strip().lower(): w for w in db.query(ProductionWorkstation).filter(ProductionWorkstation.is_deleted == False).all()}  # noqa: E712
        linked_ws = dict(db.query(MaintenanceAsset.workstation_id, MaintenanceAsset.asset_code).filter(
            MaintenanceAsset.workstation_id.isnot(None), MaintenanceAsset.is_deleted == False,  # noqa: E712
        ).all())
    except Exception:  # pragma: no cover — Production module absent
        pass

    seen_codes: dict[str, int] = {}
    seen_ws: dict[int, int] = {}
    results = []
    max_year = date.today().year + 1
    for row_no, r in rows:
        errs: list[str] = []
        code = _cell_str(r.get("asset_code")).upper()
        name = _cell_str(r.get("name"))
        if not name:
            errs.append("name is empty — every machine needs a name.")
        elif len(name) > 200:
            errs.append(f"name is {len(name)} characters — keep it under 200.")
        if code:
            if len(code) > 50:
                errs.append(f"asset_code '{code}' is longer than 50 characters.")
            if code in existing_codes:
                errs.append(f"asset_code '{code}' already exists in the asset register — use a different machine number or leave it blank to auto-number.")
            elif code in seen_codes:
                errs.append(f"asset_code '{code}' repeats row {seen_codes[code]} of this file.")
            else:
                seen_codes[code] = row_no

        cat_raw = _cell_str(r.get("category"))
        category = _category(cat_raw) if cat_raw else None
        if not cat_raw:
            errs.append(f"category is empty — use one of: {', '.join(MAINTENANCE_ASSET_CATEGORIES)}.")
        elif not category:
            errs.append(f"category '{cat_raw}' isn't valid — use one of: {', '.join(MAINTENANCE_ASSET_CATEGORIES)}.")

        plant_raw = _cell_str(r.get("plant"))
        plant = by_plant.get(plant_raw.lower()) if plant_raw else None
        if not plant_raw:
            errs.append(f"plant is empty — use one of: {plant_names}.")
        elif not plant:
            errs.append(f"plant '{plant_raw}' not found — use one of: {plant_names}.")

        dept_id = None
        dept_raw = _cell_str(r.get("department"))
        if dept_raw and plant:
            matches = [d for d in departments if d.name.strip().lower() == dept_raw.lower()]
            in_plant = [d for d in matches if d.branch_id == plant.id] or [d for d in matches if d.branch_id is None]
            if in_plant:
                dept_id = in_plant[0].id
            else:
                options = sorted({d.name for d in departments if d.branch_id in (plant.id, None)})
                errs.append(
                    f"department '{dept_raw}' isn't under plant {plant.name} — "
                    + (f"use one of: {', '.join(options)}, or leave it blank." if options else "that plant has no departments yet; leave it blank.")
                )

        crit = _cell_str(r.get("criticality")).upper()[:1] or "B"
        if crit not in MAINTENANCE_ASSET_CRITICALITIES:
            errs.append(f"criticality '{_cell_str(r.get('criticality'))}' isn't valid — use A (line stops if it fails), B or C.")

        year = None
        y_raw = _cell_str(r.get("year_of_manufacture"))
        if y_raw:
            if not y_raw.isdigit() or not (1900 <= int(y_raw) <= max_year):
                errs.append(f"year_of_manufacture '{y_raw}' isn't a 4-digit year between 1900 and {max_year}.")
            else:
                year = int(y_raw)

        commissioned = None
        try:
            commissioned = _parse_date(r.get("commissioned_on"))
            if commissioned and commissioned > date.today():
                errs.append(f"commissioned_on {commissioned:%d-%m-%Y} is in the future.")
        except ValueError as e:
            errs.append(f"commissioned_on '{e}' isn't a date — use DD-MM-YYYY (e.g. 15-06-2018) or an Excel date cell.")

        ws_id = None
        ws_raw = _cell_str(r.get("workstation_code"))
        if ws_raw:
            ws = ws_by_code.get(ws_raw.lower())
            if not ws:
                errs.append(f"workstation_code '{ws_raw}' not found in Production › Workstations — check the code or leave it blank.")
            elif ws.id in linked_ws:
                errs.append(f"workstation '{ws.code}' is already linked to asset {linked_ws[ws.id]}.")
            elif ws.id in seen_ws:
                errs.append(f"workstation '{ws.code}' is also on row {seen_ws[ws.id]} — one workstation can link to only one asset.")
            else:
                seen_ws[ws.id] = row_no
                ws_id = ws.id

        for key, limit in (("location", 200), ("make", 100), ("model", 100), ("serial_number", 100)):
            if len(_cell_str(r.get(key))) > limit:
                errs.append(f"{key} is longer than {limit} characters.")

        results.append({
            "row": row_no,
            "asset_code": code or None,
            "auto_code": not code,
            "name": name,
            "category": category or cat_raw,
            "plant": plant.name if plant else plant_raw,
            "criticality": crit,
            "errors": errs,
            "_values": {
                "name": name, "category": category, "branch_id": plant.id if plant else None, "department_id": dept_id,
                "location_text": _cell_str(r.get("location")) or None, "make": _cell_str(r.get("make")) or None,
                "model": _cell_str(r.get("model")) or None, "serial_number": _cell_str(r.get("serial_number")) or None,
                "year_of_manufacture": year, "criticality": crit, "commissioned_on": commissioned,
                "workstation_id": ws_id, "remarks": _cell_str(r.get("remarks")) or None,
            },
        })

    # Number the blank-code rows EQ-NNNN in file order, skipping codes the file uses.
    taken = existing_codes | set(seen_codes)
    nxt = int(suggest_asset_code(db)[3:])
    for res in results:
        if res["auto_code"]:
            while f"EQ-{nxt:04d}" in taken:
                nxt += 1
            res["asset_code"] = f"EQ-{nxt:04d}"
            taken.add(res["asset_code"])
            nxt += 1
    return results


def run_import(db: Session, filename: str, raw: bytes, dry_run: bool, user_id: int) -> dict:
    rows = read_rows(filename, raw)
    if not rows:
        raise HTTPException(status_code=422, detail="The sheet has no data rows — add your machines under the header row (one per row).")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=422, detail=f"The file has {len(rows)} rows — import at most {MAX_ROWS} at a time by splitting it.")
    results = validate_rows(db, rows)
    bad = [r for r in results if r["errors"]]
    flat = [f"Row {r['row']}: {e}" for r in bad for e in r["errors"]]
    if not dry_run:
        if bad:
            raise HTTPException(
                status_code=422,
                detail=f"Nothing was imported — {len(bad)} of {len(results)} row(s) have errors. Fix them and validate again. First problem: {flat[0]}",
            )
        for r in results:
            db.add(MaintenanceAsset(**r["_values"], asset_code=r["asset_code"], status="operational", created_by_id=user_id))
        try:
            db.commit()
        except Exception as e:  # pragma: no cover — e.g. a concurrent create took a code
            db.rollback()
            raise HTTPException(status_code=409, detail=f"Nothing was imported — the database rejected the batch ({type(e).__name__}). Someone may have added an asset with the same code meanwhile; validate again and retry.")
    return {
        "dry_run": dry_run,
        "total_rows": len(results),
        "valid_rows": len(results) - len(bad),
        "error_rows": len(bad),
        "created": 0 if dry_run else len(results),
        "rows": [{k: v for k, v in r.items() if k != "_values"} for r in results],
        "errors": flat,
    }


def build_template(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    branches = db.query(Branch).order_by(Branch.name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append([c for c, *_ in COLUMNS])
    for cell, (_, req, *_r) in zip(ws[1], COLUMNS):
        cell.font = Font(bold=True, color="B91C1C" if req else "000000")
        cell.fill = PatternFill("solid", fgColor="FDF1E6")
    example = [ex for *_x, ex in COLUMNS]
    example[3] = branches[0].name if branches else "Faridabad Plant"
    ws.append(example)
    for i, (c, *_r) in enumerate(COLUMNS):
        ws.column_dimensions[chr(ord("A") + i)].width = max(14, len(c) + 4)
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "A2"

    ins = wb.create_sheet("Instructions")
    ins.column_dimensions["A"].width = 24
    ins.column_dimensions["B"].width = 12
    ins.column_dimensions["C"].width = 100

    def head(text):
        ins.append([])
        ins.append([text])
        ins.cell(row=ins.max_row, column=1).font = Font(bold=True, size=12)

    ins.append(["Asset register import"])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append(["Fill one machine per row on the 'Assets' sheet. Delete the example row before importing. "
                "Validate first — nothing is saved unless every row is correct."])
    head("Columns")
    ins.append(["Column", "Required", "What to enter"])
    for cell in ins[ins.max_row]:
        cell.font = Font(bold=True)
    for c, req, help_text, _ex in COLUMNS:
        ins.append([c, "Yes" if req else "No", help_text])
        ins.cell(row=ins.max_row, column=3).alignment = Alignment(wrap_text=True)
    head("Allowed categories")
    for code in MAINTENANCE_ASSET_CATEGORIES:
        ins.append([code, "", CATEGORY_LABELS.get(code, code)])
    head("Criticality")
    ins.append(["A", "", "Line stops if it fails — breakdowns become Urgent"])
    ins.append(["B", "", "Important (default when blank)"])
    ins.append(["C", "", "Low impact"])
    head("Plants (use the name or the code)")
    if branches:
        for b in branches:
            ins.append([b.name, b.code, ""])
    else:
        ins.append(["(no plants yet — add them in Organization › Plants first)"])
    depts = db.query(Department).order_by(Department.name).all()
    if depts:
        names = {b.id: b.name for b in branches}
        head("Departments")
        for d in depts:
            ins.append([d.name, "", f"Plant: {names.get(d.branch_id, 'All plants')}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
