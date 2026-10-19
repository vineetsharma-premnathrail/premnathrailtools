"""Statutory registers per plant (Factories Act 1948 + State Factories Rules).

  form12          Register of adult workers (s.62, Form 12) — as of a date
  form25          Muster roll (Form 25) — one month, daily status codes,
                  totals and APPROVED overtime hours (s.59 OT register data)
  leave_register  Register of leave with wages (s.79/80, Form 14) — one year

Each register is built as a "spec" (quality/reports/render.py layout) for the
managed Word document — the document of record, edited in Word Online and
printed/signed from there — and as an .xlsx for the inspector's soft copy.

Who is on a plant's register: employees whose org assignment (time slices,
services/assignments.py) on the period's last day is that plant; employees
without history use their current plant. Exited before the period = left out.
"""
from __future__ import annotations

import io
from calendar import monthrange
from datetime import date
from decimal import Decimal

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.hr.models.attendance import HrAttendance
from app.modules.hr.models.employee_profile import HrEmployeeProfile
from app.modules.hr.models.leave import HrLeaveBalance, HrLeaveType
from app.modules.hr.models.masters import HrDesignation
from app.modules.hr.models.overtime import HrOvertime
from app.modules.hr.models.statutory_register import REGISTER_TYPES
from app.modules.hr.services import assignments as asg
from app.modules.hr.services.attendance import day_view
from app.modules.hr.services.calendar import CalendarCache, daterange, today_ist
from app.modules.main.models.user import User
from app.modules.organization.models.branch import Branch

CODES = {"present": "P", "absent": "A", "half_day": "HD", "on_leave": "L", "holiday": "H", "weekly_off": "WO",
         "on_duty": "OD", "work_from_home": "WFH", "unmarked": "-"}
CODE_LEGEND = "P present · A absent · HD half day · L leave · H holiday · WO weekly off · OD on duty · WFH work from home · - not marked"


def _fmt(d: date | None) -> str:
    return d.strftime("%d-%m-%Y") if d else ""


def _branch(db: Session, branch_id: int) -> Branch:
    b = db.get(Branch, branch_id)
    if b is None:
        raise HTTPException(status_code=404, detail=f"Plant #{branch_id} doesn't exist. Pick a plant from the list.")
    return b


def _company(db: Session, branch: Branch):
    from app.modules.organization.models.company import Company

    return db.get(Company, branch.company_id) if getattr(branch, "company_id", None) else db.query(Company).first()


def plant_employees(db: Session, branch_id: int, start: date, end: date) -> list[tuple[User, HrEmployeeProfile | None]]:
    slices = asg.as_of(db, end)
    rows = db.query(User, HrEmployeeProfile).outerjoin(HrEmployeeProfile, HrEmployeeProfile.user_id == User.id).all()
    out = []
    for u, p in rows:
        if p is not None and p.employment_status == "exited" and p.date_of_exit and p.date_of_exit < start:
            continue
        if p is None and not u.is_active:
            continue
        if u.date_of_joining and u.date_of_joining > end:
            continue
        s = slices.get(u.id)
        plant = s.branch_id if s is not None else ((p.branch_id if p and p.branch_id else None) or u.branch_id)
        if plant == branch_id:
            out.append((u, p))
    out.sort(key=lambda up: ((up[1].employee_code if up[1] and up[1].employee_code else "~"), up[0].name or ""))
    return out


def _age(dob: date | None, on: date) -> str:
    if not dob:
        return ""
    return str(on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day)))


def _facts(db: Session, branch: Branch, period: str) -> list[tuple[str, str | None]]:
    company = _company(db, branch)
    return [("Name of the factory", (getattr(company, "legal_name", None) or getattr(company, "name", None)) if company else None),
            ("Plant", branch.name), ("Period", period), ("Generated", _fmt(today_ist()))]


# ───────────────────────────── Form 12 ─────────────────────────────
def form12(db: Session, branch_id: int, as_of: date) -> dict:
    branch = _branch(db, branch_id)
    emps = plant_employees(db, branch_id, as_of, as_of)
    desig = {d.id: d.name for d in db.query(HrDesignation).all()}
    cache = CalendarCache(db)
    rows = []
    for i, (u, p) in enumerate(emps, start=1):
        ctx = cache.context(u, p, load_profile=False)
        shift = cache.shift_on(ctx, as_of)
        rows.append([
            str(i), u.name or u.email, (p.employee_code if p else None) or "", "",  # father's/husband's name: not held by the portal
            _fmt(p.date_of_birth) if p else "", _age(p.date_of_birth if p else None, as_of),
            ((p.gender or "").title() if p else ""), (p.permanent_address or p.current_address or "") if p else "",
            _fmt(u.date_of_joining), (desig.get(p.designation_id) if p and p.designation_id else None) or (u.designation or ""),
            shift.code if shift else "", "",  # number of relay — single-relay plant
        ])
    return {
        "title": "FORM 12 — REGISTER OF ADULT WORKERS", "number": f"F12-{branch.code or branch.id}-{as_of:%Y%m%d}",
        "intro": "Prescribed under section 62 of the Factories Act, 1948 and the Factories Rules of the State. "
                 "Father's / husband's name is not held in the portal — fill it in Word before signing.",
        "facts": _facts(db, branch, f"As on {_fmt(as_of)}"),
        "sections": [{"heading": f"Adult workers ({len(rows)})",
                      "columns": ["Sl", "Name", "Emp code", "Father's / husband's name", "Date of birth", "Age", "Sex",
                                  "Residential address", "Date of joining", "Nature of work", "Group / shift", "Relay no."],
                      "rows": rows, "widths": [0.3, 1.0, 0.55, 0.8, 0.6, 0.35, 0.4, 1.0, 0.6, 0.7, 0.45, 0.42],
                      "empty": "No employees are on this plant's rolls on that date."}],
        "statement": "Certified that the particulars above are correct as per the records of the factory.",
        "signatures": [("Prepared by (HR)", ""), ("Manager / Occupier", "")],
    }


# ───────────────────────────── Form 25 ─────────────────────────────
def muster_data(db: Session, branch_id: int, year: int, month: int) -> tuple[Branch, list[int], list[dict]]:
    if not (1 <= month <= 12):
        raise HTTPException(status_code=400, detail=f"Month must be 1–12, got {month}. Pick a month from the list.")
    branch = _branch(db, branch_id)
    start, end = date(year, month, 1), date(year, month, monthrange(year, month)[1])
    emps = plant_employees(db, branch_id, start, end)
    ids = [u.id for u, _ in emps]
    att = {(r.user_id, r.attendance_date): r for r in db.query(HrAttendance).filter(
        HrAttendance.user_id.in_(ids or [-1]), HrAttendance.attendance_date >= start, HrAttendance.attendance_date <= end).all()}
    ot: dict[int, int] = {}
    for o in db.query(HrOvertime).filter(HrOvertime.user_id.in_(ids or [-1]), HrOvertime.status == "approved",
                                         HrOvertime.work_date >= start, HrOvertime.work_date <= end).all():
        ot[o.user_id] = ot.get(o.user_id, 0) + (o.approved_minutes or 0)
    cache = CalendarCache(db)
    cache.prefetch_roster(ids, start, end)
    today = today_ist()
    out = []
    for u, p in emps:
        ctx = cache.context(u, p, load_profile=False)
        codes, totals = [], {"P": 0.0, "A": 0, "L": 0, "H": 0, "WO": 0}
        for d in daterange(start, end):
            v = day_view(cache, ctx, d, att.get((u.id, d)), today)
            c = CODES.get(v.status or "", "") if v.status else ""
            codes.append(c)
            if c in ("P", "OD", "WFH"):
                totals["P"] += 1
            elif c == "HD":
                totals["P"] += 0.5
            elif c in totals:
                totals[c] += 1
        out.append({"user_id": u.id, "name": u.name or u.email, "code": (p.employee_code if p else None) or "",
                    "codes": codes, "present": totals["P"], "absent": totals["A"], "leave": totals["L"],
                    "holidays": totals["H"] + totals["WO"], "ot_hours": round(ot.get(u.id, 0) / 60, 2)})
    return branch, list(range(1, end.day + 1)), out


def form25(db: Session, branch_id: int, year: int, month: int) -> dict:
    branch, days, data = muster_data(db, branch_id, year, month)
    label = date(year, month, 1).strftime("%B %Y")
    halves = [(1, 16), (17, len(days))]
    sections = []
    for a, b in halves:
        n = b - a + 1
        w = round((7.17 - 1.4) / max(n, 1), 3)
        sections.append({
            "heading": f"Attendance, days {a}–{b}",
            "columns": ["Employee"] + [str(d) for d in range(a, b + 1)],
            "rows": [[f"{r['code']} {r['name']}".strip()] + r["codes"][a - 1:b] for r in data],
            "widths": [1.4] + [w] * n, "empty": "No employees on this plant's rolls in the month.",
        })
    sections.append({
        "heading": "Totals for the month",
        "columns": ["Employee", "Days worked", "Absent", "Leave", "Holidays / weekly offs", "OT hours (approved)"],
        "rows": [[f"{r['code']} {r['name']}".strip(), f"{r['present']:g}", str(r["absent"]), str(r["leave"]),
                  str(r["holidays"]), f"{r['ot_hours']:.2f}"] for r in data],
        "widths": [2.2, 0.9, 0.8, 0.8, 1.3, 1.17],
    })
    return {
        "title": "FORM 25 — MUSTER ROLL", "number": f"F25-{branch.code or branch.id}-{year}{month:02d}",
        "intro": "Muster roll under the Factories Rules of the State. " + CODE_LEGEND
                 + ". Overtime = hours approved for payment at twice the ordinary rate (s.59).",
        "facts": _facts(db, branch, label), "sections": sections,
        "signatures": [("Prepared by (HR)", ""), ("Manager", "")],
    }


# ───────────────────────────── Leave register ─────────────────────────────
def leave_data(db: Session, branch_id: int, year: int) -> tuple[Branch, list[HrLeaveType], list[dict]]:
    branch = _branch(db, branch_id)
    start, end = date(year, 1, 1), date(year, 12, 31)
    emps = plant_employees(db, branch_id, start, end)
    ids = [u.id for u, _ in emps]
    types = db.query(HrLeaveType).order_by(HrLeaveType.code).all()
    bal = {(b.user_id, b.leave_type_id): b for b in db.query(HrLeaveBalance).filter(
        HrLeaveBalance.user_id.in_(ids or [-1]), HrLeaveBalance.year == year).all()}
    worked: dict[int, float] = {}
    for r in db.query(HrAttendance).filter(HrAttendance.user_id.in_(ids or [-1]), HrAttendance.attendance_date >= start,
                                           HrAttendance.attendance_date <= end).all():
        if r.status in ("present", "on_duty", "work_from_home"):
            worked[r.user_id] = worked.get(r.user_id, 0) + 1
        elif r.status == "half_day":
            worked[r.user_id] = worked.get(r.user_id, 0) + 0.5
    out = []
    for u, p in emps:
        per = {}
        for t in types:
            b = bal.get((u.id, t.id))
            if b is None:
                continue
            credit = Decimal(b.opening or 0) + Decimal(b.allotted or 0) + Decimal(b.adjusted or 0)
            per[t.code] = {"opening": float(b.opening or 0), "earned": float((b.allotted or 0) + (b.adjusted or 0)),
                           "availed": float(b.used or 0), "balance": float(credit - Decimal(b.used or 0))}
        out.append({"user_id": u.id, "name": u.name or u.email, "code": (p.employee_code if p else None) or "",
                    "days_worked": worked.get(u.id, 0), "leave": per})
    return branch, types, out


def leave_register(db: Session, branch_id: int, year: int) -> dict:
    branch, types, data = leave_data(db, branch_id, year)
    rows = []
    for r in data:
        for code, v in r["leave"].items():
            rows.append([f"{r['code']} {r['name']}".strip(), f"{r['days_worked']:g}", code, f"{v['opening']:g}",
                         f"{v['earned']:g}", f"{v['availed']:g}", f"{v['balance']:g}"])
        if not r["leave"]:
            rows.append([f"{r['code']} {r['name']}".strip(), f"{r['days_worked']:g}", "—", "", "", "", ""])
    return {
        "title": "FORM 14 — REGISTER OF LEAVE WITH WAGES", "number": f"F14-{branch.code or branch.id}-{year}",
        "intro": "Sections 79 and 80 of the Factories Act, 1948: days worked in the calendar year, leave earned, "
                 "availed and balance per leave type. Wages for leave are paid by payroll (ADP).",
        "facts": _facts(db, branch, f"Calendar year {year}"),
        "sections": [{"heading": "Leave account", "columns": ["Employee", "Days worked", "Leave type", "Opening",
                                                               "Earned / credited", "Availed", "Balance"],
                      "rows": rows, "widths": [2.1, 0.8, 0.8, 0.75, 1.0, 0.82, 0.9],
                      "empty": "No employees on this plant's rolls in the year."}],
        "signatures": [("Prepared by (HR)", ""), ("Manager", "")],
    }


def build_spec(db: Session, register_type: str, branch_id: int, year: int, month: int) -> dict:
    if register_type == "form12":
        on = date(year, month, monthrange(year, month)[1]) if month else date(year, 12, 31)
        return form12(db, branch_id, min(on, today_ist()))
    if register_type == "form25":
        return form25(db, branch_id, year, month)
    if register_type == "leave_register":
        return leave_register(db, branch_id, year)
    raise HTTPException(status_code=400, detail=f"Unknown register '{register_type}'. Use one of: {', '.join(REGISTER_TYPES)}.")


def spec_to_xlsx(spec: dict) -> bytes:
    """Excel soft copy: one sheet per section with the facts on top."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for i, sec in enumerate(spec.get("sections") or [], start=1):
        ws = wb.create_sheet(title=(sec["heading"][:28] or f"Sheet{i}").replace("/", "-").replace(":", ""))
        ws.append([spec["title"]])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([spec["number"]])
        for label, value in spec.get("facts") or []:
            ws.append([label, value or ""])
        ws.append([])
        ws.append(sec["columns"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for r in sec.get("rows") or []:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
