"""Biometric punch import and day evaluation (SAP PT_ERL00 / time events).

Import: rows of (employee_code | email, timestamp | date + time, direction,
device). A punch already stored (same employee, same minute-exact
timestamp) is a duplicate and skipped.

derive_day(user, d) builds hr_attendance from the punches inside the day's
shift window (shift start − 4 h … shift end + 6 h; a night shift's window
runs into the next morning; no shift = the calendar day):
- check-in = first IN (or first punch when the device gives no direction),
  check-out = last OUT (or last punch);
- status from worked minutes: >= full-day minutes → present, >= half-day
  minutes → half_day, else absent (HR settings, Time evaluation);
- a repeat within the duplicate window is ignored and reported;
- one punch only → check-in (or check-out) alone, reported as a missing
  punch (time evaluation then shows it as "missing check-out");
- it NEVER overwrites leave, regularization, manual or self check-in rows,
  and it respects the payroll period lock — those days are reported, not changed.
"""
from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.hr.models.attendance import HrAttendance
from app.modules.hr.models.employee_profile import HrEmployeeProfile
from app.modules.hr.models.punch import HrPunch
from app.modules.hr.services import period_lock
from app.modules.hr.services.attendance import parse_date_cell, parse_time_cell, read_csv
from app.modules.hr.services.calendar import IST, CalendarCache
from app.modules.hr.services.time_evaluation import TimeRules, as_ist, worked_minutes
from app.modules.main.models.user import User

MAX_ROWS = 100_000
PROTECTED_SOURCES = ("leave", "regularization", "manual", "self")
PUNCH_TEMPLATE = ("employee_code,email,timestamp,direction,device_id\n"
                  "E001,,2026-11-02 08:55,in,GATE-1\nE001,,2026-11-02 17:41,out,GATE-1\n"
                  ",name@premnathrail.com,02-11-2026 22:04,in,PLANT2\n")
_DT_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M:%S",
               "%d/%m/%Y %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%d.%m.%Y %H:%M")
DIRECTION_ALIASES = {"in": "in", "i": "in", "entry": "in", "0": "in", "checkin": "in", "check-in": "in",
                     "out": "out", "o": "out", "exit": "out", "1": "out", "checkout": "out", "check-out": "out"}


@dataclass
class PunchRules:
    full_day: int = 360
    half_day: int = 240
    dup_window: int = 2

    @classmethod
    def load(cls, db: Session) -> "PunchRules":
        from app.modules.hr.settings import get_setting

        def val(key, default):
            try:
                v = get_setting(db, key)
                return int(v) if v is not None else default
            except Exception:  # noqa: BLE001 — settings unavailable: defaults
                return default

        d = cls()
        return cls(full_day=val("punch_full_day_minutes", d.full_day), half_day=val("punch_half_day_minutes", d.half_day),
                   dup_window=val("punch_duplicate_window_minutes", d.dup_window))


def parse_timestamp(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=IST) if value.tzinfo is None else value.astimezone(IST)
    value = str(value or "").strip()
    for fmt in _DT_FORMATS:
        try:
            return datetime.strptime(value, fmt).replace(tzinfo=IST)
        except ValueError:
            continue
    return None


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        _, rows = read_csv(content)
        return rows
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 — real reason
            raise HTTPException(status_code=400, detail=f"Couldn't read '{filename}' as an Excel file ({e}). Save it again as .xlsx or .csv.")
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return []
        keys = [str(h or "").strip().lower() for h in header]
        out = []
        for row in it:
            if row is None or all(c in (None, "") for c in row):
                continue
            out.append({k: (v if isinstance(v, (datetime, date, time)) else ("" if v is None else str(v).strip()))
                        for k, v in zip(keys, row)})
        return out
    raise HTTPException(status_code=400, detail=f"'{filename}' isn't a CSV or Excel (.xlsx) file. Export the punch log from the biometric software as CSV or XLSX.")


def _row_timestamp(r: dict) -> datetime | None:
    ts = r.get("timestamp") or r.get("punch_time") or r.get("datetime")
    if ts:
        return parse_timestamp(ts)
    d, t = r.get("date"), r.get("time")
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, str):
        d = parse_date_cell(d)
    if isinstance(t, str):
        t = parse_time_cell(t)
        if t == "invalid":
            return None
    if isinstance(d, date) and isinstance(t, time):
        return datetime.combine(d, t.replace(tzinfo=None)).replace(tzinfo=IST)
    return None


@dataclass
class ImportResult:
    batch: str
    total_rows: int = 0
    inserted: int = 0
    duplicates: int = 0
    errors: list[dict] = field(default_factory=list)
    days: set = field(default_factory=set)


def import_punches(db: Session, rows: list[dict], actor: User, *, dry_run: bool, in_scope: set[int] | None) -> ImportResult:
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"The file has {len(rows):,} rows; the limit is {MAX_ROWS:,}. Split it by date range.")
    res = ImportResult(batch=f"PB-{datetime.now(IST):%Y%m%d%H%M%S}-{uuid.uuid4().hex[:4]}", total_rows=len(rows))
    by_email = {u.email.lower(): u for u in db.query(User).all() if u.email}
    users = {u.id: u for u in by_email.values()}
    by_code = {p.employee_code.strip().lower(): p.user_id
               for p in db.query(HrEmployeeProfile).filter(HrEmployeeProfile.employee_code.isnot(None)).all()}
    seen: set[tuple[int, datetime]] = set()
    for idx, r in enumerate(rows, start=2):
        code = str(r.get("employee_code") or r.get("emp_code") or r.get("code") or "").strip()
        email = str(r.get("email") or "").strip().lower()
        ident = code or email

        def err(msg):
            res.errors.append({"row": idx, "identifier": ident, "message": msg})

        user = None
        if code:
            uid = by_code.get(code.lower())
            user = users.get(uid) if uid else None
            if user is None:
                err(f"No employee has the code '{code}'. Check it against HR > Employees (the biometric enrolment number must equal the employee code).")
                continue
        elif email:
            user = by_email.get(email)
            if user is None:
                err(f"No portal user has the email '{email}'.")
                continue
        else:
            err("employee_code and email are both empty. Fill one of them.")
            continue
        if in_scope is not None and user.id not in in_scope:
            err(f"{user.name} is outside your HR data access. Remove the row, or ask an admin to widen it in Organization › Role & Permissions.")
            continue
        ts = _row_timestamp(r)
        if ts is None:
            err(f"Timestamp '{r.get('timestamp') or ((r.get('date') or '') and str(r.get('date')) + ' ' + str(r.get('time') or ''))}' isn't valid. Use YYYY-MM-DD HH:MM or DD-MM-YYYY HH:MM (24-hour).")
            continue
        raw_dir = str(r.get("direction") or r.get("in_out") or "").strip().lower()
        direction = DIRECTION_ALIASES.get(raw_dir) if raw_dir else None
        if raw_dir and direction is None:
            err(f"Direction '{raw_dir}' isn't recognised. Use in / out (or leave it blank when the device doesn't record it).")
            continue
        ts = ts.replace(second=0, microsecond=0)
        key = (user.id, ts)
        if key in seen or db.query(HrPunch.id).filter(HrPunch.user_id == user.id, HrPunch.punched_at == ts).first():
            res.duplicates += 1
            continue
        seen.add(key)
        res.inserted += 1
        res.days.add((user.id, ts.date()))
        res.days.add((user.id, ts.date() - timedelta(days=1)))  # may close a night shift that began the day before
        if not dry_run:
            db.add(HrPunch(user_id=user.id, punched_at=ts, direction=direction,
                           device_id=str(r.get("device_id") or r.get("device") or "").strip() or None,
                           source="biometric", import_batch=res.batch, created_by_id=actor.id))
    if not dry_run:
        db.flush()
    return res


def _window(d: date, shift) -> tuple[datetime, datetime]:
    if shift is None:
        start = datetime.combine(d, time(0, 0)).replace(tzinfo=IST)
        return start, start + timedelta(days=1)
    s = datetime.combine(d, shift.start_time.replace(tzinfo=None)).replace(tzinfo=IST)
    e = datetime.combine(d, shift.end_time.replace(tzinfo=None)).replace(tzinfo=IST)
    if e <= s:
        e += timedelta(days=1)
    return s - timedelta(hours=4), e + timedelta(hours=6)


@dataclass
class DayResult:
    user_id: int
    date: date
    action: str            # created / updated / skipped / none
    status: str | None = None
    exceptions: list[str] = field(default_factory=list)


def derive_day(db: Session, cache: CalendarCache, user: User, d: date, rules: PunchRules, trules: TimeRules,
               actor: User) -> DayResult:
    ctx = cache.context(user)
    shift = cache.shift_on(ctx, d)
    w_start, w_end = _window(d, shift)
    punches = (db.query(HrPunch).filter(HrPunch.user_id == user.id, HrPunch.punched_at >= w_start, HrPunch.punched_at < w_end)
               .order_by(HrPunch.punched_at.asc()).all())
    res = DayResult(user_id=user.id, date=d, action="none")
    if not punches:
        return res
    # Debounce double swipes.
    kept: list[HrPunch] = []
    for p in punches:
        if kept and (as_ist(p.punched_at) - as_ist(kept[-1].punched_at)) < timedelta(minutes=max(rules.dup_window, 0) or 0) \
                and (p.direction == kept[-1].direction or p.direction is None or kept[-1].direction is None):
            res.exceptions.append(f"Duplicate punch at {as_ist(p.punched_at):%H:%M} ignored (within {rules.dup_window} min of the previous).")
            continue
        kept.append(p)
    ins = [p for p in kept if p.direction == "in"]
    outs = [p for p in kept if p.direction == "out"]
    undirected = [p for p in kept if p.direction is None]
    ci = (ins[0] if ins else (undirected[0] if undirected else None))
    co = (outs[-1] if outs else (undirected[-1] if len(undirected) > 1 else None))
    ci_t = as_ist(ci.punched_at) if ci else None
    co_t = as_ist(co.punched_at) if co else None
    if ci_t and co_t and co_t <= ci_t:
        res.exceptions.append(f"The last OUT ({co_t:%H:%M}) is before the first IN ({ci_t:%H:%M}) — check the device direction setting.")
        co_t = None
    if ci_t and not co_t:
        res.exceptions.append("Missing OUT punch — only an IN was recorded. Regularize the day or add the punch.")
    if co_t and not ci_t:
        res.exceptions.append("Missing IN punch — only an OUT was recorded. Regularize the day or add the punch.")

    row = db.query(HrAttendance).filter(HrAttendance.user_id == user.id, HrAttendance.attendance_date == d).first()
    if row is not None and row.source in PROTECTED_SOURCES:
        res.action = "skipped"
        res.exceptions.append(f"Kept the existing {row.source} entry ({row.status}); punches don't overwrite it.")
        return res
    if period_lock.is_locked(db, user.id, d):
        res.action = "skipped"
        res.exceptions.append("The month is locked for payroll for this plant — the day wasn't changed.")
        return res
    kind, _ = cache.day_kind(ctx, d)
    worked = worked_minutes(ci_t, co_t, trules) if ci_t and co_t else None
    if worked is None:
        status = "present"  # one punch: present, flagged; time evaluation reports the missing punch
    elif kind != "working":
        status = "present"
    elif worked >= rules.full_day:
        status = "present"
    elif worked >= rules.half_day:
        status = "half_day"
    else:
        status = "absent"
        res.exceptions.append(f"Only {worked} minutes between the first IN and last OUT — marked Absent (below the half-day {rules.half_day} min).")
    created = row is None
    if created:
        row = HrAttendance(user_id=user.id, attendance_date=d)
        db.add(row)
    row.status = status
    row.source = "import"
    row.check_in = ci_t
    row.check_out = co_t
    row.shift_id = shift.id if shift else None
    row.marked_by_id = actor.id
    row.remarks = ("Biometric punches" + (": " + " ".join(res.exceptions) if res.exceptions else ""))[:1000]
    res.action = "created" if created else "updated"
    res.status = status
    return res


def derive_days(db: Session, pairs, actor: User) -> list[DayResult]:
    cache = CalendarCache(db)
    rules, trules = PunchRules.load(db), TimeRules.load(db)
    users: dict[int, User] = {}
    out = []
    for uid, d in sorted(pairs):
        u = users.get(uid) or db.get(User, uid)
        users[uid] = u
        if u is None:
            continue
        r = derive_day(db, cache, u, d, rules, trules, actor)
        if r.action != "none":
            out.append(r)
    db.flush()
    return out
