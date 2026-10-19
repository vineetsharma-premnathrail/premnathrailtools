"""DOCX builder for the RRV Production Handover Certificate — same content
as rrv_handover_pdf.py (vehicle identity, customer & order, handover facts,
tests, stage sign-off, signature blocks), on the company letterhead, as an
editable Word document of record (see main/services/managed_documents.py).
"""

from __future__ import annotations

import io
from typing import Any, Mapping

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt

from app.utils.docx_letterhead import (
    ACCENT, MUTED, USABLE_IN, tight, write, shade, cell_margins, set_widths, new_letterhead_document,
)

SHADE_HEX = "F3F4F6"


def _dash(value: Any) -> str:
    return "—" if value in (None, "") else str(value)


def _grid(doc: Any, rows: list[list[str]], widths: list[float], header: bool = False, bold_cols: tuple[int, ...] = ()) -> Any:
    table = doc.add_table(rows=0, cols=len(widths))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(table)
    for r, values in enumerate(rows):
        cells = table.add_row().cells
        for c, (cell, value) in enumerate(zip(cells, values)):
            is_head = header and r == 0
            write(cell, value, size=8 if (is_head or c in bold_cols) else 8.5, bold=is_head or c in bold_cols, first=True)
            if is_head or c in bold_cols:
                shade(cell, SHADE_HEX)
    set_widths(table, widths)
    return table


def build_rrv_handover_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    """Same ctx keys as rrv_handover_pdf.build_rrv_handover_pdf (printed_at /
    printed_by are not used — a document of record shouldn't change each
    time it is opened)."""
    doc = new_letterhead_document()

    title = doc.add_table(rows=1, cols=2)
    cell_margins(title, left=0, right=0, top=0, bottom=40)
    left, right = title.rows[0].cells
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    write(left, "PRODUCTION HANDOVER CERTIFICATE", size=14, bold=True, color=ACCENT, first=True)
    write(right, ctx["build_number"], size=12, bold=True, align="right", first=True)
    set_widths(title, [USABLE_IN - 2.0, 2.0])

    intro = write(
        doc,
        "This certifies that the Rail-cum-Road Vehicle described below has been manufactured, inspected and tested, "
        f"and was handed over to {ctx.get('handed_to') or 'the customer'}"
        f"{' of ' + ctx['handed_to_org'] if ctx.get('handed_to_org') else ''} on {ctx['handover_date']}.",
        size=9.5, after=10,
    )
    tight(intro, before=8, after=10)

    facts = [
        ("RRV model", ctx["rrv_model"]), ("Customer", ctx.get("customer")),
        ("Vehicle serial", ctx.get("serial")), ("Customer PO", ctx.get("po")),
        ("Chassis no.", ctx.get("chassis")), ("Engine no.", ctx.get("engine")),
        ("Year of manufacture", ctx.get("year")), ("Warranty", ctx.get("warranty")),
        ("Handover date", ctx["handover_date"]), ("Commissioning date", ctx.get("commissioning_date")),
        ("Handover location", ctx.get("location")), ("Customer acceptance ref.", ctx.get("acceptance_ref")),
        ("Final inspection", ctx.get("final_inspection")), ("Handed over by", ctx.get("handed_over_by")),
    ]
    half = USABLE_IN / 2
    rows = [[facts[i][0], _dash(facts[i][1]), facts[i + 1][0], _dash(facts[i + 1][1])] for i in range(0, len(facts), 2)]
    _grid(doc, rows, [1.3, half - 1.3, 1.3, half - 1.3], bold_cols=(0, 2))

    tight(write(doc, "Vehicle tests", size=10, bold=True), before=12, after=4)
    tests = [["Test", "Result"]] + [[label, (res or "pending").upper()] for label, res in ctx.get("tests") or []]
    _grid(doc, tests, [USABLE_IN - 1.4, 1.4], header=True)

    tight(write(doc, "Production stage sign-off", size=10, bold=True), before=12, after=4)
    stages = [["Stage", "Status", "Signed off by", "Date"]] + [
        [label, _dash(status), _dash(by), _dash(when)] for label, status, by, when in ctx.get("stages") or []
    ]
    _grid(doc, stages, [USABLE_IN - 4.1, 1.2, 1.9, 1.0], header=True)

    if ctx.get("remarks"):
        tight(write(doc, "Remarks", size=9, bold=True), before=10, after=2)
        write(doc, ctx["remarks"])

    tight(doc.add_paragraph(), before=14)
    sign = doc.add_table(rows=3, cols=2)
    sign.style = "Table Grid"
    cell_margins(sign)
    labels = [("For Premnath Rail", "Received by (customer)"), ("Signature & stamp", "Signature & stamp"),
              (f"Name: {ctx.get('handed_over_by') or ''}", f"Name: {ctx.get('handed_to') or ''}")]
    for r, (a, b) in enumerate(labels):
        for cell, text in zip(sign.rows[r].cells, (a, b)):
            write(cell, text, size=9 if r == 0 else 8.5, bold=r == 0, color=MUTED if r == 1 else None, first=True)
            if r == 1:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    sign.rows[1].height = Pt(62)
    set_widths(sign, [half, half])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
