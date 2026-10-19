"""Coded column-mapping templates for bank statement files (CSV / XLSX).

Each bank exports its statement with its own headers, date format and
preamble rows. A template lists, per field, the header names that bank uses
(matched case- and punctuation-insensitively); the parser finds the header
row within the first 40 rows, so bank preambles (account name, period...)
are skipped. Templates are code, not admin metadata — a new bank layout is a
new entry here.

Fields: value_date (required), txn_date, description, reference,
cheque_number, debit, credit, balance; or `amount` + `drcr` for banks that
print one amount column with a Dr/Cr marker."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

_DATE_FORMATS = ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d %b %Y",
                 "%d-%b-%Y", "%d-%b-%y", "%d %b %y", "%d %B %Y", "%d/%b/%Y", "%Y/%m/%d")


@dataclass(frozen=True)
class StatementFormat:
    key: str
    label: str
    columns: dict[str, tuple[str, ...]]
    date_formats: tuple[str, ...] = _DATE_FORMATS
    notes: str = ""
    required: tuple[str, ...] = field(default=("value_date",))


_GENERIC = {
    "value_date": ("value date", "value dt", "valuedate", "date", "txn date", "transaction date", "tran date", "posting date"),
    "txn_date": ("txn date", "transaction date", "tran date", "posting date", "date"),
    "description": ("description", "narration", "particulars", "transaction remarks", "remarks", "details", "transaction details"),
    "reference": ("reference", "ref no", "reference no", "utr", "utr no", "ref no cheque no", "chq ref no", "transaction id", "ref"),
    "cheque_number": ("cheque no", "cheque number", "chq no", "chqno", "instrument no", "cheque"),
    "debit": ("debit", "withdrawal", "withdrawals", "withdrawal amt", "withdrawal amount", "dr", "debit amount", "withdrawal amount inr"),
    "credit": ("credit", "deposit", "deposits", "deposit amt", "deposit amount", "cr", "credit amount", "deposit amount inr"),
    "balance": ("balance", "closing balance", "running balance", "bal", "balance inr", "available balance"),
    "amount": ("amount", "transaction amount", "amount inr"),
    "drcr": ("dr cr", "drcr", "type", "cr dr", "debit credit"),
}

FORMATS: dict[str, StatementFormat] = {f.key: f for f in (
    StatementFormat("generic", "Generic (Date, Description, Reference, Debit, Credit, Balance)", _GENERIC,
                    notes="Any CSV/XLSX with a header row; one Amount column + a Dr/Cr column also works."),
    StatementFormat("hdfc", "HDFC Bank — NetBanking statement", {
        "txn_date": ("date",), "description": ("narration",), "reference": ("chq ref no", "chq./ref.no."),
        "value_date": ("value dt", "value date"), "debit": ("withdrawal amt", "withdrawal amt."),
        "credit": ("deposit amt", "deposit amt."), "balance": ("closing balance",),
    }, date_formats=("%d/%m/%y", "%d/%m/%Y"), notes="Chq./Ref.No. is matched as both the UTR and the cheque number."),
    StatementFormat("sbi", "State Bank of India — account statement", {
        "txn_date": ("txn date",), "value_date": ("value date",), "description": ("description",),
        "reference": ("ref no cheque no", "ref no./cheque no."), "debit": ("debit",), "credit": ("credit",), "balance": ("balance",),
    }, date_formats=("%d %b %Y", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y")),
    StatementFormat("icici", "ICICI Bank — detailed statement", {
        "value_date": ("value date",), "txn_date": ("transaction date",), "cheque_number": ("cheque number",),
        "description": ("transaction remarks",), "debit": ("withdrawal amount inr", "withdrawal amount"),
        "credit": ("deposit amount inr", "deposit amount"), "balance": ("balance inr", "balance"),
    }, date_formats=("%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y")),
    StatementFormat("axis", "Axis Bank — account statement", {
        "txn_date": ("tran date",), "value_date": ("value date", "tran date"), "cheque_number": ("chqno", "chq no"),
        "description": ("particulars",), "debit": ("dr", "debit"), "credit": ("cr", "credit"), "balance": ("bal", "balance"),
    }, date_formats=("%d-%m-%Y", "%d/%m/%Y")),
    StatementFormat("canara", "Canara Bank / PNB / BoB (Date, Particulars, Cheque, Debit, Credit, Balance)", {
        "txn_date": ("txn date", "transaction date", "date"), "value_date": ("value date", "txn date", "date"),
        "description": ("particulars", "description", "narration"), "cheque_number": ("cheque no", "chq no", "instrument id"),
        "reference": ("ref no", "reference"), "debit": ("debit", "withdrawal"), "credit": ("credit", "deposit"),
        "balance": ("balance",),
    }),
)}


def format_options() -> list[dict]:
    return [{"value": f.key, "label": f.label, "notes": f.notes} for f in FORMATS.values()]


def _norm_header(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def norm_ref(s) -> str:
    """Upper-case alphanumerics only — for comparing UTRs / cheque numbers."""
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def norm_cheque(s) -> str:
    return norm_ref(s).lstrip("0")


@dataclass
class ParsedLine:
    row_number: int
    value_date: date
    txn_date: date | None
    description: str | None
    reference: str | None
    cheque_number: str | None
    debit: float
    credit: float
    balance: float | None

    def dedup_key(self, bank_account_id: int) -> str:
        raw = "|".join([str(bank_account_id), self.value_date.isoformat(), f"{self.debit:.2f}", f"{self.credit:.2f}",
                        norm_ref(self.reference), norm_ref(self.cheque_number), norm_ref(self.description)[:80],
                        "" if self.balance is None else f"{self.balance:.2f}"])
        return hashlib.sha256(raw.encode()).hexdigest()


def _read_rows(content: bytes, file_name: str) -> list[list]:
    name = (file_name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ValueError(f"'{file_name}' couldn't be opened as an Excel workbook — save it again as .xlsx (or export the statement as CSV) and upload that.")
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        raise ValueError(f"'{file_name}' is an old .xls workbook — open it in Excel and save it as .xlsx or CSV, then upload again.")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise ValueError(f"'{file_name}' isn't readable text — export the statement as CSV or XLSX.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _parse_date(v, formats: tuple[str, ...]) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    s = s.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2} \d", s) else s
    for fmt in formats + _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(s)


def _parse_amount(v) -> tuple[float | None, str | None]:
    """'1,23,456.50 Cr' -> (123456.5, 'cr')."""
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip()
    if not s or s in ("-", "--"):
        return None, None
    marker = None
    low = s.lower()
    if low.endswith("cr"):
        marker, s = "cr", s[:-2]
    elif low.endswith("dr"):
        marker, s = "dr", s[:-2]
    neg = s.strip().startswith("(") and s.strip().endswith(")")
    s = re.sub(r"[^0-9.\-]", "", s)
    if s in ("", "-", "."):
        return None, marker
    val = float(s)
    return (-val if neg else val), marker


def _find_header(rows: list[list], fmt: StatementFormat) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows[:40]):
        headers = [_norm_header(c) for c in row]
        if not any(headers):
            continue
        mapping: dict[str, int] = {}
        used: set[int] = set()
        for fld, aliases in fmt.columns.items():
            for alias in aliases:
                a = _norm_header(alias)
                pos = next((i for i, h in enumerate(headers) if h == a and i not in used), None)
                if pos is not None:
                    mapping[fld] = pos
                    if fld not in ("txn_date",):
                        used.add(pos)
                    break
        has_amounts = ("debit" in mapping or "credit" in mapping) or ("amount" in mapping)
        if "value_date" not in mapping and "txn_date" in mapping:
            mapping["value_date"] = mapping["txn_date"]
        if "value_date" in mapping and has_amounts:
            return idx, mapping
    expected = ", ".join(fmt.columns.get("value_date", ())[:3])
    raise ValueError(
        f"No header row matching the '{fmt.label}' layout was found in the first 40 rows. The file needs a date column "
        f"(e.g. {expected}) and Debit/Credit (or Amount) columns. Check you picked the right bank format, or use 'Generic'.")


def parse_statement(content: bytes, file_name: str, format_key: str) -> list[ParsedLine]:
    fmt = FORMATS.get(format_key)
    if fmt is None:
        raise ValueError(f"Unknown statement format '{format_key}'. Pick one of: {', '.join(FORMATS)}.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    rows = _read_rows(content, file_name)
    header_idx, cols = _find_header(rows, fmt)

    def cell(row, fld):
        i = cols.get(fld)
        return row[i] if i is not None and i < len(row) else None

    out: list[ParsedLine] = []
    errors: list[str] = []
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row or not any(str(c).strip() for c in row if c is not None):
            continue
        raw_date = cell(row, "value_date")
        try:
            vdate = _parse_date(raw_date, fmt.date_formats)
        except ValueError as bad:
            # Footer rows ("Closing balance", "*** End of statement") carry no date.
            text = " ".join(str(c) for c in row if c is not None).lower()
            if any(k in text for k in ("total", "closing", "opening", "end of", "statement summary", "balance b/f")):
                continue
            errors.append(f"row {n}: '{bad}' isn't a date")
            continue
        if vdate is None:
            continue
        try:
            tdate = _parse_date(cell(row, "txn_date"), fmt.date_formats)
        except ValueError:
            tdate = None
        debit = credit = 0.0
        if "debit" in cols or "credit" in cols:
            d, _ = _parse_amount(cell(row, "debit"))
            c, _ = _parse_amount(cell(row, "credit"))
            debit, credit = abs(d or 0.0), abs(c or 0.0)
        else:
            amt, marker = _parse_amount(cell(row, "amount"))
            flag = str(cell(row, "drcr") or marker or "").strip().lower()
            amt = amt or 0.0
            if flag.startswith("d") or (not flag and amt < 0):
                debit = abs(amt)
            else:
                credit = abs(amt)
        if debit and credit:
            errors.append(f"row {n}: both debit ({debit:,.2f}) and credit ({credit:,.2f}) are filled")
            continue
        if not debit and not credit:
            continue  # balance-only / info rows
        bal, marker = _parse_amount(cell(row, "balance"))
        if bal is not None and marker == "dr":
            bal = -abs(bal)
        desc = cell(row, "description")
        ref = cell(row, "reference")
        chq = cell(row, "cheque_number")
        out.append(ParsedLine(
            row_number=n, value_date=vdate, txn_date=tdate,
            description=(str(desc).strip() or None) if desc is not None else None,
            reference=(str(ref).strip()[:120] or None) if ref not in (None, "") else None,
            cheque_number=(str(chq).strip()[:50] or None) if chq not in (None, "") else None,
            debit=round(debit, 2), credit=round(credit, 2), balance=round(bal, 2) if bal is not None else None,
        ))
    if errors:
        shown = "; ".join(errors[:5]) + (f"; and {len(errors) - 5} more" if len(errors) > 5 else "")
        raise ValueError(f"The statement couldn't be read: {shown}. Fix those rows (or pick the right bank format) and upload again.")
    if not out:
        raise ValueError("The statement has a header row but no transaction lines with a debit or credit amount.")
    return out
