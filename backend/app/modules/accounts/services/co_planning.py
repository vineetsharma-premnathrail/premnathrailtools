"""CO cost-centre planning (SAP KP06) and plan vs actual (S_ALR_87013611).

A plan version belongs to one financial year and holds amounts per cost
centre x P&L ledger x month. Versions are 'original' or 'revised'; one per
FY is primary (the default for plan vs actual). A locked version can't be
changed — unlocking needs a reason (audit trail in the route).

Grid and Excel use the same shape: one row per cost centre x ledger, one
column per month of the FY (Apr..Mar), amounts in rupees (costs positive,
revenue negative, like debit-minus-credit actuals)."""
from __future__ import annotations

import io
from datetime import date, datetime, timezone

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.modules.accounts.models.co_planning import CO_PLAN_VERSION_TYPES, CoPlanLine, CoPlanVersion
from app.modules.accounts.models.gl_account import GLAccount
from app.modules.accounts.models.journal_entry import JE_LEDGER_STATUSES, JournalEntry, JournalEntryLine
from app.modules.accounts.services import fiscal
from app.modules.accounts.services.controlling import cost_center_tree_ids
from app.modules.organization.models.cost_center import CostCenter

_NON_PL_POSTINGS = ("year_end",)
_PL_TYPES = ("revenue", "expense")


def _r(x) -> float:
    return round(float(x or 0), 2)


def _month_label(period: str) -> str:
    y, m = period.split("-")
    return date(int(y), int(m), 1).strftime("%b-%Y")


def fy_bounds(db: Session, financial_year: str) -> tuple[date, date]:
    """'2026-27' -> (start, end) of that FY (Company Info FYs, else Apr-Mar)."""
    try:
        start_year = int((financial_year or "")[:4])
    except ValueError:
        raise ValueError(f"'{financial_year}' isn't a financial year — use the form 2026-27.")
    from app.modules.organization.models.company_financial_year import CompanyFinancialYear

    for fy in db.query(CompanyFinancialYear).all():
        if fiscal.fy_code(fy.start_date, fy.end_date) == financial_year:
            return fy.start_date, fy.end_date
    derived = fiscal.derived_fiscal_year(date(start_year, 4, 1))
    if derived.code != financial_year:
        raise ValueError(f"'{financial_year}' isn't a financial year — use the form {derived.code}.")
    return derived.start_date, derived.end_date


def fy_periods(db: Session, financial_year: str) -> list[str]:
    start, end = fy_bounds(db, financial_year)
    out, y, m = [], start.year, start.month
    while date(y, m, 1) <= end:
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _version(db: Session, version_id: int) -> CoPlanVersion:
    v = db.query(CoPlanVersion).filter(CoPlanVersion.id == version_id).first()
    if not v:
        raise ValueError("Plan version not found — it may have been deleted. Reload the list.")
    return v


def _ensure_draft(v: CoPlanVersion) -> None:
    if v.status == "locked":
        raise ValueError(f"Plan version {v.code} ({v.financial_year}) is locked — unlock it first (needs a reason), "
                         f"or copy it into a new revised version.")


def version_row(db: Session, v: CoPlanVersion) -> dict:
    total, count = db.query(func.coalesce(func.sum(CoPlanLine.amount), 0), func.count(CoPlanLine.id)).filter(
        CoPlanLine.version_id == v.id).one()
    return {
        "id": v.id, "code": v.code, "name": v.name, "financial_year": v.financial_year, "version_type": v.version_type,
        "status": v.status, "is_primary": v.is_primary, "copied_from_version_id": v.copied_from_version_id,
        "notes": v.notes, "locked_at": v.locked_at, "locked_by_id": v.locked_by_id,
        "total_amount": _r(total), "line_count": int(count or 0),
    }


def list_versions(db: Session, financial_year: str | None = None) -> list[dict]:
    q = db.query(CoPlanVersion)
    if financial_year:
        q = q.filter(CoPlanVersion.financial_year == financial_year)
    return [version_row(db, v) for v in q.order_by(CoPlanVersion.financial_year.desc(), CoPlanVersion.code.asc()).all()]


def _set_primary(db: Session, v: CoPlanVersion) -> None:
    db.query(CoPlanVersion).filter(CoPlanVersion.financial_year == v.financial_year, CoPlanVersion.id != v.id).update(
        {CoPlanVersion.is_primary: False}, synchronize_session=False)
    v.is_primary = True


def _actual_cells(db: Session, start: date, end: date, cc_ids: list[int] | None = None) -> dict[tuple, float]:
    """{(cc_id, gl_id, period): debit-credit} of posted P&L lines carrying a cost centre."""
    q = (
        db.query(JournalEntryLine.cost_center_id, JournalEntryLine.gl_account_id, JournalEntry.accounting_period,
                 func.coalesce(func.sum(JournalEntryLine.debit_amount), 0), func.coalesce(func.sum(JournalEntryLine.credit_amount), 0))
        .join(JournalEntry, JournalEntryLine.journal_entry_id == JournalEntry.id)
        .join(GLAccount, JournalEntryLine.gl_account_id == GLAccount.id)
        .filter(JournalEntry.status.in_(JE_LEDGER_STATUSES), JournalEntry.posting_type.notin_(_NON_PL_POSTINGS),
                JournalEntry.posting_date >= start, JournalEntry.posting_date <= end,
                JournalEntryLine.cost_center_id.isnot(None), GLAccount.account_type.in_(_PL_TYPES))
    )
    if cc_ids is not None:
        q = q.filter(JournalEntryLine.cost_center_id.in_(cc_ids))
    out: dict[tuple, float] = {}
    for cc, gl, period, dr, cr in q.group_by(JournalEntryLine.cost_center_id, JournalEntryLine.gl_account_id,
                                             JournalEntry.accounting_period).all():
        out[(cc, gl, period)] = _r(out.get((cc, gl, period), 0) + float(dr) - float(cr))
    return out


def create_version(db: Session, *, code: str, name: str, financial_year: str, version_type: str = "original",
                   is_primary: bool = False, notes: str | None = None, copy_from_version_id: int | None = None,
                   copy_actuals_from_fy: str | None = None, uplift_pct: float = 0.0, created_by_id: int | None = None) -> CoPlanVersion:
    code, name = (code or "").strip(), (name or "").strip()
    if not code or not name:
        raise ValueError("Give the version a code (e.g. 0 or R1) and a name.")
    if version_type not in CO_PLAN_VERSION_TYPES:
        raise ValueError("Version type must be 'original' or 'revised'.")
    periods = fy_periods(db, financial_year)
    if db.query(CoPlanVersion.id).filter(CoPlanVersion.financial_year == financial_year, CoPlanVersion.code == code).first():
        raise ValueError(f"Plan version {code} already exists for {financial_year} — pick another code.")
    if copy_from_version_id and copy_actuals_from_fy:
        raise ValueError("Copy either from another plan version or from a year's actuals, not both.")
    v = CoPlanVersion(code=code, name=name, financial_year=financial_year, version_type=version_type, notes=notes,
                      status="draft", is_primary=False, copied_from_version_id=copy_from_version_id, created_by_id=created_by_id)
    db.add(v)
    db.flush()
    factor = 1 + float(uplift_pct or 0) / 100
    if copy_from_version_id:
        src = _version(db, copy_from_version_id)
        src_periods = fy_periods(db, src.financial_year)
        shift = dict(zip(src_periods, periods))
        for line in src.lines:
            if line.accounting_period in shift:
                db.add(CoPlanLine(version_id=v.id, cost_center_id=line.cost_center_id, gl_account_id=line.gl_account_id,
                                  accounting_period=shift[line.accounting_period], amount=_r(float(line.amount) * factor)))
    elif copy_actuals_from_fy:
        src_periods = fy_periods(db, copy_actuals_from_fy)
        start, end = fy_bounds(db, copy_actuals_from_fy)
        shift = dict(zip(src_periods, periods))
        for (cc, gl, period), amount in _actual_cells(db, start, end).items():
            if period in shift and abs(amount) >= 0.005:
                db.add(CoPlanLine(version_id=v.id, cost_center_id=cc, gl_account_id=gl,
                                  accounting_period=shift[period], amount=_r(amount * factor)))
    if is_primary or not db.query(CoPlanVersion.id).filter(
            CoPlanVersion.financial_year == financial_year, CoPlanVersion.is_primary == True, CoPlanVersion.id != v.id).first():  # noqa: E712
        _set_primary(db, v)
    db.flush()
    return v


def update_version(db: Session, version_id: int, *, name: str | None = None, notes: str | None = None,
                   version_type: str | None = None, is_primary: bool | None = None) -> CoPlanVersion:
    v = _version(db, version_id)
    if name is not None or version_type is not None:
        _ensure_draft(v)
    if name is not None:
        if not name.strip():
            raise ValueError("The version name can't be blank.")
        v.name = name.strip()
    if version_type is not None:
        if version_type not in CO_PLAN_VERSION_TYPES:
            raise ValueError("Version type must be 'original' or 'revised'.")
        v.version_type = version_type
    if notes is not None:
        v.notes = notes
    if is_primary:
        _set_primary(db, v)
    db.flush()
    return v


def delete_version(db: Session, version_id: int) -> None:
    v = _version(db, version_id)
    _ensure_draft(v)
    if db.query(CoPlanVersion.id).filter(CoPlanVersion.copied_from_version_id == v.id).first():
        db.query(CoPlanVersion).filter(CoPlanVersion.copied_from_version_id == v.id).update(
            {CoPlanVersion.copied_from_version_id: None}, synchronize_session=False)
    db.delete(v)
    db.flush()


def set_locked(db: Session, version_id: int, *, locked: bool, user_id: int | None) -> CoPlanVersion:
    v = _version(db, version_id)
    if locked:
        if v.status == "locked":
            raise ValueError(f"Plan version {v.code} is already locked.")
        if not v.lines:
            raise ValueError(f"Plan version {v.code} has no plan lines yet — enter or upload the plan before locking it.")
        v.status, v.locked_by_id, v.locked_at = "locked", user_id, datetime.now(timezone.utc)
    else:
        if v.status != "locked":
            raise ValueError(f"Plan version {v.code} isn't locked.")
        v.status, v.locked_by_id, v.locked_at = "draft", None, None
    db.flush()
    return v


def grid(db: Session, version_id: int) -> dict:
    v = _version(db, version_id)
    periods = fy_periods(db, v.financial_year)
    cells: dict[tuple, dict] = {}
    for line in v.lines:
        cells.setdefault((line.cost_center_id, line.gl_account_id), {})[line.accounting_period] = _r(line.amount)
    ccs = {c.id: c for c in db.query(CostCenter).filter(CostCenter.id.in_({k[0] for k in cells})).all()} if cells else {}
    gls = {a.id: a for a in db.query(GLAccount).filter(GLAccount.id.in_({k[1] for k in cells})).all()} if cells else {}
    rows = []
    for (cc_id, gl_id), amounts in cells.items():
        cc, gl = ccs.get(cc_id), gls.get(gl_id)
        rows.append({
            "cost_center_id": cc_id, "cost_center_code": cc.code if cc else None, "cost_center_name": cc.name if cc else None,
            "gl_account_id": gl_id, "gl_account_code": gl.code if gl else None, "gl_account_name": gl.name if gl else None,
            "amounts": {p: amounts.get(p, 0.0) for p in periods}, "total": _r(sum(amounts.values())),
        })
    rows.sort(key=lambda r: (r["cost_center_code"] or "", r["gl_account_code"] or ""))
    return {
        "version": version_row(db, v), "periods": periods, "period_labels": [_month_label(p) for p in periods], "rows": rows,
        "period_totals": {p: _r(sum(r["amounts"][p] for r in rows)) for p in periods},
        "grand_total": _r(sum(r["total"] for r in rows)),
    }


def _check_objects(db: Session, cc_ids: set[int], gl_ids: set[int]) -> tuple[dict, dict]:
    ccs = {c.id: c for c in db.query(CostCenter).filter(CostCenter.id.in_(cc_ids)).all()} if cc_ids else {}
    gls = {a.id: a for a in db.query(GLAccount).filter(GLAccount.id.in_(gl_ids)).all()} if gl_ids else {}
    for cc_id in cc_ids - set(ccs):
        raise ValueError(f"Cost centre #{cc_id} doesn't exist — pick it again from the list.")
    for gl_id in gl_ids:
        a = gls.get(gl_id)
        if a is None:
            raise ValueError(f"Ledger #{gl_id} doesn't exist — pick it again from the list.")
        if a.account_type not in _PL_TYPES or not a.is_posting_account:
            raise ValueError(f"Ledger '{a.code} — {a.name}' isn't a P&L posting ledger — cost-centre plans are made on "
                             f"expense and income ledgers only.")
    return ccs, gls


def save_grid(db: Session, version_id: int, rows: list[dict]) -> dict:
    """Replaces the whole plan of a draft version with `rows`:
    [{cost_center_id, gl_account_id, amounts: {"YYYY-MM": amount}}]."""
    v = _version(db, version_id)
    _ensure_draft(v)
    periods = set(fy_periods(db, v.financial_year))
    seen: set[tuple] = set()
    for i, r in enumerate(rows, start=1):
        key = (r.get("cost_center_id"), r.get("gl_account_id"))
        if not key[0] or not key[1]:
            raise ValueError(f"Row {i}: pick both the cost centre and the ledger.")
        if key in seen:
            raise ValueError(f"Row {i}: this cost centre and ledger appear twice — keep one row per pair.")
        seen.add(key)
        for p in (r.get("amounts") or {}):
            if p not in periods:
                raise ValueError(f"Row {i}: {p} isn't a month of {v.financial_year}.")
    _check_objects(db, {k[0] for k in seen}, {k[1] for k in seen})
    db.query(CoPlanLine).filter(CoPlanLine.version_id == v.id).delete(synchronize_session=False)
    for r in rows:
        for p, amount in (r.get("amounts") or {}).items():
            amt = _r(amount)
            if abs(amt) >= 0.005:
                db.add(CoPlanLine(version_id=v.id, cost_center_id=r["cost_center_id"], gl_account_id=r["gl_account_id"],
                                  accounting_period=p, amount=amt))
    db.flush()
    db.expire(v)
    return grid(db, v.id)


# ── Excel ─────────────────────────────────────────────────────────────────

_FIXED_HEADERS = ["Cost centre code", "Cost centre name", "Ledger code", "Ledger name"]


def export_xlsx(db: Session, version_id: int) -> tuple[bytes, str]:
    """The version's grid as an .xlsx; an empty version gives a template."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    g = grid(db, version_id)
    v = g["version"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(_FIXED_HEADERS + g["period_labels"] + ["Total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in g["rows"]:
        ws.append([r["cost_center_code"], r["cost_center_name"], r["gl_account_code"], r["gl_account_name"]]
                  + [r["amounts"][p] for p in g["periods"]] + [r["total"]])
    ws.freeze_panes = "E2"
    notes = wb.create_sheet("Read me")
    for line in (
        f"Plan version {v['code']} — {v['name']} ({v['financial_year']})",
        "One row per cost centre and ledger. Codes must match Organization › Cost centres and Finance › Ledgers.",
        "Amounts in rupees: costs positive, income negative. Blank = 0. The names and Total columns are ignored on upload.",
        "Uploading replaces the whole plan of this (draft) version.",
    ):
        notes.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"CO-plan-{v['financial_year']}-{v['code']}.xlsx"


def import_xlsx(db: Session, version_id: int, content: bytes) -> dict:
    """Reads an .xlsx in the export layout and replaces the version's plan.
    Every problem is reported with its row number; nothing is saved unless
    the whole sheet is clean."""
    from openpyxl import load_workbook

    v = _version(db, version_id)
    _ensure_draft(v)
    periods = fy_periods(db, v.financial_year)
    labels = [_month_label(p) for p in periods]
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("That file isn't a readable .xlsx workbook — download the template from this version and fill it in.")
    ws = wb["Plan"] if "Plan" in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in (next(rows_iter, None) or [])]
    expected = _FIXED_HEADERS + labels
    if [h.lower() for h in header[:len(expected)]] != [h.lower() for h in expected]:
        raise ValueError(f"The header row doesn't match this version's layout. Expected: {', '.join(expected)}. "
                         f"Download the template from this version ({v.financial_year}) and paste the figures into it.")
    ccs = {c.code.strip().upper(): c for c in db.query(CostCenter).all()}
    gls = {a.code.strip().upper(): a for a in db.query(GLAccount).all()}
    errors: list[str] = []
    out: dict[tuple, dict] = {}
    for n, row in enumerate(rows_iter, start=2):
        row = list(row or [])
        if not any(c not in (None, "") for c in row[:4 + len(periods)]):
            continue
        cc_code = str(row[0] or "").strip().upper()
        gl_code = str(row[2] or "").strip().upper() if len(row) > 2 else ""
        cc, gl = ccs.get(cc_code), gls.get(gl_code)
        if not cc:
            errors.append(f"Row {n}: cost centre code '{row[0] or ''}' isn't in Organization › Cost centres.")
        if not gl:
            errors.append(f"Row {n}: ledger code '{row[2] if len(row) > 2 else ''}' isn't in Finance › Ledgers.")
        elif gl.account_type not in _PL_TYPES or not gl.is_posting_account:
            errors.append(f"Row {n}: ledger '{gl.code} — {gl.name}' isn't a P&L posting ledger.")
        amounts = {}
        for i, p in enumerate(periods):
            raw = row[4 + i] if len(row) > 4 + i else None
            if raw in (None, ""):
                continue
            try:
                amounts[p] = _r(float(str(raw).replace(",", "")))
            except ValueError:
                errors.append(f"Row {n}: '{raw}' under {labels[i]} isn't a number.")
        if cc and gl:
            if (cc.id, gl.id) in out:
                errors.append(f"Row {n}: {cc.code} / {gl.code} appears twice — keep one row per cost centre and ledger.")
            out[(cc.id, gl.id)] = amounts
    if errors:
        more = f" (and {len(errors) - 25} more)" if len(errors) > 25 else ""
        raise ValueError("Nothing was imported — fix these rows and upload again: " + " ".join(errors[:25]) + more)
    return save_grid(db, v.id, [{"cost_center_id": k[0], "gl_account_id": k[1], "amounts": a} for k, a in out.items()])


# ── Plan vs actual ───────────────────────────────────────────────────────

def plan_vs_actual(db: Session, *, version_id: int | None = None, financial_year: str | None = None,
                   from_period: str | None = None, to_period: str | None = None,
                   cost_center_id: int | None = None, include_children: bool = True) -> dict:
    """Plan (the version) vs actual (posted lines) per cost centre x ledger
    for a run of months inside the version's FY. Variance = actual - plan
    (positive = overspent on a cost, short on income)."""
    if version_id:
        v = _version(db, version_id)
    else:
        fy = financial_year or fiscal.fiscal_year_for(db, date.today()).code
        v = db.query(CoPlanVersion).filter(CoPlanVersion.financial_year == fy, CoPlanVersion.is_primary == True).first()  # noqa: E712
        if not v:
            raise ValueError(f"There is no primary plan version for {fy} — create one under Finance › Controlling › Planning.")
    periods = fy_periods(db, v.financial_year)
    from_period = from_period or periods[0]
    to_period = to_period or periods[-1]
    if from_period not in periods or to_period not in periods:
        raise ValueError(f"Pick months inside {v.financial_year} ({_month_label(periods[0])} to {_month_label(periods[-1])}).")
    if from_period > to_period:
        raise ValueError("The 'from' month is after the 'to' month — swap them.")
    sel = [p for p in periods if from_period <= p <= to_period]
    start, _ = fiscal.period_bounds(sel[0])
    _, end = fiscal.period_bounds(sel[-1])

    cc_ids = None
    if cost_center_id:
        if not db.query(CostCenter.id).filter(CostCenter.id == cost_center_id).first():
            raise ValueError("Cost centre not found — it may have been deleted. Reload the list.")
        cc_ids = cost_center_tree_ids(db, cost_center_id, include_children)

    plan: dict[tuple, float] = {}
    pq = db.query(CoPlanLine).filter(CoPlanLine.version_id == v.id, CoPlanLine.accounting_period.in_(sel))
    if cc_ids is not None:
        pq = pq.filter(CoPlanLine.cost_center_id.in_(cc_ids))
    for line in pq.all():
        key = (line.cost_center_id, line.gl_account_id)
        plan[key] = _r(plan.get(key, 0) + float(line.amount))
    actual: dict[tuple, float] = {}
    for (cc, gl, _p), amt in _actual_cells(db, start, end, cc_ids).items():
        actual[(cc, gl)] = _r(actual.get((cc, gl), 0) + amt)

    keys = set(plan) | set(actual)
    ccs = {c.id: c for c in db.query(CostCenter).filter(CostCenter.id.in_({k[0] for k in keys})).all()} if keys else {}
    gls = {a.id: a for a in db.query(GLAccount).filter(GLAccount.id.in_({k[1] for k in keys})).all()} if keys else {}
    rows = []
    for key in keys:
        cc, gl = ccs.get(key[0]), gls.get(key[1])
        p, a = plan.get(key, 0.0), actual.get(key, 0.0)
        rows.append({
            "cost_center_id": key[0], "cost_center_code": cc.code if cc else None, "cost_center_name": cc.name if cc else None,
            "gl_account_id": key[1], "gl_account_code": gl.code if gl else None, "gl_account_name": gl.name if gl else None,
            "account_type": gl.account_type if gl else None,
            "plan": p, "actual": a, "variance": _r(a - p), "variance_pct": _r((a - p) / abs(p) * 100) if p else None,
        })
    rows.sort(key=lambda r: (r["cost_center_code"] or "", 0 if r["account_type"] == "expense" else 1, r["gl_account_code"] or ""))
    tp, ta = _r(sum(r["plan"] for r in rows)), _r(sum(r["actual"] for r in rows))
    return {
        "version": version_row(db, v), "from_period": sel[0], "to_period": sel[-1], "from_date": start, "to_date": end,
        "cost_center_id": cost_center_id, "include_children": include_children, "rows": rows,
        "total_plan": tp, "total_actual": ta, "total_variance": _r(ta - tp),
        "total_variance_pct": _r((ta - tp) / abs(tp) * 100) if tp else None,
    }
