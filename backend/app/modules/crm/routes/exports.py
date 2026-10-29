"""CRM Excel exports for the Organizations and Inquiries & Tenders lists.

Same two modes as the Service registers: exactly the ticked records (ids),
or everything matching the list's search, optionally narrowed to a From–To
range on the date the record was created (both ends inclusive, days in IST).
Under /crm/exports/… so the paths never collide with /crm/organizations/{id}."""
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.core.permissions import require_app_access, require_tab_action
from app.core.data_scope import apply_scope
import app.modules.crm.access  # noqa: F401 — registers the CRM data-scope models
from app.modules.crm.access import sees_prices
from app.db.session import get_db
from app.modules.crm.models.inquiry import Inquiry
from app.modules.crm.models.organization import Organization, OrgContact
from app.modules.crm.models.tender import Tender
from app.modules.erp.importing import check_date_range, in_range, ist_date, parse_export_ids
from app.modules.main.models.user import User

router = APIRouter(prefix="/crm/exports", tags=["CRM - Export"])


def _xlsx(sheets: list[tuple[str, list[str], list[list]]], filename: str) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FDF1E6")
        for row in rows:
            ws.append(row)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 4))
        for col in ws.iter_cols(min_row=2):
            for cell in col:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "DD-MM-YYYY"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _names(db: Session, ids: set) -> dict[int, str]:
    ids = {i for i in ids if i}
    return {u.id: u.name or u.email for u in db.query(User).filter(User.id.in_(ids)).all()} if ids else {}


def _v(x):
    return "" if x is None else x


@router.get("/organizations")
def export_organizations(
    search: str | None = Query(None),
    ids: str | None = Query(None, description="Comma-separated organization ids — export only these."),
    date_from: date | None = Query(None),
    date_to: date | None = Query(None),
    db: Session = Depends(get_db),
    _user: User = Depends(require_app_access("crm")),
    _tab: User = Depends(require_tab_action("crm", "organizations", "export")),
):
    selected = parse_export_ids(ids, "organizations")
    check_date_range(date_from, date_to)
    q = apply_scope(db.query(Organization).filter(Organization.is_deleted == False), _user, Organization)  # noqa: E712
    if selected is not None:
        orgs = q.filter(Organization.id.in_(selected)).order_by(Organization.name).all()
        if not orgs:
            raise HTTPException(status_code=404, detail="None of the selected customer organizations can be exported — they were deleted, or are outside your CRM data access. Refresh the list.")
    else:
        if search and search.strip():
            like = f"%{search.strip()}%"
            q = q.filter(or_(Organization.name.ilike(like), Organization.org_code.ilike(like), Organization.city.ilike(like),
                             Organization.gst_number.ilike(like),
                             Organization.tax_vat_id.ilike(like)))
        orgs = [o for o in q.order_by(Organization.name).all() if in_range(ist_date(o.created_at), date_from, date_to)]

    contacts: dict[int, list[OrgContact]] = {}
    if orgs:
        for c in db.query(OrgContact).filter(OrgContact.org_id.in_([o.id for o in orgs])).order_by(OrgContact.id).all():
            contacts.setdefault(c.org_id, []).append(c)
    names = _names(db, {o.created_by_id for o in orgs})
    headers = ["Org Code", "Customer Organization", "Type", "Parent", "Railway Zone", "Division / Workshop", "Address", "City", "State",
               "PIN", "Country", "Scope", "GST No.", "PAN", "GST State Code", "Tax / VAT ID", "Company Registration No", "Phone", "Email", "Website", "Contacts", "Primary Contact", "Primary Contact Mobile",
               "Credit Limit", "Credit Days", "Created On", "Created By"]
    rows = []
    for o in orgs:
        cs = contacts.get(o.id, [])
        first = cs[0] if cs else None
        rows.append([_v(o.org_code), o.name, _v(o.org_type), _v(o.parent_org), _v(o.railway_zone), _v(o.division_workshop),
                     _v(o.address), _v(o.city), _v(o.state), _v(o.pin_code), _v(o.country), (o.org_scope or "domestic").title(), _v(o.gst_number),
                     _v(o.pan), _v(o.gst_state_code), _v(o.tax_vat_id), _v(o.company_registration_no),
                     _v(o.official_phone), _v(o.official_email), _v(o.website), len(cs),
                     _v(first.name if first else None), _v(getattr(first, "mobile", None) if first else None),
                     _v(o.credit_limit), _v(o.credit_days), _v(ist_date(o.created_at)), names.get(o.created_by_id, "")])
    return _xlsx([("Customer Organizations", headers, rows)], "organizations.xlsx")


@router.get("/inquiries-tenders")
def export_inquiries_tenders(
    search: str | None = Query(None),
    type: str = Query("all", description="all | inquiry | tender"),
    inquiry_ids: str | None = Query(None),
    tender_ids: str | None = Query(None),
    date_from: date | None = Query(None),
    date_to: date | None = Query(None),
    db: Session = Depends(get_db),
    _user: User = Depends(require_app_access("crm")),
    _tab: User = Depends(require_tab_action("crm", "inquiries_tenders", "export")),
):
    """One workbook, an Inquiries sheet and a Tenders sheet."""
    if type not in ("all", "inquiry", "tender"):
        raise HTTPException(status_code=422, detail="Type must be all, inquiry or tender.")
    check_date_range(date_from, date_to)
    sel_inq = parse_export_ids(inquiry_ids, "inquiries") if inquiry_ids else None
    sel_tnd = parse_export_ids(tender_ids, "tenders") if tender_ids else None
    selective = sel_inq is not None or sel_tnd is not None
    term = f"%{search.strip()}%" if search and search.strip() else None

    def pick(model, selected, search_cols):
        q = apply_scope(db.query(model).filter(model.is_deleted == False), _user, model)  # noqa: E712
        if selective:
            return q.filter(model.id.in_(selected)).all() if selected else []
        if term:
            q = q.outerjoin(Organization, Organization.id == model.org_id).filter(or_(*[c.ilike(term) for c in search_cols], Organization.name.ilike(term)))
        return [r for r in q.all() if in_range(ist_date(r.created_at), date_from, date_to)]

    inquiries = pick(Inquiry, sel_inq, [Inquiry.universal_id, Inquiry.product, Inquiry.product_category]) if type in ("all", "inquiry") or selective else []
    tenders = pick(Tender, sel_tnd, [Tender.universal_id, Tender.tender_name, Tender.tender_number]) if type in ("all", "tender") or selective else []
    if selective and not inquiries and not tenders:
        raise HTTPException(status_code=404, detail="None of the selected records can be exported — they were deleted, or are outside your CRM data access. Refresh the list.")
    inquiries.sort(key=lambda r: r.created_at or datetime.min, reverse=True)
    tenders.sort(key=lambda r: r.created_at or datetime.min, reverse=True)

    org_ids = {r.org_id for r in [*inquiries, *tenders]}
    orgs = {o.id: o.name for o in db.query(Organization).filter(Organization.id.in_(org_ids)).all()} if org_ids else {}
    names = _names(db, {r.created_by_id for r in [*inquiries, *tenders]})
    # Inquiries & Tenders › Prices & margins: without it the value columns are left empty.
    show_money = sees_prices(_user)

    def money(v):
        return _v(v) if show_money else ""

    inq_headers = ["Inquiry ID", "Customer Organization", "Product",
                   "Specification", "Status", "Priority", "Lead Source", "BD Owner",
                   "Sales Engineer", "Expected Value", "Probability %", "Expected Order Date", "Next Follow-up",
                   "Technical Offer No.", "Sent to R&D On", "Created On", "Created By"]
    inq_rows = [[i.universal_id, orgs.get(i.org_id, ""), _v(i.product),
                 _v(i.product_spec),
                 _v(i.status), _v(i.priority), _v(i.lead_source), _v(i.bd_owner), _v(i.sales_engineer),
                 money(i.expected_value), _v(i.probability), _v(i.expected_order_date), _v(i.next_followup_date),
                 _v(i.technical_offer_number), _v(ist_date(i.technical_offer_sent_at)), _v(ist_date(i.created_at)),
                 names.get(i.created_by_id, "")] for i in inquiries]
    tnd_headers = ["Tender ID", "Customer Organization", "Tender Name", "Tender No.", "Authority", "Portal", "Type", "Category",
                   "Tender Value", "Currency", "Status", "Stage", "Tender Details", "Priority", "BD Owner", "Railway Zone", "Division",
                   "Publish Date", "Pre-bid Meeting", "Submission Date", "Closing Date & Time", "Opening Date", "Participate", "Awarded To",
                   "Contract Value", "Technical Offer No.", "Created On", "Created By"]
    tnd_rows = [[t.universal_id, orgs.get(t.org_id, ""), _v(t.tender_name), _v(t.tender_number), _v(t.tender_authority),
                 _v(t.tender_portal), _v(t.tender_type), _v(t.tender_category), money(t.tender_value), _v(t.currency),
                 _v(t.status), _v(t.current_stage), _v(t.tender_details), _v(t.priority), _v(t.bd_owner), _v(t.railway_zone), _v(t.division),
                 _v(t.publish_date), _v(t.pre_bid_meeting_date), _v(t.submission_date),
                 t.closing_at.strftime("%d-%m-%Y %H:%M") if t.closing_at else "", _v(t.opening_date),
                 "" if t.participate is None else ("Yes" if t.participate else "No"), _v(t.awarded_to), money(t.contract_value),
                 _v(t.technical_offer_number), _v(ist_date(t.created_at)), names.get(t.created_by_id, "")] for t in tenders]

    sheets = []
    if inquiries or type in ("all", "inquiry") and not selective:
        sheets.append(("Inquiries", inq_headers, inq_rows))
    if tenders or type in ("all", "tender") and not selective:
        sheets.append(("Tenders", tnd_headers, tnd_rows))
    return _xlsx(sheets, "inquiries_tenders.xlsx")
