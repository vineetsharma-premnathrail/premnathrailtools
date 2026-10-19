"""DOCX builder for the CRM Quotation — the same layout as quotation_pdf.py,
as an editable Word file (edited in Word Online as a managed document, see
main/services/managed_documents.py).

Mirrors the PDF section by section: right-aligned "QUOTATION" title,
client block + quote meta, navy-header items grid, tax/discount summary
with grand-total band, amount in words, terms, signatory. Amount/words/
sentence-case helpers are imported from quotation_pdf.py so both files
always print identical numbers; the letterhead comes from
utils/docx_letterhead.py (same logo.JPG / logo-1.JPG as the PDF).
"""

from __future__ import annotations

import io
from typing import Any, Mapping

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.modules.crm.reports.quotation_pdf import (
    _sentence_case, _split_numbered_points, _fmt_amount, _fmt_quantity, _amount_in_words,
    LUMP_SUM_TITLE, LUMP_SUM_TERMS_HEADING, LUMP_SUM_SIGNATORY, _is_lump_sum, _particular_lines, _lump_sum_price_line,
    _detail_lines, _terms_sections,
)
from app.utils.docx_letterhead import (
    BODY, HEADER, ACCENT, ACCENT_HEX, ACCENT_LIGHT_HEX, MUTED, WHITE, USABLE_IN,
    fmt_run, tight, write, shade, cell_margins, set_widths, top_border, new_letterhead_document,
)


def build_quotation_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    """Same ctx keys as quotation_pdf.build_quotation_pdf."""
    doc = new_letterhead_document()

    currency = ctx.get("currency_symbol") or "Rs."
    is_export = (ctx.get("quotation_type") or "Domestic") == "Export"
    lump_sum = _is_lump_sum(ctx)

    # ── Title ──
    title = tight(doc.add_paragraph())
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_run(title.add_run(LUMP_SUM_TITLE if lump_sum else "QUOTATION"), size=HEADER + 12, bold=True, color=ACCENT)
    write(doc, f"Quotation type: {ctx.get('quotation_type') or 'Domestic'}", align="right", color=MUTED, after=14)

    # ── Client (left) + quote meta (right) ──
    top = doc.add_table(rows=1, cols=2)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(top, left=0, right=0, top=0, bottom=0)
    set_widths(top, [USABLE_IN * 0.55, USABLE_IN * 0.45])
    left, right = top.rows[0].cells
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    write(left, "CLIENT", size=BODY - 1, bold=True, color=ACCENT, first=True, after=3)
    write(left, ctx.get("client_name") or "-", size=HEADER + 1, bold=True)
    for key, transform in (
        ("client_address", _sentence_case), ("client_gst_number", lambda v: f"GSTIN: {v}"),
        ("client_contact_name", _sentence_case), ("client_contact_email", str), ("client_contact_phone", str),
    ):
        if ctx.get(key):
            write(left, transform(ctx[key]))

    meta_pairs = [("Quotation #", ctx.get("quot_number") or "-"), ("Quotation date", ctx.get("quote_date") or "-")]
    if ctx.get("technical_offer_number"):
        meta_pairs.append(("Technical offer #", ctx["technical_offer_number"]))
    if ctx.get("technical_offer_date"):
        meta_pairs.append(("Technical offer date", ctx["technical_offer_date"]))
    if ctx.get("valid_until"):
        meta_pairs.append(("Valid Until", ctx["valid_until"]))
    if ctx.get("delivery_time"):
        meta_pairs.append(("Delivery Time", _sentence_case(ctx["delivery_time"])))
    for i, (label, value) in enumerate(meta_pairs):
        p = right.paragraphs[0] if i == 0 else right.add_paragraph()
        tight(p, after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fmt_run(p.add_run(f"{label}   "), bold=True, color=ACCENT)
        fmt_run(p.add_run(str(value)))

    write(doc, "", after=10)

    if lump_sum:
        _lump_sum_body(doc, ctx, currency, is_export)
    else:
        _itemwise_body(doc, ctx, currency, is_export)
    _terms_and_signatory(doc, ctx, lump_sum)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _header_row(grid, headers: list[str]) -> None:
    for cell, h in zip(grid.rows[0].cells, headers):
        shade(cell, ACCENT_HEX)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        write(cell, h, bold=True, align="center", color=WHITE, first=True)
    # Repeat the header row on every page, like the PDF's repeatRows=1.
    tr_pr = grid.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _lump_sum_body(doc, ctx: Mapping[str, Any], currency: str, is_export: bool) -> None:
    grid = doc.add_table(rows=1, cols=2)
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(grid)
    _header_row(grid, ["S.No.", "Particulars"])
    for i, item in enumerate(ctx.get("items") or []):
        num_cell, text_cell = grid.add_row().cells
        num_cell.vertical_alignment = text_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        write(num_cell, str(i + 1), align="center", first=True)
        for n, line in enumerate(_particular_lines(item)):
            write(text_cell, line, bold=(n == 0), first=(n == 0))
    set_widths(grid, [0.7, USABLE_IN - 0.7])

    price = write(doc, "", after=0)
    tight(price, before=16, after=6)
    fmt_run(price.add_run(_lump_sum_price_line(ctx, currency, is_export)), size=HEADER + 1, bold=True, color=ACCENT)
    words = write(doc, "", after=0)
    fmt_run(words.add_run(f"Amount in words: {_amount_in_words(ctx.get('price') or 0, is_export)}"), bold=True)


def _itemwise_body(doc, ctx: Mapping[str, Any], currency: str, is_export: bool) -> None:
    # ── Items grid ──
    headers = ["#", "Description", "Model no.", "Qty", f"Unit price ({currency})", f"Taxable value ({currency})"]
    widths = [0.3, 2.5, 0.7, 0.5, 1.3, 1.3]
    if not is_export:
        headers.append("GST %")
        widths.append(0.5)
    headers.append(f"Total ({currency})")
    widths.append(1.4)
    scale = USABLE_IN / sum(widths)
    widths = [w * scale for w in widths]

    items = ctx.get("items") or []
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(grid)
    _header_row(grid, headers)

    taxable_total = 0.0
    gst_total = 0.0
    for i, item in enumerate(items):
        raw_subtotal = item.get("subtotal")
        subtotal = float(raw_subtotal) if raw_subtotal is not None else 0.0
        gst_amt = subtotal * float(item.get("gst_percent") or 0) / 100
        taxable_total += subtotal
        gst_total += gst_amt
        values = [
            (str(i + 1), "center"),
            (_sentence_case(item.get("description")) or "-", "left"),
            (item.get("model_number") or "-", "center"),
            (_fmt_quantity(item.get("quantity")), "center"),
            (_fmt_amount(item.get("unit_price"), is_export), "right"),
            (_fmt_amount(raw_subtotal, is_export), "right"),
        ]
        if not is_export:
            values.append((_fmt_amount(item.get("gst_percent"), is_export), "center"))
        values.append((_fmt_amount(item.get("total"), is_export), "right"))
        cells = grid.add_row().cells
        for cell, (text, align) in zip(cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            write(cell, text, align=align, first=True)
        for line in _detail_lines(item):
            write(cells[1], line, size=BODY - 1, color=MUTED)
    set_widths(grid, widths)
    write(doc, "", after=6)

    # ── Summary ──
    grand_total = sum(float(item.get("total") or 0) for item in items)
    rows: list[tuple[str, str]] = [("Taxable value", _fmt_amount(taxable_total, is_export))]
    if not is_export:
        if (ctx.get("gst_type") or "CGST_SGST") == "IGST":
            rows.append(("IGST", _fmt_amount(gst_total, is_export)))
        else:
            rows.append(("CGST", _fmt_amount(gst_total / 2, is_export)))
            rows.append(("SGST", _fmt_amount(gst_total / 2, is_export)))
    if ctx.get("discount"):
        discount = float(ctx["discount"])
        discount_amt = grand_total * discount / 100 if ctx.get("discount_type") == "percent" else discount
        grand_total -= discount_amt
        label = f"Discount ({discount:g}%)" if ctx.get("discount_type") == "percent" else "Discount"
        rows.append((label, f"-{_fmt_amount(discount_amt, is_export)}"))
    rows.append(("Grand total", _fmt_amount(grand_total, is_export)))

    summary = doc.add_table(rows=len(rows), cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(summary, left=0, right=90, top=40, bottom=40)
    last = len(rows) - 1
    for r, (label, value) in enumerate(rows):
        lc, vc = summary.rows[r].cells
        size = HEADER + 1 if r == last else BODY
        write(lc, label, size=size, bold=True, align="right", first=True)
        write(vc, value, size=size, bold=(r == last), align="right", first=True)
        if r == last:
            for c in (lc, vc):
                shade(c, ACCENT_LIGHT_HEX)
                top_border(c, ACCENT_HEX)
    set_widths(summary, [USABLE_IN - 1.5, 1.5])

    words = write(doc, "", after=0)
    tight(words, before=12)
    fmt_run(words.add_run(f"Amount in words: {_amount_in_words(grand_total, is_export)}"), bold=True)


def _terms_and_signatory(doc, ctx: Mapping[str, Any], lump_sum: bool) -> None:
    # ── Terms ──
    terms_sections = _terms_sections(ctx, lump_sum)
    if terms_sections:
        heading = write(doc, LUMP_SUM_TERMS_HEADING if lump_sum else "TERMS AND CONDITIONS", size=BODY + 1, bold=True, color=ACCENT, after=8)
        tight(heading, before=20, after=8)
        heading.paragraph_format.keep_with_next = True
        for label, text in terms_sections:
            if label:
                lp = write(doc, label, bold=True, color=MUTED, after=3)
                lp.paragraph_format.keep_with_next = True
            for point in _split_numbered_points(text):
                write(doc, point, after=5)
            write(doc, "", after=0)

    # ── Signatory ──
    if lump_sum:
        for n, line in enumerate(LUMP_SUM_SIGNATORY):
            p = write(doc, line, bold=(n > 0))
            if n == 0:
                tight(p, before=24)
            p.paragraph_format.keep_with_next = n < len(LUMP_SUM_SIGNATORY) - 1
        return
    sign_line = write(doc, "_" * 34, align="right")
    tight(sign_line, before=24)
    sign_line.paragraph_format.keep_with_next = True
    write(doc, "Authorized Signatory", align="right", color=MUTED)
