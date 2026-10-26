"""DOCX builder for the Final Project Report — the closure document of a
Project Management project: identity, scope, schedule performance, WBS
outcome, effort by department, delays, closure note, lessons learned and
sign-off. An editable Word document of record (see
main/services/managed_documents.py); the PM finishes the narrative in Word.
"""

from __future__ import annotations

import io
from typing import Any, Mapping

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt

from app.utils.docx_letterhead import (
    ACCENT, MUTED, USABLE_IN, tight, write, shade, cell_margins, set_widths, new_letterhead_document,
)

SHADE_HEX = "F3F4F6"


def _dash(value: Any) -> str:
    return "—" if value in (None, "") else str(value)


def _grid(doc: Any, rows: list[list[str]], widths: list[float], header: bool = False, bold_cols: tuple[int, ...] = ()) -> Any:
    table = doc.add_table(rows=0, cols=len(widths))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(table)
    for r, values in enumerate(rows):
        cells = table.add_row().cells
        for c, (cell, value) in enumerate(zip(cells, values)):
            is_head = header and r == 0
            write(cell, value, size=8 if (is_head or c in bold_cols) else 8.5, bold=is_head or c in bold_cols, first=True)
            if is_head or c in bold_cols:
                shade(cell, SHADE_HEX)
    set_widths(table, widths)
    return table


def _heading(doc: Any, text: str) -> None:
    tight(write(doc, text, size=10.5, bold=True, color=ACCENT), before=12, after=4)


def _text(doc: Any, value: str | None, placeholder: str) -> None:
    if value:
        write(doc, value, size=9, after=2)
    else:
        write(doc, placeholder, size=9, color=MUTED, after=2)


def build_closure_report_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    doc = new_letterhead_document()

    title = doc.add_table(rows=1, cols=2)
    cell_margins(title, left=0, right=0, top=0, bottom=40)
    left, right = title.rows[0].cells
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    write(left, "FINAL PROJECT REPORT", size=14, bold=True, color=ACCENT, first=True)
    write(right, ctx["code"], size=12, bold=True, align="right", first=True)
    set_widths(title, [USABLE_IN - 2.0, 2.0])

    half = USABLE_IN / 2
    facts = [
        ("Project", ctx["name"]), ("Customer", ctx.get("client")),
        ("Customer PO no.", ctx.get("po")), ("Sales order (SO) no.", ctx.get("so")),
        ("Project manager", ctx.get("pm")), ("Start date", ctx["start"]),
        ("Target handover", ctx["target"]), ("Technically complete", ctx["completed"]),
        ("Schedule variance", ctx["variance"]), ("Closed", ctx["closed"]),
    ]
    rows = [[facts[i][0], _dash(facts[i][1]), facts[i + 1][0], _dash(facts[i + 1][1])] for i in range(0, len(facts), 2)]
    tight(doc.add_paragraph(), before=4)
    _grid(doc, rows, [1.35, half - 1.35, 1.35, half - 1.35], bold_cols=(0, 2))

    n = iter(range(1, 20))
    _heading(doc, f"{next(n)}. Objective and scope")
    _text(doc, ctx.get("scope"), "[Describe the project objective and the agreed scope.]")

    # The WBS sections only apply to older projects that still have WBS
    # elements — the WBS tab is no longer in the UI.
    if ctx["wbs"]:
        _heading(doc, f"{next(n)}. Work breakdown structure — outcome")
        wbs = [["No.", "WBS", "Department", "Lead", "Planned", "Actual", "Slip", "Status", "Files"]] + [
            [str(r["no"]), r["title"], r["department"], r["lead"], r["planned"], r["actual"], r["slip"], r["status"], str(r["evidence"])]
            for r in ctx["wbs"]
        ]
        _grid(doc, wbs, [0.35, USABLE_IN - 5.75, 1.0, 1.0, 0.75, 0.75, 0.6, 0.85, 0.45], header=True)

        _heading(doc, f"{next(n)}. Deliverables")
        for r in ctx["wbs"]:
            if r["cancel_reason"]:
                line = f"WBS {r['no']} {r['title']}: cancelled — {r['cancel_reason']}"
            else:
                line = f"WBS {r['no']} {r['title']}: {r['deliverable'] or 'no completion criteria recorded'}" + (
                    f" Handover note: {r['handover_note']}" if r["handover_note"] else "")
            write(doc, line, size=8.5, after=2)

        _heading(doc, f"{next(n)}. Effort by department")
        effort = [["Department", "Tasks", "Estimated h", "Logged h"]] + [
            [dept, str(k), f"{est:g}", f"{logged:g}"] for dept, k, est, logged in ctx["effort"]
        ]
        _grid(doc, effort, [USABLE_IN - 3.3, 1.0, 1.15, 1.15], header=True)

    _heading(doc, f"{next(n)}. Delays and issues")
    if ctx["late"]:
        for r in ctx["late"]:
            write(doc, f"WBS {r['no']} {r['title']} ({r['department']}) finished {r['slip'].lstrip('+')} late: planned {r['planned']}, actual {r['actual']}.", size=8.5, after=2)
        write(doc, "[Add the cause of each delay and what was done about it.]", size=9, color=MUTED, after=2)
    elif ctx["wbs"]:
        write(doc, "Every WBS was accepted on or before its planned date.", size=9, after=2)
    else:
        write(doc, "[Describe any delays, their cause and what was done about them.]", size=9, color=MUTED, after=2)

    _heading(doc, f"{next(n)}. Closure and settlement")
    _text(doc, ctx.get("closure_note"), "[Final billing, retention, warranty start and any open commitments to the customer.]")

    _heading(doc, f"{next(n)}. Lessons learned")
    _text(doc, ctx.get("lessons"), "[What went well, what to do differently on the next project.]")

    tight(doc.add_paragraph(), before=14)
    sign = doc.add_table(rows=3, cols=2)
    sign.style = "Table Grid"
    cell_margins(sign)
    labels = [("Prepared by — Project Manager", "Approved by — Partner"), ("Signature & date", "Signature & date"),
              (f"Name: {ctx.get('pm') or ''}", "Name: ")]
    for r, (a, b) in enumerate(labels):
        for cell, text in zip(sign.rows[r].cells, (a, b)):
            write(cell, text, size=9 if r == 0 else 8.5, bold=r == 0, color=MUTED if r == 1 else None, first=True)
            if r == 1:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    sign.rows[1].height = Pt(56)
    set_widths(sign, [half, half])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
