"""Job work / subcontracting (SAP item category L + ME2O stock at vendor,
GST ITC-04 and Rule 45 return limits). See models/job_work.py for the flow.

Challan = a Material Issue of type "job_work" — its lines post jobwork_out
(Free → With vendor). Everything here reads the challan lines and what came
back against them (approved Material Returns, Job Work Receipt components).
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, timedelta

from fastapi import HTTPException
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.modules.store.models.item import StoreItem
from app.modules.store.models.job_work import StoreJobWorkReceipt, StoreJobWorkReceiptComponent
from app.modules.store.models.location import StoreLocation
from app.modules.store.models.material_issue import StoreMaterialIssue, StoreMaterialIssueItem
from app.modules.store.models.material_return import StoreMaterialReturn, StoreMaterialReturnItem
from app.modules.store.models.stock_transaction import StoreStockTransaction

JOB_WORK_TYPE = "job_work"
RECEIPT_REFERENCE = "job_work_receipt"
# GST Rule 45 / s.143: inputs back within 1 year, capital goods within 3 years.
INPUT_LIMIT_DAYS = 365
CAPITAL_LIMIT_DAYS = 3 * 365
CAPITAL_ITEM_TYPES = {"asset", "tool_equipment"}
# Reminder steps (days out), sent once each by tasks/store_jobwork_reminders.py.
INPUT_REMINDER_DAYS = (300, 330, 360)
CAPITAL_REMINDER_DAYS = (1000, 1080)
_EPS = 1e-9


def is_job_work(issue: StoreMaterialIssue | None) -> bool:
    return bool(issue is not None and issue.issue_type == JOB_WORK_TYPE)


def posted_to_vendor(db: Session, issue: StoreMaterialIssue | None) -> bool:
    """True when the challan's stock went to With-vendor stock (jobwork_out).
    Job-work issues posted before 2026-10-13 took the stock out completely
    ('issue'), so their returns keep posting back to Free as before."""
    if not is_job_work(issue):
        return False
    return db.query(StoreStockTransaction.id).filter(
        StoreStockTransaction.reference_type == "material_issue", StoreStockTransaction.reference_number == issue.issue_number,
        StoreStockTransaction.transaction_type == "jobwork_out").first() is not None


def is_capital(item: StoreItem | None) -> bool:
    return bool(item is not None and item.item_type in CAPITAL_ITEM_TYPES)


def limit_days(item: StoreItem | None) -> int:
    return CAPITAL_LIMIT_DAYS if is_capital(item) else INPUT_LIMIT_DAYS


def generate_receipt_number(db: Session) -> str:
    """JWR-[YEAR]-[NUMBER]."""
    from app.modules.store.service import _lock_number_series
    prefix = f"JWR-{date.today().year}-"
    _lock_number_series(db, prefix)
    last = db.query(func.max(StoreJobWorkReceipt.receipt_number)).filter(StoreJobWorkReceipt.receipt_number.like(f"{prefix}%")).scalar()
    return f"{prefix}{(int(last.rsplit('-', 1)[-1]) + 1) if last else 1:04d}"


# ---------------------------------------------------------------------------
# Challan lines — sent / back / consumed / outstanding
# ---------------------------------------------------------------------------

def challan_lines(db: Session, *, issue_ids: list[int] | None = None, vendor_id: int | None = None,
                  location_ids: set[int] | None = None, today: date | None = None) -> list[dict]:
    """One row per job-work challan line with what is still at the vendor."""
    today = today or date.today()
    q = db.query(StoreMaterialIssue).filter(StoreMaterialIssue.issue_type == JOB_WORK_TYPE)
    if issue_ids is not None:
        q = q.filter(StoreMaterialIssue.id.in_(issue_ids or [-1]))
    if vendor_id is not None:
        q = q.filter(StoreMaterialIssue.vendor_id == vendor_id)
    if location_ids is not None:
        q = q.filter(StoreMaterialIssue.location_id.in_(list(location_ids) or [-1]))
    issues = q.order_by(StoreMaterialIssue.issue_date, StoreMaterialIssue.id).all()
    if not issues:
        return []
    ids = [i.id for i in issues]
    lines = db.query(StoreMaterialIssueItem).filter(StoreMaterialIssueItem.issue_id.in_(ids)).order_by(StoreMaterialIssueItem.id).all()
    items = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_({l.item_id for l in lines} or {-1})).all()}
    locs = {l.id: l.name for l in db.query(StoreLocation).all()}

    # Returns per (issue, item) — approved (stock back) and pending (raised, waiting approval).
    back: dict[tuple[int, int], float] = defaultdict(float)
    pending: dict[tuple[int, int], float] = defaultdict(float)
    for iid, item_id, status, qty in db.query(StoreMaterialReturn.source_issue_id, StoreMaterialReturnItem.item_id, StoreMaterialReturn.status,
                                              func.sum(StoreMaterialReturnItem.quantity)).join(
            StoreMaterialReturn, StoreMaterialReturnItem.return_id == StoreMaterialReturn.id).filter(
            StoreMaterialReturn.source_issue_id.in_(ids), StoreMaterialReturn.status.in_(("approved", "pending_approval"))).group_by(
            StoreMaterialReturn.source_issue_id, StoreMaterialReturnItem.item_id, StoreMaterialReturn.status).all():
        (back if status == "approved" else pending)[(iid, item_id)] += float(qty or 0)
    consumed: dict[int, float] = defaultdict(float)
    scrap: dict[int, float] = defaultdict(float)
    for line_id, c, s in db.query(StoreJobWorkReceiptComponent.issue_item_id, func.sum(StoreJobWorkReceiptComponent.consumed_quantity),
                                  func.sum(StoreJobWorkReceiptComponent.scrap_quantity)).filter(
            StoreJobWorkReceiptComponent.issue_item_id.in_([l.id for l in lines] or [-1])).group_by(StoreJobWorkReceiptComponent.issue_item_id).all():
        consumed[line_id] += float(c or 0)
        scrap[line_id] += float(s or 0)
    # Value sent, from the jobwork_out ledger rows (valuation on), per (issue number, item).
    sent_value: dict[tuple[str, int], float] = defaultdict(float)
    sent_qty: dict[tuple[str, int], float] = defaultdict(float)
    numbers = {i.issue_number for i in issues}
    for ref, item_id, qty, value in db.query(StoreStockTransaction.reference_number, StoreStockTransaction.item_id,
                                             StoreStockTransaction.quantity, StoreStockTransaction.value).filter(
            StoreStockTransaction.reference_type == "material_issue", StoreStockTransaction.reference_number.in_(numbers),
            StoreStockTransaction.transaction_type.in_(("jobwork_out", "issue"))).all():
        sent_qty[(ref, item_id)] += float(qty or 0)
        if value is not None:
            sent_value[(ref, item_id)] += float(value)

    by_issue = {i.id: i for i in issues}
    out = []
    # Returns are recorded per item, so they are spread over the issue's lines
    # of that item in line order.
    left_back = dict(back)
    left_pending = dict(pending)
    for l in lines:
        issue = by_issue[l.issue_id]
        item = items.get(l.item_id)
        key = (issue.id, l.item_id)
        b = min(l.quantity, left_back.get(key, 0.0))
        left_back[key] = left_back.get(key, 0.0) - b
        p = min(l.quantity - b, left_pending.get(key, 0.0))
        left_pending[key] = left_pending.get(key, 0.0) - p
        c, s = consumed.get(l.id, 0.0), scrap.get(l.id, 0.0)
        outstanding = max(0.0, l.quantity - b - c - s)
        vkey = (issue.issue_number, l.item_id)
        rate = None
        if sent_qty.get(vkey) and vkey in sent_value:
            rate = sent_value[vkey] / sent_qty[vkey]
        elif item is not None:
            rate = item.moving_average_cost or item.standard_cost
        days_out = (today - issue.issue_date).days
        lim = limit_days(item)
        deadline = issue.issue_date + timedelta(days=lim)
        status = "closed" if outstanding <= _EPS else ("overdue" if today > deadline else (
            "due_back" if issue.expected_return_date and today > issue.expected_return_date else "open"))
        out.append({
            "issue_id": issue.id, "issue_item_id": l.id, "challan_number": issue.challan_number, "issue_number": issue.issue_number,
            "challan_date": issue.issue_date, "expected_return_date": issue.expected_return_date,
            "vendor_id": issue.vendor_id, "vendor_name": issue.vendor_name, "process": issue.process,
            "output_item_id": issue.output_item_id,
            "location_id": issue.location_id, "location_name": locs.get(issue.location_id),
            "item_id": l.item_id, "item_code": item.item_code if item else None, "item_name": item.item_name if item else None,
            "uom": item.uom if item else None, "hsn_code": item.hsn_sac_code if item else None,
            "gst_rate": item.gst_rate if item else None, "goods_type": "capital_goods" if is_capital(item) else "inputs",
            "batch_number": l.batch_number,
            "quantity_sent": l.quantity, "quantity_returned": round(b, 4), "quantity_return_pending": round(p, 4),
            "quantity_consumed": round(c, 4), "quantity_scrap": round(s, 4), "quantity_outstanding": round(outstanding, 4),
            "rate": round(rate, 4) if rate is not None else None,
            "taxable_value": round(l.quantity * rate, 2) if rate is not None else None,
            "outstanding_value": round(outstanding * rate, 2) if rate is not None else None,
            "days_out": days_out, "limit_days": lim, "return_deadline": deadline,
            "days_to_deadline": (deadline - today).days, "status": status,
        })
    return out


def outstanding_for_receipt(db: Session, issue_item: StoreMaterialIssueItem) -> float:
    """What can still be consumed from a challan line: sent − returned
    (approved or pending) − consumed − scrap."""
    row = next((r for r in challan_lines(db, issue_ids=[issue_item.issue_id]) if r["issue_item_id"] == issue_item.id), None)
    if row is None:
        return 0.0
    return max(0.0, row["quantity_outstanding"] - row["quantity_return_pending"])


def vendor_stock(db: Session, *, location_ids: set[int] | None = None, vendor_id: int | None = None) -> list[dict]:
    """ME2O: what each vendor holds of ours, per item."""
    agg: dict[tuple, dict] = {}
    for r in challan_lines(db, location_ids=location_ids, vendor_id=vendor_id):
        if r["quantity_outstanding"] <= _EPS:
            continue
        k = (r["vendor_id"], r["vendor_name"], r["item_id"])
        a = agg.setdefault(k, {"vendor_id": r["vendor_id"], "vendor_name": r["vendor_name"], "item_id": r["item_id"],
                               "item_code": r["item_code"], "item_name": r["item_name"], "uom": r["uom"], "quantity": 0.0,
                               "value": 0.0, "challans": [], "oldest_challan_date": r["challan_date"], "overdue": False})
        a["quantity"] = round(a["quantity"] + r["quantity_outstanding"], 4)
        a["value"] = round(a["value"] + (r["outstanding_value"] or 0), 2)
        if r["challan_number"] not in a["challans"]:
            a["challans"].append(r["challan_number"])
        a["oldest_challan_date"] = min(a["oldest_challan_date"], r["challan_date"])
        a["overdue"] = a["overdue"] or r["status"] == "overdue"
    return sorted(agg.values(), key=lambda a: ((a["vendor_name"] or "").lower(), a["item_name"] or ""))


# ---------------------------------------------------------------------------
# ITC-04
# ---------------------------------------------------------------------------

def itc04(db: Session, *, date_from: date, date_to: date, location_ids: set[int] | None = None) -> dict:
    """ITC-04 (goods sent to / received back from a job worker) for a period.
    Table 4 = challans issued in the period; Table 5A = goods received back
    (returns approved / receipts posted) in the period against any challan."""
    lines = challan_lines(db, location_ids=location_ids)
    sent = [r for r in lines if date_from <= r["challan_date"] <= date_to]
    by_line = {r["issue_item_id"]: r for r in lines}
    received: list[dict] = []
    issue_ids = list({r["issue_id"] for r in lines})
    if issue_ids:
        rets = db.query(StoreMaterialReturn, StoreMaterialReturnItem).join(
            StoreMaterialReturnItem, StoreMaterialReturnItem.return_id == StoreMaterialReturn.id).filter(
            StoreMaterialReturn.source_issue_id.in_(issue_ids), StoreMaterialReturn.status == "approved",
            StoreMaterialReturn.return_date >= date_from, StoreMaterialReturn.return_date <= date_to).all()
        first_line = {}
        for r in lines:
            first_line.setdefault((r["issue_id"], r["item_id"]), r)
        for ret, it in rets:
            ch = first_line.get((ret.source_issue_id, it.item_id))
            if not ch:
                continue
            received.append({"challan_number": ch["challan_number"], "challan_date": ch["challan_date"], "vendor_name": ch["vendor_name"],
                             "document": ret.return_number, "document_date": ret.return_date, "nature": "Returned as is / processed",
                             "item_code": ch["item_code"], "item_name": ch["item_name"], "hsn_code": ch["hsn_code"], "uom": ch["uom"],
                             "quantity": it.quantity, "taxable_value": round(it.quantity * ch["rate"], 2) if ch["rate"] is not None else None})
        comps = db.query(StoreJobWorkReceipt, StoreJobWorkReceiptComponent).join(
            StoreJobWorkReceiptComponent, StoreJobWorkReceiptComponent.receipt_id == StoreJobWorkReceipt.id).filter(
            StoreJobWorkReceipt.issue_id.in_(issue_ids), StoreJobWorkReceipt.receipt_date >= date_from,
            StoreJobWorkReceipt.receipt_date <= date_to).all()
        out_items = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_({r.output_item_id for r, _ in comps} or {-1})).all()}
        for rec, c in comps:
            ch = by_line.get(c.issue_item_id)
            if not ch:
                continue
            oi = out_items.get(rec.output_item_id)
            for qty, nature in ((c.consumed_quantity, f"Used in {oi.item_name if oi else 'processed item'}"), (c.scrap_quantity, "Process scrap / loss at job worker")):
                if qty and qty > _EPS:
                    received.append({"challan_number": ch["challan_number"], "challan_date": ch["challan_date"], "vendor_name": ch["vendor_name"],
                                     "document": rec.receipt_number, "document_date": rec.receipt_date, "nature": nature,
                                     "item_code": ch["item_code"], "item_name": ch["item_name"], "hsn_code": ch["hsn_code"], "uom": ch["uom"],
                                     "quantity": qty, "taxable_value": round(qty * ch["rate"], 2) if ch["rate"] is not None else None})
    received.sort(key=lambda r: (r["document_date"], r["document"]))
    open_lines = [r for r in lines if r["quantity_outstanding"] > _EPS and r["challan_date"] <= date_to]
    return {"from": date_from, "to": date_to, "sent": sent, "received": received, "outstanding": open_lines,
            "overdue": [r for r in open_lines if r["status"] == "overdue"]}


def itc04_workbook(data: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FDF1E6")
        for r in rows:
            ws.append(r)
        for i, _ in enumerate(header):
            ws.column_dimensions[chr(65 + i) if i < 26 else "A"].width = 16

    def d(v):
        return v.strftime("%d-%m-%Y") if v else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "4 - Sent to job worker"
    sheet(ws, ["Challan no.", "Challan date", "Job worker", "Item code", "Description", "HSN", "UQC", "Quantity", "Taxable value",
               "GST rate %", "Type of goods", "Process"],
          [[r["challan_number"], d(r["challan_date"]), r["vendor_name"], r["item_code"], r["item_name"], r["hsn_code"], r["uom"],
            r["quantity_sent"], r["taxable_value"], r["gst_rate"], "Capital goods" if r["goods_type"] == "capital_goods" else "Inputs",
            r["process"]] for r in data["sent"]])
    ws = wb.create_sheet("5A - Received back")
    sheet(ws, ["Original challan no.", "Original challan date", "Job worker", "Our document", "Document date", "Nature of job / received as",
               "Item code", "Description", "HSN", "UQC", "Quantity", "Taxable value"],
          [[r["challan_number"], d(r["challan_date"]), r["vendor_name"], r["document"], d(r["document_date"]), r["nature"], r["item_code"],
            r["item_name"], r["hsn_code"], r["uom"], r["quantity"], r["taxable_value"]] for r in data["received"]])
    ws = wb.create_sheet("Outstanding at job worker")
    sheet(ws, ["Challan no.", "Challan date", "Job worker", "Item code", "Description", "Sent", "Back", "Consumed", "Scrap",
               "Outstanding", "Outstanding value", "Days out", "Return by (Rule 45)", "Status"],
          [[r["challan_number"], d(r["challan_date"]), r["vendor_name"], r["item_code"], r["item_name"], r["quantity_sent"],
            r["quantity_returned"], r["quantity_consumed"], r["quantity_scrap"], r["quantity_outstanding"], r["outstanding_value"],
            r["days_out"], d(r["return_deadline"]), r["status"]] for r in data["outstanding"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Job work receipt (processed item back, components consumed at the vendor)
# ---------------------------------------------------------------------------

def create_receipt(db: Session, payload, user) -> StoreJobWorkReceipt:
    """Posts: jobwork_consume per component (With vendor → used up, at the
    average), then the processed item into Free (or Quality, under a job work
    inspection lot) at (component value + job charge) ÷ quantity."""
    from app.modules.store.services.stock_ledger import post_stock_transaction
    from app.modules.store.services.valuation import ledger_sort_key, valuation_enabled
    from app.modules.store.services import batches as batch_service

    issue = db.query(StoreMaterialIssue).filter(StoreMaterialIssue.id == payload.issue_id).with_for_update().first()
    if not issue:
        raise HTTPException(status_code=404, detail="Job-work challan not found — it may have been removed. Reload the page.")
    if not is_job_work(issue):
        raise HTTPException(status_code=422, detail=f"Issue {issue.issue_number} isn't a job-work challan — a job work receipt can only "
                                                    "be made against an issue of type Job work.")
    if not posted_to_vendor(db, issue):
        raise HTTPException(status_code=409, detail=f"Challan {issue.challan_number or issue.issue_number} was posted before job-work stock "
                                                    "tracking started, so its material isn't in With-vendor stock — receive the processed item "
                                                    "through a GRN / Return instead.")
    receipt_date = payload.receipt_date or date.today()
    if receipt_date > date.today():
        raise HTTPException(status_code=422, detail=f"Receipt date {receipt_date:%d-%m-%Y} is in the future — record it on the day the goods arrive.")
    if receipt_date < issue.issue_date:
        raise HTTPException(status_code=422, detail=f"Receipt date {receipt_date:%d-%m-%Y} is before the challan date ({issue.issue_date:%d-%m-%Y}).")
    location_id = payload.location_id or issue.location_id
    loc = db.get(StoreLocation, location_id)
    if not loc:
        raise HTTPException(status_code=404, detail="Store not found — pick the store the processed goods are received into.")
    out_item = db.get(StoreItem, payload.output_item_id)
    if not out_item:
        raise HTTPException(status_code=422, detail="Pick the processed item from the Item Master (create it there first if it's new).")
    if payload.output_quantity is None or payload.output_quantity <= 0:
        raise HTTPException(status_code=422, detail=f"Quantity of '{out_item.item_name}' received must be greater than zero.")
    if not payload.components:
        raise HTTPException(status_code=422, detail="Say which challan material the vendor used — at least one component line.")

    lines = {l.id: l for l in db.query(StoreMaterialIssueItem).filter(StoreMaterialIssueItem.issue_id == issue.id).all()}
    comps = []
    for c in payload.components:
        line = lines.get(c.issue_item_id)
        if not line:
            raise HTTPException(status_code=422, detail=f"Challan line #{c.issue_item_id} isn't on {issue.challan_number or issue.issue_number} — reload the page.")
        consumed, scrap = c.consumed_quantity or 0, c.scrap_quantity or 0
        if consumed < 0 or scrap < 0:
            raise HTTPException(status_code=422, detail="Consumed and scrap quantities can't be negative.")
        if consumed + scrap <= _EPS:
            continue
        avail = outstanding_for_receipt(db, line)
        name = db.query(StoreItem.item_name).filter(StoreItem.id == line.item_id).scalar()
        if consumed + scrap - avail > 1e-6:
            raise HTTPException(status_code=409, detail=f"Only {avail:g} of '{name}' is still with {issue.vendor_name or 'the vendor'} on "
                                                        f"challan {issue.challan_number or issue.issue_number}; {consumed + scrap:g} entered as consumed + scrap.")
        comps.append((line, consumed, scrap))
    if not comps:
        raise HTTPException(status_code=422, detail="Enter the quantity of challan material used (consumed and/or process scrap).")

    rate = payload.job_charge_rate
    po_id = None
    if payload.purchase_order_item_id:
        from app.modules.p2p.models.purchase_order import P2PPurchaseOrder, P2PPurchaseOrderItem
        po_item = db.get(P2PPurchaseOrderItem, payload.purchase_order_item_id)
        if not po_item:
            raise HTTPException(status_code=422, detail="The job-work PO line wasn't found — pick it again.")
        po = db.get(P2PPurchaseOrder, po_item.purchase_order_id)
        if issue.vendor_id and po.vendor_id and po.vendor_id != issue.vendor_id:
            raise HTTPException(status_code=422, detail=f"PO {po.po_number} is with {po.vendor_name}, but the challan went to {issue.vendor_name} — pick that vendor's job-work PO.")
        if po.status in ("draft", "cancelled"):
            raise HTTPException(status_code=409, detail=f"PO {po.po_number} is {po.status} — the job charge must come from an issued PO.")
        po_id = po.id
        rate = po_item.unit_price if rate is None else rate
    if rate is not None and rate < 0:
        raise HTTPException(status_code=422, detail="Job charge rate can't be negative.")

    rec = StoreJobWorkReceipt(
        receipt_number=generate_receipt_number(db), issue_id=issue.id, location_id=location_id, vendor_id=issue.vendor_id,
        vendor_name=issue.vendor_name, receipt_date=receipt_date, vendor_challan_number=(payload.vendor_challan_number or "").strip() or None,
        output_item_id=out_item.id, output_quantity=payload.output_quantity, purchase_order_id=po_id,
        purchase_order_item_id=payload.purchase_order_item_id, job_charge_rate=rate, remarks=(payload.remarks or "").strip() or None,
        created_by_id=user.id,
    )
    db.add(rec)
    db.flush()
    ref = dict(reference_type=RECEIPT_REFERENCE, reference_number=rec.receipt_number, transaction_date=receipt_date,
               created_by_id=user.id, vendor_name=issue.vendor_name)
    valued = valuation_enabled(db)
    component_value = 0.0
    try:
        for line, consumed, scrap in sorted(comps, key=lambda x: ledger_sort_key(x[0].item_id, issue.location_id)):
            value = 0.0
            for qty, what in ((consumed, "used in"), (scrap, "process scrap at vendor for")):
                if qty <= _EPS:
                    continue
                txn = post_stock_transaction(db, item_id=line.item_id, location_id=issue.location_id, transaction_type="jobwork_consume",
                                             quantity=qty, batch_number=line.batch_number,
                                             remarks=f"Challan {issue.challan_number or issue.issue_number}: {what} {out_item.item_name} ({rec.receipt_number})", **ref)
                value += txn.value or 0
            component_value += value
            rec.components.append(StoreJobWorkReceiptComponent(issue_item_id=line.id, item_id=line.item_id, consumed_quantity=consumed,
                                                               scrap_quantity=scrap, batch_number=line.batch_number,
                                                               value=round(value, 2) if valued else None))
        db.flush()
        charge = (rate or 0) * payload.output_quantity
        unit_cost = None
        if valued:
            rec.component_value = round(component_value, 2)
            unit_cost = round((component_value + charge) / payload.output_quantity, 4)
            rec.output_unit_cost = unit_cost
        # Output batch: a batch-tracked processed item gets its own batch,
        # carrying the components' heat for traceability.
        if batch_service.is_batch_managed(out_item):
            from app.modules.store.models.batch import StoreBatch
            heats = [b.heat_number for b in db.query(StoreBatch).filter(
                StoreBatch.batch_number.in_([l.batch_number for l, _, _ in comps if l.batch_number] or ["\x00"])).all() if b.heat_number]
            ob = StoreBatch(item_id=out_item.id, batch_number=batch_service.generate_batch_number(db),
                            heat_number=", ".join(dict.fromkeys(heats))[:100] or None, vendor_id=issue.vendor_id, vendor_name=issue.vendor_name,
                            location_id=location_id, received_qty=payload.output_quantity, received_date=receipt_date,
                            remarks=f"Job work receipt {rec.receipt_number} (challan {issue.challan_number or issue.issue_number})",
                            status="active", created_by_id=user.id)
            db.add(ob)
            db.flush()
            rec.output_batch_number = ob.batch_number
        from app.modules.quality.services.lots import needs_job_work_inspection, create_lot, LotError
        if needs_job_work_inspection(db, out_item):
            try:
                lot = create_lot(db, inspection_type="job_work", source_type="job_work_receipt", item=out_item,
                                 quantity=payload.output_quantity, location_id=location_id, user_id=user.id, source_id=rec.id,
                                 source_reference=f"{rec.receipt_number} (job work {issue.issue_number})", batch_number=rec.output_batch_number,
                                 vendor_name=issue.vendor_name, receive_into_quality=True, lot_date=receipt_date, unit_cost=unit_cost)
            except LotError as e:
                raise ValueError(str(e))
            rec.inspection_lot_id = lot.id
        else:
            post_stock_transaction(db, item_id=out_item.id, location_id=location_id, transaction_type="receipt",
                                   quantity=payload.output_quantity, batch_number=rec.output_batch_number,
                                   remarks=f"Processed by {issue.vendor_name or 'job worker'} on challan {issue.challan_number or issue.issue_number}",
                                   unit_cost=unit_cost, cost_source="job_work" if unit_cost is not None else None,
                                   **ref)
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=409, detail=f"The job work receipt couldn't be posted: {e}")
    db.flush()
    return rec


def receipt_dict(db: Session, rec: StoreJobWorkReceipt) -> dict:
    issue = db.get(StoreMaterialIssue, rec.issue_id)
    out_item = db.get(StoreItem, rec.output_item_id)
    items = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_({c.item_id for c in rec.components} or {-1})).all()}
    lot_number = None
    if rec.inspection_lot_id:
        from app.modules.quality.models.inspection_lot import QualityInspectionLot
        lot_number = db.query(QualityInspectionLot.lot_number).filter(QualityInspectionLot.id == rec.inspection_lot_id).scalar()
    return {
        "id": rec.id, "receipt_number": rec.receipt_number, "issue_id": rec.issue_id,
        "issue_number": issue.issue_number if issue else None, "challan_number": issue.challan_number if issue else None,
        "location_id": rec.location_id, "location_name": db.query(StoreLocation.name).filter(StoreLocation.id == rec.location_id).scalar(),
        "vendor_id": rec.vendor_id, "vendor_name": rec.vendor_name, "receipt_date": rec.receipt_date,
        "vendor_challan_number": rec.vendor_challan_number, "output_item_id": rec.output_item_id,
        "output_item_code": out_item.item_code if out_item else None, "output_item_name": out_item.item_name if out_item else None,
        "uom": out_item.uom if out_item else None, "output_quantity": rec.output_quantity, "output_batch_number": rec.output_batch_number,
        "purchase_order_id": rec.purchase_order_id, "job_charge_rate": rec.job_charge_rate, "component_value": rec.component_value,
        "output_unit_cost": rec.output_unit_cost, "inspection_lot_id": rec.inspection_lot_id, "inspection_lot_number": lot_number,
        "remarks": rec.remarks, "created_at": rec.created_at,
        "components": [{"id": c.id, "issue_item_id": c.issue_item_id, "item_id": c.item_id,
                        "item_code": items[c.item_id].item_code if c.item_id in items else None,
                        "item_name": items[c.item_id].item_name if c.item_id in items else None,
                        "consumed_quantity": c.consumed_quantity, "scrap_quantity": c.scrap_quantity,
                        "batch_number": c.batch_number, "value": c.value} for c in rec.components],
    }


def reminder_step(days_out: int, capital: bool, already: int | None) -> int | None:
    """The highest reminder step reached and not yet sent, else None."""
    steps = CAPITAL_REMINDER_DAYS if capital else INPUT_REMINDER_DAYS
    due = [s for s in steps if days_out >= s and (already or 0) < s]
    return max(due) if due else None

