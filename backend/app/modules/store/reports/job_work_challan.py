"""Job work delivery challan — the managed Word document of record
`store_job_work_challan` (company letterhead; "Save & Edit in Word").

GST: goods sent to a job worker without tax under Section 143 of the CGST
Act travel on a delivery challan (Rule 55) with the consignor / consignee
GSTIN, HSN, quantity and taxable value, and must come back within one year
(inputs) / three years (capital goods). The same challan number is what the
ITC-04 return reports."""
from __future__ import annotations

import io
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.main.models.user import User
from app.modules.main.services import managed_documents
from app.modules.store.models.material_issue import StoreMaterialIssue
from app.utils.docx_letterhead import BODY, HEADER, MUTED, USABLE_IN, cell_margins, set_widths, tight, write

CHALLAN_ENTITY = "store_job_work_challan"


def challan_context(db: Session, issue: StoreMaterialIssue) -> dict[str, Any]:
    from app.modules.accounts.models.vendor import Vendor
    from app.modules.p2p.reports.purchase_order import _company, _state
    from app.modules.store.issue_forms import JOB_WORK_PROCESSES
    from app.modules.store.services.job_work import challan_lines, is_capital
    from app.modules.store.models.item import StoreItem

    company, caddr = _company(db)
    vendor = db.get(Vendor, issue.vendor_id) if issue.vendor_id else None
    c_state, c_code = _state(company.gstin if company else None, caddr.state if caddr else None)
    v_state, v_code = _state(vendor.gstin if vendor else None, vendor.state if vendor else None)
    rows = challan_lines(db, issue_ids=[issue.id])
    items = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_({r["item_id"] for r in rows} or {-1})).all()}
    capital = any(is_capital(items.get(r["item_id"])) for r in rows)
    lines = [{"name": r["item_name"] or "", "code": r["item_code"] or "", "hsn": r["hsn_code"] or "", "qty": r["quantity_sent"],
              "uom": r["uom"] or "", "rate": r["rate"], "value": r["taxable_value"], "batch": r["batch_number"]} for r in rows]
    issued_by = db.get(User, issue.issued_by_id) if issue.issued_by_id else None
    return {
        "challan_number": issue.challan_number or issue.issue_number, "challan_date": issue.issue_date.strftime("%d-%m-%Y"),
        "issue_number": issue.issue_number,
        "company_name": (company.legal_name or company.name) if company else "Premnath Engineering Works",
        "company_address": [x for x in ((caddr.address_line1 if caddr else None), (caddr.address_line2 if caddr else None),
                                        " ".join(x for x in ((caddr.city if caddr else None), (caddr.pincode if caddr else None)) if x)) if x],
        "company_gstin": company.gstin if company else None, "company_state": c_state, "company_state_code": c_code,
        "vendor_name": issue.vendor_name or (vendor.name if vendor else ""),
        "vendor_address": [x for x in ((vendor.address if vendor else None),
                                       " ".join(x for x in ((vendor.city if vendor else None), (vendor.pin_code if vendor else None)) if x)) if x],
        "vendor_gstin": vendor.gstin if vendor else None, "vendor_state": v_state, "vendor_state_code": v_code,
        "process": JOB_WORK_PROCESSES.get(issue.process, issue.process or ""),
        "expected_return": issue.expected_return_date.strftime("%d-%m-%Y") if issue.expected_return_date else "",
        "vehicle": issue.vehicle_number or "", "lines": lines,
        "total_value": round(sum(ln["value"] or 0 for ln in lines), 2),
        "limit_text": "three years (capital goods)" if capital else "one year (inputs)",
        "remarks": issue.remarks or "", "issued_by": (issued_by.name or issued_by.email) if issued_by else "",
    }


def build_challan_docx(ctx: dict[str, Any]) -> bytes:
    from app.modules.p2p.reports.purchase_order import _borders, _label_value, _party, _n

    doc = new_doc()
    tight(write(doc, "DELIVERY CHALLAN — JOB WORK", size=HEADER + 4, bold=True, align="center"), after=2)
    tight(write(doc, "Goods sent for job work without payment of tax under Section 143 of the CGST Act, 2017 (Rule 55)",
                size=BODY - 1, color=MUTED, align="center"), after=6)

    head = doc.add_table(rows=3, cols=3)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(head)
    cell_margins(head, left=80, right=80, top=40, bottom=40)
    set_widths(head, [USABLE_IN * 0.5, USABLE_IN * 0.25, USABLE_IN * 0.25])
    _label_value(head.cell(0, 1), "Challan No.", ctx["challan_number"])
    _label_value(head.cell(0, 2), "Dated", ctx["challan_date"])
    _label_value(head.cell(1, 1), "Nature of process", ctx["process"])
    _label_value(head.cell(1, 2), "Expected back by", ctx["expected_return"])
    _label_value(head.cell(2, 1), "Vehicle / courier no.", ctx["vehicle"])
    _label_value(head.cell(2, 2), "Store reference", ctx["issue_number"])
    consignor = head.cell(0, 0).merge(head.cell(1, 0))
    _party(consignor, "Consignor (Principal)", ctx["company_name"], ctx["company_address"], ctx["company_gstin"],
           ctx["company_state"], ctx["company_state_code"], first=True)
    _party(head.cell(2, 0), "Consignee (Job worker)", ctx["vendor_name"], ctx["vendor_address"], ctx["vendor_gstin"],
           ctx["vendor_state"], ctx["vendor_state_code"], first=True)
    for c in (consignor, head.cell(2, 0)):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    write(doc, "", after=4)
    widths = [0.35, 2.75, 0.75, 0.75, 0.5, 0.85, 1.05]
    scale = USABLE_IN / sum(widths)
    grid = doc.add_table(rows=1, cols=7)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(grid)
    cell_margins(grid, left=60, right=60, top=30, bottom=30)
    set_widths(grid, [w * scale for w in widths])
    for cell, h in zip(grid.rows[0].cells, ["Sl", "Description of goods", "HSN", "Quantity", "Unit", "Rate", "Taxable value"]):
        write(cell, h, size=BODY - 0.5, bold=True, align="center", first=True)
    for n, ln in enumerate(ctx["lines"], start=1):
        cells = grid.add_row().cells
        write(cells[0], str(n), align="center", first=True)
        write(cells[1], ln["name"], bold=True, first=True)
        if ln["code"] or ln["batch"]:
            write(cells[1], " · ".join(x for x in (ln["code"], f"Batch {ln['batch']}" if ln["batch"] else None) if x), size=BODY - 1, color=MUTED)
        write(cells[2], ln["hsn"], align="center", first=True)
        write(cells[3], f"{float(ln['qty']):g}", align="right", first=True)
        write(cells[4], ln["uom"], align="center", first=True)
        write(cells[5], _n(ln["rate"]) if ln["rate"] is not None else "", align="right", first=True)
        write(cells[6], _n(ln["value"]) if ln["value"] is not None else "", align="right", first=True)
    total = grid.add_row().cells
    write(total[1], "Total taxable value", bold=True, align="right", first=True)
    write(total[6], _n(ctx["total_value"]), bold=True, align="right", first=True)

    write(doc, "", after=4)
    for i, text in enumerate([
        "The goods are sent for job work and remain the property of the consignor. No sale is involved.",
        f"The goods (or the processed goods) must be returned within {ctx['limit_text']} of this challan, "
        "failing which the supply is deemed to have been made on the day they were sent (Section 143).",
        "Please quote this challan number and date on your delivery challan / invoice when returning the goods.",
    ], start=1):
        write(doc, f"{i}. {text}", size=BODY - 0.5)
    if ctx["remarks"]:
        write(doc, f"Remarks: {ctx['remarks']}", size=BODY - 0.5, after=2)

    write(doc, "", after=14)
    sign = doc.add_table(rows=1, cols=2)
    set_widths(sign, [USABLE_IN / 2, USABLE_IN / 2])
    write(sign.cell(0, 0), "Received the above goods in good condition", size=BODY - 0.5, first=True)
    write(sign.cell(0, 0), "", after=18)
    write(sign.cell(0, 0), "Job worker's signature & stamp", size=BODY - 0.5, color=MUTED)
    write(sign.cell(0, 1), f"for {ctx['company_name']}", bold=True, align="right", first=True)
    write(sign.cell(0, 1), "", after=18)
    write(sign.cell(0, 1), "Authorised Signatory" + (f" (prepared by {ctx['issued_by']})" if ctx["issued_by"] else ""),
          size=BODY - 0.5, color=MUTED, align="right")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def new_doc():
    from app.utils.docx_letterhead import new_letterhead_document
    return new_letterhead_document()


def challan_document_target(db: Session, issue_id: int, user: User, modify: bool) -> managed_documents.DocumentTarget:
    issue = db.get(StoreMaterialIssue, issue_id)
    if not issue or issue.issue_type != "job_work":
        raise HTTPException(status_code=404, detail="Job-work challan not found — only issues of type Job work have a challan document.")
    from app.core.permission_registry import can_view_tab
    from app.core.permissions import ensure_tab_action
    from app.modules.store.services.plant_access import ensure_location_allowed
    if modify:
        ensure_tab_action(user, "store", "issues", "edit")
    elif not can_view_tab(user, "store", "issues"):
        raise HTTPException(status_code=403, detail="You can't open job-work challans — your role has no access to Store › Issues. "
                                                    "Ask an admin to add it in Organization › Role & Permissions.")
    ensure_location_allowed(db, user, issue.location_id)
    number = (issue.challan_number or issue.issue_number).replace("/", "-")
    return managed_documents.DocumentTarget(
        entity_id=issue.id, revision=0, title=f"Job work challan {issue.challan_number or issue.issue_number}",
        folder=managed_documents.sharepoint_folder("Store-JobWorkChallans", number),
        filename_stem=managed_documents.safe_filename(f"JW-Challan-{number}"),
        build_docx=lambda: build_challan_docx(challan_context(db, issue)),
    )


managed_documents.register(managed_documents.DocumentAdapter(entity_type=CHALLAN_ENTITY, app="store", load=challan_document_target))
