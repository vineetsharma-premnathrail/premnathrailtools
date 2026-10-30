from datetime import date
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy.orm import Session, selectinload

from app.core.data_scope import apply_scope, ensure_in_scope
from app.core.permissions import ensure_feature_action, require_app_access, require_tab_action
from app.db.session import get_db
from app.modules.main.models.user import User
from app.modules.store.access import ensure_any_store_action
from app.modules.organization.models.department import Department
from app.modules.store.services.plant_access import ensure_location_allowed
from app.modules.store.models.item import StoreItem
from app.modules.store.models.location import StoreLocation
from app.modules.store.models.material_issue import StoreMaterialIssue, StoreMaterialIssueItem
from app.modules.store.schemas.material_issue import StoreMaterialIssueCreate, StoreMaterialIssueResponse
from app.modules.store.service import generate_material_issue_number
from app.modules.store.services.stock_ledger import post_stock_transaction
from app.modules.store.routes.doc_types import doc_type_labels, get_doc_type
from app.modules.store.issue_forms import ISSUE_FORMS, public_standard, validate_issue
from app.modules.store.services import project_stock

ISSUE_LINE_COLUMNS = ["item_code", "quantity", "batch_number", "remarks"]

router = APIRouter(prefix="/store/material-issues", tags=["Store"])


def _to_response(db: Session, issue: StoreMaterialIssue) -> StoreMaterialIssueResponse:
    location = db.query(StoreLocation).filter(StoreLocation.id == issue.location_id).first()
    department = db.query(Department).filter(Department.id == issue.department_id).first() if issue.department_id else None
    user_ids = {issue.requested_by_id, issue.issued_by_id} - {None}
    users = {u.id: u for u in db.query(User).filter(User.id.in_(user_ids)).all()} if user_ids else {}
    items_by_id = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_([it.item_id for it in issue.items])).all()} if issue.items else {}

    resp = StoreMaterialIssueResponse.model_validate(issue)
    resp.location_name = location.name if location else None
    resp.issue_type_label = doc_type_labels(db, "issue").get(issue.issue_type) if issue.issue_type else None
    resp.department_name = department.name if department else None
    if issue.p2p_request_id:
        from app.modules.p2p.models.p2p_request import P2PRequest
        pr = db.query(P2PRequest).filter(P2PRequest.id == issue.p2p_request_id).first()
        if pr:
            resp.p2p_number = pr.p2p_number
            resp.department_name = resp.department_name or pr.department
    if issue.requested_by_id and issue.requested_by_id in users:
        resp.requested_by_name = users[issue.requested_by_id].name or users[issue.requested_by_id].email
    if issue.issued_by_id and issue.issued_by_id in users:
        resp.issued_by_name = users[issue.issued_by_id].name or users[issue.issued_by_id].email
    from app.modules.store.models.bin import StoreBin
    bin_ids = {l.bin_id for l in issue.items} - {None}
    bin_codes = {b.id: b.code for b in db.query(StoreBin).filter(StoreBin.id.in_(bin_ids)).all()} if bin_ids else {}
    for line, item_resp in zip(issue.items, resp.items):
        item = items_by_id.get(line.item_id)
        item_resp.bin_code = bin_codes.get(line.bin_id)
        item_resp.item_code = item.item_code if item else None
        item_resp.item_name = item.item_name if item else None
        item_resp.uom = item.uom if item else None
        item_resp.element_display = project_stock.element_label(db, line.order_project_element_id)
    info = project_stock.element_info(db, issue.order_project_element_id, issue.order_project_id)
    if info:
        resp.op_number, resp.element_no, resp.element_label = info["op_number"], info.get("element_no"), info.get("element_label")   # project-level issues carry the project code only
    elif issue.order_project_id:
        from app.modules.projects.models.project import PmProject
        p = db.get(PmProject, issue.order_project_id)
        resp.op_number = p.project_code if p else None
    if issue.erp_project_id:
        from app.modules.erp.models.project import Project
        project = db.query(Project).filter(Project.id == issue.erp_project_id).first()
        resp.erp_project_label = _project_label(project) if project else None
    resp.against_label = _against_label(db, issue, resp)
    return resp


def _against_label(db: Session, issue: StoreMaterialIssue, resp) -> str | None:
    """One line saying what the issue is against — shown on lists, the
    detail view and the challan."""
    from app.modules.crm.models.organization import Organization
    from app.modules.maintenance.models.asset import MaintenanceAsset
    from app.modules.maintenance.models.work_order import MaintenanceWorkOrder
    from app.modules.quality.models.inspection_lot import QualityInspectionLot
    from app.modules.quality.models.ncr import QualityNcr
    from app.modules.store.issue_forms import JOB_WORK_PROCESSES
    parts: list[str] = []
    if issue.maintenance_work_order_id and (w := db.get(MaintenanceWorkOrder, issue.maintenance_work_order_id)):
        parts.append(f"Work order {w.wo_number}")
    if issue.asset_id and (a := db.get(MaintenanceAsset, issue.asset_id)):
        parts.append(f"Machine {a.asset_code} {a.name}")
    if issue.ncr_id and (n := db.get(QualityNcr, issue.ncr_id)):
        parts.append(f"NCR {n.ncr_number}")
    if issue.inspection_lot_id and (lot := db.get(QualityInspectionLot, issue.inspection_lot_id)):
        parts.append(f"Lot {lot.lot_number}")
    if issue.customer_id and (o := db.get(Organization, issue.customer_id)):
        parts.append(f"Customer {o.name}")
    if issue.vendor_name:
        parts.append(f"Vendor {issue.vendor_name}" + (f" · {JOB_WORK_PROCESSES.get(issue.process, issue.process)}" if issue.process else ""))
    if getattr(issue, "service_request_id", None):
        from app.modules.erp.models.service_request import ServiceRequest
        if (sr := db.get(ServiceRequest, issue.service_request_id)):
            parts.append(f"Service request {sr.request_number}")
    if resp.erp_project_label:
        parts.append(resp.erp_project_label)
    if issue.order_project_element_id:
        parts.append(project_stock.element_label(db, issue.order_project_element_id))
    elif resp.op_number:
        parts.append(f"Project {resp.op_number}")
    if issue.issue_type == "general" and resp.department_name:
        parts.append(f"Department {resp.department_name}")
    return " · ".join(parts) or None


def _project_label(p) -> str:
    return " — ".join(x for x in (p.serial_number, p.model_name, p.client_company, f"PO {p.po_number}" if p.po_number else None) if x)


@router.get("", response_model=list[StoreMaterialIssueResponse])
async def list_material_issues(
    location_id: int | None = Query(None),
    department_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("store")),
):
    # The New Return form picks the issue being returned from this list too.
    ensure_any_store_action(user, ("issues", "view"), ("returns", "create"))
    query = db.query(StoreMaterialIssue).options(selectinload(StoreMaterialIssue.items))
    query = apply_scope(query, user, StoreMaterialIssue)
    if location_id is not None:
        query = query.filter(StoreMaterialIssue.location_id == location_id)
    if department_id is not None:
        query = query.filter(StoreMaterialIssue.department_id == department_id)
    issues = query.order_by(StoreMaterialIssue.created_at.desc()).all()
    return [_to_response(db, i) for i in issues]


@router.get("/form-standard")
async def issue_form_standard(_user: User = Depends(require_tab_action("store", "issues", "create"))):
    """What each issue type asks for — the New Issue form renders from this."""
    return public_standard()


@router.get("/lookups/{kind}")
async def issue_lookups(kind: str, db: Session = Depends(get_db), _user: User = Depends(require_tab_action("store", "issues", "create"))):
    """Pick-lists for the "against" fields, readable by Store users without
    the owning module (departments, maintenance work orders / assets, NCRs,
    inspection lots, customers, vendors)."""
    from app.modules.accounts.models.vendor import Vendor
    from app.modules.crm.models.organization import Organization
    from app.modules.maintenance.models.asset import MaintenanceAsset
    from app.modules.maintenance.models.work_order import MaintenanceWorkOrder
    from app.modules.quality.models.inspection_lot import QualityInspectionLot
    from app.modules.quality.models.ncr import QualityNcr
    if kind == "departments":
        return [{"id": d.id, "label": d.name} for d in db.query(Department).order_by(Department.name).all()]
    if kind == "maintenance_work_orders":
        rows = db.query(MaintenanceWorkOrder).filter(MaintenanceWorkOrder.status.notin_(("closed", "cancelled"))).order_by(MaintenanceWorkOrder.id.desc()).limit(500).all()
        return [{"id": w.id, "label": f"{w.wo_number} — {w.title}" if getattr(w, "title", None) else w.wo_number} for w in rows]
    if kind == "assets":
        rows = db.query(MaintenanceAsset).filter(MaintenanceAsset.is_deleted == False).order_by(MaintenanceAsset.asset_code).all()  # noqa: E712
        return [{"id": a.id, "label": f"{a.asset_code} — {a.name}"} for a in rows]
    if kind == "ncrs":
        rows = db.query(QualityNcr).filter(QualityNcr.is_deleted == False, QualityNcr.status != "closed").order_by(QualityNcr.id.desc()).limit(500).all()  # noqa: E712
        return [{"id": n.id, "label": f"{n.ncr_number} — {getattr(n, 'item_name', None) or ''}".rstrip(" —")} for n in rows]
    if kind == "inspection_lots":
        rows = db.query(QualityInspectionLot).filter(QualityInspectionLot.status.notin_(("closed", "cancelled"))).order_by(QualityInspectionLot.id.desc()).limit(500).all()
        return [{"id": l.id, "label": f"{l.lot_number} — {l.item_name}"} for l in rows]
    if kind == "customers":
        rows = db.query(Organization).filter(Organization.is_deleted == False).order_by(Organization.name).all()  # noqa: E712
        return [{"id": o.id, "label": o.name} for o in rows]
    if kind == "vendors":
        rows = db.query(Vendor).filter(Vendor.status == "active").order_by(Vendor.name).all()
        return [{"id": v.id, "label": f"{v.name} ({v.code})"} for v in rows]
    if kind == "service_requests":
        from app.modules.erp.models.project import Project
        from app.modules.erp.models.service_request import ServiceRequest
        rows = db.query(ServiceRequest, Project).join(Project, Project.id == ServiceRequest.project_id).filter(
            ServiceRequest.is_deleted == False, ServiceRequest.status.notin_(("closed", "cancelled")),  # noqa: E712
        ).order_by(ServiceRequest.id.desc()).limit(500).all()
        return [{"id": sr.id, "label": " — ".join(x for x in (sr.request_number, p.serial_number, p.client_company, sr.issue_title) if x)}
                for sr, p in rows]
    raise HTTPException(status_code=404, detail=f"Unknown pick-list '{kind}'.")


@router.get("/dispatch-projects")
async def dispatch_project_options(
    db: Session = Depends(get_db),
    _user: User = Depends(require_tab_action("store", "issues", "create")),
):
    """Customer machines / orders (ERP) a dispatch issue can go against,
    with their Quality inspection requirement (rule 8)."""
    from app.modules.erp.models.project import Project
    rows = db.query(Project).filter(Project.is_deleted == False).order_by(Project.id.desc()).limit(1000).all()  # noqa: E712
    return [{"id": p.id, "label": _project_label(p), "third_party_inspection_required": bool(p.third_party_inspection_required),
             "inspection_agency": p.inspection_agency} for p in rows]


@router.get("/import-lines/template")
async def issue_lines_template(_user: User = Depends(require_tab_action("store", "issues", "create"))):
    """Excel template for loading a material issue's item lines."""
    import io
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Issue lines"
    ws.append(ISSUE_LINE_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FDF1E6")
    ws.append(["MT-ELE-0017", 5, "", "Example line — replace with your items"])
    for col, w in zip("ABCD", (20, 12, 18, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="material_issue_lines_template.xlsx"'})


@router.post("/import-lines")
async def import_issue_lines(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _user: User = Depends(require_tab_action("store", "issues", "create")),
):
    """Reads item lines (item_code, quantity, batch_number, remarks) from an
    Excel / CSV file into the New Issue form — nothing is posted. Rows with
    the same item + batch are added together. Unknown codes / bad
    quantities come back per row so the user can fix the file."""
    from app.modules.store.routes.item_import import _read_rows
    rows = [r for r in _read_rows(file.filename, await file.read()) if any(r.values())]
    if not rows:
        raise HTTPException(status_code=422, detail="The file has no item rows — fill in item_code and quantity under the header row (download the template).")
    if len(rows) > 1000:
        raise HTTPException(status_code=422, detail=f"The file has {len(rows)} rows — one issue can hold at most 1000 lines; split the file.")
    missing = [c for c in ("item_code", "quantity") if c not in rows[0]]
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing column(s) {', '.join(missing)} — download the template and keep its header row.")
    codes = {(r.get("item_code") or "").strip().lower() for r in rows}
    items = {i.item_code.strip().lower(): i for i in db.query(StoreItem).filter(StoreItem.item_code.isnot(None)).all()
             if i.item_code and i.item_code.strip().lower() in codes}
    lines: dict[tuple[int, str], dict] = {}
    errors = []
    for idx, r in enumerate(rows, start=2):
        code = (r.get("item_code") or "").strip()
        item = items.get(code.lower())
        if not code:
            errors.append({"row": idx, "item_code": "", "reason": "Item code is empty."}); continue
        if not item:
            errors.append({"row": idx, "item_code": code, "reason": f"'{code}' isn't in the Item Master — check the code."}); continue
        try:
            qty = float(str(r.get("quantity") or "").replace(",", ""))
        except ValueError:
            errors.append({"row": idx, "item_code": code, "reason": f"Quantity '{r.get('quantity')}' isn't a number."}); continue
        if qty <= 0:
            errors.append({"row": idx, "item_code": code, "reason": "Quantity must be greater than zero."}); continue
        batch = (r.get("batch_number") or r.get("batch") or "").strip()
        key = (item.id, batch)
        if key in lines:
            lines[key]["quantity"] += qty
        else:
            lines[key] = {"item_id": item.id, "item_code": item.item_code, "item_name": item.item_name, "uom": item.uom,
                          "quantity": qty, "batch_number": batch or None, "remarks": (r.get("remarks") or "").strip() or None}
    return {"lines": list(lines.values()), "errors": errors}


@router.get("/{issue_id}", response_model=StoreMaterialIssueResponse)
async def get_material_issue(
    issue_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("store")),
):
    issue = db.query(StoreMaterialIssue).options(selectinload(StoreMaterialIssue.items)).filter(StoreMaterialIssue.id == issue_id).first()
    if not issue:
        raise HTTPException(status_code=404, detail="Material issue not found")
    ensure_any_store_action(user, ("issues", "view"), ("returns", "create"))
    ensure_in_scope(db, user, issue)
    return _to_response(db, issue)


@router.post("", response_model=StoreMaterialIssueResponse)
async def create_material_issue(
    payload: StoreMaterialIssueCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_tab_action("store", "issues", "create")),
):
    if not payload.items:
        raise HTTPException(status_code=422, detail="At least one item is required")
    location = db.query(StoreLocation).filter(StoreLocation.id == payload.location_id).first()
    if not location:
        raise HTTPException(status_code=404, detail="Store not found")
    ensure_location_allowed(db, user, location.id)
    items_by_id = {i.id: i for i in db.query(StoreItem).filter(StoreItem.id.in_([p.item_id for p in payload.items])).all()}
    for p in payload.items:
        if p.item_id not in items_by_id:
            raise HTTPException(status_code=422, detail=f"Item {p.item_id} not found")
        if p.quantity <= 0:
            raise HTTPException(status_code=422, detail=f"Quantity for '{items_by_id[p.item_id].item_name}' must be greater than zero")
    # Optional at the API level (issues raised from a PR have no type);
    # the Store issue form always sends one. The type decides which fields
    # the issue must carry (store/issue_forms.py — the dynamic form standard).
    issue_type_row = get_doc_type(db, "issue", payload.issue_type, field="issue type") if payload.issue_type else None
    from app.modules.store.services.job_work import JOB_WORK_TYPE
    if issue_type_row and issue_type_row.value == JOB_WORK_TYPE:
        # Sending material out to a job worker: Issues › Job work › Create.
        ensure_feature_action(user, "store", "issues", "job_work", "create")
    issue_date = payload.issue_date or date.today()
    if issue_date > date.today():
        raise HTTPException(status_code=422, detail=f"Issue date {issue_date.isoformat()} is in the future — stock can only be issued for today or earlier.")
    form = validate_issue(db, issue_type_row.value, issue_type_row.label, payload.model_dump(), issue_date) if issue_type_row else {}
    spec = ISSUE_FORMS.get(issue_type_row.value, {}) if issue_type_row else {}
    if spec.get("remarks_label") and not (payload.remarks or "").strip():
        raise HTTPException(status_code=422, detail=f"Issue type '{issue_type_row.label}' needs Remarks: {spec['remarks_label'].lower()}.")
    if spec.get("manager_only") and not (user.role == "admin" or getattr(user, "is_azure_admin", False)
                                         or getattr(user, "is_store_manager", False)
                                         or (location.manager_user_id and location.manager_user_id == user.id)):
        raise HTTPException(status_code=403, detail=f"'{issue_type_row.label}' writes stock off, so only a Store Manager (or {location.name}'s "
                                                    "manager) or an admin can post it — ask them, or ask an admin to tick 'Store Manager' on your user.")
    if form.get("challan_number") and len(form["challan_number"]) > 50:
        raise HTTPException(status_code=422, detail="Challan number is too long — keep it to 50 characters or fewer.")
    # Job work: the vendor comes from the Vendor master; its name is kept on
    # the issue (a job-work Return reads it to raise the job-work inspection).
    vendor_name = None
    if form.get("vendor_id"):
        from app.modules.accounts.models.vendor import Vendor
        vendor_name = db.get(Vendor, form["vendor_id"]).name
    erp_project_id = form.get("erp_project_id")
    if issue_type_row and issue_type_row.value == "job_work":
        # Items to receive back — several allowed (plate sent, flange + ring back).
        # None given = the same items come back (plating, heat treatment…).
        back: dict[int, float | None] = {}
        for b in payload.output_items or ([] if not payload.output_item_id else [{"item_id": payload.output_item_id}]):
            b = b if isinstance(b, dict) else b.model_dump()
            if b.get("quantity") is not None and b["quantity"] <= 0:
                raise HTTPException(status_code=422, detail="Items to receive back: quantity must be greater than zero, or leave it blank if not known yet.")
            if db.get(StoreItem, b["item_id"]) is None:
                raise HTTPException(status_code=422, detail=f"Item to receive back #{b['item_id']} wasn't found — pick it again from the list.")
            back[b["item_id"]] = (back.get(b["item_id"]) or 0) + b["quantity"] if b.get("quantity") is not None else back.get(b["item_id"])
        if not back:
            for line in payload.items:
                back[line.item_id] = (back.get(line.item_id) or 0) + line.quantity
        form["output_items"] = [{"item_id": i, "quantity": q} for i, q in back.items()]
        form["output_item_id"] = next(iter(back))

    # Project unit: the vehicle / lot the material goes to — mandatory for an
    # RRV order, filled from the work order / build / delivery row where it can be.
    op_link = project_stock.resolve_issue_project(db, issue_type_row.value if issue_type_row else None, payload)
    override_reason = (payload.project_stock_override_reason or "").strip() or None

    if issue_type_row and issue_type_row.value == "dispatch":
        # Quality rule 8: FG leave only with a decided final inspection and,
        # where the contract needs it, an accepted third-party certificate.
        from app.modules.quality.services.dispatch import ensure_dispatch_allowed, DispatchBlocked
        try:
            ensure_dispatch_allowed(db, lines=[p.model_dump() for p in payload.items], erp_project_id=erp_project_id,
                                    order_project_element_id=op_link["element_id"])
        except DispatchBlocked as e:
            raise HTTPException(status_code=409, detail=f"Dispatch blocked by Quality — {e}")
        # Finance credit check (FD32/UKM): over the credit limit or overdue
        # beyond credit days -> on hold until a Finance approver releases it.
        from app.modules.accounts.services.credit import ensure_dispatch_credit_ok, CreditBlocked
        try:
            ensure_dispatch_credit_ok(db, sales_order_id=payload.sales_order_id, items=[p.model_dump() for p in payload.items],
                                      erp_project_id=erp_project_id)
        except CreditBlocked as e:
            raise HTTPException(status_code=409, detail=f"Dispatch blocked by Finance — {e}")

    # WM bin + serial numbers (IQ01) per line, checked before anything posts.
    from app.modules.store.services.bins import BinError, ensure_pick_possible, validate_bin
    from app.modules.store.services import serials as serial_svc
    picked_units: list[list] = []
    taken_from: dict[tuple[int, int | None], float] = {}
    taken_ids: set[int] = set()
    try:
        for p in payload.items:
            it = items_by_id[p.item_id]
            validate_bin(db, p.bin_id, location.id, what=f"'{it.item_name}'")
            ensure_pick_possible(db, item=it, location_id=location.id, qty=p.quantity, bin_id=p.bin_id,
                                 already=taken_from.get((it.id, p.bin_id), 0.0))
            taken_from[(it.id, p.bin_id)] = taken_from.get((it.id, p.bin_id), 0.0) + p.quantity
            units = serial_svc.pick(db, it, location.id, p.quantity, p.serials, bin_id=p.bin_id, exclude_ids=taken_ids)
            taken_ids |= {u.id for u in units}
            picked_units.append(units)
    except (BinError, serial_svc.SerialError) as e:
        raise HTTPException(status_code=422, detail=str(e))

    issue = StoreMaterialIssue(
        issue_number=generate_material_issue_number(db),
        location_id=payload.location_id,
        requested_by_id=payload.requested_by_id,
        department_id=form.get("department_id", payload.department_id),
        project_or_work_order=form.get("project_or_work_order", payload.project_or_work_order),
        issue_type=payload.issue_type or None,
        vendor_name=vendor_name,
        issue_date=issue_date,
        **{k: v for k, v in form.items() if k not in ("department_id", "project_or_work_order")},
        issued_by_id=user.id,
        remarks=payload.remarks,
        order_project_id=op_link["project_id"],
        order_project_element_id=op_link["element_id"],
    )
    db.add(issue)
    db.flush()

    # Batch determination: a batch-tracked item's line names its batch (the
    # form pre-picks FIFO / FEFO) or is split over batches FIFO here — one
    # issue line per batch, so the trace shows exactly which heat went where.
    from app.modules.store.services.batches import allocate_issue
    taken: dict[tuple[int, str], dict[str, float]] = {}
    lines = []
    line_units: list[tuple[StoreMaterialIssueItem, list]] = []
    for p, units, line_el in zip(payload.items, picked_units, op_link["line_elements"]):
        left = list(units)
        for batch_no, qty in allocate_issue(db, items_by_id[p.item_id], payload.location_id, p.quantity,
                                            (p.batch_number or "").strip() or None, taken.setdefault(p.item_id, {})):
            mine, left = (left[:int(round(qty))], left[int(round(qty)):]) if units else ([], [])
            line = StoreMaterialIssueItem(issue_id=issue.id, item_id=p.item_id, quantity=qty, batch_number=batch_no, remarks=p.remarks,
                                          bin_id=p.bin_id, serial_numbers=[u.serial_no for u in mine] or None,
                                          order_project_element_id=line_el)
            db.add(line)
            lines.append(line)
            if mine:
                line_units.append((line, mine))
    # Job work: the material stays ours at the vendor (With-vendor stock,
    # SAP special stock O) until it comes back or is used up there.
    txn_type = "jobwork_out" if issue.issue_type == "job_work" else "issue"
    # Project stock: an element's own holds are consumed first; stock held for
    # another element is refused (or released by a Store Manager with a reason).
    need: dict[tuple[int, int | None], float] = {}
    for ln in lines:
        need[(ln.item_id, ln.order_project_element_id)] = need.get((ln.item_id, ln.order_project_element_id), 0.0) + ln.quantity
    overridden: list[tuple[str, float, str | None]] = []
    try:
        for (item_id, el_id), qty in sorted(need.items(), key=lambda kv: (kv[0][0], kv[0][1] or 0)):
            got = project_stock.secure_issue_qty(db, item_id=item_id, location_id=location.id, qty=qty, element_id=el_id, user=user,
                                                 location=location, override_reason=override_reason,
                                                 document=f"issue {issue.issue_number}", project_id=issue.order_project_id)
            overridden += got["overridden"]
    except HTTPException:
        db.rollback()
        raise
    if overridden:
        issue.project_stock_override_reason = override_reason
        note = "Project stock used (Store Manager override): " + "; ".join(
            f"{q:g} from {num}" + (f" ({lbl})" if lbl else "") for num, q, lbl in overridden) + f" — {override_reason}"
        issue.remarks = f"{issue.remarks}\n{note}" if issue.remarks else note
    try:
        # Item-id order, so two issues listing the same items differently
        # lock the balance rows in the same sequence and can't deadlock.
        for p in sorted(lines, key=lambda x: x.item_id):
            post_stock_transaction(
                db, item_id=p.item_id, location_id=payload.location_id, transaction_type=txn_type,
                quantity=p.quantity, batch_number=p.batch_number, reference_type="material_issue",
                reference_number=issue.issue_number, transaction_date=issue.issue_date, bin_id=p.bin_id,
                created_by_id=user.id, vendor_name=vendor_name,
                remarks=f"Job work challan {issue.challan_number} to {vendor_name}" if txn_type == "jobwork_out" else None,
            )
    except ValueError as e:
        db.rollback()
        raise HTTPException(status_code=409, detail=str(e))

    if line_units:
        # Serialised units issued: record what each went against (WO / machine / SR).
        db.flush()
        db.refresh(issue)
        against = _to_response(db, issue).against_label
        if issue.project_or_work_order:
            against = " · ".join(x for x in (f"Work order {issue.project_or_work_order}", against) if x)
        for line, units in line_units:
            serial_svc.issue_units(db, units, issue=issue, line=line, against=against, user_id=user.id)

    if issue_type_row and issue_type_row.value == "service":
        # Spares to a customer site: book them on the service request (ERP).
        from app.modules.erp.spares import book_store_issue
        db.flush()
        db.refresh(issue)
        book_store_issue(db, issue)

    if issue_type_row and issue_type_row.value == "dispatch" and issue.order_project_id:
        # A dispatch against a PM project moves the project's lifecycle stage (replaces the old sales-order dispatch posting).
        from app.modules.projects import so_projects
        from app.modules.projects.models.project import PmProject
        proj = db.get(PmProject, issue.order_project_id)
        if proj:
            so_projects.refresh_project_lifecycle(db, proj, user)

    db.commit()
    db.refresh(issue)
    issue = db.query(StoreMaterialIssue).options(selectinload(StoreMaterialIssue.items)).filter(StoreMaterialIssue.id == issue.id).first()
    return _to_response(db, issue)
