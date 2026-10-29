"""Job work / subcontracting — challans out at vendors, stock with vendor
(ME2O), Job Work Receipts (processed item in, components consumed) and the
GST ITC-04 report. See store/services/job_work.py."""
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session, selectinload

from app.core.permissions import ensure_feature_action, require_feature_action
from app.db.session import get_db
from app.modules.main.models.audit_log import AuditLog
from app.modules.main.models.user import User
from app.modules.store.models.item import StoreItem
from app.modules.store.models.job_work import StoreJobWorkReceipt
from app.modules.store.models.material_issue import StoreMaterialIssue
from app.modules.store.schemas.job_work import StoreJobWorkReceiptCreate
from app.modules.store.services import job_work as jw
from app.modules.store.services.stock_values import mask_for
from app.modules.store.services.plant_access import ensure_location_allowed, filter_by_location, visible_location_ids

router = APIRouter(prefix="/store/job-work", tags=["Store"])


def _challans(lines: list[dict]) -> list[dict]:
    out: dict[int, dict] = {}
    for r in lines:
        c = out.setdefault(r["issue_id"], {
            "issue_id": r["issue_id"], "issue_number": r["issue_number"], "challan_number": r["challan_number"],
            "challan_date": r["challan_date"], "expected_return_date": r["expected_return_date"], "vendor_id": r["vendor_id"],
            "vendor_name": r["vendor_name"], "process": r["process"], "location_name": r["location_name"], "lines": [],
            "output_item_id": r.get("output_item_id"), "output_items": r.get("output_items") or [],
            "outstanding_value": 0.0, "days_out": r["days_out"], "return_deadline": r["return_deadline"], "status": "closed",
        })
        c["lines"].append(r)
        c["outstanding_value"] = round(c["outstanding_value"] + (r["outstanding_value"] or 0), 2)
        c["return_deadline"] = min(c["return_deadline"], r["return_deadline"]) if r["quantity_outstanding"] > 0 else c["return_deadline"]
        rank = {"closed": 0, "open": 1, "due_back": 2, "overdue": 3}
        if rank[r["status"]] > rank[c["status"]]:
            c["status"] = r["status"]
    return list(out.values())


@router.get("/challans")
async def list_challans(
    status: str = Query("open"),  # open (anything still at the vendor) | all
    vendor_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_feature_action("store", "issues", "job_work", "view")),
):
    """Job-work challans with what is still at the vendor, per line, and the
    GST Rule 45 return deadline (1 year inputs / 3 years capital goods)."""
    rows = _challans(jw.challan_lines(db, vendor_id=vendor_id, location_ids=visible_location_ids(db, user)))
    if status == "open":
        rows = [c for c in rows if c["status"] != "closed"]
    rows.sort(key=lambda c: (c["status"] != "overdue", c["return_deadline"], c["challan_date"]))
    return mask_for(user, rows)


@router.get("/challans/{issue_id}")
async def get_challan(issue_id: int, db: Session = Depends(get_db), user: User = Depends(require_feature_action("store", "issues", "job_work", "view"))):
    issue = db.get(StoreMaterialIssue, issue_id)
    if not issue or not jw.is_job_work(issue):
        raise HTTPException(status_code=404, detail="Job-work challan not found — only issues of type Job work are challans.")
    ensure_location_allowed(db, user, issue.location_id)
    rows = _challans(jw.challan_lines(db, issue_ids=[issue_id]))
    challan = rows[0] if rows else None
    receipts = db.query(StoreJobWorkReceipt).options(selectinload(StoreJobWorkReceipt.components), selectinload(StoreJobWorkReceipt.outputs)).filter(
        StoreJobWorkReceipt.issue_id == issue_id).order_by(StoreJobWorkReceipt.id).all()
    from app.core.permission_registry import can_feature
    sees_receipts = can_feature(user, "store", "issues", "job_work_receipts", "view")
    return mask_for(user, {**(challan or {}), "tracks_vendor_stock": jw.posted_to_vendor(db, issue),
                           "receipts": [jw.receipt_dict(db, r) for r in receipts] if sees_receipts else []})


@router.get("/vendor-stock")
async def vendor_stock(vendor_id: int | None = Query(None), db: Session = Depends(get_db),
                       user: User = Depends(require_feature_action("store", "issues", "job_work", "view"))):
    """Stock with vendors (SAP ME2O / special stock O): still ours, still
    valued, physically at the job worker."""
    return mask_for(user, jw.vendor_stock(db, location_ids=visible_location_ids(db, user), vendor_id=vendor_id))


@router.get("/itc04")
async def itc04_report(
    date_from: date = Query(..., alias="from"),
    date_to: date = Query(..., alias="to"),
    format: str = Query("json"),
    db: Session = Depends(get_db),
    user: User = Depends(require_feature_action("store", "issues", "job_work", "view")),
):
    """GST ITC-04: goods sent to job workers (table 4) and received back
    (table 5A) in the period, plus what is still outstanding. format=xlsx
    downloads the workbook."""
    if date_to < date_from:
        raise HTTPException(status_code=422, detail=f"'To' ({date_to:%d-%m-%Y}) is before 'From' ({date_from:%d-%m-%Y}) — pick the quarter / half-year again.")
    data = mask_for(user, jw.itc04(db, date_from=date_from, date_to=date_to, location_ids=visible_location_ids(db, user)))
    if format == "xlsx":
        ensure_feature_action(user, "store", "issues", "job_work", "export")
        return Response(content=jw.itc04_workbook(data),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="ITC-04_{date_from:%Y%m%d}_{date_to:%Y%m%d}.xlsx"'})
    return data


@router.get("/po-lines")
async def job_work_po_lines(vendor_id: int = Query(...), db: Session = Depends(get_db),
                            _user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "create"))):
    """The vendor's issued PO lines a job charge can be taken from."""
    from app.modules.p2p.models.purchase_order import P2PPurchaseOrder, P2PPurchaseOrderItem
    rows = db.query(P2PPurchaseOrderItem, P2PPurchaseOrder).join(P2PPurchaseOrder, P2PPurchaseOrderItem.purchase_order_id == P2PPurchaseOrder.id).filter(
        P2PPurchaseOrder.vendor_id == vendor_id, P2PPurchaseOrder.status.notin_(("draft", "cancelled"))).order_by(P2PPurchaseOrder.id.desc()).limit(200).all()
    return [{"id": it.id, "label": f"{po.po_number} — {it.item_name}" + (f" @ ₹{it.unit_price:g}/{it.unit or 'unit'}" if it.unit_price is not None else ""),
             "unit_price": it.unit_price, "po_number": po.po_number} for it, po in rows]


@router.get("/receipts")
async def list_receipts(db: Session = Depends(get_db), user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "view"))):
    q = filter_by_location(db.query(StoreJobWorkReceipt).options(selectinload(StoreJobWorkReceipt.components), selectinload(StoreJobWorkReceipt.outputs)),
                           StoreJobWorkReceipt.location_id, visible_location_ids(db, user))
    return mask_for(user, [jw.receipt_dict(db, r) for r in q.order_by(StoreJobWorkReceipt.id.desc()).limit(500).all()])


OUTPUT_COLUMNS = ["item_code", "quantity", "job_charge_rate", "cost_share_pct"]


@router.get("/receipts/import-outputs/template")
async def receipt_outputs_template(_user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "create"))):
    """Excel template for the processed items on a job work receipt."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Processed items"
    ws.append(OUTPUT_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FDF1E6")
    ws.append(["SF-FLG-0001", 4, 250, ""])
    ws.append(["SF-RNG-0002", 8, 120, ""])
    for col, w in zip("ABCD", (20, 12, 18, 16)):
        ws.column_dimensions[col].width = w
    notes = wb.create_sheet("How to fill")
    for line in ("item_code — the processed item's code from the Item Master (required).",
                 "quantity — how many came back (required, more than 0).",
                 "job_charge_rate — vendor's charge per unit of this item (optional).",
                 "cost_share_pct — share of the sent material's value this item carries. Leave blank on every row to split "
                 "by standard cost × quantity, or fill every row so they total 100.",
                 "One row per item; the same item on two rows is added together."):
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="job_work_receipt_items_template.xlsx"'})


@router.post("/receipts/import-outputs")
async def import_receipt_outputs(file: UploadFile = File(...), db: Session = Depends(get_db),
                                 _user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "create"))):
    """Reads processed-item lines (item_code, quantity, job_charge_rate,
    cost_share_pct) from Excel / CSV into the Receive form — nothing is
    posted. Same item on two rows is added together; bad rows come back
    with the reason so the file can be fixed."""
    from app.modules.store.routes.item_import import _read_rows
    rows = [r for r in _read_rows(file.filename, await file.read()) if any(r.values())]
    if not rows:
        raise HTTPException(status_code=422, detail="The file has no item rows — fill in item_code and quantity under the header row (download the template).")
    if len(rows) > 200:
        raise HTTPException(status_code=422, detail=f"The file has {len(rows)} rows — one receipt can hold at most 200 processed items; split the file.")
    missing = [c for c in ("item_code", "quantity") if c not in rows[0]]
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing column(s) {', '.join(missing)} — download the template and keep its header row.")
    codes = {(r.get("item_code") or "").strip().lower() for r in rows}
    items = {i.item_code.strip().lower(): i for i in db.query(StoreItem).filter(StoreItem.item_code.isnot(None)).all()
             if i.item_code and i.item_code.strip().lower() in codes}

    def num(v):
        v = str(v if v is not None else "").replace(",", "").strip()
        return float(v) if v else None

    lines: dict[int, dict] = {}
    errors = []
    for idx, r in enumerate(rows, start=2):
        code = (r.get("item_code") or "").strip()
        item = items.get(code.lower())
        if not code:
            errors.append({"row": idx, "item_code": "", "reason": "Item code is empty."}); continue
        if not item:
            errors.append({"row": idx, "item_code": code, "reason": f"'{code}' isn't in the Item Master — create it there first or check the code."}); continue
        try:
            qty, rate, share = num(r.get("quantity")), num(r.get("job_charge_rate")), num(r.get("cost_share_pct"))
        except ValueError:
            errors.append({"row": idx, "item_code": code, "reason": "quantity, job_charge_rate and cost_share_pct must be numbers."}); continue
        if not qty or qty <= 0:
            errors.append({"row": idx, "item_code": code, "reason": "Quantity must be greater than zero."}); continue
        if (rate is not None and rate < 0) or (share is not None and share < 0):
            errors.append({"row": idx, "item_code": code, "reason": "Job charge and cost share can't be negative."}); continue
        if item.id in lines:
            lines[item.id]["quantity"] += qty
        else:
            lines[item.id] = {"item_id": item.id, "item_code": item.item_code, "item_name": item.item_name, "uom": item.uom,
                              "quantity": qty, "job_charge_rate": rate, "cost_share_pct": share}
    return {"lines": list(lines.values()), "errors": errors}


@router.get("/receipts/{receipt_id}")
async def get_receipt(receipt_id: int, db: Session = Depends(get_db), user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "view"))):
    rec = db.get(StoreJobWorkReceipt, receipt_id)
    if not rec:
        raise HTTPException(status_code=404, detail="Job work receipt not found")
    ensure_location_allowed(db, user, rec.location_id)
    return mask_for(user, jw.receipt_dict(db, rec))


@router.post("/receipts")
async def create_receipt(payload: StoreJobWorkReceiptCreate, db: Session = Depends(get_db),
                         user: User = Depends(require_feature_action("store", "issues", "job_work_receipts", "create"))):
    """Processed item back from the job worker: the challan material used is
    consumed out of With-vendor stock and the processed item is received at
    component value + job charge (into Quality under a job work lot when the
    item is inspection-controlled)."""
    issue = db.get(StoreMaterialIssue, payload.issue_id)
    if issue:
        ensure_location_allowed(db, user, issue.location_id)
    if payload.location_id:
        ensure_location_allowed(db, user, payload.location_id)
    rec = jw.create_receipt(db, payload, user)
    got = ", ".join(f"{o['quantity']:g} {o['item_name'] or 'item'}" for o in jw.receipt_dict(db, rec)["outputs"])
    db.add(AuditLog(entity_type="store_job_work_receipt", entity_id=rec.id, action="created", performed_by_id=user.id,
                    summary=f"{user.name or user.email} received {got} on {rec.receipt_number} "
                            f"against job-work challan {issue.challan_number or issue.issue_number}."))
    db.commit()
    return mask_for(user, jw.receipt_dict(db, db.get(StoreJobWorkReceipt, rec.id)))
