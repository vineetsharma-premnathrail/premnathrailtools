"""R&D Technical Offer (see models/technical_offer.py): numbering, the Word
document of record, and who may do what.

The Word document is a managed document (main/services/managed_documents.py):
R&D creates it from the offer's fields, edits it in Word Online, and the R&D
head releases it. Once released the document is frozen — a change is a new
revision (…-R1), never an edit of what CRM may already have quoted against.
"""
from __future__ import annotations

import io
import logging
from typing import Any, Mapping

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.crm.models.inquiry import Inquiry
from app.modules.crm.models.organization import Organization, OrgContact
from app.modules.crm.models.tender import Tender
from app.modules.main.models.user import User
from app.modules.main.services import managed_documents
from app.modules.rnd.models.technical_offer import RndTechnicalOffer
from app.modules.rnd.services import requirement_access as ra
from app.utils.docx_letterhead import (
    BODY, HEADER, ACCENT, ACCENT_LIGHT_HEX, MUTED, USABLE_IN,
    fmt_run, tight, write, shade, cell_margins, set_widths, new_letterhead_document,
)

logger = logging.getLogger(__name__)

ENTITY = "rnd_technical_offer"
SIGNATORY = ["FOR & ON BEHALF OF", "Premnath Rail", "C/o PREMNATH ENGINEERING WORKS", "", "Research & Development"]
SECTIONS = (
    ("scope", "1. Scope of Supply"),
    ("technical_details", "2. Technical Specification"),
    ("deviations", "3. Deviations from Customer Specification"),
    ("exclusions", "4. Exclusions"),
)
# What the Word template asks for in a section R&D hasn't written yet.
PROMPTS = {
    "scope": "[Write the scope of supply — e.g. 1 no. Rail Grinding Machine RGM-25 with standard tool kit, operation & maintenance manual]",
    "technical_details": "[Write the technical specification, one point per line — gauge, engine, capacity, speed, standards complied with…]",
    "deviations": "Nil",
    "exclusions": "Nil",
}


def record_for(db: Session, source: str, record_id: int):
    model = Inquiry if source == "inquiry" else Tender
    return db.query(model).filter(model.id == record_id, model.is_deleted == False).first()  # noqa: E712


def number_base(source: str, record) -> str:
    """INQ-2026-0042 → TO-INQ-2026-0042 (the source prefix keeps inquiry and
    tender series apart)."""
    return f"TO-{record.universal_id or ('INQ' if source == 'inquiry' else 'TND') + f'-{record.id}'}"


def number_for(base: str, revision: int) -> str:
    return base if revision == 0 else f"{base}-R{revision}"


def ensure_creator(db: Session, user: User) -> None:
    """Preparing a Technical Offer is R&D's formal answer: R&D heads, and the
    people with R&D › Requirements › Technical Offer team › Create."""
    if not ra.is_rnd_side(user):
        raise HTTPException(status_code=403, detail="Only R&D prepares the Technical Offer. Design / Electrical / Hydraulic & Pneumatic give their input through tasks and documents.")
    if not ra.can_create_technical_offer(db, user):
        raise HTTPException(status_code=403, detail=f"You can't create Technical Offers. Ask an admin to tick R&D › Requirements › "
                                                    f"'Technical Offer team (every requirement)' › Create for you (or your role) in Organization › "
                                                    f"Role & Permissions. R&D head: {ra._head_names(db)}.")


def ensure_can_prepare(db: Session, user: User, source: str, record_id: int) -> dict:
    ensure_creator(db, user)
    return ra.ensure(db, user, source, record_id, "view")


def prefill(source: str, record) -> dict:
    """Starting text for a new offer's Word template, taken from the
    requirement CRM sent — the rest is written in Word."""
    if source == "inquiry":
        qty = f"{record.quantity:g} no. " if record.quantity is not None else ""
        product = record.product or record.product_category
        return {
            "title": product or record.universal_id or "Technical Offer",
            "scope": f"{qty}{product}" if product else None,
            "technical_details": record.product_spec or None,
        }
    return {"title": record.tender_name or record.tender_number or record.universal_id or "Technical Offer", "scope": None, "technical_details": None}


def build_context(db: Session, offer: RndTechnicalOffer) -> dict:
    record = record_for(db, offer.source, offer.record_id)
    org = db.query(Organization).filter(Organization.id == record.org_id).first() if record else None
    contact = db.query(OrgContact).filter(OrgContact.id == record.org_contact_id).first() if record and record.org_contact_id else None
    return {
        "offer_number": offer.offer_number, "offer_date": offer.offer_date.strftime("%d-%m-%Y"), "revision": offer.revision,
        "revision_reason": offer.revision_reason, "title": offer.title,
        "requirement_ref": record.universal_id if record else None,
        "tor_number": record.technical_offer_number if record else None,
        "org_name": org.name if org else None, "org_address": org.address if org else None,
        "contact_name": contact.name if contact else None, "contact_designation": contact.designation if contact else None,
        "contact_email": contact.email if contact else None, "contact_mobile": contact.mobile if contact else None,
        **{key: getattr(offer, key) for key, _ in SECTIONS},
    }


def build_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    doc = new_letterhead_document()

    title = tight(doc.add_paragraph())
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_run(title.add_run("TECHNICAL OFFER"), size=HEADER + 12, bold=True, color=ACCENT)
    write(doc, "", after=10)

    # ── Customer (left) + offer meta (right) ──
    top = doc.add_table(rows=1, cols=2)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(top, left=0, right=0, top=0, bottom=0)
    set_widths(top, [USABLE_IN * 0.55, USABLE_IN * 0.45])
    left, right = top.rows[0].cells
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    write(left, "TO", size=BODY - 1, bold=True, color=ACCENT, first=True, after=3)
    write(left, ctx.get("org_name") or "-", size=HEADER + 1, bold=True)
    for key in ("org_address", "contact_name", "contact_designation", "contact_email", "contact_mobile"):
        if ctx.get(key):
            write(left, str(ctx[key]))

    meta = [("Offer #", ctx["offer_number"]), ("Offer date", ctx["offer_date"])]
    if ctx.get("revision"):
        meta.append(("Revision", f"R{ctx['revision']}"))
    if ctx.get("requirement_ref"):
        meta.append(("Your enquiry ref", ctx["requirement_ref"]))
    for i, (label, value) in enumerate(meta):
        p = right.paragraphs[0] if i == 0 else right.add_paragraph()
        tight(p, after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fmt_run(p.add_run(f"{label}   "), bold=True, color=ACCENT)
        fmt_run(p.add_run(str(value)))

    subject = write(doc, "", after=0)
    tight(subject, before=14, after=8)
    fmt_run(subject.add_run("Subject: "), bold=True, color=ACCENT)
    fmt_run(subject.add_run(ctx.get("title") or ""), bold=True)
    write(doc, "Dear Sir / Madam,", after=6)
    write(doc, "With reference to your enquiry, we are pleased to submit our technical offer as below. "
               "The commercial offer is submitted separately.", after=6)
    if ctx.get("revision") and ctx.get("revision_reason"):
        write(doc, f"This revision supersedes our earlier offer. Reason: {ctx['revision_reason']}", color=MUTED, after=6)

    for key, heading in SECTIONS:
        box = doc.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(box)
        set_widths(box, [USABLE_IN])
        shade(box.rows[0].cells[0], ACCENT_LIGHT_HEX)
        write(box.rows[0].cells[0], heading, size=HEADER, bold=True, color=ACCENT, first=True)
        text = (ctx.get(key) or "").strip()
        lines = text.splitlines() or [PROMPTS[key]]
        for n, line in enumerate(lines):
            p = write(doc, line, after=0, color=None if text or PROMPTS[key] == "Nil" else MUTED)
            tight(p, before=6 if n == 0 else 0, after=2)
        write(doc, "", after=8)

    write(doc, "We trust the above is in line with your requirement and look forward to your valued order.", after=18)
    for n, line in enumerate(SIGNATORY):
        write(doc, line, bold=n in (0, 1), color=ACCENT if n == 1 else None)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _renumber_docx(data: bytes, old: str, new: str) -> bytes:
    """The previous revision's Word file with its offer number swapped for
    the new one (runs, so the letterhead formatting is kept)."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def carry_forward_document(db: Session, prev: RndTechnicalOffer, new: RndTechnicalOffer, user: User) -> bool:
    """A revision starts from what R&D wrote in Word on the previous one, not
    from a fresh template. False when there is nothing to copy (or SharePoint
    failed) — the new draft then gets the template on 'Save & Edit in Word'."""
    source = managed_documents.get_active(db, ENTITY, prev.id)
    if not source:
        return False
    try:
        data = _renumber_docx(await managed_documents.current_docx(source), prev.offer_number, new.offer_number)
        target = _load(db, new.id, user, True)
        await managed_documents.generate(
            db, entity_type=ENTITY, entity_id=new.id, source_revision=target.revision, title=target.title,
            folder=target.folder, filename=f"{target.filename_stem}.docx", docx_bytes=data, user=user,
        )
        return True
    except Exception:  # noqa: BLE001 — the revision itself is saved; Word falls back to the template
        logger.warning("Could not copy the Word document of %s into %s", prev.offer_number, new.offer_number, exc_info=True)
        db.rollback()
        return False


async def create_template_document(db: Session, offer: RndTechnicalOffer, user: User) -> str | None:
    """Create the new offer's Word template right away. Returns why it
    couldn't (the offer is saved either way; 'Create Word' retries)."""
    try:
        target = _load(db, offer.id, user, True)
        await managed_documents.generate(
            db, entity_type=ENTITY, entity_id=offer.id, source_revision=target.revision, title=target.title,
            folder=target.folder, filename=f"{target.filename_stem}.docx", docx_bytes=await managed_documents.build_bytes(target), user=user,
        )
        return None
    except HTTPException as exc:
        db.rollback()
        return str(exc.detail)
    except Exception as exc:  # noqa: BLE001 — SharePoint/network failure must not lose the offer
        logger.warning("Could not create the Word template for %s", offer.offer_number, exc_info=True)
        db.rollback()
        return f"SharePoint did not accept the file ({exc.__class__.__name__}). Click 'Create Word' to try again."


def _load(db: Session, offer_id: int, user: User, modify: bool) -> managed_documents.DocumentTarget:
    offer = db.query(RndTechnicalOffer).filter(RndTechnicalOffer.id == offer_id).first()
    if not offer or not record_for(db, offer.source, offer.record_id):
        raise HTTPException(status_code=404, detail="This Technical Offer (or the requirement behind it) was deleted — refresh the page.")
    if modify:
        ensure_can_prepare(db, user, offer.source, offer.record_id)
        if offer.status != "draft":
            if offer.status == "submitted":
                raise HTTPException(status_code=409, detail=f"Technical Offer {offer.offer_number} was sent to the R&D head for approval — its document is frozen. Ask the R&D head to return it for changes if something must be corrected.")
            raise HTTPException(status_code=409, detail=f"Technical Offer {offer.offer_number} is {'already sent to the BD team' if offer.status == 'released' else offer.status} — its document is frozen. Use 'Ask for revision'; once the R&D head approves, a new revision opens for editing.")
    else:
        ra.ensure(db, user, offer.source, offer.record_id, "view")
    return managed_documents.DocumentTarget(
        entity_id=offer.id, revision=0, title=f"Technical Offer {offer.offer_number}",
        folder=managed_documents.sharepoint_folder("RnD-TechnicalOffers", offer.offer_number_base),
        filename_stem=managed_documents.safe_filename(offer.offer_number),
        build_docx=lambda: build_docx(build_context(db, offer)).getvalue(),
    )


managed_documents.register(managed_documents.DocumentAdapter(entity_type=ENTITY, app="rnd", load=_load))
