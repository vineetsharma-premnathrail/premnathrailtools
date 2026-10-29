"""Purchase Order — the managed Word document of record `p2p_purchase_order`
(company letterhead header/footer; "Save & Edit in Word" in Word Online,
versioned).

Laid out like the company's official Tally purchase order: a boxed header
(Invoice To / Supplier on the left; Voucher No., Dated, Reference, Mode of
Payment, Other References, Dispatched through, Destination, Terms of Delivery
on the right), the Description of Goods grid (Sl / Description / GST Rate /
Quantity / Rate / per / Amount) with CGST+SGST or IGST input lines, round off
and total, amount in words, Company's PAN, the numbered Declaration and the
"for <company> / Authorised Signatory" block."""
from __future__ import annotations

import io
import re
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.accounts.models.vendor import Vendor
from app.modules.main.models.user import User
from app.modules.main.services import managed_documents
from app.modules.p2p.models.p2p_request import P2PRequest
from app.modules.p2p.models.purchase_order import P2PPurchaseOrder
from app.modules.p2p.models.rfq import RFQ
from app.utils.docx_letterhead import (
    BODY, HEADER, MUTED, USABLE_IN, cell_margins, fmt_run, new_letterhead_document, set_widths, tight, write,
)

PO_ENTITY = "p2p_purchase_order"

# "Company terms:" / "Vendor terms (<name>):" are the section labels rfq._po_terms
# writes into the PO's terms box — the printed Declaration is one numbered list.
_TERMS_SECTION_LABEL = re.compile(r"^(company terms|vendor terms\b.*):$", re.IGNORECASE)
_LEADING_NUMBER = re.compile(r"^\s*(?:\d{1,2}[.)]|[-•*])\s*")


def declaration_points(terms: str | None) -> list[str]:
    """The PO's terms box as numbered Declaration points: one point per line (or per
    inline "1. … 2. …" number), section labels and the user's own numbering dropped
    so the document numbers them 1, 2, 3 … itself. Empty box -> the company default."""
    from app.modules.crm.reports.quotation_pdf import _split_numbered_points
    from app.modules.p2p.settings import DEFAULT_COMPANY_TERMS
    points = []
    for line in (terms or "").splitlines():
        line = line.strip()
        if not line or _TERMS_SECTION_LABEL.match(line):
            continue
        points += [_LEADING_NUMBER.sub("", p).strip() for p in _split_numbered_points(line)]
    return [p for p in points if p] or list(DEFAULT_COMPANY_TERMS)

# GST state codes (first two digits of a GSTIN) — prints "State Name : Haryana, Code : 06"
# and decides CGST+SGST (same state) vs IGST (inter-state).
GST_STATES = {
    "01": "Jammu & Kashmir", "02": "Himachal Pradesh", "03": "Punjab", "04": "Chandigarh", "05": "Uttarakhand",
    "06": "Haryana", "07": "Delhi", "08": "Rajasthan", "09": "Uttar Pradesh", "10": "Bihar", "11": "Sikkim",
    "12": "Arunachal Pradesh", "13": "Nagaland", "14": "Manipur", "15": "Mizoram", "16": "Tripura", "17": "Meghalaya",
    "18": "Assam", "19": "West Bengal", "20": "Jharkhand", "21": "Odisha", "22": "Chhattisgarh", "23": "Madhya Pradesh",
    "24": "Gujarat", "26": "Dadra & Nagar Haveli and Daman & Diu", "27": "Maharashtra", "29": "Karnataka", "30": "Goa",
    "31": "Lakshadweep", "32": "Kerala", "33": "Tamil Nadu", "34": "Puducherry", "35": "Andaman & Nicobar Islands",
    "36": "Telangana", "37": "Andhra Pradesh", "38": "Ladakh",
}
_STATE_CODES = {name.lower(): code for code, name in GST_STATES.items()}


def _state(gstin: str | None, state_name: str | None) -> tuple[str | None, str | None]:
    """(state name, GST state code) from the GSTIN prefix, else from the typed state name."""
    code = gstin[:2] if gstin and gstin[:2] in GST_STATES else None
    if code:
        return GST_STATES[code], code
    if state_name:
        return state_name, _STATE_CODES.get(state_name.strip().lower())
    return None, None


def _d(v) -> str:
    return v.strftime("%d-%b-%y").lstrip("0") if v else ""


def _n(v) -> str:
    from app.modules.crm.reports.quotation_pdf import _fmt_amount
    return _fmt_amount(v)


def _qty(v: float) -> str:
    return f"{float(v or 0):.2f}"


def _company(db: Session):
    from app.modules.organization.models.company import Company
    from app.modules.organization.models.company_address import CompanyAddress
    c = db.query(Company).order_by(Company.id).first()
    addr = None
    if c:
        addr = db.query(CompanyAddress).filter(CompanyAddress.company_id == c.id).order_by(
            CompanyAddress.is_primary.desc(), CompanyAddress.id).first()
    return c, addr


def _ship_to(db: Session, po: P2PPurchaseOrder) -> str:
    """ME21N receiving plant / ship-to store printed as the destination."""
    from app.modules.organization.models.branch import Branch
    from app.modules.store.models.location import StoreLocation

    loc = db.get(StoreLocation, po.ship_to_location_id) if getattr(po, "ship_to_location_id", None) else None
    plant = db.get(Branch, po.plant_branch_id) if getattr(po, "plant_branch_id", None) else None
    parts = [x for x in ((plant.name if plant else None), (loc.name if loc else None), (loc.address if loc else None)) if x]
    return ", ".join(parts)


def po_context(db: Session, po: P2PPurchaseOrder) -> dict[str, Any]:
    """Everything the PO document prints, resolved from the PO, its requisition,
    RFQ, vendor and Company Info."""
    from app.modules.crm.reports.quotation_pdf import _amount_in_words

    pr = db.get(P2PRequest, po.p2p_request_id) if po.p2p_request_id else None
    rfq = db.query(RFQ).filter(RFQ.p2p_request_id == pr.id).order_by(RFQ.id.desc()).first() if pr else None
    vendor = db.get(Vendor, po.vendor_id) if po.vendor_id else None
    company, caddr = _company(db)

    buyer_state, buyer_code = _state(company.gstin if company else None, caddr.state if caddr else None)
    vendor_state, vendor_code = _state(vendor.gstin if vendor else None, vendor.state if vendor else None)
    # Same state -> CGST + SGST; different -> IGST; unknown -> one plain "GST" line.
    tax_split = "unknown" if not (buyer_code and vendor_code) else ("intra" if buyer_code == vendor_code else "inter")

    lines, taxable, gst = [], 0.0, 0.0
    for i in po.items:
        amount = float(i.quantity or 0) * float(i.unit_price or 0)
        taxable += amount
        gst += amount * float(i.tax_rate or 0) / 100
        sub = [ln.strip() for ln in (i.specification or "").splitlines() if ln.strip()]
        if i.make:
            sub.append(f"Make: {i.make}")
        if i.part_code:
            sub.append(f"({i.part_code})")
        if getattr(i, "delivery_date", None):
            sub.append(f"Delivery: {i.delivery_date.strftime('%d/%m/%Y')}")
        lines.append({
            "name": i.item_name, "sub": sub, "gst_rate": i.tax_rate, "qty": i.quantity, "unit": i.unit or "",
            "rate": i.unit_price, "amount": amount,
        })
    gst = round(gst, 2)
    exact = taxable + gst
    total = round(exact)
    units = {ln["unit"] for ln in lines}

    approver = (pr.partner_approved_by_name if pr.partner_approved_at else (pr.po_approved_by_name or "")) if pr else ""
    prepared = ""
    if po.created_by_id:
        u = db.get(User, po.created_by_id)
        prepared = (u.name or u.email) if u else ""

    delivery = [f"DELIVERY DATE- {po.expected_delivery.strftime('%d/%m/%Y')}"] if po.expected_delivery else []
    if po.delivery_terms:
        delivery.append(po.delivery_terms)

    return {
        "draft": po.status == "draft",
        # ME22N: an approved amendment reissues the PO as "Amendment N".
        "amendment_no": getattr(po, "amendment_no", 0) or 0,
        "company_name": (company.legal_name or company.name) if company else "Premnath Engineering Works",
        "company_address": [x for x in ((caddr.address_line1 if caddr else None), (caddr.address_line2 if caddr else None),
                                        " ".join(x for x in ((caddr.city if caddr else None), (caddr.pincode if caddr else None)) if x)) if x],
        "company_gstin": company.gstin if company else None,
        "company_state": buyer_state, "company_state_code": buyer_code,
        "company_pan": company.pan if company else None,
        "supplier_name": po.vendor_name or (vendor.name if vendor else ""),
        "supplier_address": [x for x in ((vendor.address if vendor else None),
                                         " ".join(x for x in ((vendor.city if vendor else None), (vendor.pin_code if vendor else None)) if x)) if x],
        "supplier_gstin": vendor.gstin if vendor else None,
        "supplier_state": vendor_state, "supplier_state_code": vendor_code,
        "po_number": po.po_number, "po_date": _d(po.po_date),
        "reference": " ".join(x for x in ((pr.p2p_number if pr else None), (_d(pr.request_date) if pr else None)) if x) or po.po_number,
        "payment_terms": (rfq.payment_terms if rfq and rfq.payment_terms else "") or "",
        "other_references": f"FOR {pr.project_label}" if pr and pr.project_label else "",
        "dispatched_through": "", "destination": _ship_to(db, po),
        "terms_of_delivery": delivery,
        "lines": lines, "tax_split": tax_split, "gst": gst, "round_off": round(total - exact, 2), "total": total,
        "total_qty": f"{_qty(sum(float(ln['qty'] or 0) for ln in lines))} {units.pop()}".strip() if len(units) == 1 else "",
        "amount_in_words": f"INR {_amount_in_words(total, False).replace(' Rupees', '')}",
        "declaration": declaration_points(po.terms_conditions),
        "prepared_by": prepared, "approved_by": approver,
    }


# ── Word layout ────────────────────────────────────────────────────────────

def _borders(table, inside_h: bool = True) -> None:
    """Single black outer + vertical borders; horizontal inner rules optional
    (the goods grid has none between lines, like the Tally print)."""
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideV", "insideH"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "nil" if side == "insideH" and not inside_h else "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:color"), "000000")
        el.append(b)
    tbl_pr.append(el)


def _rule(cell, side: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:color"), "000000")
    borders.append(b)


def _label_value(cell, label: str, value: str | list[str]) -> None:
    write(cell, label, size=BODY - 1, color=MUTED, first=True)
    for v in (value if isinstance(value, list) else [value]) or [""]:
        write(cell, v, bold=True)


def _party(cell, heading: str, name: str, address: list[str], gstin: str | None, state: str | None, code: str | None, first: bool) -> None:
    write(cell, heading, size=BODY - 1, color=MUTED, first=first)
    write(cell, name or "", size=HEADER, bold=True)
    for line in address:
        write(cell, line)
    if gstin:
        write(cell, f"GSTIN/UIN : {gstin}")
    if state:
        write(cell, f"State Name : {state}" + (f", Code : {code}" if code else ""))


def build_po_docx(ctx: dict[str, Any]) -> io.BytesIO:
    doc = new_letterhead_document()
    title = write(doc, "PURCHASE ORDER" + (f" — AMENDMENT {ctx['amendment_no']}" if ctx.get("amendment_no") else "")
                  + (" (DRAFT)" if ctx["draft"] else ""), size=HEADER + 4, bold=True, align="center")
    tight(title, after=6)

    # ── Boxed header: parties left, voucher details right ──
    left_w, pair_w = USABLE_IN * 0.5, USABLE_IN * 0.25
    head = doc.add_table(rows=5, cols=3)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(head)
    cell_margins(head, left=80, right=80, top=40, bottom=40)
    set_widths(head, [left_w, pair_w, pair_w])
    _label_value(head.cell(0, 1), "Voucher No.", ctx["po_number"])
    _label_value(head.cell(0, 2), "Dated", ctx["po_date"])
    _label_value(head.cell(1, 1), "Reference No. & Date", ctx["reference"])
    _label_value(head.cell(1, 2), "Mode/Terms of Payment", ctx["payment_terms"])
    _label_value(head.cell(2, 1), "Dispatched through", ctx["dispatched_through"])
    _label_value(head.cell(2, 2), "Other References", ctx["other_references"])
    dest = head.cell(3, 1).merge(head.cell(3, 2))
    _label_value(dest, "Destination (Ship To)", ctx["destination"])
    delivery = head.cell(4, 1).merge(head.cell(4, 2))
    _label_value(delivery, "Terms of Delivery", ctx["terms_of_delivery"])
    invoice_to = head.cell(0, 0).merge(head.cell(2, 0))
    _party(invoice_to, "Invoice To", ctx["company_name"], ctx["company_address"], ctx["company_gstin"],
           ctx["company_state"], ctx["company_state_code"], first=True)
    supplier = head.cell(3, 0).merge(head.cell(4, 0))
    _party(supplier, "Supplier (Bill from)", ctx["supplier_name"], ctx["supplier_address"], ctx["supplier_gstin"],
           ctx["supplier_state"], ctx["supplier_state_code"], first=True)
    for c in (invoice_to, supplier):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    # ── Description of Goods ──
    widths = [0.35, 3.05, 0.55, 0.85, 0.85, 0.45, 1.05]
    scale = USABLE_IN / sum(widths)
    widths = [w * scale for w in widths]
    grid = doc.add_table(rows=1, cols=7)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(grid, inside_h=False)
    cell_margins(grid, left=60, right=60, top=30, bottom=30)
    for cell, h in zip(grid.rows[0].cells, ["Sl No.", "Description of Goods", "GST Rate", "Quantity", "Rate", "per", "Amount"]):
        write(cell, h, size=BODY - 0.5, bold=True, align="center", first=True)
        _rule(cell, "bottom")
    tr_pr = grid.rows[0]._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    tr_pr.append(rep)

    def row(values: list[tuple[str, str, bool]]):
        cells = grid.add_row().cells
        for cell, (text, align, bold) in zip(cells, values):
            write(cell, text, align=align, bold=bold, first=True)
        return cells

    for n, ln in enumerate(ctx["lines"], start=1):
        cells = row([
            (str(n), "center", False), (ln["name"], "left", True),
            (f"{float(ln['gst_rate']):g} %" if ln["gst_rate"] is not None else "", "center", False),
            (f"{_qty(ln['qty'])} {ln['unit']}".strip(), "right", True),
            (_n(ln["rate"]) if ln["rate"] is not None else "", "right", False),
            (ln["unit"], "center", False), (_n(ln["amount"]), "right", True),
        ])
        for s in ln["sub"]:
            p = write(cells[1], s, size=BODY - 0.5)
            for r in p.runs:
                r.italic = True
        write(cells[1], "", after=2)

    if ctx["tax_split"] == "intra":
        taxes = [("CGST INPUT", ctx["gst"] / 2), ("SGST INPUT", ctx["gst"] / 2)]
    elif ctx["tax_split"] == "inter":
        taxes = [("IGST INPUT", ctx["gst"])]
    else:
        taxes = [("GST INPUT", ctx["gst"])]
    if ctx["round_off"]:
        taxes.append(("ROUND OFF", ctx["round_off"]))
    for label, amount in taxes:
        cells = row([("", "left", False), (label, "right", True), ("", "left", False), ("", "left", False),
                     ("", "left", False), ("", "left", False), (_n(amount), "right", True)])
        for r in cells[1].paragraphs[0].runs:
            r.italic = True

    cells = row([("", "left", False), ("Total", "right", False), ("", "left", False), (ctx["total_qty"], "right", True),
                 ("", "left", False), ("", "left", False), (f"₹ {_n(ctx['total'])}", "right", True)])
    for c in cells:
        _rule(c, "top")
    set_widths(grid, widths)

    # ── Amount in words, PAN, declaration, signatory ──
    foot = doc.add_table(rows=2, cols=2)
    foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(foot, inside_h=False)
    cell_margins(foot, left=80, right=80, top=50, bottom=50)
    set_widths(foot, [USABLE_IN * 0.6, USABLE_IN * 0.4])
    words = foot.cell(0, 0).merge(foot.cell(0, 1))
    write(words, "Amount Chargeable (in words)", size=BODY - 1, color=MUTED, first=True)
    p = write(words, ctx["amount_in_words"], bold=True)
    fmt_run(p.add_run("\tE. & O.E"), size=BODY - 1)
    p.paragraph_format.tab_stops.add_tab_stop(int(USABLE_IN * 914400 * 0.97), alignment=2)  # right tab

    decl, sign = foot.cell(1, 0), foot.cell(1, 1)
    _rule(decl, "top")
    _rule(sign, "top")
    if ctx["company_pan"]:
        p = write(decl, "Company's PAN : ", first=True, after=4)
        fmt_run(p.add_run(ctx["company_pan"]), bold=True)
    write(decl, "Declaration", bold=True, first=not ctx["company_pan"], after=2)
    for n, text in enumerate(ctx["declaration"], start=1):
        write(decl, f"{n}. {text}", size=BODY - 1, after=1)

    write(sign, f"for {ctx['company_name']}", bold=True, align="right", first=True)
    for _ in range(4):
        write(sign, "")
    if ctx["prepared_by"] or ctx["approved_by"]:
        write(sign, " · ".join(x for x in ((f"Prepared by: {ctx['prepared_by']}" if ctx["prepared_by"] else ""),
                                           (f"Approved by: {ctx['approved_by']}" if ctx["approved_by"] else "")) if x),
              size=BODY - 1.5, color=MUTED, align="right")
    write(sign, "Authorised Signatory", align="right")
    sign.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    end = write(doc, "This is a Computer Generated Document", size=BODY - 1, align="center")
    tight(end, before=4)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def po_document_target(db: Session, po_id: int, user: User, modify: bool) -> managed_documents.DocumentTarget:
    po = db.get(P2PPurchaseOrder, po_id)
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found — it may have been cancelled. Refresh the page.")
    from app.core.data_scope import ensure_in_scope
    from app.core.permission_registry import can_view_tab
    from app.core.permissions import ensure_tab_action
    import app.modules.p2p.access  # noqa: F401 — registers the Procurement data scopes
    if modify:
        ensure_tab_action(user, "p2p", "purchase_orders", "edit")
    elif not can_view_tab(user, "p2p", "purchase_orders"):
        raise HTTPException(status_code=403, detail="You can't open purchase order documents — your role has no access to "
                                                    "Procurement › Purchase Orders. Ask an admin to add it in Organization › "
                                                    "Role & Permissions.")
    ensure_in_scope(db, user, po)
    # An approved ME22N amendment (amendment_no > 0) reissues the PO — the
    # buyer regenerates the "Amendment N" document after it is approved.
    if modify and po.status != "draft" and not (getattr(po, "amendment_no", 0) or 0):
        raise HTTPException(status_code=409, detail=f"PO {po.po_number} has been sent for approval, so its document can't be changed — "
                                                    "it is the version the approvers signed off.")
    return managed_documents.DocumentTarget(
        entity_id=po.id, revision=po.revision or 0,
        title=f"Purchase order {po.po_number}",
        folder=managed_documents.sharepoint_folder("P2P-PurchaseOrders", po.po_number.replace("/", "-")),
        filename_stem=managed_documents.safe_filename(f"PO-{po.po_number}"),
        build_docx=lambda: build_po_docx(po_context(db, po)).getvalue(),
    )


managed_documents.register(managed_documents.DocumentAdapter(entity_type=PO_ENTITY, app="purchase", load=po_document_target))
