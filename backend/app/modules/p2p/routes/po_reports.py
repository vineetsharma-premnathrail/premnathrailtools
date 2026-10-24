"""PO reporting (SAP ME2N by PO / ME2L by vendor / ME2M by material /
ME80FN general analysis) and PO history (ME23N 'PO history' tab).

Line-level, computed from the PO lines, their GRN lines (not reversed), Return
to Vendor lines, accepted service entries and vendor invoice lines
(VendorInvoiceLine.po_item_id) — nothing is stored twice.

Per PO line:
  ordered    PO quantity
  received   physically received on GRNs that aren't reversed
  rejected   rejected on those GRNs' inspections
  returned   accepted quantity sent back on Return to Vendor (a replacement is owed)
  delivered  what counts against the PO: accepted − returned once inspected,
             received while still in inspection; service lines: accepted service entries
  invoiced   on vendor invoices that aren't cancelled
  open       ordered − delivered (never below 0); open value at the PO rate
  overdue    days past the PO's expected delivery while something is still open

Own prefix (/p2p/po-reports) so it can never be shadowed by /p2p/purchase-orders/{id}."""
import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, selectinload

from app.core.data_scope import apply_scope
from app.core.permissions import GRANT_HINT, ensure_feature_action, require_app_access, require_tab_access, require_tab_action
import app.modules.p2p.access  # noqa: F401 — registers the Procurement data scopes
from app.modules.p2p.pricing import mask_for, sees_prices
from app.db.session import get_db
from app.modules.main.models.user import User
from app.modules.p2p.models.goods_receipt import P2PGoodsReceipt, P2PGoodsReceiptItem
from app.modules.p2p.models.p2p_request import P2PRequest
from app.modules.p2p.models.purchase_order import P2PPurchaseOrder, P2PPurchaseOrderItem

router = APIRouter(prefix="/p2p/po-reports", tags=["P2P"])

_EPS = 1e-9
_DECIDED = ("passed", "failed", "partial")
_REPORT_PO_STATUSES = ("issued", "acknowledged", "partially_fulfilled", "fulfilled")


def _invoiced_by_po_item(db: Session, po_item_ids: list[int]) -> dict[int, float]:
    if not po_item_ids:
        return {}
    from app.modules.accounts.models.vendor_invoice import VendorInvoice, VendorInvoiceLine
    rows = (db.query(VendorInvoiceLine.po_item_id, func.coalesce(func.sum(VendorInvoiceLine.quantity), 0))
            .join(VendorInvoice, VendorInvoice.id == VendorInvoiceLine.vendor_invoice_id)
            .filter(VendorInvoiceLine.po_item_id.in_(po_item_ids), VendorInvoice.status != "cancelled")
            .group_by(VendorInvoiceLine.po_item_id).all())
    return {pid: float(q or 0) for pid, q in rows}


def _service_by_po_item(db: Session, po_item_ids: list[int]) -> dict[int, float]:
    if not po_item_ids:
        return {}
    from app.modules.p2p.models.service_entry import P2PServiceEntry
    rows = (db.query(P2PServiceEntry.po_item_id, func.coalesce(func.sum(P2PServiceEntry.quantity), 0))
            .filter(P2PServiceEntry.po_item_id.in_(po_item_ids), P2PServiceEntry.status == "accepted")
            .group_by(P2PServiceEntry.po_item_id).all())
    return {pid: float(q or 0) for pid, q in rows}


def _grn_lines_by_po_item(db: Session, po_item_ids: list[int]) -> dict[int, list[tuple[P2PGoodsReceipt, P2PGoodsReceiptItem]]]:
    out: dict[int, list] = {}
    if not po_item_ids:
        return out
    rows = (db.query(P2PGoodsReceiptItem, P2PGoodsReceipt)
            .join(P2PGoodsReceipt, P2PGoodsReceipt.id == P2PGoodsReceiptItem.goods_receipt_id)
            .filter(P2PGoodsReceiptItem.po_item_id.in_(po_item_ids)).order_by(P2PGoodsReceipt.received_date, P2PGoodsReceipt.id).all())
    for line, grn in rows:
        out.setdefault(line.po_item_id, []).append((grn, line))
    return out


def item_lines(db: Session, pairs: list[tuple[P2PPurchaseOrderItem, P2PPurchaseOrder]], today: date) -> list[dict]:
    """One dict per PO line (ordered / received / delivered / open ...), in the
    order of `pairs`. Batched: a fixed number of queries however many lines.
    Pure arithmetic on the lines — whether the PO is released or short-closed
    is applied by the caller (p2p/po_tracking.py)."""
    ids = [i.id for i, _ in pairs]
    grns = _grn_lines_by_po_item(db, ids)
    invoiced = _invoiced_by_po_item(db, ids)
    service = _service_by_po_item(db, ids)
    pr_numbers = dict(db.query(P2PRequest.id, P2PRequest.p2p_number).filter(
        P2PRequest.id.in_({po.p2p_request_id for _, po in pairs} - {None})).all()) if pairs else {}
    rows = []
    for it, po in pairs:
        received = rejected = returned = delivered = 0.0
        for grn, line in grns.get(it.id, []):
            if grn.status == "reversed":
                continue
            received += line.received_quantity or 0
            rejected += line.rejected_quantity or 0
            returned += line.returned_quantity or 0
            if line.quality_status in _DECIDED:
                delivered += max((line.accepted_quantity or 0) - (line.returned_quantity or 0), 0.0)
            else:
                delivered += line.received_quantity or 0
        if it.category == "service":
            delivered += service.get(it.id, 0.0)
            received += service.get(it.id, 0.0)
        ordered = it.quantity or 0
        open_qty = max(ordered - delivered, 0.0)
        rate = float(it.unit_price) if it.unit_price else 0.0
        overdue = (today - po.expected_delivery).days if po.expected_delivery and open_qty > _EPS and po.expected_delivery < today else 0
        rows.append({
            "po_id": po.id, "po_number": po.po_number, "po_date": po.po_date, "po_status": po.status,
            "p2p_request_id": po.p2p_request_id, "p2p_number": pr_numbers.get(po.p2p_request_id),
            "vendor_id": po.vendor_id, "vendor_name": po.vendor_name, "po_item_id": it.id, "store_item_id": it.store_item_id,
            "item_code": it.store_item.item_code if it.store_item is not None else None, "item_name": it.item_name,
            "unit": it.unit, "category": it.category, "unit_price": rate, "ordered_qty": ordered,
            "received_qty": round(received, 6), "rejected_qty": round(rejected, 6), "returned_qty": round(returned, 6),
            "delivered_qty": round(delivered, 6), "invoiced_qty": round(invoiced.get(it.id, 0.0), 6),
            "open_qty": round(open_qty, 6), "open_value": round(open_qty * rate, 2),
            "to_invoice_qty": round(max(delivered - invoiced.get(it.id, 0.0), 0.0), 6),
            "expected_delivery": po.expected_delivery, "days_overdue": overdue,
        })
    return rows


def open_items(db: Session, *, vendor_id: int | None = None, vendor_name: str | None = None, store_item_id: int | None = None,
               po_number: str | None = None, overdue_only: bool = False, open_only: bool = True, released_only: bool = True,
               date_from: date | None = None, date_to: date | None = None, today: date | None = None) -> list[dict]:
    """The line-level report. Only released POs (final approval given) count —
    a PO still in approval is not an order yet — and a short-closed or closed
    PO has no open balance: what the vendor never supplied is written off, so
    it leaves open qty / open value / overdue (`short_closed_qty` shows it).
    `released_only=False` also lists POs that are not released (PO history)."""
    from app.modules.p2p import po_tracking

    today = today or date.today()
    q = (db.query(P2PPurchaseOrderItem, P2PPurchaseOrder)
         .join(P2PPurchaseOrder, P2PPurchaseOrder.id == P2PPurchaseOrderItem.purchase_order_id)
         .filter(P2PPurchaseOrder.status.in_(_REPORT_PO_STATUSES)))
    if vendor_id is not None:
        q = q.filter(P2PPurchaseOrder.vendor_id == vendor_id)
    if vendor_name and vendor_name.strip():
        q = q.filter(P2PPurchaseOrder.vendor_name.ilike(f"%{vendor_name.strip()}%"))
    if store_item_id is not None:
        q = q.filter(P2PPurchaseOrderItem.store_item_id == store_item_id)
    if po_number and po_number.strip():
        q = q.filter(P2PPurchaseOrder.po_number.ilike(f"%{po_number.strip()}%"))
    if date_from:
        q = q.filter(P2PPurchaseOrder.po_date >= date_from)
    if date_to:
        q = q.filter(P2PPurchaseOrder.po_date <= date_to)
    pairs = q.options(selectinload(P2PPurchaseOrderItem.store_item)).order_by(P2PPurchaseOrder.po_date.desc(), P2PPurchaseOrder.id.desc(),
                                                                               P2PPurchaseOrderItem.id).all()
    stages = po_tracking.stages_for(db, list({po.id: po for _, po in pairs}.values()))
    pairs = [(it, po) for it, po in pairs if not released_only or stages[po.id] in po_tracking.RELEASED_STAGES]
    rows = []
    for row in item_lines(db, pairs, today):
        stage = stages[row["po_id"]]
        row["po_stage"] = stage
        row["po_stage_label"] = po_tracking.STAGES[stage]
        row["short_closed_qty"] = 0.0
        if stage not in po_tracking.OPEN_STAGES:
            # Delivered in full, short-closed, closed or not yet released: nothing is awaited from the vendor.
            row["short_closed_qty"] = row["open_qty"] if stage == "short_closed" else 0.0
            row.update(open_qty=0.0, open_value=0.0, days_overdue=0)
        if open_only and row["open_qty"] <= _EPS:
            continue
        if overdue_only and row["days_overdue"] <= 0:
            continue
        rows.append(row)
    return rows


def _summary(rows: list[dict], key: str, label: str) -> list[dict]:
    out: dict = {}
    for r in rows:
        k = r[key] if r[key] is not None else r[label]
        s = out.setdefault(k, {key: r[key], label: r[label], "lines": 0, "open_qty": 0.0, "open_value": 0.0, "overdue_lines": 0,
                               "max_days_overdue": 0})
        s["lines"] += 1
        s["open_qty"] += r["open_qty"]
        s["open_value"] = round(s["open_value"] + r["open_value"], 2)
        if r["days_overdue"] > 0:
            s["overdue_lines"] += 1
            s["max_days_overdue"] = max(s["max_days_overdue"], r["days_overdue"])
    return sorted(out.values(), key=lambda s: -s["open_value"])


def _parse(d: str | None, field: str) -> date | None:
    if not d:
        return None
    try:
        return date.fromisoformat(d)
    except ValueError:
        raise HTTPException(status_code=422, detail=f"'{d}' is not a valid {field.replace('_', ' ')} — use YYYY-MM-DD.")


def _report(db, vendor_id, vendor_name, store_item_id, po_number, overdue_only, open_only, date_from, date_to) -> list[dict]:
    return open_items(db, vendor_id=vendor_id, vendor_name=vendor_name, store_item_id=store_item_id, po_number=po_number,
                      overdue_only=overdue_only, open_only=open_only, date_from=_parse(date_from, "date_from"),
                      date_to=_parse(date_to, "date_to"))


@router.get("/tracking")
async def po_tracking_list(
    stage: str | None = Query(None, description="draft, in_approval, with_vendor, partly_received, received, short_closed, closed or cancelled"),
    search: str | None = Query(None, description="PO no., vendor, requisition, project or item"),
    vendor_id: int | None = Query(None),
    buyer_id: int | None = Query(None),
    overdue_only: bool = Query(False),
    due_soon_only: bool = Query(False, description="Delivery due within the next 7 days"),
    date_from: str | None = Query(None, description="PO date from (YYYY-MM-DD)"),
    date_to: str | None = Query(None, description="PO date to (YYYY-MM-DD)"),
    order_project_id: int | None = Query(None),
    order_project_element_id: int | None = Query(None),
    sort: str = Query("code", description="code, urgency, due, date, vendor, stage, value or progress"),
    desc: bool = Query(False),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("purchase")),
    _tab: User = Depends(require_tab_access("p2p", "purchase_orders")),
):
    """Every PO with its true lifecycle stage, delivery progress and days to its expected delivery —
    filtered, sorted and paged on the server over everything the viewer may see."""
    from app.modules.p2p import po_tracking

    if stage and stage not in po_tracking.STAGES:
        raise HTTPException(status_code=422, detail=f"'{stage}' is not a PO stage — use one of: {', '.join(po_tracking.STAGES)}.")
    if sort not in po_tracking.SORTS:
        raise HTTPException(status_code=422, detail=f"Cannot sort by '{sort}' — use one of: {', '.join(po_tracking.SORTS)}.")
    query = apply_scope(db.query(P2PPurchaseOrder), user, P2PPurchaseOrder)
    data = po_tracking.tracking(
        db, query, stage=stage, search=search, vendor_id=vendor_id, buyer_id=buyer_id, overdue_only=overdue_only,
        due_soon_only=due_soon_only, date_from=_parse(date_from, "date_from"), date_to=_parse(date_to, "date_to"),
        order_project_id=order_project_id, order_project_element_id=order_project_element_id, sort=sort, desc=desc,
        skip=skip, limit=limit)
    return mask_for(user, "purchase_orders", data)


@router.get("/tracking/filters")
async def po_tracking_filters(
    db: Session = Depends(get_db),
    user: User = Depends(require_app_access("purchase")),
    _tab: User = Depends(require_tab_access("p2p", "purchase_orders")),
):
    """Vendors and buyers on the viewer's POs, for the tracking filter pickers."""
    from app.modules.p2p import po_tracking

    return po_tracking.filter_options(db, apply_scope(db.query(P2PPurchaseOrder), user, P2PPurchaseOrder))


@router.get("/open-items")
async def po_open_items(
    vendor_id: int | None = Query(None),
    vendor_name: str | None = Query(None),
    store_item_id: int | None = Query(None),
    po_number: str | None = Query(None),
    overdue_only: bool = Query(False),
    open_only: bool = Query(True, description="False = every PO line (ME80FN), not only open ones"),
    date_from: str | None = Query(None, description="PO date from"),
    date_to: str | None = Query(None, description="PO date to"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_tab_access("p2p", "purchase_orders")),
):
    rows = _report(db, vendor_id, vendor_name, store_item_id, po_number, overdue_only, open_only, date_from, date_to)
    return mask_for(_user, "purchase_orders", {
        "rows": rows,
        "by_vendor": _summary(rows, "vendor_id", "vendor_name"),
        "by_item": _summary(rows, "store_item_id", "item_name"),
        "total_open_value": round(sum(r["open_value"] for r in rows), 2),
        "overdue_lines": sum(1 for r in rows if r["days_overdue"] > 0),
    })


@router.get("/open-items/export")
async def po_open_items_export(
    vendor_id: int | None = Query(None),
    vendor_name: str | None = Query(None),
    store_item_id: int | None = Query(None),
    po_number: str | None = Query(None),
    overdue_only: bool = Query(False),
    open_only: bool = Query(True),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    db: Session = Depends(get_db),
    _user: User = Depends(require_tab_action("p2p", "purchase_orders", "export")),
):
    if not sees_prices(_user, "purchase_orders"):
        raise HTTPException(status_code=403, detail=(
            "The open-items export carries PO rates and values, and your role doesn't show prices on Purchase Orders "
            f"(p2p:purchase_orders.prices:view). Use the on-screen report instead. {GRANT_HINT}"))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = _report(db, vendor_id, vendor_name, store_item_id, po_number, overdue_only, open_only, date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "PO lines"
    head = ["PO no.", "PO date", "PO stage", "PR no.", "Vendor", "Item code", "Item", "Unit", "Category", "Rate", "Ordered",
            "Received", "Rejected", "Returned", "Delivered (net)", "Invoiced", "To invoice", "Open qty", "Open value",
            "Expected delivery", "Days overdue"]
    ws.append(head)
    for r in rows:
        ws.append([r["po_number"], r["po_date"], r["po_stage_label"], r["p2p_number"], r["vendor_name"], r["item_code"], r["item_name"],
                   r["unit"], r["category"], r["unit_price"], r["ordered_qty"], r["received_qty"], r["rejected_qty"], r["returned_qty"],
                   r["delivered_qty"], r["invoiced_qty"], r["to_invoice_qty"], r["open_qty"], r["open_value"],
                   r["expected_delivery"], r["days_overdue"]])
    for sheet_name, data, key in (("By vendor", _summary(rows, "vendor_id", "vendor_name"), "vendor_name"),
                                  ("By item", _summary(rows, "store_item_id", "item_name"), "item_name")):
        s = wb.create_sheet(sheet_name)
        s.append(["Vendor" if key == "vendor_name" else "Item", "Lines", "Open qty", "Open value", "Overdue lines", "Max days overdue"])
        for d in data:
            s.append([d[key], d["lines"], d["open_qty"], d["open_value"], d["overdue_lines"], d["max_days_overdue"]])
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FDF1E6")
        sheet.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", (18, 12, 16, 16, 28, 16, 34, 7, 10, 10, 10, 10, 10, 10, 13, 10, 10, 10, 13, 15, 12)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="po_open_items.xlsx"'})


@router.get("/{po_id}/history")
async def po_history(po_id: int, db: Session = Depends(get_db), _user: User = Depends(require_tab_access("p2p", "purchase_orders"))):
    """ME23N PO history: per PO line, every GRN, Return to Vendor, service
    entry and vendor invoice booked against it."""
    ensure_feature_action(_user, "p2p", "purchase_orders", "history", "view")
    po = db.query(P2PPurchaseOrder).options(selectinload(P2PPurchaseOrder.items)).filter(P2PPurchaseOrder.id == po_id).first()
    if not po:
        raise HTTPException(status_code=404, detail=f"Purchase order #{po_id} not found — it may have been deleted. Reload the list.")
    ids = [i.id for i in po.items]
    grns = _grn_lines_by_po_item(db, ids)
    from app.modules.accounts.models.vendor_invoice import VendorInvoice, VendorInvoiceLine
    from app.modules.p2p.models.service_entry import P2PServiceEntry
    from app.modules.p2p.models.vendor_return import P2PVendorReturn, P2PVendorReturnItem
    grn_line_ids = [l.id for pairs in grns.values() for _, l in pairs]
    rtv_rows = (db.query(P2PVendorReturnItem, P2PVendorReturn).join(P2PVendorReturn, P2PVendorReturn.id == P2PVendorReturnItem.vendor_return_id)
                .filter(P2PVendorReturnItem.grn_item_id.in_(grn_line_ids or [-1])).order_by(P2PVendorReturn.return_date).all())
    grn_line_to_po_item = {l.id: l.po_item_id for pairs in grns.values() for _, l in pairs}
    inv_rows = (db.query(VendorInvoiceLine, VendorInvoice).join(VendorInvoice, VendorInvoice.id == VendorInvoiceLine.vendor_invoice_id)
                .filter(VendorInvoiceLine.po_item_id.in_(ids or [-1])).order_by(VendorInvoice.invoice_date).all())
    se_rows = db.query(P2PServiceEntry).filter(P2PServiceEntry.po_item_id.in_(ids or [-1])).order_by(P2PServiceEntry.id).all()
    summary = {r["po_item_id"]: r for r in open_items(db, po_number=po.po_number, open_only=False, released_only=False) if r["po_id"] == po.id}

    lines = []
    for it in po.items:
        events = []
        for grn, l in grns.get(it.id, []):
            events.append({"type": "grn", "label": "Goods receipt", "document_id": grn.id, "document_number": grn.grn_number,
                           "date": grn.received_date, "status": grn.status, "quantity": l.received_quantity,
                           "accepted": l.accepted_quantity, "rejected": l.rejected_quantity, "invoice_quantity": l.invoice_quantity,
                           "receipt_variance": l.receipt_variance, "vendor_invoice_number": grn.vendor_invoice_number,
                           "delivery_challan_number": grn.delivery_challan_number, "quality_status": l.quality_status})
        for ri, rtv in rtv_rows:
            if grn_line_to_po_item.get(ri.grn_item_id) == it.id:
                events.append({"type": "rtv", "label": "Return to vendor", "document_id": rtv.id, "document_number": rtv.rtv_number,
                               "date": rtv.return_date, "status": ri.source, "quantity": ri.quantity, "value": ri.value})
        for se in se_rows:
            if se.po_item_id == it.id:
                events.append({"type": "service_entry", "label": "Service entry", "document_id": se.id,
                               "document_number": getattr(se, "entry_number", None) or f"SE #{se.id}",
                               "date": getattr(se, "entry_date", None) or (se.created_at.date() if se.created_at else None),
                               "status": se.status, "quantity": se.quantity})
        for il, inv in inv_rows:
            if il.po_item_id == it.id:
                events.append({"type": "invoice", "label": "Vendor invoice", "document_id": inv.id, "document_number": inv.invoice_number,
                               "date": inv.invoice_date, "status": inv.status, "quantity": il.quantity,
                               "amount": il.taxable_amount, "match_status": il.match_status})
        events.sort(key=lambda e: (e["date"] or date.min))
        lines.append({"po_item_id": it.id, "item_name": it.item_name, "item_code": it.store_item.item_code if it.store_item is not None else None,
                      "unit": it.unit, "ordered_qty": it.quantity, "summary": summary.get(it.id), "events": events})
    return mask_for(_user, "purchase_orders", {"po_id": po.id, "po_number": po.po_number, "vendor_name": po.vendor_name,
                                               "po_date": po.po_date, "expected_delivery": po.expected_delivery,
                                               "status": po.status, "lines": lines})
