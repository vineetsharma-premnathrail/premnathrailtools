"""Two renderers for one quality document "spec" — the managed Word
document (company letterhead, edited in Word Online; the document of
record) and a reportlab PDF used only until a Word draft exists.

spec = {
  "title": "INSPECTION REPORT", "number": "IL-2026-0001",
  "intro": str | None,
  "facts": [(label, value), ...],            # 2-column grid
  "sections": [{"heading": str, "columns": [...], "rows": [[...]], "widths": [inches]}],
  "statement": str | None,                   # certification / decision text
  "remarks": str | None,
  "signatures": [(role, name), ...],
}
"""
from __future__ import annotations

import io
import os
from typing import Any
from xml.sax.saxutils import escape


def _dash(v: Any) -> str:
    return "—" if v in (None, "") else str(v)


def render_docx(spec: dict) -> bytes:
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Pt
    from app.utils.docx_letterhead import (
        ACCENT, MUTED, USABLE_IN, tight, write, shade, cell_margins, set_widths, new_letterhead_document,
    )
    shade_hex = "F3F4F6"

    def grid(doc, rows, widths, header=False, bold_cols=()):
        table = doc.add_table(rows=0, cols=len(widths))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(table)
        for r, values in enumerate(rows):
            cells = table.add_row().cells
            for c, (cell, value) in enumerate(zip(cells, values)):
                head = header and r == 0
                write(cell, _dash(value), size=8 if (head or c in bold_cols) else 8.5, bold=head or c in bold_cols, first=True)
                if head or c in bold_cols:
                    shade(cell, shade_hex)
        set_widths(table, widths)
        return table

    doc = new_letterhead_document()
    title = doc.add_table(rows=1, cols=2)
    cell_margins(title, left=0, right=0, top=0, bottom=40)
    left, right = title.rows[0].cells
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    write(left, spec["title"], size=14, bold=True, color=ACCENT, first=True)
    write(right, spec["number"], size=12, bold=True, align="right", first=True)
    set_widths(title, [USABLE_IN - 2.2, 2.2])
    if spec.get("intro"):
        tight(write(doc, spec["intro"], size=9.5), before=8, after=8)

    facts = list(spec.get("facts") or [])
    if len(facts) % 2:
        facts.append(("", ""))
    half = USABLE_IN / 2
    if facts:
        rows = [[facts[i][0], _dash(facts[i][1]), facts[i + 1][0], _dash(facts[i + 1][1])] for i in range(0, len(facts), 2)]
        grid(doc, rows, [1.35, half - 1.35, 1.35, half - 1.35], bold_cols=(0, 2))

    for sec in spec.get("sections") or []:
        tight(write(doc, sec["heading"], size=10, bold=True), before=12, after=4)
        if sec.get("rows"):
            grid(doc, [sec["columns"]] + sec["rows"], sec["widths"], header=True)
        else:
            write(doc, sec.get("empty") or "None.", color=MUTED)

    if spec.get("statement"):
        tight(write(doc, spec["statement"], size=9.5, bold=True), before=12, after=4)
    if spec.get("remarks"):
        tight(write(doc, "Remarks", size=9, bold=True), before=10, after=2)
        write(doc, spec["remarks"])

    sigs = spec.get("signatures") or []
    if sigs:
        tight(doc.add_paragraph(), before=14)
        sign = doc.add_table(rows=3, cols=len(sigs))
        sign.style = "Table Grid"
        cell_margins(sign)
        for c, (role, name) in enumerate(sigs):
            write(sign.rows[0].cells[c], role, size=9, bold=True, first=True)
            write(sign.rows[1].cells[c], "Signature & date", size=8.5, color=MUTED, first=True)
            sign.rows[1].cells[c].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
            write(sign.rows[2].cells[c], f"Name: {name or ''}", size=8.5, first=True)
        sign.rows[1].height = Pt(54)
        set_widths(sign, [USABLE_IN / len(sigs)] * len(sigs))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(spec: dict, printed: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.lib.colors import HexColor, black
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Image, Spacer
        from reportlab.lib.styles import ParagraphStyle
    except ImportError:
        raise RuntimeError("reportlab is not installed on the server (pip install reportlab).")
    grid_c, shade_c = HexColor("#9ca3af"), HexColor("#f3f4f6")

    def p(text, size=8.5, bold=False):
        return Paragraph(escape(_dash(text)), ParagraphStyle(
            "c", fontName="Helvetica-Bold" if bold else "Helvetica", fontSize=size, leading=size + 2))

    def grid(rows, widths, header=False):
        t = Table(rows, colWidths=widths)
        style = [("GRID", (0, 0), (-1, -1), 0.4, grid_c), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if header:
            style.append(("BACKGROUND", (0, 0), (-1, 0), shade_c))
        t.setStyle(TableStyle(style))
        return t

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=10 * mm, bottomMargin=10 * mm, leftMargin=12 * mm,
                            rightMargin=12 * mm, title=f"{spec['title']} {spec['number']}")
    width = A4[0] - 24 * mm
    logo_path = os.path.join(os.path.dirname(__file__), "..", "..", "..", "utils", "templates", "premnath_logo_mark.png")
    logo = Image(logo_path, width=24 * mm, height=9 * mm) if os.path.exists(logo_path) else p("")
    head = Table([[logo, p(spec["title"], 14, True), p(spec["number"], 12, True)]],
                 colWidths=[30 * mm, width - 80 * mm, 50 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (2, 0), (2, 0), "RIGHT"),
                              ("LINEBELOW", (0, 0), (-1, 0), 1.2, black), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story: list[Any] = [head, Spacer(1, 3 * mm)]
    if spec.get("intro"):
        story += [p(spec["intro"], 9.5), Spacer(1, 3 * mm)]
    facts = list(spec.get("facts") or [])
    if len(facts) % 2:
        facts.append(("", ""))
    if facts:
        rows = [[p(facts[i][0], 8, True), p(facts[i][1]), p(facts[i + 1][0], 8, True), p(facts[i + 1][1])] for i in range(0, len(facts), 2)]
        info = grid(rows, [32 * mm, width / 2 - 32 * mm, 32 * mm, width / 2 - 32 * mm])
        info.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), shade_c), ("BACKGROUND", (2, 0), (2, -1), shade_c)]))
        story += [info, Spacer(1, 4 * mm)]
    for sec in spec.get("sections") or []:
        story.append(p(sec["heading"], 10, True))
        story.append(Spacer(1, 1.5 * mm))
        if sec.get("rows"):
            total = sum(sec["widths"])
            widths = [w / total * width for w in sec["widths"]]
            story.append(grid([[p(c, 8, True) for c in sec["columns"]]] + [[p(v) for v in r] for r in sec["rows"]], widths, header=True))
        else:
            story.append(p(sec.get("empty") or "None."))
        story.append(Spacer(1, 4 * mm))
    if spec.get("statement"):
        story += [p(spec["statement"], 9.5, True), Spacer(1, 3 * mm)]
    if spec.get("remarks"):
        story += [p("Remarks", 9, True), p(spec["remarks"]), Spacer(1, 3 * mm)]
    sigs = spec.get("signatures") or []
    if sigs:
        sw = width / len(sigs)
        sign = Table([[p(r, 9, True) for r, _ in sigs], [p("\n\nSignature & date") for _ in sigs],
                      [p(f"Name: {n or ''}") for _, n in sigs]], colWidths=[sw] * len(sigs), rowHeights=[None, 18 * mm, None])
        sign.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, grid_c), ("VALIGN", (0, 0), (-1, -1), "BOTTOM")]))
        story.append(sign)
    story += [Spacer(1, 3 * mm), p(printed, 7)]
    doc.build(story)
    return buf.getvalue()
