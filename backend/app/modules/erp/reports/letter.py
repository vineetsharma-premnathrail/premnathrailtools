"""PDF + Word builders for a machine's free-form letter (ERP Projects →
Documents → Created).

Both reuse the CRM letterhead helpers so a letter carries exactly the same
header (logo.JPG) and footer (logo-1.JPG) banners as quotations and
technical offers: the PDF via quotation_pdf's canvas-drawn letterhead, the
.docx via technical_offer_docx's full-page-width floating banners. The PDF
is the version to send; the .docx exists for when someone needs to tweak
wording in Word before sending.
"""

from __future__ import annotations

import io
from typing import Any, Mapping
from xml.sax.saxutils import escape as _xml_escape

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether

from app.modules.crm.reports.quotation_pdf import (
    _draw_letterhead, _p, FONT_NAME, FONT_NAME_BOLD, MARGIN, USABLE_WIDTH, HEADER_HEIGHT, FOOTER_HEIGHT,
)
from app.modules.crm.reports.technical_offer_docx import (
    Document, _cropped_content_bytes, _add_full_width_banner, _image_aspect,
    HEADER_LOGO_PATH, FOOTER_LOGO_PATH, PAGE_WIDTH_EMU,
)

LETTER_SIZE = 10.5


def _body_paragraphs(text: str) -> list[str]:
    """Blank lines separate paragraphs; single newlines stay as line breaks."""
    blocks = [b.strip("\n") for b in (text or "").replace("\r\n", "\n").split("\n\n")]
    return [b for b in blocks if b.strip()]


def _lines(text: str | None) -> list[str]:
    return [ln.rstrip() for ln in (text or "").replace("\r\n", "\n").split("\n")]


# Text width inside the letter margins (PDF).
LETTER_WIDTH = USABLE_WIDTH - 0.3 * inch


def letterhead_pdf_doc(buf: io.BytesIO, title: str) -> SimpleDocTemplate:
    """A4 doc whose frame clears the letterhead header/footer banners; build it
    with onFirstPage/onLaterPages=_draw_letterhead."""
    return SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=MARGIN + HEADER_HEIGHT + 0.2 * inch,
        bottomMargin=MARGIN + FOOTER_HEIGHT + 0.15 * inch,
        leftMargin=MARGIN + 0.15 * inch, rightMargin=MARGIN + 0.15 * inch,
        title=title,
    )


def letterhead_docx():
    """New A4 python-docx Document with the letterhead banners and Calibri body."""
    if Document is None:
        raise Exception("python-docx library is not installed. Please run: pip install python-docx")
    from docx.shared import Pt, Inches, Emu

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.75)
    header_bytes = _cropped_content_bytes(HEADER_LOGO_PATH)
    footer_bytes = _cropped_content_bytes(FOOTER_LOGO_PATH)
    # Body clears the banners exactly as the PDF frame does (letterhead_pdf_doc):
    # a fixed 1.35in left the first line (the copy label) riding over the header line.
    header_h_in = 8.27 * _image_aspect(header_bytes) if header_bytes else HEADER_HEIGHT / 72
    footer_h_in = 8.27 * _image_aspect(footer_bytes) if footer_bytes else FOOTER_HEIGHT / 72
    section.top_margin = Inches(MARGIN / 72 + header_h_in + 0.2)
    section.bottom_margin = Inches(MARGIN / 72 + footer_h_in + 0.15)
    if header_bytes:
        _add_full_width_banner(section.header.paragraphs[0], header_bytes, offset_y_emu=0)
    if footer_bytes:
        footer_height_emu = int(int(PAGE_WIDTH_EMU) * _image_aspect(footer_bytes))
        _add_full_width_banner(
            section.footer.paragraphs[0], footer_bytes,
            offset_y_emu=int(Emu(Inches(11.69))) - footer_height_emu,
        )

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(LETTER_SIZE)
    style.paragraph_format.space_after = Pt(0)
    return doc


def build_letter_pdf(ctx: Mapping[str, Any]) -> io.BytesIO:
    """Expected keys: letter_no, letter_date (display string), recipient, subject,
    body, signatory_name, signatory_designation, company_name (all strings; the
    last three optional). Templated delivery documents (letter_type other than
    "general") are laid out by letter_templates instead."""
    if ctx.get("letter_type", "general") != "general":
        from app.modules.erp.reports.letter_templates import build_template_pdf
        return build_template_pdf(ctx)
    buf = io.BytesIO()
    doc = letterhead_pdf_doc(buf, ctx.get("subject") or "Letter")
    width = LETTER_WIDTH
    body_style = ParagraphStyle(
        "body", fontName=FONT_NAME, fontSize=LETTER_SIZE, leading=LETTER_SIZE + 4.5,
        alignment=TA_JUSTIFY, spaceAfter=8,
    )

    story: list[Any] = []
    ref_row = Table(
        [[_p(f"Ref: {ctx.get('letter_no', '')}", size=LETTER_SIZE, bold=True),
          _p(f"Date: {ctx.get('letter_date', '')}", size=LETTER_SIZE, bold=True, align="right")]],
        colWidths=[width * 0.6, width * 0.4],
    )
    ref_row.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [ref_row, Spacer(1, 16)]

    story.append(_p("To,", size=LETTER_SIZE))
    for line in _lines(ctx.get("recipient")):
        story.append(_p(line or " ", size=LETTER_SIZE))
    story.append(Spacer(1, 14))

    subject = _xml_escape(ctx.get("subject") or "")
    story.append(Paragraph(
        f"<font name='{FONT_NAME_BOLD}'>Subject: </font><font name='{FONT_NAME_BOLD}'><u>{subject}</u></font>",
        ParagraphStyle("subj", fontName=FONT_NAME, fontSize=LETTER_SIZE, leading=LETTER_SIZE + 4),
    ))
    story.append(Spacer(1, 14))

    for block in _body_paragraphs(ctx.get("body") or ""):
        story.append(Paragraph(_xml_escape(block).replace("\n", "<br/>"), body_style))

    sign: list[Any] = [Spacer(1, 18)]
    if ctx.get("company_name"):
        sign.append(_p(f"For {ctx['company_name']}", size=LETTER_SIZE, bold=True))
    sign.append(Spacer(1, 42))
    if ctx.get("signatory_name"):
        sign.append(_p(ctx["signatory_name"], size=LETTER_SIZE, bold=True))
    if ctx.get("signatory_designation"):
        sign.append(_p(ctx["signatory_designation"], size=LETTER_SIZE))
    story.append(KeepTogether(sign))

    doc.build(story, onFirstPage=_draw_letterhead, onLaterPages=_draw_letterhead)
    buf.seek(0)
    return buf


def build_letter_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    """Same layout and keys as build_letter_pdf, as an editable .docx."""
    if ctx.get("letter_type", "general") != "general":
        from app.modules.erp.reports.letter_templates import build_template_docx
        return build_template_docx(ctx)
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

    doc = letterhead_docx()

    def para(text: str = "", bold: bool = False, underline: bool = False, space_after: float = 0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            run = p.add_run(text)
            run.bold, run.underline = bold, underline
        return p

    ref = doc.add_paragraph()
    ref.paragraph_format.tab_stops.add_tab_stop(Inches(6.77), WD_TAB_ALIGNMENT.RIGHT)
    ref.add_run(f"Ref: {ctx.get('letter_no', '')}\tDate: {ctx.get('letter_date', '')}").bold = True
    ref.paragraph_format.space_after = Pt(14)

    para("To,")
    for line in _lines(ctx.get("recipient")):
        para(line)
    para(space_after=6)

    subj = para(space_after=12)
    subj.add_run("Subject: ").bold = True
    subj_run = subj.add_run(ctx.get("subject") or "")
    subj_run.bold = subj_run.underline = True

    for block in _body_paragraphs(ctx.get("body") or ""):
        p = para(space_after=8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for i, line in enumerate(block.split("\n")):
            run = p.add_run(line)
            if i < len(block.split("\n")) - 1:
                run.add_break()

    para(space_after=6)
    if ctx.get("company_name"):
        para(f"For {ctx['company_name']}", bold=True, space_after=36)
    if ctx.get("signatory_name"):
        para(ctx["signatory_name"], bold=True)
    if ctx.get("signatory_designation"):
        para(ctx["signatory_designation"])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
