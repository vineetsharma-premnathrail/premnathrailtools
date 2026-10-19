"""Fixed-format delivery documents for a machine (ERP Projects → Documents →
Created): Delivery Cover Letter, Delivery Confirmation and Guarantee
Certificate for Quality.

Each document is described once as a list of layout blocks (`_blocks`) and
rendered by both the PDF and the Word renderer below, so the two formats can
never drift apart. Every document prints twice: the "ORIGINAL FOR RECIPIENT"
copy that goes to the customer, and a "DUPLICATE FOR SUPPLIER" copy that is
signed/stamped by the customer and comes back as proof of delivery.

Context keys (from routes/projects.py `_letter_context`): letter_type,
letter_no, letter_date (display string), serial_number, company_name,
signatory_name, signatory_designation, and `data` — the template's fields.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Mapping

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Table, TableStyle, Paragraph, Spacer, KeepTogether, PageBreak
from xml.sax.saxutils import escape as _xml_escape

from app.modules.crm.reports.quotation_pdf import _draw_letterhead, _p, FONT_NAME, FONT_NAME_BOLD
from app.modules.erp.reports.letter import (
    LETTER_SIZE, LETTER_WIDTH, letterhead_pdf_doc, letterhead_docx, _lines,
)

TITLES = {
    "delivery_cover": "DELIVERY COVER LETTER",
    "delivery_confirmation": "DELIVERY CONFIRMATION",
    "guarantee_certificate": "GUARANTEE CERTIFICATE FOR QUALITY",
}
COPIES = ("(ORIGINAL FOR RECIPIENT)", "(DUPLICATE FOR SUPPLIER)")

DEFAULT_CONFIRMATION_STATEMENT = (
    "We have delivered the vehicle, as stated in the Letter of Agreement, to your site as per the tender, "
    "against the delivery challans listed below, and it is in good condition."
)
DEFAULT_COVER_STATEMENT = (
    "We have delivered the vehicle, as stated in the Letter of Agreement, to your site as per the tender, "
    "with the documents listed below."
)

GAP = ("", False)  # signature space inside a signature column


def _s(data: Mapping[str, Any], key: str) -> str:
    v = data.get(key)
    return str(v).strip() if v is not None else ""


def _fmt_invoice_date(value: str) -> str:
    try:
        return date.fromisoformat(value).strftime("%d-%b-%Y")
    except (TypeError, ValueError):
        return value or ""


def _particulars(ctx: Mapping[str, Any], data: Mapping[str, Any]) -> list[tuple[str, str]]:
    invoice = _s(data, "invoice_no")
    if invoice and _s(data, "invoice_date"):
        invoice += f" (Dt.: {_fmt_invoice_date(_s(data, 'invoice_date'))})"
    rows = [
        ("Project", _s(data, "project_title")),
        ("Client", _s(data, "client")),
        ("Application", _s(data, "application")),
        ("Vehicle", _s(data, "vehicle")),
        ("Model No.", _s(data, "model_no")),
        ("Serial No.", ctx.get("serial_number") or ""),
        ("LOA No.", _s(data, "loa_ref")),
        ("Invoice No.", invoice),
    ]
    return [(k, v) for k, v in rows if v]


def _supplier_sign(ctx: Mapping[str, Any]) -> list[tuple[str, bool]]:
    lines: list[tuple[str, bool]] = []
    if ctx.get("company_name"):
        lines.append((f"For {ctx['company_name']}", True))
    lines += [GAP, GAP]
    if ctx.get("signatory_name"):
        lines.append((ctx["signatory_name"], True))
    if ctx.get("signatory_designation"):
        lines.append((f"({ctx['signatory_designation']})", False))
    return lines


def _receiver_sign(data: Mapping[str, Any]) -> list[tuple[str, bool]]:
    org = _s(data, "receiver_org") or _s(data, "client")
    return [("RECEIVED FOR & ON BEHALF OF", True), GAP, GAP, (org, True), ("Name & Designation:", False), ("Date & Stamp:", False)]


def _blocks(ctx: Mapping[str, Any], copy_label: str) -> list[tuple]:
    kind = ctx["letter_type"]
    data: Mapping[str, Any] = ctx.get("data") or {}
    date_txt = f"Date: {ctx.get('letter_date', '')}"
    b: list[tuple] = [("copy", copy_label)]

    if kind == "delivery_confirmation":
        b += [("title", TITLES[kind]), ("space", 10), ("kv", _particulars(ctx, data)), ("space", 10),
              ("split", f"Delivery Note: {_s(data, 'delivery_note_no')}", date_txt), ("space", 8),
              ("para", _s(data, "statement") or DEFAULT_CONFIRMATION_STATEMENT), ("space", 4),
              ("table", ["S.No", "Delivery Challan No.", "Item Delivered"],
               [[str(i), it.get("challan_no", ""), it.get("description", "")] for i, it in enumerate(data.get("items") or [], 1)],
               [0.1, 0.3, 0.6]),
              ("space", 14), ("lines", ["Thank you", "Regards"]), ("space", 6),
              ("sign", _supplier_sign(ctx), _receiver_sign(data))]
    elif kind == "delivery_cover":
        client = _s(data, "receiver_org") or _s(data, "client")
        b += [("split", f"Ref: {ctx.get('letter_no', '')}", date_txt), ("space", 6),
              ("title", TITLES[kind]), ("space", 10), ("kv", _particulars(ctx, data)), ("space", 10),
              ("lines", ["To,", client]), ("space", 8),
              ("para", _s(data, "statement") or DEFAULT_COVER_STATEMENT),
              ("para", f"The vehicle bears Model No.: {_s(data, 'model_no') or '—'} and Serial No.: {ctx.get('serial_number') or '—'}."),
              ("numbered", list(data.get("documents") or [])),
              ("space", 14), ("lines", ["Thank you", "Regards"]), ("space", 6),
              ("sign", _supplier_sign(ctx), _receiver_sign(data))]
    else:  # guarantee_certificate
        loa = " ".join(ln.strip() for ln in _lines(_s(data, "loa_ref")) if ln.strip())
        company = ctx.get("company_name") or "the undersigned"
        text = (
            f"We, {company}, hereby guarantee the satisfactory operation and quality of the "
            f"{_s(data, 'supply_description')}, supplied by us to you against your LOA No.: {loa}, "
            f"for a period of {data.get('months_commissioning')} months from the date of commissioning or "
            f"{data.get('months_supply')} months from the date of supply, whichever is later. We shall be "
            "responsible for any failure of the machine to conform to the standard performance, and shall remedy "
            "at our cost any defects that may develop under proper use arising from faulty materials, design or "
            "workmanship in the supply made."
        )
        b += [("split", f"Ref: {ctx.get('letter_no', '')}", date_txt), ("space", 6),
              ("title", TITLES[kind]), ("space", 12),
              ("lines", ["To,"] + [ln for ln in _lines(_s(data, "recipient"))]), ("space", 10),
              ("lines", ["Dear Sir/Madam,"]), ("space", 6),
              ("para", text),
              ("space", 14), ("lines", ["Thank you", "Regards"]), ("space", 6),
              ("sign", _supplier_sign(ctx), None)]
    return b


# ── PDF ──────────────────────────────────────────────────────────────────────

def _pdf_sign_col(lines: list[tuple[str, bool]]) -> list[Any]:
    return [Spacer(1, 14) if not text else _p(text, size=LETTER_SIZE, bold=bold) for text, bold in lines]


def build_template_pdf(ctx: Mapping[str, Any]) -> io.BytesIO:
    buf = io.BytesIO()
    doc = letterhead_pdf_doc(buf, TITLES.get(ctx["letter_type"], "Document").title())
    width = LETTER_WIDTH - 12  # frame's 6pt side padding, so tables line up with body text
    para_style = ParagraphStyle("tbody", fontName=FONT_NAME, fontSize=LETTER_SIZE, leading=LETTER_SIZE + 4.5, alignment=TA_JUSTIFY, spaceAfter=8)
    no_pad = [("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("VALIGN", (0, 0), (-1, -1), "TOP")]

    story: list[Any] = []
    for copy_no, label in enumerate(COPIES):
        if copy_no:
            story.append(PageBreak())
        for block in _blocks(ctx, label):
            kind = block[0]
            if kind == "copy":
                # Aptos has no italic file bundled — the built-in oblique keeps this label italic.
                story.append(Paragraph(_xml_escape(block[1]), ParagraphStyle(
                    "copy", fontName="Helvetica-Oblique", fontSize=8.5, leading=12, alignment=1, spaceAfter=6)))
            elif kind == "title":
                story.append(Paragraph(_xml_escape(block[1]), ParagraphStyle(
                    "ttl", fontName=FONT_NAME_BOLD, fontSize=11, leading=15, alignment=1)))
            elif kind == "space":
                story.append(Spacer(1, block[1]))
            elif kind == "kv":
                rows = [[_p(f"{k}:", size=LETTER_SIZE, bold=True),
                         Paragraph(_xml_escape(v).replace("\n", "<br/>"), ParagraphStyle("kv", fontName=FONT_NAME, fontSize=LETTER_SIZE, leading=LETTER_SIZE + 3))]
                        for k, v in block[1]]
                if rows:
                    t = Table(rows, colWidths=[width * 0.18, width * 0.82])
                    t.setStyle(TableStyle(no_pad + [("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 1)]))
                    story.append(t)
            elif kind == "split":
                t = Table([[_p(block[1], size=LETTER_SIZE, bold=True), _p(block[2], size=LETTER_SIZE, bold=True, align="right")]],
                          colWidths=[width * 0.65, width * 0.35])
                t.setStyle(TableStyle(no_pad))
                story.append(t)
            elif kind == "para":
                story.append(Paragraph(_xml_escape(block[1]), para_style))
            elif kind == "lines":
                story += [_p(ln or " ", size=LETTER_SIZE) for ln in block[1]]
            elif kind == "numbered":
                for i, item in enumerate(block[1], 1):
                    t = Table([[_p(f"{i}.", size=LETTER_SIZE), _p(item, size=LETTER_SIZE)]], colWidths=[0.35 * inch, width - 0.35 * inch])
                    t.setStyle(TableStyle(no_pad + [("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
                    story.append(t)
            elif kind == "table":
                headers, rows, fracs = block[1], block[2], block[3]
                cell = ParagraphStyle("cell", fontName=FONT_NAME, fontSize=LETTER_SIZE - 0.5, leading=LETTER_SIZE + 2.5)
                data = [[_p(h, size=LETTER_SIZE - 0.5, bold=True) for h in headers]]
                data += [[Paragraph(_xml_escape(c).replace("\n", "<br/>"), cell) for c in r] for r in rows]
                t = Table(data, colWidths=[width * f for f in fracs], repeatRows=1)
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#555555")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2f2f2")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
            elif kind == "sign":
                left, right = block[1], block[2]
                cols = [_pdf_sign_col(left), _pdf_sign_col(right) if right else ""]
                t = Table([cols], colWidths=[width * 0.5, width * 0.5])
                t.setStyle(TableStyle(no_pad))
                story.append(KeepTogether([t]))

    doc.build(story, onFirstPage=_draw_letterhead, onLaterPages=_draw_letterhead)
    buf.seek(0)
    return buf


# ── Word ─────────────────────────────────────────────────────────────────────

def build_template_docx(ctx: Mapping[str, Any]) -> io.BytesIO:
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

    doc = letterhead_docx()
    text_width = 6.77  # inches between the 0.75in margins on A4

    def para(text: str = "", bold: bool = False, space_after: float = 0, container=None, align=None):
        p = (container or doc).add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        if text:
            p.add_run(text).bold = bold
        return p

    def borderless(rows: int, widths: list[float]):
        t = doc.add_table(rows=rows, cols=len(widths))
        t.autofit = False
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        return t

    for copy_no, label in enumerate(COPIES):
        if copy_no:
            doc.add_page_break()
        for block in _blocks(ctx, label):
            kind = block[0]
            if kind == "copy":
                p = para(block[1], space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
                p.runs[0].italic = True
                p.runs[0].font.size = Pt(8.5)
            elif kind == "title":
                p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
                run = p.add_run(block[1])
                run.bold = True
                run.font.size = Pt(11)
            elif kind == "space":
                para(space_after=max(0, block[1] - LETTER_SIZE))
            elif kind == "kv":
                t = borderless(len(block[1]), [1.25, text_width - 1.25])
                for row, (k, v) in zip(t.rows, block[1]):
                    row.cells[0].paragraphs[0].add_run(f"{k}:").bold = True
                    row.cells[1].paragraphs[0].add_run(v)
            elif kind == "split":
                p = para()
                p.paragraph_format.tab_stops.add_tab_stop(Inches(text_width), WD_TAB_ALIGNMENT.RIGHT)
                p.add_run(f"{block[1]}\t{block[2]}").bold = True
            elif kind == "para":
                para(block[1], space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif kind == "lines":
                for ln in block[1]:
                    para(ln)
            elif kind == "numbered":
                for i, item in enumerate(block[1], 1):
                    p = para(f"{i}.\t{item}", space_after=3)
                    p.paragraph_format.left_indent = Inches(0.35)
                    p.paragraph_format.first_line_indent = Inches(-0.35)
            elif kind == "table":
                headers, rows, fracs = block[1], block[2], block[3]
                t = doc.add_table(rows=1 + len(rows), cols=len(headers))
                t.style = "Table Grid"
                for r_idx, values in enumerate([headers] + rows):
                    for c_idx, value in enumerate(values):
                        cell = t.rows[r_idx].cells[c_idx]
                        cell.width = Inches(text_width * fracs[c_idx])
                        cell.paragraphs[0].add_run(value).bold = r_idx == 0
            elif kind == "sign":
                left, right = block[1], block[2]
                t = borderless(1, [text_width / 2, text_width / 2])
                for cell, lines in zip(t.rows[0].cells, (left, right or [])):
                    first = True
                    for text, bold in lines:
                        p = cell.paragraphs[0] if first else cell.add_paragraph()
                        first = False
                        if text:
                            p.add_run(text).bold = bold
                        else:
                            p.paragraph_format.space_after = Pt(6)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
