"""Shared helpers for the ERP bulk import / Excel export (Machines and
Service Requests) — same contract as the Store Item Master import:
rows that already exist are NOT saved and come back in `existing` with the
record they matched; rows with a problem come back in `errors` with the
reason and how to fix it; valid rows are created.

Headers are matched after normalising ("Serial Number", "serial number" and
"serial_number" are the same column), so an exported Excel can be edited and
imported straight back."""
import csv
import io
import re
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from pydantic import ValidationError

MAX_ROWS = 5000

IST = ZoneInfo("Asia/Kolkata")


def parse_export_ids(ids: str | None, what: str) -> list[int] | None:
    """`ids` query param ("3,7,12") for "export the rows I ticked"."""
    if ids is None or not ids.strip():
        return None
    out: list[int] = []
    for part in ids.split(","):
        part = part.strip()
        if not part:
            continue
        if not part.isdigit():
            raise HTTPException(status_code=422, detail=f"Selected {what} list has an invalid id '{part}' — clear the selection and tick the rows again.")
        out.append(int(part))
    if not out:
        raise HTTPException(status_code=422, detail=f"No {what} selected — tick at least one row, or export by date range instead.")
    return out


def check_date_range(date_from: date | None, date_to: date | None) -> None:
    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail=f"From date ({date_from.strftime('%d-%m-%Y')}) is after To date ({date_to.strftime('%d-%m-%Y')}) — swap them or pick a later To date.",
        )


def ist_date(ts: datetime | None) -> date | None:
    """A stored timestamp as the calendar day in India (what users see)."""
    if ts is None:
        return None
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=timezone.utc)
    return ts.astimezone(IST).date()


def in_range(d: date | None, date_from: date | None, date_to: date | None) -> bool:
    """Inclusive on both ends. A row with no date is outside any range."""
    if date_from is None and date_to is None:
        return True
    if d is None:
        return False
    return (date_from is None or d >= date_from) and (date_to is None or d <= date_to)

_HEADER_JUNK = re.compile(r"[^a-z0-9]+")


def norm_header(h: object) -> str:
    return _HEADER_JUNK.sub("_", str(h or "").strip().lower()).strip("_")


def norm_key(v: str | None) -> str:
    """Comparison key for identifiers: case- and whitespace-insensitive, so
    'SN 001', 'sn  001' and 'SN 001 ' are one machine."""
    return " ".join((v or "").split()).lower()


def read_rows(filename: str, raw: bytes) -> list[dict]:
    """Rows as {normalised header: cell}. Excel dates stay `datetime`, every
    other cell is a stripped string ('' when blank)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        except Exception:
            raise HTTPException(status_code=422, detail="Could not open the Excel file — save it as .xlsx (or CSV) and try again.")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [norm_header(h) for h in rows[0]]
        out = []
        for r in rows[1:]:
            row = {}
            for i, v in enumerate(r):
                if i >= len(headers) or not headers[i]:
                    continue
                if isinstance(v, (datetime, date)):
                    row[headers[i]] = v
                elif isinstance(v, float) and v.is_integer():
                    row[headers[i]] = str(int(v))  # 12345.0 → "12345" (serials, phones, pincodes)
                else:
                    row[headers[i]] = "" if v is None else str(v).strip()
            out.append(row)
        return out
    if not name.endswith(".csv"):
        raise HTTPException(status_code=422, detail="Upload a .csv or .xlsx file — download the template to start from the right columns.")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=422, detail="Could not read the file as UTF-8 text — save it as CSV (UTF-8) from Excel and try again.")
    reader = csv.DictReader(io.StringIO(text))
    return [{norm_header(k): (v or "").strip() for k, v in r.items() if k} for r in reader]


def check_rows(rows: list[dict], required: list[str]) -> None:
    if not rows:
        raise HTTPException(status_code=422, detail="The file has no data rows — fill in at least one row under the header row.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=422, detail=f"The file has {len(rows)} rows — import at most {MAX_ROWS} at a time by splitting it.")
    missing = [c for c in required if c not in rows[0]]
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing the {', '.join(missing)} column(s) — download the template and keep its header row.")


def text(r: dict, key: str) -> str | None:
    v = r.get(key)
    if isinstance(v, (datetime, date)):
        return v.strftime("%d-%m-%Y")
    return (v or "").strip() or None


_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y")


def parse_date(r: dict, key: str, label: str) -> date | None:
    v = r.get(key)
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    v = (v or "").strip()
    if not v:
        return None
    v = v.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2} ", v) else v
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{label} '{v}' isn't a date — use DD-MM-YYYY (e.g. 15-08-2026).")


def parse_bool(r: dict, key: str, label: str) -> bool:
    v = (text(r, key) or "").lower()
    if v in ("", "no", "n", "false", "0"):
        return False
    if v in ("yes", "y", "true", "1"):
        return True
    raise ValueError(f"{label} '{v}' isn't valid — write Yes or No.")


def parse_number(r: dict, key: str, label: str, *, minimum: float = 0, maximum: float | None = None) -> float | None:
    v = (text(r, key) or "").replace(",", "")
    if not v:
        return None
    try:
        n = float(v)
    except ValueError:
        raise ValueError(f"{label} '{v}' isn't a number — write digits only (e.g. 1500).")
    if n < minimum or (maximum is not None and n > maximum):
        rng = f"{minimum:g}–{maximum:g}" if maximum is not None else f"{minimum:g} or more"
        raise ValueError(f"{label} {n:g} is out of range — it must be {rng}.")
    return n


def enum_label(value: str) -> str:
    return value.replace("_", " ").title()


def parse_choice(r: dict, key: str, label: str, choices: dict[str, str], default: str | None = None) -> str | None:
    """Accepts the stored value (repair_breakdown) or its label (Repair / Breakdown)."""
    v = (text(r, key) or "").strip()
    if not v:
        return default
    k = norm_header(v)
    for value, lbl in choices.items():
        if k in (norm_header(value), norm_header(lbl)):
            return value
    raise ValueError(f"{label} '{v}' isn't valid — use one of: {', '.join(choices.values())}.")


def validation_reason(exc: ValidationError, labels: dict[str, str]) -> str:
    """Pydantic errors as one readable sentence, using column titles."""
    parts = []
    for e in exc.errors():
        field = str(e["loc"][0]) if e.get("loc") else ""
        msg = str(e.get("msg", "")).removeprefix("Value error, ")
        parts.append(f"{labels.get(field, field)}: {msg}" if field else msg)
    return " ".join(parts)


def csv_template(columns: list[tuple[str, str]], samples: list[list[str]], filename: str) -> StreamingResponse:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([key for key, _ in columns])
    for s in samples:
        w.writerow(s)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def xlsx_response(sheet: str, headers: list[str], rows: list[list], filename: str, widths: list[int] | None = None) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FDF1E6")
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        width = widths[i - 1] if widths and i <= len(widths) else max(12, min(40, len(h) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "DD-MM-YYYY"
            elif isinstance(cell.value, str) and cell.value[:1] in _FORMULA_STARTS:
                # Spreadsheet formula injection: openpyxl writes a string starting with "=" as a
                # formula. Forcing the string type stores it as plain text (no visible apostrophe,
                # so an exported file still re-imports unchanged); numbers and dates are untouched.
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})
