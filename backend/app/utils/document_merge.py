"""Join two documents into one, the first followed by the second on a new page.

Used for the quotation that goes to the customer: R&D's Technical Offer on top,
the commercial quotation below (owner, 2026-10-27). Word files are joined with
docxcompose so images, headers / footers, tables and styles of both survive and
the result stays editable in Word; PDFs with pypdf."""
import io

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer
from pypdf import PdfWriter


def merge_docx(first: bytes, second: bytes) -> bytes:
    master = Document(io.BytesIO(first))
    master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    composer = Composer(master)
    composer.append(Document(io.BytesIO(second)))
    out = io.BytesIO()
    composer.save(out)
    return out.getvalue()


def merge_pdf(first: bytes, second: bytes) -> bytes:
    writer = PdfWriter()
    for part in (first, second):
        writer.append(io.BytesIO(part))
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()
