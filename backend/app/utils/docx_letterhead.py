"""Shared building blocks for system-generated Word documents on the
company letterhead (same logo.JPG header / logo-1.JPG footer banners as the
PDFs): a ready-to-fill A4 Document plus small paragraph/table helpers.
Used by the quotation, RRV handover certificate, etc.
"""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Emu, RGBColor

from app.modules.crm.reports.technical_offer_docx import (
    _cropped_content_bytes, _add_full_width_banner, _image_aspect,
    HEADER_LOGO_PATH, FOOTER_LOGO_PATH, PAGE_WIDTH_EMU,
)

FONT = "Aptos"
BODY = 9
HEADER = 10
ACCENT = RGBColor(0x1C, 0x21, 0x3E)
ACCENT_HEX = "1C213E"
ACCENT_LIGHT_HEX = "EEF0F6"
MUTED = RGBColor(0x57, 0x53, 0x4E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

PAGE_W_IN, PAGE_H_IN = 8.27, 11.69
MARGIN_IN = 0.55
USABLE_IN = PAGE_W_IN - 2 * MARGIN_IN

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


def fmt_run(run: Any, size: float = BODY, bold: bool = False, color: RGBColor | None = None) -> Any:
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def tight(p: Any, before: float = 0, after: float = 0) -> Any:
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    return p


def write(cell_or_doc: Any, text: str, size: float = BODY, bold: bool = False, align: str = "left",
           color: RGBColor | None = None, first: bool = False, after: float = 0) -> Any:
    """Add a paragraph to a cell/document. first=True reuses a cell's built-in
    empty paragraph instead of leaving a blank line above the text."""
    p = cell_or_doc.paragraphs[0] if first else cell_or_doc.add_paragraph()
    tight(p, after=after)
    p.alignment = ALIGN[align]
    fmt_run(p.add_run(text or ""), size=size, bold=bold, color=color)
    return p


def shade(cell: Any, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def cell_margins(table: Any, left: int = 70, right: int = 70, top: int = 60, bottom: int = 60) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_widths(table: Any, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def top_border(cell: Any, color_hex: str, size_eighths: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size_eighths))
    top.set(qn("w:color"), color_hex)
    borders.append(top)
    tc_pr.append(borders)



def new_letterhead_document(margin_in: float = MARGIN_IN) -> Any:
    """A4 Document with the letterhead banners edge-to-edge in the header and
    footer, body margins set so text starts below / stops above them."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_W_IN), Inches(PAGE_H_IN)
    section.left_margin = section.right_margin = Inches(margin_in)
    section.header_distance = section.footer_distance = Inches(0)

    header_bytes = _cropped_content_bytes(HEADER_LOGO_PATH)
    footer_bytes = _cropped_content_bytes(FOOTER_LOGO_PATH)
    header_h_in = PAGE_W_IN * _image_aspect(header_bytes) if header_bytes else 0.85
    footer_h_in = PAGE_W_IN * _image_aspect(footer_bytes) if footer_bytes else 0.55
    # Body starts below the header banner and stops above the footer banner,
    # matching the PDFs' MARGIN + banner height + 0.15in.
    section.top_margin = Inches(MARGIN_IN + header_h_in + 0.15)
    section.bottom_margin = Inches(MARGIN_IN + footer_h_in + 0.15)
    if header_bytes:
        _add_full_width_banner(section.header.paragraphs[0], header_bytes, offset_y_emu=0)
    if footer_bytes:
        footer_height_emu = int(int(PAGE_WIDTH_EMU) * _image_aspect(footer_bytes))
        _add_full_width_banner(section.footer.paragraphs[0], footer_bytes, offset_y_emu=int(Emu(Inches(PAGE_H_IN))) - footer_height_emu)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal.paragraph_format.space_after = Pt(0)
    return doc
