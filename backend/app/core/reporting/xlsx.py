"""The designed report workbook (Summary, grouped summary, records, definitions).

Styling is lifted from the P2P M.I.S export (`modules/p2p/routes/mis.py`, left
unchanged): dark title band, orange accent row, brand header, zebra body, thin
borders, data bars and native charts. Everything a customer or engineer typed
is untrusted text: a cell that starts with = + - @ (or tab / CR) is written as
a string so Excel never runs it as a formula."""
from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings
from app.core.record_ref import encode_ref
from app.core.reporting import dates as rdates

BRAND = "FF7A45"
INK = "0F172A"
MUTED = "64748B"
ZEBRA = "F8FAFC"
GRID = "E2E8F0"
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
HEADER_ROW = 7
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")

FORMATS = {
    "int": "#,##0", "decimal": "#,##0.00", "money": '"₹" #,##,##0.00', "percent": "0.0%",
    "hours": "0.0", "days": "0.0", "date": "DD-MM-YYYY", "datetime": "DD-MM-YYYY HH:MM", "time": "HH:MM",
}
STATUS_COLORS = {
    "open": "2563EB", "acknowledged": "2563EB", "assigned": "2563EB", "scheduled": "2563EB",
    "in_progress": "F59E0B", "pending_parts": "F59E0B", "on_hold": "F59E0B", "work_completed": "16A34A",
    "review": "16A34A", "closed": "64748B", "cancelled": "64748B", "urgent": "DC2626", "standard": "64748B",
    "paid": "16A34A", "invoiced": "2563EB", "to_bill": "F59E0B", "not_billable": "64748B",
}


def argb(hex6: str) -> str:
    """openpyxl cell colours need 8-digit ARGB (a bare RGB gets alpha 00 = invisible)."""
    return f"FF{hex6}"


def _fill(hex6: str) -> PatternFill:
    return PatternFill(start_color=argb(hex6), end_color=argb(hex6), fill_type="solid")


_THIN = Side(style="thin", color=argb(GRID))
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
TOTAL_BORDER = Border(left=_THIN, right=_THIN, top=Side(style="double", color=argb(INK)), bottom=_THIN)
BODY_FONT = Font(size=10.5, color=argb("1E293B"))


def put(cell, value: Any) -> None:
    """Write a value; text that could run as a formula is forced to a string."""
    if isinstance(value, str) and value.startswith(FORMULA_STARTS):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def safe_sheet_name(label: str, taken: set[str], suffix: str = "") -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", label).strip() or "Report"
    name = (base[: 31 - len(suffix)] + suffix).strip()
    n, cand = 2, name
    while cand.lower() in taken:
        tail = f" {n}"
        cand = name[: 31 - len(tail)] + tail
        n += 1
    taken.add(cand.lower())
    return cand


def safe_filename(module_label: str, report_label: str, d_from: str | None, d_to: str | None) -> str:
    def clean(s: str) -> str:
        return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")
    mod = clean(module_label.split("&")[0]) or "Report"
    rng = f"_{d_from}_to_{d_to}" if d_from and d_to else "_all_dates"
    return f"{mod}_{clean(report_label) or 'Report'}{rng}.xlsx"


def _value_for(col_type: str, v: Any) -> Any:
    if v is None:
        return None
    if col_type == "bool":
        return "Yes" if v else "No"
    if col_type == "datetime" and isinstance(v, datetime):
        return rdates.to_ist(v).replace(tzinfo=None)
    if col_type == "datetime" and isinstance(v, str):
        try:
            return rdates.to_ist(datetime.fromisoformat(v)).replace(tzinfo=None)
        except ValueError:
            return v
    if col_type == "date" and isinstance(v, str):
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    if col_type == "status" and isinstance(v, str):
        return v.replace("_", " ").title()
    return v


def _style_cell(cell, col_type: str, zebra: bool, wrap: bool = False) -> None:
    cell.font = BODY_FONT
    cell.border = BORDER
    if zebra:
        cell.fill = _fill(ZEBRA)
    fmt = FORMATS.get(col_type)
    if fmt:
        cell.number_format = fmt
    numeric = col_type in ("int", "decimal", "money", "percent", "hours", "days")
    cell.alignment = Alignment(horizontal="right" if numeric else "left", vertical="top" if wrap else "center", wrap_text=wrap)


def _title_block(ws: Worksheet, company: str, title: str, lines: list[str], n_cols: int) -> None:
    n_cols = max(n_cols, 6)
    last = get_column_letter(n_cols)
    ws.sheet_view.showGridLines = False
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    put(c, f"{company} - {title}")
    c.font = Font(color=argb("FFFFFF"), bold=True, size=16)
    c.fill = _fill(INK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    for i in range(4):  # rows 2-5
        r = 2 + i
        ws.merge_cells(f"A{r}:{last}{r}")
        cell = ws[f"A{r}"]
        put(cell, lines[i] if i < len(lines) else "")
        cell.font = Font(color=argb("CBD5E1"), size=10.5)
        cell.fill = _fill(INK)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 17
    for col in range(1, n_cols + 1):
        ws.cell(row=6, column=col).fill = _fill(BRAND)
    ws.row_dimensions[6].height = 4


def _print_setup(ws: Worksheet, title: str, header_row: int | None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddHeader.center.text = title
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i)
        put(c, lab)
        c.fill = _fill(BRAND)
        c.font = Font(color=argb("FFFFFF"), bold=True, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        c.border = BORDER
    ws.row_dimensions[row].height = 20


def _widths(ws: Worksheet, headers: list[str], rows: list[list], specs: list[dict]) -> None:
    for i, h in enumerate(headers, start=1):
        spec = specs[i - 1] if i - 1 < len(specs) else {}
        if spec.get("width"):
            w = spec["width"]
        elif spec.get("cut"):
            w = 50
        else:
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r) and r[i - 1] is not None])
            w = min(max(longest + 3, 10), 45)
        ws.column_dimensions[get_column_letter(i)].width = w


def _unit_label(label: str, t: str) -> str:
    if t == "hours" and "(h" not in label and " h" not in label.lower():
        return f"{label} (h)"
    if t == "days" and "day" not in label.lower():
        return f"{label} (days)"
    return label


def describe_lines(result: dict, meta: dict) -> list[str]:
    p = result.get("period") or {}
    if p.get("mode") == "snapshot":
        per = f"As of {rdates.fmt(rdates.today_ist())} (position today, not a period)"
        fa = result.get("filters_applied") or []
        flt = "Filters: " + " | ".join(f"{x['label']}: {x['value']}" for x in fa) if fa else "Filters: none"
        return [f"Period: {per}", flt, f"Data: your access - {result.get('scope_label') or meta.get('scope_label', '')}",
                f"Generated {meta.get('generated', '')} by {meta.get('user_name', '')}"]
    if p.get("from"):
        f, t = date.fromisoformat(p["from"]), date.fromisoformat(p["to"])
        per = f"{p.get('date_field_label', 'Period')} - {p.get('label', '')} ({rdates.fmt(f)} to {rdates.fmt(t)})"
        if p.get("compare"):
            c = p["compare"]
            per += f" compared with {rdates.fmt(date.fromisoformat(c['from']))} to {rdates.fmt(date.fromisoformat(c['to']))}"
    else:
        per = f"{p.get('date_field_label', 'Dates')}: all dates"
    fa = result.get("filters_applied") or []
    flt = "Filters: " + " | ".join(f"{x['label']}: {x['value']}{' (fixed)' if x.get('locked') else ''}" for x in fa) if fa else "Filters: none"
    if result.get("q"):
        flt += f" | Search: {result['q']}"
    return [f"Period: {per}", flt, f"Data: your access - {result.get('scope_label') or meta.get('scope_label', '')}",
            f"Generated {meta.get('generated', '')} by {meta.get('user_name', '')}"]


def _kpi_sheet(ws: Worksheet, kpis: list[dict], row0: int) -> int:
    for col in "ABCD":
        ws.column_dimensions[col].width = 28
    r = row0
    for i, k in enumerate(kpis):
        r = row0 + (i // 2) * 4
        c = 1 if i % 2 == 0 else 3
        lab = ws.cell(row=r, column=c)
        put(lab, k["label"])
        lab.font = Font(size=10, color=argb(MUTED), bold=True)
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        val = ws.cell(row=r + 1, column=c)
        put(val, _value_for(k["type"], k.get("value")))
        val.font = Font(size=22, color=argb(BRAND), bold=True)
        val.alignment = Alignment(horizontal="left", vertical="center")
        if FORMATS.get(k["type"]):
            val.number_format = FORMATS[k["type"]]
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
        cmp_cell = ws.cell(row=r + 2, column=c)
        if k.get("compare") is not None:
            d = k.get("delta_pct")
            txt = f"vs {k['compare']}" + (f" ({d:+.1f}%)" if d is not None else "")
            good = None
            if d is not None and k.get("better"):
                good = (d >= 0) == (k["better"] == "up")
            put(cmp_cell, txt)
            cmp_cell.font = Font(size=10, color=argb("16A34A" if good else "DC2626" if good is False else MUTED))
        ws.merge_cells(start_row=r + 2, start_column=c, end_row=r + 2, end_column=c + 1)
        for cc in (c, c + 1):
            for rr in (r, r + 1, r + 2):
                ws.cell(row=rr, column=cc).fill = _fill(ZEBRA)
                ws.cell(row=rr, column=cc).border = BORDER
    return r + 4


def _add_chart(ws: Worksheet, series: dict, anchor_row: int, anchor_col: int = 1) -> int:
    """Write a visible 'Chart data' block at anchor_row and a native chart beside/under it.
    Returns the next free row."""
    labels, datasets = series.get("labels") or [], series.get("datasets") or []
    if not labels or not datasets:
        return anchor_row
    r0 = anchor_row
    t = ws.cell(row=r0, column=1)
    put(t, f"Chart data - {series.get('label', '')}")
    t.font = Font(bold=True, size=10, color=argb(MUTED))
    hdr = r0 + 1
    _header(ws, hdr, [""] + [str(d.get("label", "")) for d in datasets])
    for i, lab in enumerate(labels):
        put(ws.cell(row=hdr + 1 + i, column=1), str(lab))
        for j, d in enumerate(datasets):
            data = d.get("data") or []
            ws.cell(row=hdr + 1 + i, column=2 + j, value=data[i] if i < len(data) else None)
    n = len(labels)
    kind = series.get("kind", "bar")
    if kind == "doughnut":
        chart = DoughnutChart()
    elif kind == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if kind == "hbar" else "col"
        if kind == "stacked_bar":
            chart.grouping = "stacked"
            chart.overlap = 100
    data_ref = Reference(ws, min_col=2, max_col=1 + len(datasets), min_row=hdr, max_row=hdr + n)
    cats = Reference(ws, min_col=1, min_row=hdr + 1, max_row=hdr + n)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.title = series.get("label")
    chart.height, chart.width = 8, 16
    if kind != "doughnut":
        for i, s in enumerate(chart.series):
            s.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            s.graphicalProperties.line.solidFill = PALETTE[i % len(PALETTE)]
    ws.add_chart(chart, f"{get_column_letter(max(4, 3 + len(datasets)))}{r0}")
    return hdr + n + max(2, 17 - n)


def _write_summary_table(wb: Workbook, name: str, report_label: str, company: str, result: dict, meta: dict) -> None:
    ws = wb.create_sheet(name)
    measures = result.get("measures") or []
    cmp_on = bool((result.get("period") or {}).get("compare"))
    heads = [result.get("group_label") or ("Group" if result.get("group_by") else "Total")]
    specs: list[dict] = [{}]
    types = ["text"]
    for m in measures:
        lab = _unit_label(m["label"], m["type"])
        if cmp_on:
            heads += [f"{lab} - this period", f"{lab} - previous", f"{lab} - change %"]
            specs += [{}, {}, {}]
            types += [m["type"], m["type"], "decimal"]
        else:
            heads.append(lab)
            specs.append({})
            types.append(m["type"])
    _title_block(ws, company, f"{meta['module_label']} Reports - {report_label}", describe_lines(result, meta), len(heads))
    _header(ws, HEADER_ROW, heads)
    r = HEADER_ROW + 1
    body: list[list] = []

    def emit(row: dict, level: int) -> None:
        nonlocal r
        label = next(iter((row.get("label") or {}).values()), "")
        vals = [label]
        for m in measures:
            v = (row.get("values") or {}).get(m["key"])
            if cmp_on:
                vals += [v, (row.get("compare") or {}).get(m["key"]), (row.get("delta_pct") or {}).get(m["key"])]
            else:
                vals.append(v)
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=ci)
            put(cell, _value_for(types[ci - 1], v))
            _style_cell(cell, types[ci - 1], (r - HEADER_ROW) % 2 == 0)
            if level == 0 and result.get("group_by2") is not None:
                cell.font = Font(size=10.5, bold=True, color=argb("1E293B"))
            if ci == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=2 * level)
        if level:
            ws.row_dimensions[r].outline_level = 1
        body.append(vals)
        r += 1
        for ch in row.get("children") or []:
            emit(ch, level + 1)

    for row in result.get("rows") or []:
        emit(row, 0)
    ws.sheet_properties.outlinePr.summaryBelow = False
    first_body = HEADER_ROW + 1
    if result.get("totals", {}).get("values") is not None and measures:
        tv = result["totals"]
        vals = ["Total"]
        for m in measures:
            v = (tv.get("values") or {}).get(m["key"])
            if cmp_on:
                vals += [v, (tv.get("compare") or {}).get(m["key"]), (tv.get("delta_pct") or {}).get(m["key"])]
            else:
                vals.append(v)
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=ci)
            put(cell, _value_for(types[ci - 1], v))
            _style_cell(cell, types[ci - 1], False)
            cell.font = Font(size=10.5, bold=True)
            cell.border = TOTAL_BORDER
        r += 1
    last_body = r - 2 if measures else r - 1
    # data bar on the first count column
    for ci, t in enumerate(types[1:], start=2):
        if t == "int" and last_body >= first_body:
            L = get_column_letter(ci)
            ws.conditional_formatting.add(f"{L}{first_body}:{L}{last_body}", DataBarRule(start_type="min", end_type="max", color=argb(BRAND)))
            break
    _widths(ws, heads, body, specs)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 24)
    ws.freeze_panes = f"B{HEADER_ROW + 1}"
    if body:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(heads))}{HEADER_ROW + len(body)}"
    nxt = r + 2
    for s in result.get("series") or []:
        nxt = _add_chart(ws, s, nxt)
    for note in result.get("notes") or []:
        put(ws.cell(row=nxt, column=1), note)
        nxt += 1
    _print_setup(ws, f"{report_label} - summary", HEADER_ROW)


def _write_records(wb: Workbook, name: str, report_label: str, company: str, detail: dict, meta: dict) -> None:
    ws = wb.create_sheet(name)
    cols = detail["columns"]
    heads = [_unit_label(c["label"], c["type"]) for c in cols]
    _title_block(ws, company, f"{meta['module_label']} Reports - {report_label} (records)", describe_lines(detail, meta), len(heads))
    _header(ws, HEADER_ROW, heads)
    base = settings.FRONTEND_URL.rstrip("/")
    rows = detail["rows"]
    plain: list[list] = []
    for ri, row in enumerate(rows):
        r = HEADER_ROW + 1 + ri
        out = []
        for ci, c in enumerate(cols, start=1):
            v = _value_for(c["type"], row.get(c["key"]))
            cell = ws.cell(row=r, column=ci)
            put(cell, v)
            _style_cell(cell, c["type"], ri % 2 == 1, wrap=bool(c.get("cut")))
            if c["type"] == "status" and row.get(c["key"]) is not None:
                col = STATUS_COLORS.get(str(row.get(c["key"])).lower(), MUTED)
                cell.fill = _fill(col)
                cell.font = Font(color=argb("FFFFFF"), bold=True, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c.get("link") and row.get(c["id_key"]) is not None and v is not None:
                cell.hyperlink = f"{base}{c['link']}/{encode_ref(int(row[c['id_key']]))}"
                cell.font = Font(size=10.5, color=argb("2563EB"), underline="single")
            out.append(v)
        plain.append(out)
    last = HEADER_ROW + len(rows)
    tot_row = last + 1
    if rows and detail.get("totals"):
        for ci, c in enumerate(cols, start=1):
            cell = ws.cell(row=tot_row, column=ci)
            if ci == 1:
                put(cell, "Total")
            if c["key"] in detail["totals"]:
                L = get_column_letter(ci)
                cell.value = f"=SUBTOTAL(109,{L}{HEADER_ROW + 1}:{L}{last})"
                _style_cell(cell, c["type"], False)
            cell.font = Font(size=10.5, bold=True)
            cell.border = TOTAL_BORDER
    _widths(ws, heads, plain, [{"cut": c.get("cut"), "width": c.get("width")} for c in cols])
    ws.freeze_panes = f"B{HEADER_ROW + 1}"
    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(cols))}{last}"
    _print_setup(ws, f"{report_label} - records", HEADER_ROW)


def build_workbook(parts: list[dict], meta: dict) -> bytes:
    """parts: [{"label", "summary": result dict | None, "detail": result dict | None, "definitions": [(metric, text)]}].
    meta: company, module_label, user_name, generated, scope_label, sheets (subset of summary/detail/definitions)."""
    company = meta.get("company") or "Premnath Rail Tools"
    sheets = set(meta.get("sheets") or ("summary", "detail", "definitions"))
    wb = Workbook()
    first = wb.active
    first.title = "Summary"
    main = parts[0]
    ref = main.get("summary") or main.get("detail") or {}
    _title_block(first, company, f"{meta['module_label']} Reports - {main['label']}", describe_lines(ref, meta), 4)
    grp = ref.get("group_by")
    if grp:
        by = ref.get("group_label") or grp
        put(first["A5"], f"Grouped by {by}" + (f", then {ref['group_by2']}" if ref.get("group_by2") else ""))
        first["A5"].font = Font(color=argb("CBD5E1"), size=10.5)
    kpis = (main.get("summary") or {}).get("kpis") or []
    nxt = _kpi_sheet(first, kpis, HEADER_ROW + 1) if kpis else HEADER_ROW + 1
    if len(parts) > 1:
        put(first.cell(row=nxt, column=1), "Reports in this file")
        first.cell(row=nxt, column=1).font = Font(bold=True, color=argb(MUTED))
        for i, p in enumerate(parts):
            put(first.cell(row=nxt + 1 + i, column=1), p["label"])
    _print_setup(first, f"{meta['module_label']} Reports", None)
    taken = {"summary", "definitions"}
    for p in parts:
        if p.get("summary") is not None and "summary" in sheets:
            _write_summary_table(wb, safe_sheet_name(p["label"], taken, " - summary"), p["label"], company, p["summary"], meta)
        if p.get("detail") is not None and "detail" in sheets:
            _write_records(wb, safe_sheet_name(p["label"], taken, " - records"), p["label"], company, p["detail"], meta)
    if "definitions" in sheets:
        ws = wb.create_sheet("Definitions")
        _title_block(ws, company, f"{meta['module_label']} Reports - How the figures are calculated", [], 3)
        _header(ws, HEADER_ROW, ["Report", "Metric", "How it is calculated"])
        r = HEADER_ROW + 1
        for p in parts:
            for metric, text in p.get("definitions") or []:
                for ci, v in enumerate((p["label"], metric, text), start=1):
                    cell = ws.cell(row=r, column=ci)
                    put(cell, v)
                    _style_cell(cell, "text", r % 2 == 0, wrap=(ci == 3))
                r += 1
        for L, w in (("A", 30), ("B", 34), ("C", 90)):
            ws.column_dimensions[L].width = w
        _print_setup(ws, "Definitions", HEADER_ROW)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _table_chart(ws: Worksheet, kind: str, title: str, hdr: int, n: int, n_series: int, anchor_col: int) -> None:
    if n < 1 or n_series < 1:
        return
    if kind == "donut":
        chart = DoughnutChart()
    elif kind == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if kind == "hbar" else "col"
        if kind == "stacked_bar":
            chart.grouping = "stacked"
            chart.overlap = 100
    chart.add_data(Reference(ws, min_col=2, max_col=1 + n_series, min_row=hdr, max_row=hdr + n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=hdr + 1, max_row=hdr + n))
    chart.title = title
    chart.height, chart.width = 9, 17
    if kind != "donut":
        for i, s in enumerate(chart.series):
            s.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            s.graphicalProperties.line.solidFill = PALETTE[i % len(PALETTE)]
    ws.add_chart(chart, f"{get_column_letter(anchor_col)}{HEADER_ROW}")


def _write_visual(wb: Workbook, name: str, company: str, spec: dict, res: dict, page_result: dict, meta: dict) -> None:
    lines = describe_lines(page_result, meta)
    if res.get("kind") == "table" and res.get("columns") and "rows" in res and res["rows"] and "label" not in res["rows"][0]:
        _write_records(wb, name, spec["title"], company, {**res, "totals": {}, **{k: page_result.get(k) for k in ("period", "filters_applied", "scope_label")}}, meta)
        return
    ws = wb.create_sheet(name)
    kind = res.get("kind", "column")
    if kind == "table":
        heads = ["Group"] + [_unit_label(c["label"], c["type"]) for c in res["columns"]]
        types = ["text"] + [c["type"] for c in res["columns"]]
        body = [[r["label"]] + [r["values"].get(c["key"]) for c in res["columns"]] for r in res["rows"]]
    else:
        ds = res.get("datasets") or []
        mtypes = {m["key"]: m["type"] for m in res.get("measures", [])}
        heads = ["Group"] + [str(d["label"]) for d in ds]
        types = ["text"] + [mtypes.get(d.get("measure"), (res.get("measure") or {}).get("type", "int")) for d in ds]
        body = [[lab] + [d["data"][i] if i < len(d["data"]) else None for d in ds] for i, lab in enumerate(res.get("labels", []))]
        if res.get("compare"):
            heads.append("Previous period")
            types.append(types[1] if len(types) > 1 else "int")
            for i, row in enumerate(body):
                row.append(res["compare"][i])
    _title_block(ws, company, f"{meta['module_label']} Reports - {spec['title']}", lines, max(len(heads), 6))
    _header(ws, HEADER_ROW, heads)
    for ri, row in enumerate(body):
        for ci, v in enumerate(row, start=1):
            cell = ws.cell(row=HEADER_ROW + 1 + ri, column=ci)
            put(cell, _value_for(types[ci - 1], v))
            _style_cell(cell, types[ci - 1], ri % 2 == 1)
    tot = res.get("totals") or {}
    if body and kind != "table" and len(res.get("datasets") or []) == 1 and (res.get("measure") or {}).get("key") in tot:
        r = HEADER_ROW + 1 + len(body)
        put(ws.cell(row=r, column=1), "Total (all records)")
        c = ws.cell(row=r, column=2)
        put(c, _value_for(types[1], tot[res["measure"]["key"]]))
        _style_cell(c, types[1], False)
        for cc in (ws.cell(row=r, column=1), c):
            cc.font = Font(size=10.5, bold=True)
            cc.border = TOTAL_BORDER
    _widths(ws, heads, body, [{}] * len(heads))
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 26)
    ws.freeze_panes = f"B{HEADER_ROW + 1}"
    if body:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(heads))}{HEADER_ROW + len(body)}"
    if kind != "table":
        _table_chart(ws, kind, spec["title"], HEADER_ROW, len(body), len(res.get("datasets") or []), len(heads) + 2)
    if res.get("truncated"):
        put(ws.cell(row=HEADER_ROW + len(body) + 3, column=1), "Only the largest groups are shown on the dashboard.")
    _print_setup(ws, spec["title"], HEADER_ROW)


def build_dashboard_workbook(page_result: dict, specs: list[dict], records: dict | None, meta: dict) -> bytes:
    """page_result: dashboards.evaluate output. specs: [{id, kind, title, definition}] in page order
    (KPIs first). records: Records sheet content (detail dict) or None."""
    company = meta.get("company") or "Premnath Rail Tools"
    wb = Workbook()
    first = wb.active
    first.title = "Summary"
    _title_block(first, company, f"{meta['module_label']} Reports - {page_result['label']}", describe_lines(page_result, meta), 4)
    kpis = []
    for sp in specs:
        k = page_result["kpis"].get(sp["id"])
        if k:
            kpis.append({"label": sp["title"], "type": k["type"], "value": k.get("value"), "compare": k.get("compare"),
                         "delta_pct": k.get("delta_pct"), "better": k.get("better")})
    nxt = _kpi_sheet(first, kpis, HEADER_ROW + 1) if kpis else HEADER_ROW + 1
    for n in page_result.get("notes") or []:
        put(first.cell(row=nxt, column=1), n)
        nxt += 1
    _print_setup(first, f"{meta['module_label']} Reports - {page_result['label']}", None)
    taken = {"summary", "definitions", "records"}
    for sp in specs:
        res = page_result["results"].get(sp["id"])
        if res:
            _write_visual(wb, safe_sheet_name(sp["title"], taken), company, sp, res, page_result, meta)
    if records is not None:
        _write_records(wb, "Records", "Records", company,
                       {**records, **{k: page_result.get(k) for k in ("period", "filters_applied", "scope_label")}}, meta)
    ws = wb.create_sheet("Definitions")
    _title_block(ws, company, f"{meta['module_label']} Reports - How the figures are calculated", [], 3)
    _header(ws, HEADER_ROW, ["Figure", "How it is calculated"])
    r = HEADER_ROW + 1
    for sp in specs:
        if sp["id"] in page_result["kpis"] or sp["id"] in page_result["results"]:
            for ci, v in enumerate((sp["title"], sp.get("definition") or ""), start=1):
                cell = ws.cell(row=r, column=ci)
                put(cell, v)
                _style_cell(cell, "text", r % 2 == 0, wrap=(ci == 2))
            r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    _print_setup(ws, "Definitions", HEADER_ROW)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
